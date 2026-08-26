from __future__ import annotations

import pytest

from datetime import datetime, timezone
from docx import Document
from openpyxl import Workbook

from app.domains.conversations.contracts import ArtifactRecord, AttachmentRecord
from app.domains.conversations.repository import ConversationRepository
from app.domains.identity import IdentityRepository
from app.domains.resources import ResourceLocator, ResourceWorkspace
from app.domains.skills import SkillDraft, SkillRegistry
from app.settings import Settings


def _repositories(tmp_path):
    settings = Settings(data_dir=tmp_path, database_path=tmp_path / "unihr.db")
    identity = IdentityRepository(settings.database_path, settings)
    conversations = ConversationRepository(
        settings.database_path, default_tenant_id=settings.default_tenant_id,
    )
    return settings, identity, conversations


def test_user_data_isolation_and_admin_history_scope(tmp_path):
    _, identity, repository = _repositories(tmp_path)
    admin = identity.login("admin", "UniHR@2026").user
    user = identity.login("hr-user", "UniHR@2026").user

    admin_conversation = repository.create_conversation("管理员会话", admin)
    user_conversation = repository.create_conversation("普通用户会话", user)

    assert [item.id for item in repository.list_conversations(user)] == [user_conversation.id]
    with pytest.raises(KeyError):
        repository.get_conversation(admin_conversation.id, user)

    history = repository.admin_list_history(admin)
    assert {item.id for item in history} == {admin_conversation.id, user_conversation.id}
    with pytest.raises(PermissionError):
        repository.admin_list_history(user)


def test_feedback_is_upserted_per_answer_and_can_be_collected_for_testing(tmp_path):
    _, identity, repository = _repositories(tmp_path)
    admin = identity.login("admin", "UniHR@2026").user
    user = identity.login("hr-user", "UniHR@2026").user
    conversation = repository.create_conversation("政策问答", user)
    repository.add_message(conversation.id, "user", "这项政策适用吗？", principal=user)
    answer = repository.add_message(conversation.id, "assistant", "请以现行文件为准。", principal=user)

    first = repository.upsert_feedback(answer.id, True, "有依据", user)
    second = repository.upsert_feedback(answer.id, False, "还缺页码", user)

    assert first.id == second.id
    assert second.helpful is False
    assert second.comment == "还缺页码"
    assert repository.admin_feedback_stats(admin).not_helpful == 1
    collected = repository.admin_list_feedback(admin)
    assert collected[0].question == "这项政策适用吗？"
    assert collected[0].answer == "请以现行文件为准。"
    reviewed = repository.admin_update_feedback_status(collected[0].id, "added_to_test", admin)
    assert reviewed.test_status == "added_to_test"


def test_user_cannot_rate_another_users_answer(tmp_path):
    _, identity, repository = _repositories(tmp_path)
    admin = identity.login("admin", "UniHR@2026").user
    user = identity.login("hr-user", "UniHR@2026").user
    conversation = repository.create_conversation("管理员会话", admin)
    answer = repository.add_message(conversation.id, "assistant", "内部回答", principal=admin)

    with pytest.raises(KeyError):
        repository.upsert_feedback(answer.id, True, principal=user)


def test_user_skills_are_private_while_system_skills_are_shared(tmp_path):
    settings, identity, _ = _repositories(tmp_path)
    registry = SkillRegistry(settings.database_path, settings.default_tenant_id)
    admin = identity.login("admin", "UniHR@2026").user
    user = identity.login("hr-user", "UniHR@2026").user
    custom = registry.create(SkillDraft(
        name="private-hr-workflow",
        display_name="个人流程",
        description="只属于当前用户",
        instructions="读取输入，检查证据，再输出草稿。",
        trigger_phrases=["个人流程"],
    ), user)

    assert custom.id in {item.id for item in registry.list(user)}
    assert custom.id not in {item.id for item in registry.list(admin)}
    system = next(item for item in registry.list(user) if item.source_type == "system")

    enabled = registry.update(custom.id, enabled=True, principal=user)
    assert enabled.enabled is True and enabled.status == "active"
    assert registry.select("请执行个人流程", principal=user)[0].skill.id == custom.id
    with pytest.raises(ValueError, match="system_skill_is_read_only"):
        registry.update(system.id, enabled=False, principal=user)


def test_resource_workspace_links_generated_artifact_to_source_files(tmp_path):
    _, identity, repository = _repositories(tmp_path)
    user = identity.login("hr-user", "UniHR@2026").user
    conversation = repository.create_conversation("表格任务", user)
    user_message = repository.add_message(conversation.id, "user", "处理工资表", principal=user)
    assistant_message = repository.add_message(conversation.id, "assistant", "已生成", principal=user)
    now = datetime.now(timezone.utc).isoformat()
    source = tmp_path / "工资表.xlsx"
    source.write_bytes(b"source")
    output = tmp_path / "结果.xlsx"
    output.write_bytes(b"result")
    repository.add_attachment(AttachmentRecord(
        id="att_source",
        conversation_id=conversation.id,
        message_id=user_message.id,
        display_name=source.name,
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        extension=".xlsx",
        byte_size=source.stat().st_size,
        saved_path=str(source),
        created_at=now,
    ), user)
    repository.add_artifact(ArtifactRecord(
        id="art_result",
        conversation_id=conversation.id,
        message_id=assistant_message.id,
        kind="spreadsheet",
        display_name=output.name,
        path=str(output),
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        byte_size=output.stat().st_size,
        metadata={
            "source_files": [source.name],
            "artifact_harness": {"passed": True, "status": "passed"},
        },
        created_at=now,
    ), user)

    manifest = ResourceWorkspace(repository).build(conversation.id, user)

    assert manifest.source_count == 1
    assert manifest.artifact_count == 1
    artifact = next(item for item in manifest.items if item.resource_type == "artifact")
    assert artifact.parent_resource_ids == ["att_source"]
    assert artifact.status == "passed"
    assert repository.get_attachment("att_source", user).display_name == source.name


def test_resource_workspace_reads_only_requested_excel_range(tmp_path):
    _, identity, repository = _repositories(tmp_path)
    user = identity.login("hr-user", "UniHR@2026").user
    conversation = repository.create_conversation("局部读取表格", user)
    message = repository.add_message(conversation.id, "user", "读取工资表", principal=user)
    source = tmp_path / "工资表.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "调整测算"
    sheet.append(["姓名", "部门", "调整额"])
    sheet.append(["甲", "研发", 1200])
    sheet.append(["乙", "人事", 800])
    workbook.save(source)
    repository.add_attachment(AttachmentRecord(
        id="att_sheet_query", conversation_id=conversation.id, message_id=message.id,
        display_name=source.name,
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        extension=".xlsx", byte_size=source.stat().st_size, saved_path=str(source),
        summary={"kind": "workbook", "sheets": [{"name": "调整测算", "rows": 3, "columns": 3}]},
        created_at=datetime.now(timezone.utc).isoformat(),
    ), user)

    excerpt = ResourceWorkspace(repository).query(
        "att_sheet_query",
        ResourceLocator(sheet="调整测算", cell_range="B2:C3"),
        user,
    )

    assert excerpt.text == "研发 | 1200\n人事 | 800"
    assert excerpt.metadata["range"] == "B2:C3"
    assert excerpt.truncated is False


def test_resource_workspace_versions_and_archives_owner_scoped_artifacts(tmp_path):
    _, identity, repository = _repositories(tmp_path)
    owner = identity.login("hr-user", "UniHR@2026").user
    other = identity.login("admin", "UniHR@2026").user
    conversation = repository.create_conversation("资源版本", owner)
    user_message = repository.add_message(conversation.id, "user", "生成两版名单", principal=owner)
    assistant_message = repository.add_message(conversation.id, "assistant", "已生成", principal=owner)
    now = datetime.now(timezone.utc).isoformat()
    for index, name in enumerate(("候选人名单-v1.xlsx", "候选人名单-v2.xlsx"), start=1):
        path = tmp_path / name
        workbook = Workbook()
        workbook.active.append(["姓名", f"版本{index}"])
        workbook.save(path)
        repository.add_artifact(ArtifactRecord(
            id=f"art_version_{index}", conversation_id=conversation.id,
            message_id=assistant_message.id, kind="spreadsheet", display_name=name,
            path=str(path), mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            byte_size=path.stat().st_size, metadata={}, created_at=f"{now}-{index}",
        ), owner)

    workspace = ResourceWorkspace(repository)
    manifest = workspace.build(conversation.id, owner)
    versions = sorted(
        (item for item in manifest.items if item.resource_type == "artifact"),
        key=lambda item: item.version_number,
    )
    assert [item.version_number for item in versions] == [1, 2]
    assert versions[0].lineage_key == versions[1].lineage_key
    assert versions[0].is_current is False
    assert versions[1].is_current is True

    archived = workspace.set_archived("art_version_1", True, owner)
    assert archived.archived_at
    assert [item.id for item in workspace.build(conversation.id, owner).items] == ["art_version_2"]
    assert workspace.build(conversation.id, owner, include_archived=True).archived_count == 1
    with pytest.raises(KeyError):
        workspace.set_archived("art_version_2", True, other)

    restored = workspace.set_archived("art_version_1", False, owner)
    assert restored.archived_at == ""


def test_resource_workspace_reads_named_document_section_and_enforces_owner(tmp_path):
    _, identity, repository = _repositories(tmp_path)
    owner = identity.login("hr-user", "UniHR@2026").user
    other = identity.login("admin", "UniHR@2026").user
    conversation = repository.create_conversation("局部读取公文", owner)
    message = repository.add_message(conversation.id, "user", "读取第二章", principal=owner)
    source = tmp_path / "政策.docx"
    document = Document()
    document.add_heading("第一章 总则", level=1)
    document.add_paragraph("第一章内容不应出现在结果中。")
    document.add_heading("第二章 调整规则", level=1)
    document.add_paragraph("公积金调整自2026年7月起执行。")
    document.add_heading("第三章 附则", level=1)
    document.add_paragraph("第三章内容不应出现在结果中。")
    document.save(source)
    repository.add_attachment(AttachmentRecord(
        id="att_doc_query", conversation_id=conversation.id, message_id=message.id,
        display_name=source.name,
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        extension=".docx", byte_size=source.stat().st_size, saved_path=str(source),
        summary={"kind": "document", "paragraphs": 6, "tables": 0},
        created_at=datetime.now(timezone.utc).isoformat(),
    ), owner)

    excerpt = ResourceWorkspace(repository).query(
        "att_doc_query", ResourceLocator(section="第二章"), owner,
    )

    assert "公积金调整自2026年7月起执行" in excerpt.text
    assert "第一章内容" not in excerpt.text
    assert "第三章内容" not in excerpt.text
    with pytest.raises(KeyError):
        ResourceWorkspace(repository).query("att_doc_query", ResourceLocator(section="第二章"), other)
