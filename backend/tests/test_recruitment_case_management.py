from datetime import datetime, timezone

import pytest
from docx import Document
from openpyxl import load_workbook

from app.domains.assistant.contracts import BusinessIntent, ExecutionMode, TurnAction
from app.domains.conversations.contracts import ArtifactRecord
from app.domains.conversations.repository import ConversationRepository
from app.domains.identity import IdentityRepository, Principal, RUN_APPROVE
from app.domains.recruitment import (
    CandidateReviewRecord,
    RecruitmentCaseAction,
    RecruitmentCaseCreate,
    RecruitmentCaseModule,
    RecruitmentEvidencePackageModule,
)
from app.domains.runs import RunRepository, WorkflowRunRequest
from app.settings import Settings


def _principal(user_id: str, tenant_id: str = "tenant_a", *, admin: bool = False) -> Principal:
    return Principal(
        user_id=user_id,
        username=user_id,
        display_name="管理员" if admin else "招聘专员",
        role="admin" if admin else "user",
        tenant_id=tenant_id,
        permissions=[RUN_APPROVE] if admin else [],
    )


def _open_case(tmp_path):
    database = tmp_path / "unihr.db"
    IdentityRepository(database, Settings(
        database_path=database, data_dir=tmp_path,
        default_tenant_id="tenant_a", default_tenant_name="测试组织",
    ))
    owner = _principal("owner")
    conversations = ConversationRepository(database)
    conversation = conversations.create_conversation("研发实习生筛选", owner)
    message = conversations.add_message(conversation.id, "assistant", "筛选完成", principal=owner)
    runs = RunRepository(database)
    runs.start(WorkflowRunRequest(
        run_id="run_recruitment",
        tenant_id=owner.tenant_id,
        owner_user_id=owner.user_id,
        conversation_id=conversation.id,
        business_intent=BusinessIntent.RECRUITMENT_SCREENING,
        turn_action=TurnAction.START_NEW,
        execution_mode=ExecutionMode.EXECUTE,
        max_iterations=2,
    ))
    now = datetime.now(timezone.utc).isoformat()
    artifact = ArtifactRecord(
        id="artifact_recruitment",
        conversation_id=conversation.id,
        message_id=message.id,
        kind="recruitment_report",
        display_name="候选人筛选报告.xlsx",
        path=str(tmp_path / "report.xlsx"),
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        byte_size=1024,
        created_at=now,
    )
    conversations.add_artifact(artifact, owner)
    module = RecruitmentCaseModule(database)
    record = module.open(RecruitmentCaseCreate(
        case_id="recruitment_case_1",
        tenant_id=owner.tenant_id,
        owner_user_id=owner.user_id,
        conversation_id=conversation.id,
        run_id="run_recruitment",
        artifact_id=artifact.id,
        job_title="研发实习生",
        rubric_hash="rubric_abc",
        selection_count=2,
        candidates=[
            CandidateReviewRecord(
                candidate_id="candidate_a", candidate_name="甲", source_name="甲.pdf",
                rank=1, ai_score=88, ai_recommendation="优先复核", evidence_coverage=.9,
            ),
            CandidateReviewRecord(
                candidate_id="candidate_b", candidate_name="乙", source_name="乙.pdf",
                rank=2, ai_score=80, ai_recommendation="重点核实", evidence_coverage=.8,
            ),
        ],
    ))
    return module, owner, record


def test_candidate_review_interview_and_final_decision_are_persisted(tmp_path):
    module, owner, record = _open_case(tmp_path)
    reviewed = module.apply(record.id, RecruitmentCaseAction(
        action="review_candidate", candidate_id="candidate_a",
        review_status="shortlisted", reviewer_note="B2B 证据充分，进入结构化面试",
    ), owner)
    assert reviewed.status == "interviewing"
    assert reviewed.candidates[0].review_status == "shortlisted"
    assert reviewed.candidates[0].reviewed_by == owner.display_name

    interviewed = module.apply(record.id, RecruitmentCaseAction(
        action="record_interview", candidate_id="candidate_a",
        interview_outcome="passed", interview_note="案例回答包含明确个人贡献和量化结果",
    ), owner)
    assert interviewed.candidates[0].interview_outcome == "passed"
    assert interviewed.candidates[0].interview_notes[0].recorded_by == owner.display_name

    finalized = module.apply(record.id, RecruitmentCaseAction(
        action="finalize", final_candidate_ids=["candidate_a"],
        final_rationale="结合简历证据与结构化面试，建议进入下一轮。",
    ), owner)
    assert finalized.status == "finalized"
    assert finalized.final_candidate_ids == ["candidate_a"]
    assert finalized.finalized_by == owner.display_name


def test_case_is_owner_scoped_but_tenant_admin_can_access(tmp_path):
    module, owner, record = _open_case(tmp_path)
    with pytest.raises(KeyError):
        module.get(record.id, _principal("another_user"))
    with pytest.raises(KeyError):
        module.get(record.id, _principal("other_tenant_admin", "tenant_b", admin=True))
    assert module.get(record.id, _principal("tenant_admin", admin=True)).id == record.id
    assert module.get(record.id, owner).id == record.id
    assert [item.id for item in module.list(owner)] == [record.id]
    assert module.list(_principal("another_user")) == []
    assert [item.id for item in module.list(_principal("tenant_admin", admin=True))] == [record.id]
    assert module.list(_principal("other_tenant_admin", "tenant_b", admin=True)) == []


def test_case_list_supports_governed_status_filter(tmp_path):
    module, owner, record = _open_case(tmp_path)
    assert [item.id for item in module.list(owner, status="reviewing")] == [record.id]
    assert module.list(owner, status="finalized") == []
    with pytest.raises(ValueError, match="recruitment_case_status_invalid"):
        module.list(owner, status="unknown")


def test_case_folders_search_rename_move_archive_and_restore(tmp_path):
    module, owner, record = _open_case(tmp_path)
    folder = module.create_folder("2026 校招", owner)
    assert [item.name for item in module.list_folders(owner)] == ["2026 校招"]

    renamed = module.apply(record.id, RecruitmentCaseAction(
        action="rename_case", case_name="研发中心实习生 Case 01",
    ), owner)
    moved = module.apply(record.id, RecruitmentCaseAction(
        action="move_case", folder_id=folder.id,
    ), owner)
    assert renamed.case_name == "研发中心实习生 Case 01"
    assert moved.folder_id == folder.id
    assert [item.id for item in module.list(owner, folder_id=folder.id)] == [record.id]
    assert [item.id for item in module.list(owner, query="Case 01")] == [record.id]

    archived = module.apply(record.id, RecruitmentCaseAction(action="archive_case"), owner)
    assert archived.archived_at
    assert module.list(owner) == []
    assert [item.id for item in module.list(owner, archived=True)] == [record.id]

    restored = module.apply(record.id, RecruitmentCaseAction(action="restore_case"), owner)
    assert restored.archived_at == ""
    module.delete_folder(folder.id, owner)
    assert module.get(record.id, owner).folder_id == ""
    assert module.list_folders(owner) == []


def test_case_folder_is_owner_scoped(tmp_path):
    module, owner, _ = _open_case(tmp_path)
    folder = module.create_folder("我的项目", owner)
    with pytest.raises(KeyError):
        module.rename_folder(folder.id, "越权修改", _principal("another_user"))
    with pytest.raises(ValueError, match="recruitment_folder_name_exists"):
        module.create_folder("我的项目", owner)


def test_final_decision_is_bounded_and_requires_explicit_rationale(tmp_path):
    module, owner, record = _open_case(tmp_path)
    with pytest.raises(ValueError, match="final_candidates_required"):
        module.apply(record.id, RecruitmentCaseAction(
            action="finalize", final_candidate_ids=[], final_rationale="有依据",
        ), owner)
    with pytest.raises(ValueError, match="final_rationale_required"):
        module.apply(record.id, RecruitmentCaseAction(
            action="finalize", final_candidate_ids=["candidate_a"], final_rationale="",
        ), owner)
    with pytest.raises(ValueError, match="candidate_not_found"):
        module.apply(record.id, RecruitmentCaseAction(
            action="finalize", final_candidate_ids=["missing"], final_rationale="有依据",
        ), owner)


def test_open_is_idempotent_for_the_same_run(tmp_path):
    module, owner, record = _open_case(tmp_path)
    existing = module.open(RecruitmentCaseCreate(
        case_id="different_case_id",
        tenant_id=owner.tenant_id,
        owner_user_id=owner.user_id,
        conversation_id=record.conversation_id,
        run_id=record.run_id,
        artifact_id=record.artifact_id,
        job_title="不应覆盖",
        rubric_hash="different",
        selection_count=1,
        candidates=[],
    ))
    assert existing.id == record.id
    assert existing.job_title == "研发实习生"


def test_multi_interviewer_schedule_feedback_permissions_and_calendar(tmp_path):
    module, owner, record = _open_case(tmp_path)
    module.apply(record.id, RecruitmentCaseAction(
        action="review_candidate", candidate_id="candidate_a",
        review_status="shortlisted", reviewer_note="进入结构化面试",
    ), owner)
    scheduled = module.apply(record.id, RecruitmentCaseAction(
        action="schedule_interview", candidate_id="candidate_a",
        interview_title="甲 · 产业合作结构化面试",
        interview_start_at="2026-08-12T10:00:00+08:00",
        interview_end_at="2026-08-12T11:00:00+08:00",
        interview_timezone="Asia/Shanghai", interview_mode="online",
        interview_meeting_url="https://meeting.example/interview-a",
        interviewer_ids=["usr_unihr_demo", "usr_unihr_admin"],
        focus_dimensions=["业务拓展", "数据分析"],
    ), owner)
    session = scheduled.interview_sessions[0]
    assert scheduled.candidates[0].interview_outcome == "scheduled"
    assert scheduled.collaborator_user_ids == ["usr_unihr_admin", "usr_unihr_demo"]
    assert [item.user_id for item in session.interviewers] == ["usr_unihr_demo", "usr_unihr_admin"]

    interviewer_a = _principal("usr_unihr_demo")
    interviewer_a.display_name = "HR 用户"
    interviewer_b = _principal("usr_unihr_admin")
    interviewer_b.display_name = "业务面试官"
    assert module.get(record.id, interviewer_a).id == record.id
    assert [item.id for item in module.list(interviewer_b)] == [record.id]
    with pytest.raises(ValueError, match="recruitment_case_owner_required"):
        module.apply(record.id, RecruitmentCaseAction(
            action="cancel_interview", interview_session_id=session.id,
        ), interviewer_a)

    first = module.apply(record.id, RecruitmentCaseAction(
        action="submit_interview_feedback", interview_session_id=session.id,
        recommendation="yes", dimension_scores={"业务拓展": 5, "数据分析": 4},
        evidence_note="能说明独立拓客步骤和本人贡献。", risk_note="行业经验仍需验证。",
    ), interviewer_a)
    assert first.interview_sessions[0].status == "scheduled"
    assert len(first.interview_sessions[0].feedback) == 1

    second = module.apply(record.id, RecruitmentCaseAction(
        action="submit_interview_feedback", interview_session_id=session.id,
        recommendation="hold", dimension_scores={"业务拓展": 4, "数据分析": 3},
        evidence_note="案例完整，但量化结果解释不足。", risk_note="建议补充背调。",
    ), interviewer_b)
    assert second.interview_sessions[0].status == "completed"
    assert second.candidates[0].interview_outcome == "passed"
    assert len(second.interview_sessions[0].feedback) == 2

    filename, calendar = module.calendar_event(record.id, session.id, interviewer_a)
    assert filename == "甲-面试安排.ics"
    assert "DTSTART:20260812T020000Z" in calendar
    assert "SUMMARY:甲 · 产业合作结构化面试" in calendar
    assert "核实重点：业务拓展、数据分析" in calendar


def test_interview_schedule_rejects_unassigned_users_and_unstructured_feedback(tmp_path):
    module, owner, record = _open_case(tmp_path)
    module.apply(record.id, RecruitmentCaseAction(
        action="review_candidate", candidate_id="candidate_a", review_status="shortlisted",
    ), owner)
    with pytest.raises(ValueError, match="interviewers_invalid"):
        module.apply(record.id, RecruitmentCaseAction(
            action="schedule_interview", candidate_id="candidate_a",
            interview_start_at="2026-08-12T10:00:00+08:00",
            interview_end_at="2026-08-12T11:00:00+08:00", interview_mode="online",
            interview_meeting_url="https://meeting.example/a", interviewer_ids=["missing"],
            focus_dimensions=["业务拓展"],
        ), owner)

    scheduled = module.apply(record.id, RecruitmentCaseAction(
        action="schedule_interview", candidate_id="candidate_a",
        interview_start_at="2026-08-12T10:00:00+08:00",
        interview_end_at="2026-08-12T11:00:00+08:00", interview_mode="phone",
        interviewer_ids=["usr_unihr_demo"], focus_dimensions=["业务拓展", "协作"],
    ), owner)
    session_id = scheduled.interview_sessions[0].id
    with pytest.raises(ValueError, match="interview_feedback_assignment_required"):
        module.apply(record.id, RecruitmentCaseAction(
            action="submit_interview_feedback", interview_session_id=session_id,
            recommendation="yes", dimension_scores={"业务拓展": 4, "协作": 4},
            evidence_note="有证据",
        ), owner)
    with pytest.raises(ValueError, match="interview_dimension_scores_invalid"):
        module.apply(record.id, RecruitmentCaseAction(
            action="submit_interview_feedback", interview_session_id=session_id,
            recommendation="yes", dimension_scores={"业务拓展": 4}, evidence_note="有证据",
        ), _principal("usr_unihr_demo"))


def _packages(tmp_path, module: RecruitmentCaseModule) -> RecruitmentEvidencePackageModule:
    return RecruitmentEvidencePackageModule(
        module.database_path,
        tmp_path / "recruitment-packages",
        module,
        ConversationRepository(module.database_path),
    )


def test_interview_evaluation_package_is_generated_persisted_and_downloadable(tmp_path):
    module, owner, record = _open_case(tmp_path)
    packages = _packages(tmp_path, module)

    preview = packages.preview(record.id, "interview_evaluation", owner)
    assert preview.can_generate is True
    assert preview.candidate_count == 2
    assert preview.final_candidate_count == 0

    generated = packages.generate(record.id, "interview_evaluation", owner)
    assert generated.package_type == "interview_evaluation"
    assert generated.harness["passed"] is True
    assert {item.format for item in generated.files} == {"docx", "xlsx"}
    assert all(item.artifact_id for item in generated.files)
    assert [item.id for item in packages.list(record.id, owner)] == [generated.id]

    docx_path, _, _ = packages.download(record.id, generated.id, "docx", owner)
    xlsx_path, _, _ = packages.download(record.id, generated.id, "xlsx", owner)
    assert "结构化面试评价包" in "\n".join(item.text for item in Document(docx_path).paragraphs)
    workbook = load_workbook(xlsx_path, data_only=False, keep_links=False)
    assert [item.title for item in workbook.worksheets if item.sheet_state == "visible"] == [
        "使用说明", "候选人总览", "面试评价模板", "多人反馈",
        "最终人工建议", "证据索引", "审计信息",
    ]
    assert workbook["PackageData"].sheet_state == "hidden"
    assert workbook["候选人总览"]["I5"].value.startswith("=IF(COUNTIF")
    assert workbook.calculation.fullCalcOnLoad is True
    workbook.close()


def test_final_decision_package_is_gated_until_human_finalization(tmp_path):
    module, owner, record = _open_case(tmp_path)
    packages = _packages(tmp_path, module)
    preview = packages.preview(record.id, "final_decision", owner)
    assert preview.can_generate is False
    assert preview.blockers == ["recruitment_case_not_finalized"]
    with pytest.raises(ValueError, match="recruitment_case_not_finalized"):
        packages.generate(record.id, "final_decision", owner)

    module.apply(record.id, RecruitmentCaseAction(
        action="finalize", final_candidate_ids=["candidate_a"],
        final_rationale=(
            "结合简历事实证据和两位面试官的结构化反馈，人工建议候选人甲进入后续。"
            "其产业合作流程经验和本人贡献证据较完整；下一轮继续核实独立谈判深度、"
            "主动拓客能力及到岗稳定性。本结论由HR人工复核形成，不以AI分数替代最终就业决定。"
        ),
    ), owner)
    generated = packages.generate(record.id, "final_decision", owner)
    xlsx_path, _, _ = packages.download(record.id, generated.id, "xlsx", owner)
    workbook = load_workbook(xlsx_path, data_only=False, keep_links=False)
    final_sheet = workbook["最终人工建议"]
    assert final_sheet["A5"].value == "candidate_a"
    assert final_sheet["C5"].value == "最终建议人选"
    assert "人工建议候选人甲" in final_sheet["D5"].value
    assert final_sheet.row_dimensions[5].height > 50
    workbook.close()


def test_package_generation_requires_case_owner_or_tenant_admin(tmp_path):
    module, owner, record = _open_case(tmp_path)
    packages = _packages(tmp_path, module)
    collaborator = _principal("usr_unihr_demo")
    module.apply(record.id, RecruitmentCaseAction(
        action="review_candidate", candidate_id="candidate_a", review_status="shortlisted",
    ), owner)
    module.apply(record.id, RecruitmentCaseAction(
        action="schedule_interview", candidate_id="candidate_a",
        interview_start_at="2026-08-12T10:00:00+08:00",
        interview_end_at="2026-08-12T11:00:00+08:00", interview_mode="phone",
        interviewer_ids=[collaborator.user_id], focus_dimensions=["业务拓展"],
    ), owner)
    preview = packages.preview(record.id, "interview_evaluation", collaborator)
    assert preview.can_generate is False
    assert preview.blockers == ["recruitment_case_owner_required"]
    with pytest.raises(ValueError, match="recruitment_case_owner_required"):
        packages.generate(record.id, "interview_evaluation", collaborator)

    admin = _principal("tenant_admin", admin=True)
    generated = packages.generate(record.id, "interview_evaluation", admin)
    assert all(not item.artifact_id for item in generated.files)
    assert packages.download(record.id, generated.id, "docx", collaborator)[0].is_file()
