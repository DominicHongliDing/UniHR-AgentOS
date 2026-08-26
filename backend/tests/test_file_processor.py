from __future__ import annotations

from docx import Document
from openpyxl import Workbook

from app.domains.files import FileProcessor, LegacyOfficeConverter


def _fake_runner(source, target, kind, _timeout):
    assert source.is_file()
    if kind == "word":
        document = Document()
        document.add_paragraph("旧版 Word 政策正文")
        document.save(target)
    else:
        workbook = Workbook()
        workbook.active.append(["工号", "金额"])
        workbook.active.append(["001", 200000])
        workbook.save(target)
    return {"application_version": "test"}


def test_legacy_doc_is_converted_to_safe_docx_before_extraction(tmp_path):
    converter = LegacyOfficeConverter(tmp_path / "converted", runner=_fake_runner)
    processor = FileProcessor(tmp_path / "uploads", 1024 * 1024, legacy_converter=converter)

    result = processor.process("旧版政策.doc", "application/msword", b"legacy-binary-placeholder")

    assert result.display_name == "旧版政策.doc"
    assert result.extension == ".docx"
    assert result.saved_path.suffix == ".docx"
    assert "旧版 Word 政策正文" in result.extracted_text
    assert result.summary["original_extension"] == ".doc"
    assert result.summary["converted_format"] == ".docx"
    assert result.summary["macros_removed"] is True
    assert not list((tmp_path / "uploads").glob("*.doc"))


def test_legacy_xls_is_converted_to_xlsx_and_enters_spreadsheet_path(tmp_path):
    converter = LegacyOfficeConverter(tmp_path / "converted", runner=_fake_runner)
    processor = FileProcessor(tmp_path / "uploads", 1024 * 1024, legacy_converter=converter)

    result = processor.process("旧版工资表.xls", "application/vnd.ms-excel", b"legacy-binary-placeholder")

    assert result.extension == ".xlsx"
    assert result.mime_type.endswith("spreadsheetml.sheet")
    assert "工号 | 金额" in result.extracted_text
    assert result.summary["original_extension"] == ".xls"
    assert result.summary["conversion_engine"] == "Microsoft Excel COM"
