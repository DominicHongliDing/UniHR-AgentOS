import re
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm
from openpyxl import Workbook, load_workbook
from openpyxl.cell.read_only import EMPTY_CELL

from app.domains.artifacts.documents import (
    DocumentArtifactBuilder,
    DocumentDraft,
    DraftSection,
    DraftTable,
)
from app.domains.documents.harness import DocumentLayoutHarness
from app.domains.artifacts.harness import ArtifactHarness
from app.domains.artifacts.excel_verification import ExcelAutomationVerifier
from app.domains.artifacts.spreadsheet_engine import SpreadsheetPlanError, SpreadsheetPlanExecutor
from app.domains.artifacts.spreadsheet_plans import SpreadsheetPlan
from app.domains.artifacts.spreadsheets import SpreadsheetArtifactBuilder
from app.domains.task_briefs import DocumentStyleSpec, SpreadsheetStyleSpec


def test_excel_formula_cache_inspection_accepts_read_only_empty_cells(tmp_path):
    formulas_path = tmp_path / "ragged.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["C1"] = "=1+1"
    workbook.save(formulas_path)
    workbook.close()

    # This is the exact object openpyxl returns for empty positions in a ragged
    # read-only row. It deliberately has no coordinate attribute.
    assert not hasattr(EMPTY_CELL, "coordinate")
    checks: dict[str, bool] = {}
    metrics: dict[str, int | float | str | bool] = {}
    issues: list[str] = []

    assert ExcelAutomationVerifier._inspect_formula_cache(
        formulas_path, checks, metrics, issues,
    ) is True
    assert metrics["formula_count_after_recalculation"] == 1
    assert "formula_result_error" not in issues


def test_document_builder_applies_selected_chinese_font_and_a4(tmp_path):
    builder = DocumentArtifactBuilder(tmp_path)
    path, metadata = builder.build(
        DocumentDraft(
            title="关于开展年度人事材料核对工作的通知",
            document_type="通知",
            addressee="各部门",
            lead="为确保年度材料准确完整，现将有关事项通知如下。",
            sections=[DraftSection("一、工作安排", ["请于规定时间内完成材料核对。"])],
        ),
        body_font="宋体",
        body_size_pt=14,
    )

    document = Document(path)
    section = document.sections[0]
    assert round(section.page_width.cm, 1) == 21.0
    assert round(section.page_height.cm, 1) == 29.7
    assert metadata["body_font"] == "宋体"
    body_run = next(run for paragraph in document.paragraphs for run in paragraph.runs if "确保年度材料" in run.text)
    assert body_run._element.rPr.rFonts.get(qn("w:eastAsia")) == "宋体"


def test_document_builder_scales_wide_table_to_printable_page_width(tmp_path):
    builder = DocumentArtifactBuilder(tmp_path)
    path, metadata = builder.build(
        DocumentDraft(
            title="Candidate interview confirmation",
            tables=[
                DraftTable(
                    title="Candidate summary",
                    headers=["Name", "Rank", "Score", "Strengths", "Risks", "Review"],
                    rows=[["Candidate A", "1", "88", "Evidence", "Verify", ""]],
                    column_widths_cm=[2.5, 2.0, 2.0, 6.0, 6.0, 4.0],
                )
            ],
        )
    )

    document = Document(path)
    section = document.sections[0]
    usable_width_cm = section.page_width.cm - section.left_margin.cm - section.right_margin.cm
    first_row_width_cm = sum(cell.width.cm for cell in document.tables[0].rows[0].cells)

    assert first_row_width_cm <= usable_width_cm + 0.05
    assert metadata["validation"]["checks"]["tables_within_page_width"] is True


def test_document_layout_harness_rejects_table_wider_than_printable_page(tmp_path):
    builder = DocumentArtifactBuilder(tmp_path)
    path, _ = builder.build(
        DocumentDraft(
            title="Wide table guard",
            tables=[
                DraftTable(
                    title="Review table",
                    headers=["A", "B", "C"],
                    rows=[["1", "2", "3"]],
                    column_widths_cm=[3.0, 3.0, 3.0],
                )
            ],
        )
    )
    document = Document(path)
    table = document.tables[0]
    table.autofit = False
    for row in table.rows:
        for cell in row.cells:
            cell.width = Cm(8.0)
    document.save(path)

    report = DocumentLayoutHarness().validate(path, DocumentStyleSpec())

    assert report.passed is False
    assert report.checks["tables_within_page_width"] is False


def test_spreadsheet_builder_preserves_formula_and_adds_audit_sheet(tmp_path):
    source = tmp_path / "2025-template.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "测算"
    sheet.append(["姓名", "基数", "比例", "结果"])
    sheet.append(["匿名A", 1000, 0.12, "=B2*C2"])
    workbook.save(source)
    workbook.close()

    result = SpreadsheetArtifactBuilder(tmp_path / "out").process(
        source, source.name, "生成 2026 年草稿"
    )

    output = load_workbook(result.path, data_only=False)
    assert output["测算"]["D2"].value == "=B2*C2"
    assert output.sheetnames[0] == "UniHR处理说明"
    assert "未覆盖原文件" in output["UniHR处理说明"]["B2"].value
    output.close()


def test_artifact_harness_opens_actual_workbook_and_reports_structure(tmp_path):
    source = tmp_path / "harness.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "数据"
    sheet.append(["金额", "税额"])
    sheet.append([100, "=A2*0.06"])
    workbook.calculation.fullCalcOnLoad = True
    workbook.save(source)
    workbook.close()

    report = ArtifactHarness().inspect(
        source,
        kind="spreadsheet",
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        metadata={"validation": {"passed": True}},
    )

    assert report.passed is True
    assert report.checks["package_opens"] is True
    assert report.metrics["sheet_count"] == 1
    assert report.metrics["formula_count"] == 1
    assert report.metrics["full_calc_on_load"] is True


def test_artifact_harness_rejects_corrupt_xlsx_package(tmp_path):
    source = tmp_path / "broken.xlsx"
    source.write_bytes(b"not-an-xlsx")

    report = ArtifactHarness().inspect(
        source,
        kind="spreadsheet",
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    assert report.passed is False
    assert report.status == "failed"
    assert report.checks["package_opens"] is False


def test_excel_automation_recalculates_formula_and_renders_pdf_when_available(tmp_path):
    verifier = ExcelAutomationVerifier(tmp_path / "previews")
    if not verifier.available:
        return
    source = tmp_path / "recalculate.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "测算"
    sheet.append(["基数", "比例", "结果"])
    sheet.append([1000, 0.12, "=A2*B2"])
    workbook.save(source)
    workbook.close()

    report = verifier.verify(source)

    assert report.passed is True, report.model_dump(mode="json")
    assert report.checks["excel_recalculation_completed"] is True
    assert report.checks["formula_cache_populated"] is True
    assert report.metrics["formula_cache_coverage"] == 1.0
    assert report.metrics["rendered_page_count"] >= 1
    assert Path(report.preview_files[0]).is_file()


def test_excel_verifier_failure_never_mutates_the_business_artifact(tmp_path, monkeypatch):
    source = tmp_path / "business-artifact.xlsx"
    workbook = Workbook()
    workbook.active.append(["employee_id", "amount"])
    workbook.active.append(["E001", 100])
    workbook.save(source)
    workbook.close()
    original_bytes = source.read_bytes()
    fake_excel = tmp_path / "EXCEL.EXE"
    fake_excel.touch()
    verifier = ExcelAutomationVerifier(
        tmp_path / "previews",
        excel_executable=fake_excel,
    )

    def destructive_failure(validation_copy: Path, _preview: Path) -> str:
        target = str(validation_copy).replace("'", "''")
        return (
            "$ErrorActionPreference='Stop'; "
            f"[IO.File]::WriteAllBytes('{target}', [byte[]](1,2,3)); "
            "exit 1"
        )

    monkeypatch.setattr(verifier, "_powershell_script", destructive_failure)

    report = verifier.verify(source)

    assert report.passed is False
    assert source.read_bytes() == original_bytes
    assert not list((tmp_path / "previews").glob("excel-validation-*.xlsx"))


def test_csv_copy_and_filter_workbook_opens_in_real_excel(tmp_path):
    """A CSV-derived sheet must remain Excel-compatible after copy/filter operations."""
    verifier = ExcelAutomationVerifier(tmp_path / "previews")
    if not verifier.available:
        return
    source = tmp_path / "contracts.csv"
    source.write_text(
        "employee_id,name,contract_end_date\n"
        "E001,Employee A,2026-08-20\n"
        "E002,Employee B,2026-09-30\n",
        encoding="utf-8",
    )
    plan = SpreadsheetPlan.model_validate({
        "objective": "Create an auditable upcoming-contract review sheet",
        "operations": [
            {"type": "copy_sheet", "source_sheet": "输入数据", "target_sheet": "合同复核"},
            {
                "type": "filter_date_range",
                "sheet": "合同复核",
                "header": "contract_end_date",
                "start_date": "2026-08-12",
                "end_date": "2026-09-11",
            },
        ],
    })

    result = SpreadsheetArtifactBuilder(tmp_path / "out").process(
        source,
        source.name,
        "筛选未来三十天到期合同",
        plan=plan,
        style=SpreadsheetStyleSpec(enable_filter=True),
    )
    report = verifier.verify(result.path)

    assert report.passed is True, report.model_dump(mode="json")
    assert report.checks["excel_recalculation_completed"] is True
    output = load_workbook(result.path)
    assert output["输入数据"].tables
    assert output["输入数据"].auto_filter.ref is None
    assert output["合同复核"].auto_filter.ref == "A1:C3"
    output.close()


def test_spreadsheet_builder_creates_workbook_from_text_plan_without_source(tmp_path):
    plan = SpreadsheetPlan.model_validate({
        "objective": "从零创建月度培训预算跟踪表",
        "operations": [
            {
                "type": "create_sheet",
                "sheet": "培训预算",
                "headers": ["部门", "预算金额", "实际金额", "执行率", "状态"],
                "rows": [
                    ["研发部", 20000, 15000, None, "进行中"],
                    ["人事部", 12000, 6000, None, "未开始"],
                ],
            },
            {
                "type": "add_formula_column",
                "sheet": "培训预算",
                "header": "执行率",
                "formula_template": "=IFERROR(C{row}/B{row},0)",
                "number_format": "0.0%",
            },
            {
                "type": "add_data_validation",
                "sheet": "培训预算",
                "header": "状态",
                "validation_type": "list",
                "values": ["未开始", "进行中", "已完成"],
            },
        ],
        "validations": {
            "required_sheets": ["培训预算"],
            "required_headers": [{
                "sheet": "培训预算",
                "headers": ["部门", "预算金额", "实际金额", "执行率", "状态"],
            }],
        },
    })

    result = SpreadsheetArtifactBuilder(tmp_path / "out").create(
        "请创建月度培训预算跟踪表",
        display_name="月度培训预算跟踪表.xlsx",
        plan=plan,
    )

    output = load_workbook(result.path, data_only=False)
    assert output.sheetnames[:2] == ["UniHR处理说明", "培训预算"]
    assert output["培训预算"]["D2"].value == "=IFERROR(C2/B2,0)"
    assert output["培训预算"]["D2"].number_format == "0.0%"
    assert output["培训预算"].freeze_panes == "A2"
    assert len(output["培训预算"].data_validations.dataValidation) == 1
    assert result.display_name == "月度培训预算跟踪表.xlsx"
    assert result.metadata["created_from_scratch"] is True
    assert result.metadata["source_files"] == []
    assert result.metadata["validation"]["passed"] is True
    assert result.metadata["workbook_diff"]["added_sheets"] == ["培训预算"]
    assert result.metadata["workbook_diff"]["formula_changes"] == 2
    output.close()


def test_spreadsheet_builder_reports_data_formula_and_style_differences(tmp_path):
    source = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "明细"
    sheet.append(["姓名", "基本工资", "奖金"])
    sheet.append(["测试甲", 8000, 1000])
    workbook.save(source)
    workbook.close()
    plan = SpreadsheetPlan.model_validate({
        "objective": "更新数据、添加公式并设置样式",
        "operations": [
            {"type": "set_cells", "sheet": "明细", "cells": [{"cell": "C2", "value": 1500}]},
            {"type": "add_formula_column", "sheet": "明细", "header": "应发", "formula_template": "=B{row}+C{row}"},
            {"type": "format_range", "sheet": "明细", "range": "B2:C2", "number_format": "¥#,##0.00", "horizontal_alignment": "right"},
        ],
    })

    result = SpreadsheetArtifactBuilder(tmp_path / "out").process(
        source, source.name, "生成应发列", plan=plan,
    )
    diff = result.metadata["workbook_diff"]

    assert diff["data_changes"] >= 1
    assert diff["formula_changes"] == 1
    assert diff["style_changes"] >= 2
    assert diff["changed_cells"] >= 3
    assert any(item["cell"] == "D2" and "formula" in item["change_types"] for item in diff["samples"])


def test_spreadsheet_create_sheet_rejects_formula_values_outside_formula_operation():
    try:
        SpreadsheetPlan.model_validate({
            "objective": "不得绕过公式白名单",
            "operations": [{
                "type": "create_sheet",
                "sheet": "测试",
                "headers": ["结果"],
                "rows": [["=WEBSERVICE(\"https://example.com\")"]],
            }],
        })
    except ValueError as exc:
        assert "spreadsheet_create_sheet_formula_not_allowed" in str(exc)
    else:  # pragma: no cover - protects the security contract.
        raise AssertionError("formula payload should be rejected")


def test_spreadsheet_inspector_supports_valid_xlsx_without_cached_dimensions(tmp_path):
    """Some standards-compliant writers omit worksheet dimension metadata."""
    sized = tmp_path / "sized.xlsx"
    source = tmp_path / "unsized.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "员工明细"
    sheet.append(["工号", "姓名"])
    sheet.append(["E1001", "测试员工"])
    workbook.save(sized)
    workbook.close()

    with ZipFile(sized, "r") as input_zip, ZipFile(source, "w", ZIP_DEFLATED) as output_zip:
        for entry in input_zip.infolist():
            content = input_zip.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                content = re.sub(br"<dimension\b[^>]*/>", b"", content, count=1)
            output_zip.writestr(entry, content)

    inspection = SpreadsheetArtifactBuilder(tmp_path / "out").inspect_source(source)

    assert inspection["sheets"][0]["rows"] == 2
    assert inspection["sheets"][0]["columns"] == 2
    assert inspection["sheets"][0]["headers"] == ["工号", "姓名"]


def test_spreadsheet_executor_supports_multi_metric_summary_and_formula_sort():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "员工明细"
    sheet.append(["工号", "部门", "基本工资", "绩效奖金"])
    sheet.append(["E1", "研发", 8000, 1000])
    sheet.append(["E2", "人事", 12000, 2000])
    sheet.append(["E3", "研发", 9000, 1500])
    plan = SpreadsheetPlan.model_validate({
        "objective": "计算、排序并生成多指标部门汇总",
        "operations": [
            {"type": "add_formula_column", "sheet": "员工明细", "header": "月度应发", "formula_template": "=C{row}+D{row}", "number_format": "¥#,##0.00"},
            {"type": "sort_rows", "sheet": "员工明细", "header": "月度应发", "direction": "descending"},
            {"type": "create_summary", "source_sheet": "员工明细", "target_sheet": "部门汇总", "group_header": "部门", "aggregation": "count", "result_header": "人数"},
            {"type": "create_summary", "source_sheet": "员工明细", "target_sheet": "部门汇总", "group_header": "部门", "value_header": "月度应发", "aggregation": "sum", "result_header": "月度应发合计"},
            {"type": "create_summary", "source_sheet": "员工明细", "target_sheet": "部门汇总", "group_header": "部门", "value_header": "月度应发", "aggregation": "average", "result_header": "月度应发平均值"},
            {"type": "add_chart", "sheet": "部门汇总", "category_header": "部门", "value_headers": ["月度应发合计"], "chart_type": "column"},
        ],
    })

    result = SpreadsheetPlanExecutor().execute(workbook, plan)

    assert result.operation_results[-1]["type"] == "add_chart"
    assert [cell.value for cell in workbook["部门汇总"][1]] == ["部门", "人数", "月度应发合计", "月度应发平均值"]
    assert workbook["员工明细"]["A2"].value == "E2"
    assert workbook["员工明细"]["E2"].value == "=C2+D2"
    assert workbook["部门汇总"]["C2"].number_format == "¥#,##0.00"
    assert len(workbook["部门汇总"]._charts) == 1


def test_spreadsheet_plan_normalizes_row_placeholder_before_column_letter():
    plan = SpreadsheetPlan.model_validate({
        "objective": "兼容模型常见的行占位符顺序",
        "operations": [{
            "type": "add_formula_column",
            "sheet": "员工明细",
            "header": "月度应发",
            "formula_template": "={row}C+{row}D",
        }],
    })

    assert plan.operations[0].formula_template == "=C{row}+D{row}"


def test_spreadsheet_executor_writes_existing_data_quality_issues_to_exception_sheet():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "员工明细"
    sheet.append(["工号", "在职状态"])
    sheet.append(["E1", "在职"])
    sheet.append(["E1", "借调"])
    sheet.append([None, "试用"])
    plan = SpreadsheetPlan.model_validate({
        "objective": "校验已有数据",
        "operations": [{
            "type": "add_data_validation", "sheet": "员工明细", "header": "在职状态",
            "validation_type": "list", "values": ["在职", "试用", "离职", "停薪留职"],
        }],
        "validations": {
            "unique_keys": [{"sheet": "员工明细", "header": "工号", "allow_blank": False}],
        },
    })

    result = SpreadsheetPlanExecutor().execute(workbook, plan)

    codes = {item["code"] for item in result.exceptions}
    assert codes == {"unique_key_duplicate", "unique_key_blank", "data_validation_existing_value_invalid"}
    assert workbook["UniHR异常清单"].max_row == 4


def test_spreadsheet_builder_sizes_chinese_headers_for_readability(tmp_path):
    source = tmp_path / "adjustment.csv"
    source.write_text("工号,姓名,2025基数,2026拟调整基数\nT001,测试甲,6200,6800\n", encoding="utf-8")

    result = SpreadsheetArtifactBuilder(tmp_path / "out").process(
        source, source.name, "生成数据复核准备稿"
    )

    output = load_workbook(result.path)
    assert output["输入数据"].column_dimensions["D"].width >= 18
    assert output["UniHR处理说明"].row_dimensions[3].height >= 24
    output.close()


def test_document_builder_applies_structured_style_spec_and_validates(tmp_path):
    builder = DocumentArtifactBuilder(tmp_path)
    path, metadata = builder.build(
        DocumentDraft(title="会议材料", lead="这是需要居中展示的正文。"),
        style=DocumentStyleSpec(
            preset="meeting",
            body_font="微软雅黑",
            body_size_pt=14,
            body_alignment="center",
            line_spacing=2,
        ),
    )

    document = Document(path)
    body = next(paragraph for paragraph in document.paragraphs if "需要居中展示" in paragraph.text)
    assert body.alignment == 1
    assert metadata["preset"] == "meeting"
    assert metadata["validation"]["passed"] is True


def test_spreadsheet_builder_applies_task_brief_style_and_validates(tmp_path):
    source = tmp_path / "people.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "人员表"
    sheet.append(["工号", "姓名", "基数"])
    sheet.append(["A001", "张三", 6200])
    workbook.save(source)
    workbook.close()

    result = SpreadsheetArtifactBuilder(tmp_path / "out").process(
        source,
        source.name,
        "统一表头并居中",
        style=SpreadsheetStyleSpec(
            target_sheet="人员表",
            header_fill="0F766E",
            horizontal_alignment="center",
            freeze_header=True,
        ),
    )

    output = load_workbook(result.path)
    assert output["人员表"]["A1"].fill.fgColor.rgb.endswith("0F766E")
    assert output["人员表"]["A2"].alignment.horizontal == "center"
    assert output["人员表"].freeze_panes == "A2"
    assert result.metadata["validation"]["passed"] is True
    output.close()


def test_spreadsheet_engine_maps_prior_year_and_builds_reviewable_annual_formulas(tmp_path):
    source = tmp_path / "annual-adjustment.xlsx"
    workbook = Workbook()
    current = workbook.active
    current.title = "2026人员名单"
    current.append(["工号", "姓名", "基本工资", "绩效部分", "实际工作月数"])
    current.append(["0000001", "匿名甲", 72000, 24000, 12])
    current.append(["0000002", "匿名乙", 36000, 6000, 6])
    prior = workbook.create_sheet("人员计算汇总2025")
    prior.append(["工号", "上年度个人缴存额"])
    prior.append(["0000001", 900])
    prior.append(["0000002", 600])
    params = workbook.create_sheet("参数说明")
    params.append(["参数", "值"])
    params.append(["基数下限", 2430])
    params.append(["基数上限", 31911])
    params.append(["个人比例", 0.12])
    params.append(["个人最低缴存额", 292])
    params.append(["个人最高缴存额", 3829])
    workbook.save(source)
    workbook.close()

    plan = SpreadsheetPlan.model_validate({
        "objective": "依据已确认参数生成2026年度公积金调整复核稿",
        "operations": [
            {"type": "copy_sheet", "source_sheet": "2026人员名单", "target_sheet": "人员计算汇总2026"},
            {
                "type": "map_column",
                "source_sheet": "人员计算汇总2025",
                "target_sheet": "人员计算汇总2026",
                "source_key_header": "工号",
                "target_key_header": "工号",
                "source_value_header": "上年度个人缴存额",
                "target_header": "上年度个人缴存额",
            },
            {
                "type": "add_formula_column", "sheet": "人员计算汇总2026", "header": "年度计入工资总额",
                "formula_template": "=C{row}+D{row}", "number_format": "0",
            },
            {
                "type": "add_formula_column", "sheet": "人员计算汇总2026", "header": "月平均工资（向上取整）",
                "formula_template": "=IF(E{row}=0,0,ROUNDUP(G{row}/E{row},0))", "number_format": "0",
            },
            {
                "type": "add_formula_column", "sheet": "人员计算汇总2026", "header": "理论个人缴存额（未限高）",
                "formula_template": "=ROUND(H{row}*'参数说明'!$B$4,0)", "number_format": "0",
            },
            {
                "type": "add_formula_column", "sheet": "人员计算汇总2026", "header": "限额后最终个人缴存额",
                "formula_template": "=MIN('参数说明'!$B$6,MAX('参数说明'!$B$5,I{row}))", "number_format": "0",
            },
            {
                "type": "add_formula_column", "sheet": "人员计算汇总2026", "header": "增加额",
                "formula_template": "=IF(F{row}=\"\",\"\",J{row}-F{row})", "number_format": "0",
            },
        ],
        "validations": {
            "required_sheets": ["人员计算汇总2026", "参数说明"],
            "required_headers": [{
                "sheet": "人员计算汇总2026",
                "headers": ["工号", "上年度个人缴存额", "限额后最终个人缴存额", "增加额"],
            }],
            "unique_keys": [{"sheet": "人员计算汇总2026", "header": "工号"}],
        },
    })
    result = SpreadsheetArtifactBuilder(tmp_path / "out").process(
        source, source.name, "生成2026年度公积金调整复核稿", plan=plan,
    )

    output = load_workbook(result.path, data_only=False)
    summary = output["人员计算汇总2026"]
    assert summary["A2"].value == "0000001"
    assert summary["F2"].value == 900
    assert summary["G2"].value == "=C2+D2"
    assert summary["H2"].value == "=IF(E2=0,0,ROUNDUP(G2/E2,0))"
    assert summary["J2"].value == "=MIN('参数说明'!$B$6,MAX('参数说明'!$B$5,I2))"
    assert summary["K2"].value == '=IF(F2="","",J2-F2)'
    assert "UniHR执行记录" in output.sheetnames
    assert "UniHR异常清单" in output.sheetnames
    assert result.metadata["changed_cell_count"] == 12
    assert result.metadata["validation"]["passed"] is True
    output.close()


def test_spreadsheet_engine_blocks_external_or_dynamic_formulas(tmp_path):
    source = tmp_path / "unsafe-formula.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "数据"
    sheet.append(["工号"])
    sheet.append(["001"])
    workbook.save(source)
    workbook.close()
    plan = SpreadsheetPlan.model_validate({
        "objective": "测试公式安全边界",
        "operations": [{
            "type": "add_formula_column",
            "sheet": "数据",
            "header": "外部数据",
            "formula_template": '=WEBSERVICE("https://example.com/{row}")',
        }],
    })

    try:
        SpreadsheetArtifactBuilder(tmp_path / "out").process(source, source.name, "测试", plan=plan)
    except SpreadsheetPlanError as exc:
        assert str(exc) == "spreadsheet_formula_unsafe"
    else:
        raise AssertionError("unsafe formula must be rejected")


def test_spreadsheet_engine_surfaces_duplicate_and_unmatched_mapping_keys(tmp_path):
    source = tmp_path / "mapping-exceptions.xlsx"
    workbook = Workbook()
    source_sheet = workbook.active
    source_sheet.title = "来源"
    source_sheet.append(["工号", "金额"])
    source_sheet.append(["001", 100])
    source_sheet.append(["001", 200])
    target_sheet = workbook.create_sheet("目标")
    target_sheet.append(["工号"])
    target_sheet.append(["001"])
    target_sheet.append(["002"])
    workbook.save(source)
    workbook.close()
    plan = SpreadsheetPlan.model_validate({
        "objective": "按工号映射金额",
        "operations": [{
            "type": "map_column", "source_sheet": "来源", "target_sheet": "目标",
            "source_key_header": "工号", "target_key_header": "工号",
            "source_value_header": "金额", "target_header": "映射金额",
        }],
    })

    result = SpreadsheetArtifactBuilder(tmp_path / "out").process(
        source, source.name, "按工号映射金额", plan=plan,
    )

    assert {item["code"] for item in result.metadata["exceptions"]} == {
        "duplicate_source_key", "unmatched_target_key",
    }
    assert result.metadata["validation"]["passed"] is False


def test_spreadsheet_engine_merges_multiple_source_workbooks_with_traceable_sheet_names(tmp_path):
    roster_path = tmp_path / "当年名单.xlsx"
    roster = Workbook()
    roster.active.title = "人员"
    roster.active.append(["工号", "姓名"])
    roster.active.append(["0000001", "匿名甲"])
    roster.save(roster_path)
    roster.close()

    prior_path = tmp_path / "上年计算.xlsx"
    prior = Workbook()
    summary = prior.active
    summary.title = "汇总"
    summary.append(["工号", "上年度个人缴存额", "参数引用"])
    summary.append(["0000001", 900, "=参数!B2"])
    params = prior.create_sheet("参数")
    params.append(["名称", "值"])
    params.append(["个人比例", 0.12])
    prior.save(prior_path)
    prior.close()

    builder = SpreadsheetArtifactBuilder(tmp_path / "out")
    inspection = builder.inspect_sources([
        (roster_path, roster_path.name),
        (prior_path, prior_path.name),
    ])
    assert [sheet["name"] for sheet in inspection["sheets"]] == [
        "人员", "上年计算-汇总", "上年计算-参数",
    ]
    plan = SpreadsheetPlan.model_validate({
        "objective": "合并来源并映射上年度金额",
        "operations": [{
            "type": "map_column", "source_sheet": "上年计算-汇总", "target_sheet": "人员",
            "source_key_header": "工号", "target_key_header": "工号",
            "source_value_header": "上年度个人缴存额", "target_header": "上年度个人缴存额",
        }],
    })
    result = builder.process(
        roster_path,
        roster_path.name,
        "合并并映射",
        plan=plan,
        additional_sources=[(prior_path, prior_path.name)],
    )

    output = load_workbook(result.path, data_only=False)
    assert output["人员"]["C2"].value == 900
    assert output["上年计算-汇总"]["C2"].value == "='上年计算-参数'!B2"
    assert len(result.metadata["source_sheet_map"]) == 3
    assert result.metadata["validation"]["passed"] is True
    output.close()


def test_spreadsheet_engine_supports_format_filter_sort_validation_summary_and_chart():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "人员数据"
    sheet.append(["部门", "姓名", "得分", "状态"])
    sheet.append(["研发", "甲", 80, "通过"])
    sheet.append(["市场", "乙", 92, "通过"])
    sheet.append(["研发", "丙", 70, "待复核"])

    plan = SpreadsheetPlan.model_validate({
        "objective": "整理人员数据并生成部门汇总",
        "operations": [
            {
                "type": "format_range", "sheet": "人员数据", "range": "A1:D1",
                "bold": True, "fill_color": "2563EB", "font_color": "FFFFFF",
                "horizontal_alignment": "center",
            },
            {"type": "set_auto_filter", "sheet": "人员数据", "range": "A1:D4"},
            {
                "type": "filter_rows", "sheet": "人员数据", "header": "状态",
                "allowed_values": ["通过"],
            },
            {
                "type": "sort_rows", "sheet": "人员数据", "header": "得分",
                "direction": "descending",
            },
            {
                "type": "add_data_validation", "sheet": "人员数据", "header": "状态",
                "validation_type": "list", "values": ["通过", "待复核", "不通过"],
                "start_row": 2, "end_row": 100,
            },
            {
                "type": "create_summary", "source_sheet": "人员数据", "target_sheet": "部门汇总",
                "group_header": "部门", "value_header": "得分", "aggregation": "average",
                "result_header": "平均得分", "number_format": "0.0",
            },
            {
                "type": "add_chart", "sheet": "部门汇总", "category_header": "部门",
                "value_headers": ["平均得分"], "chart_type": "column", "title": "部门平均得分",
            },
        ],
        "validations": {
            "required_sheets": ["人员数据", "部门汇总"],
            "required_headers": [{"sheet": "部门汇总", "headers": ["部门", "平均得分"]}],
        },
    })
    result = SpreadsheetPlanExecutor().execute(workbook, plan)

    assert result.validation["passed"] is True
    assert sheet["C2"].value == 92
    assert sheet["A1"].fill.fgColor.rgb.endswith("2563EB")
    assert sheet.auto_filter.ref == "A1:D4"
    assert len(sheet.data_validations.dataValidation) == 1
    assert workbook["部门汇总"]["B2"].value.startswith("=AVERAGEIF(")
    assert len(workbook["部门汇总"]._charts) == 1
    assert any(item["type"] == "create_summary" for item in result.operation_results)


def test_spreadsheet_engine_filters_an_inclusive_date_range_and_reports_bad_dates():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "合同台账"
    sheet.append(["工号", "姓名", "合同到期日"])
    sheet.append(["E001", "甲", "2026-08-12"])
    sheet.append(["E002", "乙", "2026/09/30"])
    sheet.append(["E003", "丙", "2026年8月20日"])
    sheet.append(["E004", "丁", "待确认"])

    plan = SpreadsheetPlan.model_validate({
        "objective": "筛选未来30天合同到期人员",
        "operations": [{
            "type": "filter_date_range",
            "sheet": "合同台账",
            "header": "合同到期日",
            "start_date": "2026-08-09",
            "end_date": "2026-09-08",
        }],
    })
    result = SpreadsheetPlanExecutor().execute(workbook, plan)

    assert sheet.row_dimensions[2].hidden is False
    assert sheet.row_dimensions[3].hidden is True
    assert sheet.row_dimensions[4].hidden is False
    assert sheet.row_dimensions[5].hidden is True
    assert sheet.auto_filter.ref == "A1:C5"
    assert result.changed_cell_count == 6
    assert "2026-08-09 至 2026-09-08" in result.operation_results[0]["summary"]
    assert "匹配 2 行" in result.operation_results[0]["summary"]
    assert result.exceptions[0]["code"] == "date_filter_invalid_date"
    assert result.validation["passed"] is False


def test_spreadsheet_plan_rejects_an_inverted_date_range():
    try:
        SpreadsheetPlan.model_validate({
            "objective": "日期筛选",
            "operations": [{
                "type": "filter_date_range",
                "sheet": "Sheet1",
                "header": "日期",
                "start_date": "2026-09-01",
                "end_date": "2026-08-01",
            }],
        })
    except ValueError as exc:
        assert "spreadsheet_date_range_invalid" in str(exc)
    else:
        raise AssertionError("an inverted date range must be rejected")


def test_spreadsheet_engine_refuses_to_sort_formula_rows():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "计算"
    sheet.append(["名称", "金额"])
    sheet.append(["甲", "=1+1"])
    plan = SpreadsheetPlan.model_validate({
        "objective": "排序",
        "operations": [{"type": "sort_rows", "sheet": "计算", "header": "金额"}],
    })
    try:
        SpreadsheetPlanExecutor().execute(workbook, plan)
    except SpreadsheetPlanError as exc:
        assert str(exc) == "spreadsheet_sort_formula_rows_not_supported"
    else:
        raise AssertionError("formula rows must not be reordered without formula translation")
