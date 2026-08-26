import json
import io
import sqlite3
from datetime import date

from langgraph.checkpoint.memory import InMemorySaver
from openpyxl import Workbook, load_workbook
from PIL import Image

from app.domains.conversations.repository import ConversationRepository
from app.domains.runs import RunPointer
from app.domains.skills import (
    SkillDraft,
    SkillExecutionPlan,
    SkillPlanNode,
    SkillRegistry,
    SkillResourceStore,
)
from app.domains.task_briefs import TaskBrief
from app.domains.workbench.service import (
    UploadPayload,
    WorkbenchService,
    _compile_deterministic_spreadsheet_plan,
    _compile_inline_records_spreadsheet_plan,
    _locked_spreadsheet_headers,
)
from app.models.gateway import ModelGatewayError, ModelReply
from app.settings import Settings
from app.domains.public_info.contracts import PublicInfoResult


class RecordingGateway:
    def __init__(self, replies):
        self.replies = list(replies)
        self.calls = []

    def complete(self, messages, **kwargs):
        self.calls.append((messages, kwargs))
        return ModelReply(self.replies.pop(0), kwargs.get("model") or "fake", 7, {})


class RecordingKnowledgeIndex:
    def __init__(self):
        self.queries = []

    def search(self, query, limit=6):
        self.queries.append(query)
        return []


class RouterUnavailableThenAnswerGateway:
    def __init__(self, answer: str):
        self.answer = answer
        self.calls = 0

    def complete(self, messages, **kwargs):
        self.calls += 1
        if self.calls == 1:
            raise ModelGatewayError("router_unavailable")
        return ModelReply(self.answer, kwargs.get("model") or "fake", 7, {})


class FailingWorkflow:
    def start(self, _request, **_kwargs):
        raise TypeError("executor exploded")

    def get(self, _pointer):
        return None

    def close(self):
        return None


class RecordingPublicInfo:
    def __init__(self):
        self.queries = []

    def answer(self, message, **_kwargs):
        self.queries.append(message)
        return PublicInfoResult(
            answer="找到两条公开职业资料线索。[1][2]",
            citations=[
                {"index": "1", "title": "公开资料A", "url": "https://example.com/a"},
                {"index": "2", "title": "公开资料B", "url": "https://example.com/b"},
            ],
            capability_status="available",
            query=message,
            trace=[{"event": "web.search", "status": "completed", "summary": "公开检索完成"}],
        )


def test_recruitment_web_sourcing_uses_web_without_downgrading_domain(tmp_path):
    gateway = RecordingGateway([route_json("recruitment_screening")])
    public_info = RecordingPublicInfo()
    settings = Settings(data_dir=tmp_path / "data", database_path=tmp_path / "unihr.db")
    repository = ConversationRepository(settings.database_path)
    service = WorkbenchService(
        settings, repository, model_gateway=gateway, public_info=public_info,
        workflow_checkpointer=InMemorySaver(),
    )
    question = (
        "请联网从公开的 GitHub 个人主页搜索适合 Python 后端工程师岗位的候选人资料，"
        "重点看 FastAPI 和 PostgreSQL 项目经验，列出来源。"
    )

    result = service.handle_turn(
        message=question, conversation_id=None, task_mode="自动选择", uploads=[],
        document_font="宋体", document_font_size=14,
    )

    assert result.interpretation.business_intent.value == "recruitment_screening"
    assert result.interpretation.retrieval_mode.value == "web"
    assert public_info.queries == [
        "公开的 GitHub 个人主页搜索适合 Python 后端工程师岗位的候选人资料，"
        "重点看 FastAPI 和 PostgreSQL 项目经验，列出来源"
    ]
    assert "候选线索" in result.assistant_message.content
    assert "自动录用或淘汰" in result.assistant_message.content
    assert result.run.status == "succeeded"
    assert result.run.graph_path == ["prepare", "execute", "quality_gate", "finalize"]


def route_json(intent="conversation", *, turn_action=None, missing_inputs=None):
    action_fields = {}
    if turn_action is not None:
        action_fields = {
            "turn_action": turn_action,
        }
    return json.dumps({
        "primary_intent": intent,
        "secondary_intents": [],
        "context_mode": "uses_conversation",
        "missing_inputs": list(missing_inputs or []),
        "risk_level": "low",
        "data_class": "public",
        "required_capabilities": [],
        "confidence": 0.92,
        "reason_summary": "上下文对话",
        **action_fields,
    }, ensure_ascii=False)


def test_product_help_explanation_is_answered_without_a_prior_result(tmp_path):
    gateway = RecordingGateway([
        route_json("product_help", turn_action="explain_existing"),
        "模型调用归属到会话和招聘 Case，便于成本分摊、数据隔离与审计追踪。",
    ])
    settings = Settings(data_dir=tmp_path / "data", database_path=tmp_path / "unihr.db")
    repository = ConversationRepository(settings.database_path)
    service = WorkbenchService(
        settings, repository, model_gateway=gateway, workflow_checkpointer=InMemorySaver(),
    )

    result = service.handle_turn(
        message="P0-F 前端验收：请用一句话说明模型调用记录为什么需要归属到会话和招聘 Case。",
        conversation_id=None,
        task_mode="自动选择",
        uploads=[],
        document_font="宋体",
        document_font_size=14,
    )

    assert "成本分摊" in result.assistant_message.content
    assert result.interpretation.turn_action.value == "start_new"
    assert result.interpretation.execution_mode.value == "execute"
    assert result.run.graph_path == ["prepare", "execute", "quality_gate", "finalize"]
    answer_system_prompt = gateway.calls[1][0][0]["content"]
    assert "唯一可用于说明上线状态的依据" in answer_system_prompt
    assert "管理员可手动选择一份失败文件并只重试这一份" in answer_system_prompt
    assert "当前没有自动定时或无限失败重试" in answer_system_prompt


def test_image_question_uses_structured_evidence_and_surfaces_review_notice(tmp_path):
    image_bytes = io.BytesIO()
    Image.new("RGB", (400, 160), "white").save(image_bytes, format="PNG")
    gateway = RecordingGateway([
        route_json("image_understanding"),
        json.dumps({
            "page_text": "会议时间：2026年8月18日\n人数：?人",
            "blocks": [
                {"text": "会议时间：2026年8月18日", "kind": "field", "bbox": [10, 20, 500, 100], "confidence": 0.96},
                {"text": "人数：?人", "kind": "field", "bbox": [10, 120, 300, 200], "confidence": 0.68},
            ],
        }, ensure_ascii=False),
        "会议时间为2026年8月18日；人数无法确认。",
    ])
    settings = Settings(data_dir=tmp_path / "data", database_path=tmp_path / "unihr.db")
    repository = ConversationRepository(settings.database_path)
    service = WorkbenchService(
        settings, repository, model_gateway=gateway, workflow_checkpointer=InMemorySaver(),
    )

    result = service.handle_turn(
        message="请读取图片中的会议时间和人数；看不清的不要猜。",
        conversation_id=None, task_mode="自动选择",
        uploads=[type("Upload", (), {
            "filename": "会议通知.png", "mime_type": "image/png", "content": image_bytes.getvalue(),
        })()],
        document_font="宋体", document_font_size=14,
    )

    assert "人数无法确认" in result.assistant_message.content
    assert "视觉证据复核提示" in result.assistant_message.content
    assert any(item[1]["purpose"] == "image_evidence_extraction" for item in gateway.calls)
    assert any(item[1]["purpose"] == "vision_understanding" for item in gateway.calls)
    assert any(item["event"] == "vision.evidence.extracted" for item in result.trace)
    assert any(item["event"] == "vision.evidence.review_required" for item in result.trace)


def test_explicit_two_column_formula_is_compiled_without_model_guessing():
    brief = TaskBrief.model_validate({
            "kind": "spreadsheet",
            "objective": "增加年度计入工资总额列",
            "mode_spec": {
                "source_sheet": "测算",
                "calculation_rules": "年度计入工资总额=基本工资+绩效部分，结果取整数",
            },
        })
    plan = _compile_deterministic_spreadsheet_plan(
        brief,
        {
            "sheets": [{
                "name": "测算", "header_row": 1,
                "header_columns": {"工号": "A", "基本工资": "B", "绩效部分": "C"},
            }],
        },
    )

    assert plan is not None
    operations = plan.model_dump(mode="json")["operations"]
    assert operations[0] == {"type": "copy_sheet", "source_sheet": "测算", "target_sheet": "测算_复核"}
    assert operations[1]["header"] == "年度计入工资总额"
    assert operations[1]["formula_template"] == "=B{row}+C{row}"
    assert operations[1]["number_format"] == "0"
    assert _locked_spreadsheet_headers(
        "任务要求：\n- 处理目标：增加年度计入工资总额公式列\n锁定约束：\n- 年度计入工资总额=基本工资+绩效部分",
        brief,
    ) == ["年度计入工资总额"]


def test_relative_date_filter_is_compiled_from_a_simple_prompt_without_model_guessing():
    brief = TaskBrief.model_validate({
        "kind": "spreadsheet",
        "objective": "查看近期合同到期人员",
    })
    plan = _compile_deterministic_spreadsheet_plan(
        brief,
        {
            "sheets": [{
                "name": "合同台账",
                "header_row": 1,
                "header_columns": {"工号": "A", "姓名": "B", "contract_end_date": "C"},
            }],
        },
        message="帮我看看未来30天谁的合同要到期",
        reference_date=date(2026, 8, 9),
    )

    assert plan is not None
    operations = plan.model_dump(mode="json")["operations"]
    assert operations[0] == {
        "type": "copy_sheet", "source_sheet": "合同台账", "target_sheet": "合同台账_日期筛选",
    }
    assert operations[1] == {
        "type": "filter_date_range",
        "sheet": "合同台账_日期筛选",
        "header": "contract_end_date",
        "start_date": "2026-08-09",
        "end_date": "2026-09-08",
        "header_row": 1,
        "invalid_date_policy": "exclude_and_report",
    }
    assert len(operations) == 2
    assert "以 2026-08-09 为基准" in plan.assumptions[0]


def test_date_filter_preserves_explicit_sort_direction_in_the_code_plan():
    brief = TaskBrief.model_validate({
        "kind": "spreadsheet",
        "objective": "筛选未来30天合同到期人员并按到期日期升序",
    })
    plan = _compile_deterministic_spreadsheet_plan(
        brief,
        {
            "sheets": [{
                "name": "合同台账",
                "header_row": 1,
                "header_columns": {"工号": "A", "contract_end_date": "B"},
            }],
        },
        message="保留原始数据，按到期日期升序，生成复核表",
        reference_date=date(2026, 8, 12),
    )

    assert plan is not None
    operations = plan.model_dump(mode="json")["operations"]
    assert [item["type"] for item in operations] == [
        "copy_sheet", "filter_date_range", "sort_rows",
    ]
    assert operations[2] == {
        "type": "sort_rows",
        "sheet": "合同台账_日期筛选",
        "header": "contract_end_date",
        "direction": "ascending",
        "header_row": 1,
        "start_row": None,
        "end_row": None,
    }


def test_this_week_uses_monday_to_sunday_and_explicit_ranges_override_it():
    brief = TaskBrief.model_validate({"kind": "spreadsheet", "objective": "合同到期筛选"})
    inspection = {
        "sheets": [{
            "name": "合同台账", "header_row": 1,
            "header_columns": {"姓名": "A", "合同到期日": "B"},
        }],
    }
    weekly = _compile_deterministic_spreadsheet_plan(
        brief, inspection, message="这周谁到期", reference_date=date(2026, 8, 9),
    )
    explicit = _compile_deterministic_spreadsheet_plan(
        brief,
        inspection,
        message="筛选 2026-08-10 到 2026-08-20 到期的合同",
        reference_date=date(2026, 8, 9),
    )

    weekly_filter = weekly.model_dump(mode="json")["operations"][1]
    explicit_filter = explicit.model_dump(mode="json")["operations"][1]
    assert (weekly_filter["start_date"], weekly_filter["end_date"]) == (
        "2026-08-03", "2026-08-09",
    )
    assert (explicit_filter["start_date"], explicit_filter["end_date"]) == (
        "2026-08-10", "2026-08-20",
    )


def test_compact_repeated_records_compile_to_a_formula_workbook_plan_without_model_guessing():
    message = (
        "做一张培训完成情况表：制度培训应到50实到42，安全培训应到40实到38，"
        "合规培训应到30实到30；计算每项完成率和总计，生成 Excel。"
    )
    brief = TaskBrief.model_validate({
        "kind": "spreadsheet", "objective": message,
    })

    plan = _compile_inline_records_spreadsheet_plan(brief, message)

    assert plan is not None
    operations = plan.model_dump(mode="json")["operations"]
    assert operations[0] == {
        "type": "create_sheet",
        "sheet": "培训完成情况",
        "headers": ["项目", "应到", "实到", "完成率"],
        "rows": [
            ["制度培训", 50, 42, None],
            ["安全培训", 40, 38, None],
            ["合规培训", 30, 30, None],
            ["总计", 120, 110, None],
        ],
    }
    assert operations[1]["type"] == "add_formula_column"
    assert operations[1]["formula_template"] == '=IF(B{row}=0,"",C{row}/B{row})'
    assert operations[1]["number_format"] == "0.0%"
    assert plan.requires_human_review is False


def test_inline_record_compiler_refuses_ambiguous_single_record():
    brief = TaskBrief.model_validate({
        "kind": "spreadsheet", "objective": "做个预算表，预算100实际80",
    })

    assert _compile_inline_records_spreadsheet_plan(brief, brief.objective) is None


def test_locked_headers_extracts_named_columns_instead_of_the_column_count():
    message = (
        "新增三列：出勤率＝实际出勤天数÷应出勤天数；"
        "月度应发＝基本工资＋绩效奖金；"
        "工号异常时，数据状态标记‘需复核’，否则标记‘正常’。"
    )

    assert _locked_spreadsheet_headers(message, None) == ["出勤率", "月度应发", "数据状态"]


def test_workbench_second_turn_sends_previous_context(tmp_path):
    gateway = RecordingGateway([
        route_json(), "好的，我记住了。",
        route_json(), "你刚才说你负责博士后管理。",
    ])
    settings = Settings(data_dir=tmp_path, database_path=tmp_path / "unihr.db")
    repository = ConversationRepository(settings.database_path)
    service = WorkbenchService(
        settings, repository, model_gateway=gateway, workflow_checkpointer=InMemorySaver(),
    )

    first = service.handle_turn(
        message="我负责博士后管理。", conversation_id=None, task_mode="自动选择",
        uploads=[], document_font="宋体", document_font_size=14,
    )
    second = service.handle_turn(
        message="我刚才说我负责什么？", conversation_id=first.conversation.id,
        task_mode="自动选择", uploads=[], document_font="宋体", document_font_size=14,
    )

    assert "博士后管理" in second.assistant_message.content
    answer_call_messages = gateway.calls[3][0]
    serialized = json.dumps(answer_call_messages, ensure_ascii=False)
    assert "我负责博士后管理" in serialized
    assert second.conversation.id == first.conversation.id


def test_read_only_attachment_question_returns_answer_without_creating_artifact(tmp_path):
    misclassified_route = json.dumps({
        "primary_intent": "document_generation",
        "secondary_intents": [], "context_mode": "standalone", "missing_inputs": [],
        "risk_level": "medium", "data_class": "internal_hr",
        "required_capabilities": ["document.compose", "artifact.create"],
        "confidence": 0.93, "reason_summary": "读取附件中的字段", "turn_action": "start_new",
        "resolved_request": "读取附件并列出目标年度和预算上限", "retrieval_queries": [],
        "priority_constraints": [],
    }, ensure_ascii=False)
    gateway = RecordingGateway([
        misclassified_route,
        "目标年度为2026年，预算上限为20万元。",
    ])
    settings = Settings(data_dir=tmp_path, database_path=tmp_path / "unihr.db")
    repository = ConversationRepository(settings.database_path)
    service = WorkbenchService(
        settings, repository, model_gateway=gateway, workflow_checkpointer=InMemorySaver(),
    )

    result = service.handle_turn(
        message="请读取附件并列出目标年度和预算上限。",
        conversation_id=None,
        task_mode="自动选择",
        uploads=[type("Upload", (), {
            "filename": "约束.txt", "mime_type": "text/plain",
            "content": "目标年度：2026年\n预算上限：20万元".encode("utf-8"),
        })()],
        document_font="宋体",
        document_font_size=14,
    )

    assert result.interpretation.business_intent.value == "conversation"
    assert result.interpretation.delivery_mode.value == "inline_response"
    assert result.artifacts == []
    assert result.run.status == "succeeded"
    assert "2026年" in result.assistant_message.content
    assert "20万元" in json.dumps(gateway.calls[1][0], ensure_ascii=False)
    assert any(
        item["event"] == "turn.interpreted" and item["business_intent"] == "conversation"
        for item in result.trace
    )


def test_inline_polish_request_does_not_create_word_artifact_when_route_overreaches(tmp_path):
    misclassified_route = json.dumps({
        "primary_intent": "document_generation",
        "secondary_intents": [],
        "context_mode": "standalone",
        "missing_inputs": [],
        "risk_level": "low",
        "data_class": "internal_hr",
        "required_capabilities": ["document.compose", "artifact.create"],
        "confidence": 0.91,
        "reason_summary": "用户提到正式通知",
        "turn_action": "start_new",
        "resolved_request": "将用户提供的话润色成正式通知语气",
        "retrieval_queries": [],
        "priority_constraints": [],
    }, ensure_ascii=False)
    polished_text = (
        "王老师，您好！主任特别提醒，本次面试要求较为严格，请提醒候选人认真做好准备。"
        "我们计划于本周四、本周五组织专家面试。如您认为相关内容需要修改或补充，烦请及时发给我。"
    )
    gateway = RecordingGateway([misclassified_route, polished_text])
    settings = Settings(data_dir=tmp_path, database_path=tmp_path / "unihr.db")
    repository = ConversationRepository(settings.database_path)
    service = WorkbenchService(
        settings, repository, model_gateway=gateway, workflow_checkpointer=InMemorySaver(),
    )
    message = (
        "王老师，这个视频我们主任说提醒候选人要认真准备下，还蛮严格的，"
        "我们准备这周四这周五找专家面试了，您看看如果有要改的要增加的可以发给我。"
        "帮我润色一下这段话，要有正式通知的感觉。"
    )

    result = service.handle_turn(
        message=message,
        conversation_id=None,
        task_mode="自动选择",
        uploads=[],
        document_font="宋体",
        document_font_size=14,
    )

    assert result.interpretation.business_intent.value == "conversation"
    assert result.interpretation.delivery_mode.value == "inline_transform"
    assert "没有授权创建文件" in result.interpretation.reason_summary
    assert result.artifacts == []
    assert "Word 草稿" not in result.assistant_message.content
    assert result.assistant_message.content == polished_text
    assert gateway.calls[1][1]["purpose"] == "inline_text_transform"
    assert "不得把‘本周四、本周五、明天’等相对日期" in gateway.calls[1][0][0]["content"]
    assert any(item["event"] == "delivery.inline_guard" for item in result.trace)


def test_latest_no_file_instruction_cannot_be_overridden_by_semantic_document_skill(tmp_path):
    router_decision = json.dumps({
        "primary_intent": "document_generation",
        "secondary_intents": [],
        "context_mode": "uses_conversation",
        "missing_inputs": [],
        "risk_level": "low",
        "data_class": "internal_hr",
        "required_capabilities": ["document.compose", "artifact.create"],
        "confidence": 0.96,
        "reason_summary": "用户撤回文件交付，改为聊天内正文",
        "turn_action": "revise_existing",
        "delivery_mode": "mutate_artifact",
        "resolved_request": "基于上一轮内容，只在聊天中给出简洁正式的通知正文",
        "retrieval_queries": [],
        "priority_constraints": ["不要创建任何文件", "只在聊天中回复"],
    }, ensure_ascii=False)
    inline_answer = "各位候选人：专家面试拟于明天下午举行，请认真做好相关准备。具体安排另行通知。"
    gateway = RecordingGateway([router_decision, inline_answer])
    settings = Settings(data_dir=tmp_path, database_path=tmp_path / "unihr.db")
    repository = ConversationRepository(settings.database_path)
    conversation = repository.create_conversation("专家面试通知")
    repository.add_message(
        conversation.id,
        "user",
        "请生成一份关于明天下午专家面试安排的正式 Word 通知。",
    )
    registry = SkillRegistry(settings.database_path)
    document_skill = registry.create(SkillDraft(
        name="formal-document",
        display_name="公文与会议材料",
        description="生成正式 Word 公文和会议材料。",
        instructions="生成结构化公文并创建 Word 文件。",
        trigger_phrases=["正式通知", "Word 通知"],
        capabilities=["document.compose", "artifact.create"],
        source_type="prompt",
        enabled=True,
        status="active",
    ))
    service = WorkbenchService(
        settings,
        repository,
        model_gateway=gateway,
        skill_registry=registry,
        workflow_checkpointer=InMemorySaver(),
    )
    service.skill_orchestrator.plan = lambda *_args, **_kwargs: SkillExecutionPlan(
        status="ready",
        nodes=[SkillPlanNode(
            skill_id=document_skill.id,
            skill_version=document_skill.version,
            display_name=document_skill.display_name,
            score=0.92,
            selection_mode="automatic",
            matched_trigger="正式通知",
            semantic_score=0.88,
            rerank_score=0.81,
            provides=["document.compose", "artifact.create"],
            outputs=["docx"],
            reason="历史任务与公文 Skill 语义相似",
        )],
        covered_capabilities=["document.compose", "artifact.create"],
        warnings=[f"semantic_intent_override:{document_skill.id}"],
    )

    def fail_if_document_created(*_args, **_kwargs):
        raise AssertionError("latest no-file instruction must prevent document execution")

    service._create_document = fail_if_document_created
    result = service.handle_turn(
        message="先别生成 Word，也不要创建任何文件，只在聊天里给我一段简洁正式的通知正文。",
        conversation_id=conversation.id,
        task_mode="自动选择",
        uploads=[],
        document_font="宋体",
        document_font_size=14,
    )

    assert result.interpretation.business_intent.value == "conversation"
    assert result.interpretation.delivery_mode.value == "inline_response"
    assert result.artifacts == []
    assert result.assistant_message.content == inline_answer
    assert not any(item["event"] == "document.compose" for item in result.trace)


def test_workbench_marks_run_failed_when_executor_raises(tmp_path):
    gateway = RecordingGateway([route_json()])
    settings = Settings(data_dir=tmp_path / "data", database_path=tmp_path / "unihr.db")
    repository = ConversationRepository(settings.database_path)
    service = WorkbenchService(
        settings, repository, model_gateway=gateway, workflow_checkpointer=InMemorySaver(),
    )
    service.workflow = FailingWorkflow()

    result = service.handle_turn(
        message="执行一个会抛出异常的测试任务",
        conversation_id=None,
        task_mode="自动选择",
        uploads=[],
        document_font="宋体",
        document_font_size=14,
    )

    run = service.runs.list_for_conversation(
        result.conversation.id,
        tenant_id=repository.default_tenant_id,
        owner_user_id=repository.default_owner_user_id,
    )[0]
    assert result.run.status == "failed"
    assert run.status.value == "failed"
    assert run.error_code == "TypeError"
    assert "未完成" in result.assistant_message.content
    assert any(item["event"] == "graph.run.failed" for item in result.trace)


def test_task_brief_only_turn_is_compiled_persisted_and_traced(tmp_path):
    gateway = RecordingGateway([route_json(), "已收到结构化任务要求。"])
    settings = Settings(data_dir=tmp_path, database_path=tmp_path / "unihr.db")
    repository = ConversationRepository(settings.database_path)
    service = WorkbenchService(
        settings, repository, model_gateway=gateway, workflow_checkpointer=InMemorySaver(),
    )

    result = service.handle_turn(
        message="",
        conversation_id=None,
        task_mode="政策问答",
        task_brief={
            "version": "1.0",
            "kind": "policy",
            "objective": "查询浙江省博士后进站备案材料",
            "constraints": ["适用地区：浙江省"],
            "deliverables": ["带证据的政策回答"],
            "mode_spec": {"topic": "博士后进站备案", "scope": "浙江省"},
        },
        uploads=[],
        document_font="宋体",
        document_font_size=14,
    )

    stored = repository.list_messages(result.conversation.id)
    assert result.user_message.task_brief["kind"] == "policy"
    assert stored[0].task_brief["mode_spec"]["scope"] == "浙江省"
    assert "政策主题：博士后进站备案" in stored[0].content
    assert any(item["event"] == "brief.compiled" for item in result.trace)
    router_prompt = json.dumps(gateway.calls[0][0], ensure_ascii=False)
    assert "查询浙江省博士后进站备案材料" in router_prompt


def test_workbench_recruitment_mode_creates_reviewable_excel(tmp_path):
    screening = json.dumps({
        "job_title": "研发工程师",
        "criteria": [{"id": "python", "label": "Python经验", "weight": 100, "requirement": "3年以上"}],
        "candidates": [{
            "candidate_name": "候选人甲", "source_name": "resume.txt",
            "criterion_scores": {"python": 88}, "score_evidence": {"python": "4年Python经验"},
            "strengths": ["经验匹配"], "gaps": [], "verification_questions": ["核实项目成果"],
        }],
    }, ensure_ascii=False)
    gateway = RecordingGateway([route_json(), screening])
    settings = Settings(data_dir=tmp_path, database_path=tmp_path / "unihr.db")
    repository = ConversationRepository(settings.database_path)
    service = WorkbenchService(
        settings, repository, model_gateway=gateway, workflow_checkpointer=InMemorySaver(),
    )

    result = service.handle_turn(
        message="岗位：研发工程师；必须有3年以上Python经验",
        conversation_id=None,
        task_mode="简历筛选",
        uploads=[type("Upload", (), {
            "filename": "resume.txt", "mime_type": "text/plain", "content": "候选人甲，4年Python经验".encode("utf-8"),
        })()],
        document_font="宋体",
        document_font_size=14,
    )

    assert result.artifacts[0].kind == "recruitment_report"
    assert result.artifacts[0].metadata["human_review_required"] is True
    assert any(item["event"] == "skill.selected" for item in result.trace)
    assert any(item["event"] == "skill.graph.node.business_executed" for item in result.trace)
    assert any(
        "批量简历证据化筛选" in str(item.get("summary") or "")
        for item in result.trace
        if item["event"] == "skill.graph.node.business_executed"
    )
    assert any(item["event"] == "human_review.gate" for item in result.trace)
    assert "候选人甲" in result.assistant_message.content
    assert result.run.status == "waiting_review"
    assert result.run.graph_path == [
        "prepare", "execute", "quality_gate", "human_review_gate",
    ]
    pointer = RunPointer(
        run_id=result.run.id,
        tenant_id=repository.default_tenant_id,
        owner_user_id=repository.default_owner_user_id,
    )
    persisted = service.runs.get(pointer)
    assert persisted.status.value == "waiting_review"
    assert any(item.event == "graph.run.paused" for item in service.runs.list_events(pointer))


def test_recruitment_skill_graph_pauses_as_waiting_input_when_resumes_are_missing(tmp_path):
    gateway = RecordingGateway([route_json("recruitment_screening")])
    settings = Settings(data_dir=tmp_path, database_path=tmp_path / "unihr.db")
    repository = ConversationRepository(settings.database_path)
    service = WorkbenchService(
        settings, repository, model_gateway=gateway, workflow_checkpointer=InMemorySaver(),
    )

    result = service.handle_turn(
        message="帮我选两个人进入面试",
        conversation_id=None,
        task_mode="简历筛选",
        uploads=[],
        document_font="宋体",
        document_font_size=14,
    )

    assert result.run.status == "waiting_input"
    assert "上传" in result.assistant_message.content
    assert any(item["event"] == "skill.graph.node.waiting_input" for item in result.trace)
    assert not any(item["event"] == "skill.graph.node.completed" for item in result.trace)


def test_new_recruitment_with_resume_attachments_executes_when_router_requests_clarification(tmp_path):
    screening = json.dumps({
        "job_title": "产业拓展实习生",
        "criteria": [{"id": "experience", "label": "拓展经历", "weight": 100, "requirement": "产业拓展"}],
        "candidates": [
            {
                "candidate_name": "候选人甲", "source_name": "候选人甲.txt",
                "criterion_scores": {"experience": 88}, "score_evidence": {"experience": "独立对接合作单位"},
                "strengths": ["经验匹配"], "gaps": [], "verification_questions": ["核实本人贡献"],
            },
            {
                "candidate_name": "候选人乙", "source_name": "候选人乙.txt",
                "criterion_scores": {"experience": 76}, "score_evidence": {"experience": "负责活动外联"},
                "strengths": ["有外联经验"], "gaps": [], "verification_questions": ["核实项目深度"],
            },
        ],
    }, ensure_ascii=False)
    gateway = RecordingGateway([
        route_json(
            "recruitment_screening",
            turn_action="clarify",
            missing_inputs=["筛选标准或岗位要求"],
        ),
        screening,
    ])
    settings = Settings(data_dir=tmp_path, database_path=tmp_path / "unihr.db")
    repository = ConversationRepository(settings.database_path)
    service = WorkbenchService(
        settings, repository, model_gateway=gateway, workflow_checkpointer=InMemorySaver(),
    )

    result = service.handle_turn(
        message="帮我从这些产业拓展实习简历里选两个人面试",
        conversation_id=None,
        task_mode="自动选择",
        uploads=[
            type("Upload", (), {"filename": "候选人甲.txt", "mime_type": "text/plain", "content": "候选人甲，独立对接合作单位".encode("utf-8")})(),
            type("Upload", (), {"filename": "候选人乙.txt", "mime_type": "text/plain", "content": "候选人乙，负责活动外联".encode("utf-8")})(),
        ],
        document_font="宋体",
        document_font_size=14,
    )

    assert result.interpretation.execution_mode.value == "execute"
    assert result.run.graph_path == ["prepare", "execute", "quality_gate", "human_review_gate"]
    assert result.artifacts[0].kind == "recruitment_report"
    assert "候选人甲" in result.assistant_message.content


def test_clarification_answer_asks_for_the_actual_missing_slot(tmp_path):
    gateway = RecordingGateway([route_json(
        "recruitment_screening",
        turn_action="clarify",
        missing_inputs=["岗位要求"],
    )])
    settings = Settings(data_dir=tmp_path, database_path=tmp_path / "unihr.db")
    repository = ConversationRepository(settings.database_path)
    service = WorkbenchService(
        settings, repository, model_gateway=gateway, workflow_checkpointer=InMemorySaver(),
    )

    result = service.handle_turn(
        message="帮我筛一下简历",
        conversation_id=None,
        task_mode="自动选择",
        uploads=[],
        document_font="宋体",
        document_font_size=14,
    )

    assert "岗位要求" in result.assistant_message.content
    assert "继续处理已有结果" not in result.assistant_message.content


def test_public_current_info_answers_date_then_requests_location_without_searching(tmp_path):
    gateway = RecordingGateway([route_json(
        "public_current_info",
        turn_action="start_new",
        missing_inputs=["具体城市或地区"],
    )])
    settings = Settings(data_dir=tmp_path, database_path=tmp_path / "unihr.db")
    repository = ConversationRepository(settings.database_path)
    service = WorkbenchService(
        settings, repository, model_gateway=gateway, workflow_checkpointer=InMemorySaver(),
    )

    result = service.handle_turn(
        message="任务目标：今天是礼拜几，是不是台风很大",
        conversation_id=None,
        task_mode="自动选择",
        uploads=[],
        document_font="宋体",
        document_font_size=14,
    )

    assert "星期" in result.assistant_message.content
    assert "城市或区县" in result.assistant_message.content
    assert result.interpretation.business_intent.value == "public_current_info"
    assert result.interpretation.execution_mode.value == "clarify"
    assert result.interpretation.selected_capability_id == "public_information"
    trace_events = [item["event"] for item in result.trace]
    assert "task_intent.compiled" in trace_events
    assert "capability.candidates" in trace_events
    assert "capability.selected" in trace_events
    assert "web.search" not in trace_events


def test_popular_movie_query_does_not_fall_into_weather_location_clarification(tmp_path):
    gateway = RecordingGateway([route_json(
        "public_current_info",
        turn_action="clarify",
        missing_inputs=["location"],
    )])
    public_info = RecordingPublicInfo()
    settings = Settings(data_dir=tmp_path, database_path=tmp_path / "unihr.db")
    repository = ConversationRepository(settings.database_path)
    service = WorkbenchService(
        settings, repository, model_gateway=gateway, public_info=public_info,
        workflow_checkpointer=InMemorySaver(),
    )
    question = "晚上好，我想看电影，查查最近有哪些热门大片"

    result = service.handle_turn(
        message=question, conversation_id=None, task_mode="自动选择", uploads=[],
        document_font="宋体", document_font_size=14,
    )

    assert result.interpretation.business_intent.value == "public_current_info"
    assert result.interpretation.execution_mode.value == "execute"
    assert public_info.queries == [question]
    assert "台风" not in result.assistant_message.content


def test_location_sensitive_public_clarification_stays_on_the_current_topic(tmp_path):
    gateway = RecordingGateway([route_json(
        "public_current_info",
        turn_action="clarify",
        missing_inputs=["location"],
    )])
    public_info = RecordingPublicInfo()
    settings = Settings(data_dir=tmp_path, database_path=tmp_path / "unihr.db")
    repository = ConversationRepository(settings.database_path)
    service = WorkbenchService(
        settings, repository, model_gateway=gateway, public_info=public_info,
        workflow_checkpointer=InMemorySaver(),
    )

    result = service.handle_turn(
        message="帮我查查附近最近有什么音乐会",
        conversation_id=None, task_mode="自动选择", uploads=[],
        document_font="宋体", document_font_size=14,
    )

    assert result.interpretation.execution_mode.value == "clarify"
    assert "城市或区县" in result.assistant_message.content
    assert "当前问题" in result.assistant_message.content
    assert "台风" not in result.assistant_message.content
    assert public_info.queries == []


def test_recruitment_followup_compares_existing_result_without_rescreening(tmp_path):
    screening = json.dumps({
        "job_title": "产业拓展实习生",
        "criteria": [{"id": "experience", "label": "类似工作经验", "weight": 100, "requirement": "产业拓展"}],
        "candidates": [
            {
                "candidate_name": "候选人甲", "source_name": "resume-a.txt",
                "criterion_scores": {"experience": 90}, "score_evidence": {"experience": "有B2B运营经验"},
                "strengths": ["经验匹配"], "gaps": [], "verification_questions": ["核实项目成果"],
            },
            {
                "candidate_name": "候选人乙", "source_name": "resume-b.txt",
                "criterion_scores": {"experience": 80}, "score_evidence": {"experience": "有销售经验"},
                "strengths": ["销售执行"], "gaps": [], "verification_questions": ["核实销售口径"],
            },
        ],
    }, ensure_ascii=False)
    gateway = RecordingGateway([
        route_json("recruitment_screening"), screening,
        route_json("recruitment_screening", turn_action="compare_existing"),
        "如果只能选择一人，建议候选人甲；其B2B经验与岗位更直接匹配。",
    ])
    settings = Settings(data_dir=tmp_path, database_path=tmp_path / "unihr.db")
    repository = ConversationRepository(settings.database_path)
    service = WorkbenchService(
        settings, repository, model_gateway=gateway, workflow_checkpointer=InMemorySaver(),
    )

    first = service.handle_turn(
        message="请筛选产业拓展实习生简历",
        conversation_id=None,
        task_mode="简历筛选",
        uploads=[
            type("Upload", (), {"filename": "resume-a.txt", "mime_type": "text/plain", "content": "候选人甲，有B2B运营经验".encode("utf-8")})(),
            type("Upload", (), {"filename": "resume-b.txt", "mime_type": "text/plain", "content": "候选人乙，有销售经验".encode("utf-8")})(),
        ],
        document_font="宋体",
        document_font_size=14,
    )

    def fail_if_rescreened(*args, **kwargs):
        raise AssertionError("follow-up comparison must reuse the previous screening result")

    service.recruitment.screen = fail_if_rescreened
    followup = service.handle_turn(
        message="再单独对比前两名，如果非要选一个，你会选谁？",
        conversation_id=first.conversation.id,
        task_mode="简历筛选",
        uploads=[],
        document_font="宋体",
        document_font_size=14,
    )

    assert followup.artifacts == []
    assert "建议候选人甲" in followup.assistant_message.content
    assert not any(item["event"] == "recruitment.evaluate" for item in followup.trace)
    assert followup.run.graph_path == ["prepare", "reuse_result", "quality_gate", "finalize"]


def test_exact_recruitment_followup_reuses_result_when_router_model_is_unavailable(tmp_path):
    gateway = RouterUnavailableThenAnswerGateway("最终建议候选人甲和候选人乙进入后续人工面试。")
    settings = Settings(data_dir=tmp_path, database_path=tmp_path / "unihr.db")
    repository = ConversationRepository(settings.database_path)
    conversation = repository.create_conversation("产业拓展候选人筛选")
    repository.add_message(conversation.id, "user", "请筛选六位产业拓展候选人")
    repository.add_message(
        conversation.id,
        "assistant",
        "已完成六位候选人的证据评分：候选人甲第一、候选人乙第二。",
    )
    service = WorkbenchService(
        settings, repository, model_gateway=gateway, workflow_checkpointer=InMemorySaver(),
    )

    def fail_if_rescreened(*args, **kwargs):
        raise AssertionError("this follow-up must not run recruitment screening again")

    service.recruitment.screen = fail_if_rescreened
    result = service.handle_turn(
        message="请选择两名最终的候选人",
        conversation_id=conversation.id,
        task_mode="简历筛选",
        uploads=[],
        document_font="宋体",
        document_font_size=14,
    )

    assert "候选人甲和候选人乙" in result.assistant_message.content
    assert result.interpretation.execution_mode.value == "reuse_result"
    assert result.run.graph_path == ["prepare", "reuse_result", "quality_gate", "finalize"]
    assert not any(item["event"] == "recruitment.evaluate" for item in result.trace)


def test_document_confirmation_reuses_draft_without_regenerating(tmp_path):
    draft = json.dumps({
        "title": "年度工作会议通知",
        "document_type": "通知",
        "addressee": "各部门",
        "lead": "现将有关事项通知如下。",
        "sections": [{"heading": "一、会议安排", "paragraphs": ["会议时间另行通知。"]}],
        "closing": "请遵照执行。",
        "issuer": "人事处",
        "date_text": "2026年8月6日",
    }, ensure_ascii=False)
    gateway = RecordingGateway([
        route_json("document_generation"), draft,
        route_json("document_generation", turn_action="confirm_existing"),
        "已确认沿用当前草稿；本轮没有重新生成文件。",
    ])
    settings = Settings(data_dir=tmp_path, database_path=tmp_path / "unihr.db")
    repository = ConversationRepository(settings.database_path)
    service = WorkbenchService(
        settings, repository, model_gateway=gateway, workflow_checkpointer=InMemorySaver(),
    )

    first = service.handle_turn(
        message="起草一份年度工作会议通知",
        conversation_id=None,
        task_mode="公文材料",
        uploads=[],
        document_font="宋体",
        document_font_size=14,
    )
    assert gateway.calls[1][1]["model"] == settings.document_model

    def fail_if_regenerated(*args, **kwargs):
        raise AssertionError("confirmation must not regenerate the document")

    service._create_document = fail_if_regenerated
    followup = service.handle_turn(
        message="这份草稿可以定稿，就按这个版本",
        conversation_id=first.conversation.id,
        task_mode="公文材料",
        uploads=[],
        document_font="宋体",
        document_font_size=14,
    )

    assert followup.artifacts == []
    assert any(item["event"] == "workflow.result_reused" for item in followup.trace)
    assert not any(item["event"] == "document.compose" for item in followup.trace)


def test_planner_delegated_document_step_is_new_execution_not_mutation(tmp_path):
    gateway = RecordingGateway([
        route_json("document_generation", turn_action="revise_existing"),
    ])
    settings = Settings(data_dir=tmp_path, database_path=tmp_path / "unihr.db")
    repository = ConversationRepository(settings.database_path)
    conversation = repository.create_conversation("跨 Agent 招聘协同")
    repository.add_message(conversation.id, "assistant", "招聘筛选已完成。")
    service = WorkbenchService(
        settings, repository, model_gateway=gateway, workflow_checkpointer=InMemorySaver(),
    )
    from docx import Document

    artifact_path = tmp_path / "候选人面试确认材料.docx"
    document = Document()
    document.add_heading("候选人面试确认材料", level=1)
    document.add_paragraph("张三：岗位相关优势待面试复核。")
    document.save(artifact_path)
    create_calls = []

    def create_document(*args, **kwargs):
        create_calls.append((args, kwargs))
        return (
            "已生成候选人面试确认材料。",
            [(
                "document", artifact_path, artifact_path.name,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                {},
            )],
            [],
        )

    service._create_document = create_document
    brief = TaskBrief.model_validate({
        "kind": "document",
        "objective": "根据人工确认快照生成面试确认材料",
        "constraints": ["不得写入受保护属性"],
        "deliverables": ["Word 面试确认材料"],
        "mode_spec": {
            "document_type": "确认材料",
            "title": "候选人面试确认材料",
            "issuer": "人事部门",
            "date_text": "2026年8月10日",
        },
    })

    result = service.handle_turn(
        message="人工确认候选人：张三。",
        conversation_id=conversation.id,
        task_mode="公文材料",
        task_brief=brief,
        uploads=[],
        document_font="宋体",
        document_font_size=14,
        orchestration_enabled=False,
        message_role="system",
        delegated_step_execution=True,
    )

    assert result.interpretation.turn_action.value == "start_new"
    assert result.interpretation.execution_mode.value == "execute"
    assert result.interpretation.reason_code == "planner_delegated_step"
    assert len(create_calls) == 1
    assert len(result.artifacts) == 1
    assert result.artifacts[0].mime_type.endswith("wordprocessingml.document")


def test_spreadsheet_explanation_reuses_workbook_without_reprocessing(tmp_path):
    gateway = RecordingGateway([
        route_json("spreadsheet_processing"), "已识别姓名和部门字段，请确认计算规则。",
        route_json("spreadsheet_processing", turn_action="explain_existing"),
        "当前工作簿只是准备稿，尚未执行正式金额计算。",
    ])
    settings = Settings(data_dir=tmp_path, database_path=tmp_path / "unihr.db")
    repository = ConversationRepository(settings.database_path)
    service = WorkbenchService(
        settings, repository, model_gateway=gateway, workflow_checkpointer=InMemorySaver(),
    )

    first = service.handle_turn(
        message="读取这份人员表并生成准备稿",
        conversation_id=None,
        task_mode="表格处理",
        uploads=[type("Upload", (), {
            "filename": "people.csv", "mime_type": "text/csv", "content": "姓名,部门\n张三,人事处".encode("utf-8"),
        })()],
        document_font="宋体",
        document_font_size=14,
    )

    def fail_if_reprocessed(*args, **kwargs):
        raise AssertionError("explanation must not reprocess the workbook")

    service._process_spreadsheet = fail_if_reprocessed
    followup = service.handle_turn(
        message="解释一下这个准备稿现在做了什么",
        conversation_id=first.conversation.id,
        task_mode="表格处理",
        uploads=[],
        document_font="宋体",
        document_font_size=14,
    )

    assert followup.artifacts == []
    assert any(item["event"] == "workflow.result_reused" for item in followup.trace)
    assert not any(item["event"] == "spreadsheet.transform" for item in followup.trace)


def test_workbench_executes_validated_spreadsheet_plan_and_delivers_reviewable_draft(tmp_path):
    source = tmp_path / "工资汇总.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "测算"
    sheet.append(["工号", "基本工资", "绩效部分"])
    sheet.append(["0000001", 72000, 24000])
    workbook.save(source)
    workbook.close()
    plan = json.dumps({
        "version": "1.0",
        "objective": "增加年度计入工资总额公式列",
        "operations": [{
            "type": "add_formula_column",
            "sheet": "测算",
            "header": "年度计入工资总额",
            "formula_template": "=B{row}+C{row}",
            "number_format": "0",
        }],
        "validations": {
            "required_sheets": ["测算"],
            "required_headers": [{"sheet": "测算", "headers": ["工号", "年度计入工资总额"]}],
            "unique_keys": [{"sheet": "测算", "header": "工号"}],
        },
        "assumptions": [],
        "unresolved_questions": [],
        "requires_human_review": True,
    }, ensure_ascii=False)
    gateway = RecordingGateway([route_json("spreadsheet_processing"), plan])
    settings = Settings(data_dir=tmp_path / "data", database_path=tmp_path / "unihr.db")
    repository = ConversationRepository(settings.database_path)
    service = WorkbenchService(
        settings, repository, model_gateway=gateway, workflow_checkpointer=InMemorySaver(),
    )

    result = service.handle_turn(
        message="在测算表中增加年度计入工资总额公式列",
        conversation_id=None,
        task_mode="表格处理",
        uploads=[type("Upload", (), {
            "filename": source.name,
            "mime_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "content": source.read_bytes(),
        })()],
        document_font="宋体",
        document_font_size=14,
    )

    assert result.artifacts[0].metadata["changed_cell_count"] == 1
    assert result.artifacts[0].metadata["validation"]["passed"] is True
    assert result.artifacts[0].metadata["inspection_payload_ref"]["kind"] == "spreadsheet.inspection"
    assert result.artifacts[0].metadata["validation_payload_ref"]["kind"] == "spreadsheet.validation"
    assert result.run.status == "succeeded"
    assert result.artifacts[0].metadata["review_recommended"] is True
    assert result.artifacts[0].metadata["requires_human_review"] is False
    assert any(item["event"] == "model.spreadsheet_plan" for item in result.trace)
    assert any(item["event"] == "spreadsheet.transform" for item in result.trace)


def test_spreadsheet_quality_failure_repairs_plan_and_reexecutes_from_original_source(tmp_path):
    def plan(required_header: str) -> str:
        return json.dumps({
            "version": "1.1",
            "objective": "增加月度应发列",
            "operations": [{
                "type": "add_formula_column",
                "sheet": "输入数据",
                "header": "月度应发",
                "formula_template": "=B{row}+C{row}",
                "number_format": "0.00",
            }],
            "validations": {
                "required_headers": [{
                    "sheet": "输入数据", "headers": ["姓名", "基本工资", "绩效", required_header],
                }],
            },
            "assumptions": [], "unresolved_questions": [], "requires_human_review": False,
        }, ensure_ascii=False)

    gateway = RecordingGateway([
        route_json("spreadsheet_processing"),
        plan("错误的校验列"),
        plan("月度应发"),
    ])
    settings = Settings(data_dir=tmp_path / "data", database_path=tmp_path / "unihr.db")
    repository = ConversationRepository(settings.database_path)
    service = WorkbenchService(
        settings, repository, model_gateway=gateway, workflow_checkpointer=InMemorySaver(),
    )

    result = service.handle_turn(
        message="请增加月度应发列，等于基本工资加绩效",
        conversation_id=None,
        task_mode="表格处理",
        uploads=[type("Upload", (), {
            "filename": "工资.csv", "mime_type": "text/csv",
            "content": "姓名,基本工资,绩效\n张三,8000,1200".encode("utf-8"),
        })()],
        document_font="宋体",
        document_font_size=14,
    )

    assert result.run.status == "succeeded"
    assert result.run.graph_path == [
        "prepare", "execute", "quality_gate", "repair", "execute", "quality_gate", "finalize",
    ]
    assert any(item["event"] == "spreadsheet.repair.started" for item in result.trace)
    assert any(item["event"] == "spreadsheet.repair.plan_created" for item in result.trace)
    purposes = [kwargs.get("purpose") for _, kwargs in gateway.calls]
    assert purposes == ["turn_interpretation", "spreadsheet_plan", "spreadsheet_execution_repair"]
    artifact = result.artifacts[0]
    assert artifact.metadata["validation"]["passed"] is True
    assert artifact.metadata["plan_source"].endswith("自修复规划器")
    assert artifact.metadata["repair_count"] == 1
    assert "重新执行并复验通过" in result.assistant_message.content
    workbook = load_workbook(repository.get_artifact(artifact.id).path, data_only=False)
    assert workbook["输入数据"]["D2"].value == "=B2+C2"
    workbook.close()


def test_workbench_executes_a_simple_date_request_without_calling_the_model_planner(tmp_path):
    gateway = RecordingGateway([route_json("spreadsheet_processing")])
    settings = Settings(data_dir=tmp_path / "data", database_path=tmp_path / "unihr.db")
    repository = ConversationRepository(settings.database_path)
    service = WorkbenchService(
        settings, repository, model_gateway=gateway, workflow_checkpointer=InMemorySaver(),
    )

    result = service.handle_turn(
        message="帮我筛出2026-08-09到2026-09-08合同到期的人",
        conversation_id=None,
        task_mode="表格处理",
        task_brief={"kind": "spreadsheet", "objective": "筛选合同到期人员"},
        uploads=[type("Upload", (), {
            "filename": "contracts.csv",
            "mime_type": "text/csv",
            "content": (
                "employee_id,name,contract_end_date\n"
                "E001,测试员工,2026-08-12\n"
                "E002,示例员工,2026-09-30"
            ).encode("utf-8"),
        })()],
        document_font="宋体",
        document_font_size=14,
    )

    artifact = result.artifacts[0]
    output = load_workbook(repository.get_artifact(artifact.id).path, data_only=False)
    operations = artifact.metadata["operation_plan"]["operations"]
    filtered = output[operations[1]["sheet"]]
    assert filtered.row_dimensions[2].hidden is False
    assert filtered.row_dimensions[3].hidden is True
    assert output[operations[0]["source_sheet"]]["C3"].value == "2026-09-30"
    assert all(sheet.page_setup.fitToWidth == 1 for sheet in output.worksheets)
    assert all(sheet.sheet_properties.pageSetUpPr.fitToPage for sheet in output.worksheets)
    assert all(sheet.print_area for sheet in output.worksheets)
    output.close()
    assert "日期筛选结果" in result.assistant_message.content
    assert "2026-08-09 至 2026-09-08" in result.assistant_message.content
    assert any(item["event"] == "spreadsheet.date_filter.applied" for item in result.trace)
    assert not any(kwargs.get("purpose") == "spreadsheet_plan" for _, kwargs in gateway.calls)
    assert result.run.status == "succeeded"
    assert result.artifacts[0].metadata["requires_human_review"] is False
    assert result.artifacts[0].metadata["review_recommended"] is True


def test_workbench_creates_spreadsheet_from_text_without_attachment(tmp_path):
    plan = json.dumps({
        "version": "1.1",
        "objective": "从零创建招聘进度跟踪表",
        "operations": [
            {
                "type": "create_sheet",
                "sheet": "招聘进度",
                "headers": ["岗位", "招聘人数", "已面试人数", "完成率"],
                "rows": [["研发工程师", 3, 1, None]],
            },
            {
                "type": "add_formula_column",
                "sheet": "招聘进度",
                "header": "完成率",
                "formula_template": "=IFERROR(C{row}/B{row},0)",
                "number_format": "0.0%",
            },
        ],
        "validations": {
            "required_sheets": ["招聘进度"],
            "required_headers": [{
                "sheet": "招聘进度",
                "headers": ["岗位", "招聘人数", "已面试人数", "完成率"],
            }],
        },
        "assumptions": [],
        "unresolved_questions": [],
        "requires_human_review": True,
    }, ensure_ascii=False)
    gateway = RecordingGateway([route_json("spreadsheet_processing"), plan])
    settings = Settings(data_dir=tmp_path / "data", database_path=tmp_path / "unihr.db")
    repository = ConversationRepository(settings.database_path)
    service = WorkbenchService(
        settings, repository, model_gateway=gateway, workflow_checkpointer=InMemorySaver(),
    )

    progress_events: list[tuple[str, str, dict]] = []
    result = service.handle_turn(
        message="请从零创建招聘进度跟踪表，包含完成率公式",
        conversation_id=None,
        task_mode="表格处理",
        uploads=[],
        document_font="宋体",
        document_font_size=14,
        progress=lambda event, summary, details: progress_events.append((event, summary, details)),
    )

    assert len(result.artifacts) == 1
    assert result.artifacts[0].metadata["created_from_scratch"] is True
    assert result.artifacts[0].metadata["source_files"] == []
    assert result.artifacts[0].metadata["validation_payload_ref"]["kind"] == "spreadsheet.validation"
    assert result.run.status == "succeeded"
    assert result.artifacts[0].metadata["requires_human_review"] is False
    assert result.artifacts[0].metadata["review_recommended"] is True
    assert any(item["event"] == "spreadsheet.create" for item in result.trace)
    progress_names = [event for event, _summary, _details in progress_events]
    assert progress_names.index("graph.node.execute.started") < progress_names.index(
        "tool.spreadsheet.plan.started"
    )
    assert progress_names.index("tool.spreadsheet.plan.completed") < progress_names.index(
        "tool.spreadsheet.build.started"
    )
    assert progress_names.index("tool.spreadsheet.build.completed") < progress_names.index(
        "tool.artifact.harness.started"
    )
    assert progress_names.index("tool.artifact.harness.completed") < progress_names.index(
        "graph.node.execute.completed"
    )


def test_workbench_creates_compact_inline_spreadsheet_when_planner_model_is_not_needed(tmp_path):
    message = (
        "做一张培训完成情况表：制度培训应到50实到42，安全培训应到40实到38，"
        "合规培训应到30实到30；计算每项完成率和总计，生成 Excel。"
    )
    gateway = RecordingGateway([route_json("spreadsheet_processing")])
    settings = Settings(data_dir=tmp_path / "data", database_path=tmp_path / "unihr.db")
    repository = ConversationRepository(settings.database_path)
    service = WorkbenchService(
        settings, repository, model_gateway=gateway, workflow_checkpointer=InMemorySaver(),
    )

    result = service.handle_turn(
        message=message, conversation_id=None, task_mode="自动选择", uploads=[],
        document_font="宋体", document_font_size=14,
    )

    assert result.run.status == "succeeded"
    assert len(result.artifacts) == 1
    artifact = result.artifacts[0]
    workbook = load_workbook(repository.get_artifact(artifact.id).path, data_only=False)
    sheet = workbook["培训完成情况"]
    assert sheet["A5"].value == "总计"
    assert sheet["B5"].value == 120
    assert sheet["C5"].value == 110
    assert sheet["D2"].value == '=IF(B2=0,"",C2/B2)'
    assert sheet["D5"].value == '=IF(B5=0,"",C5/B5)'
    workbook.close()
    assert any(item["event"] == "spreadsheet.plan.compiled" for item in result.trace)
    assert any(item["event"] == "skill.graph.node.business_executed" for item in result.trace)
    recorded_stages = [
        item["stage_id"] for item in result.trace
        if item["event"] == "skill.graph.stage.recorded"
    ]
    assert recorded_stages == ["plan", "execute", "validate", "harness"]
    assert any(item["event"] == "skill.graph.node.result_offloaded" for item in result.trace)
    with sqlite3.connect(repository.database_path) as db:
        payload_kinds = {
            row[0] for row in db.execute(
                "select kind from run_payloads where run_id=?", (result.run.id,)
            ).fetchall()
        }
    assert any(kind.endswith(".workbench_result") for kind in payload_kinds)
    assert any(kind.endswith(".stage_states") for kind in payload_kinds)
    assert {
        "professional.spreadsheet.stage.plan",
        "professional.spreadsheet.stage.execute",
        "professional.spreadsheet.stage.validate",
        "professional.spreadsheet.stage.harness",
    }.issubset(payload_kinds)
    atomic_stages = [
        item["stage_id"] for item in result.trace
        if item["event"] == "spreadsheet.stage.completed"
    ]
    assert atomic_stages == ["plan", "execute", "validate", "harness"]
    assert not any(kwargs.get("purpose") == "spreadsheet_plan" for _, kwargs in gateway.calls)


def test_workbench_creates_recruitment_progress_spreadsheet_from_natural_inline_records(tmp_path):
    message = (
        "帮我整理一个三部门招聘进度表：研发中心计划招12人，目前面试9人；"
        "运营中心计划招8人，目前面试6人；财务中心计划招5人，目前面试5人。"
        "请算出每个部门的面试达成率和整体达成率，给我一个能直接下载的 Excel。"
    )
    compiled_plan = _compile_inline_records_spreadsheet_plan(
        TaskBrief.model_validate({"kind": "spreadsheet", "objective": message}), message,
    )
    assert compiled_plan is not None
    gateway = RecordingGateway([route_json("spreadsheet_processing")])
    settings = Settings(data_dir=tmp_path / "data", database_path=tmp_path / "unihr.db")
    repository = ConversationRepository(settings.database_path)
    service = WorkbenchService(
        settings, repository, model_gateway=gateway, workflow_checkpointer=InMemorySaver(),
    )

    result = service.handle_turn(
        message=message, conversation_id=None, task_mode="自动选择", uploads=[],
        document_font="宋体", document_font_size=14,
    )

    assert result.run.status == "succeeded"
    artifact = result.artifacts[0]
    workbook = load_workbook(repository.get_artifact(artifact.id).path, data_only=False)
    sheet = workbook[[name for name in workbook.sheetnames if not name.startswith("UniHR")][0]]
    assert list(sheet.values)[0] == ("项目", "计划招", "目前面试", "达成率")
    assert sheet["A2"].value == "研发中心"
    assert sheet["B5"].value == 25
    assert sheet["C5"].value == 20
    assert sheet["D5"].value == '=IF(B5=0,"",C5/B5)'
    workbook.close()
    assert any(item["event"] == "spreadsheet.plan.compiled" for item in result.trace)


def test_spreadsheet_planner_repairs_synonym_that_violates_locked_task_brief_header(tmp_path):
    source = tmp_path / "工资汇总.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "测算"
    sheet.append(["工号", "基本工资", "绩效部分"])
    sheet.append(["0000001", 72000, 24000])
    workbook.save(source)
    workbook.close()

    def plan_with_header(header: str) -> str:
        return json.dumps({
            "version": "1.0",
            "objective": "增加公式列",
            "operations": [{
                "type": "add_formula_column", "sheet": "测算", "header": header,
                "formula_template": "=B{row}+C{row}", "number_format": "0",
            }],
            "validations": {}, "assumptions": [], "unresolved_questions": [],
            "requires_human_review": True,
        }, ensure_ascii=False)

    gateway = RecordingGateway([
        route_json("spreadsheet_processing", turn_action="clarify"),
        plan_with_header("应发合计"),
        plan_with_header("年度计入工资总额"),
    ])
    settings = Settings(data_dir=tmp_path / "data", database_path=tmp_path / "unihr.db")
    repository = ConversationRepository(settings.database_path)
    service = WorkbenchService(
        settings, repository, model_gateway=gateway, workflow_checkpointer=InMemorySaver(),
    )

    result = service.handle_turn(
        message="请按已填写的表格要求生成可复核工作簿",
        conversation_id=None,
        task_mode="表格处理",
        task_brief={
            "kind": "spreadsheet",
            "objective": "在测算表中增加年度计入工资总额公式列",
            "mode_spec": {
                "goal": "在测算表中增加年度计入工资总额公式列",
                "source_sheet": "测算",
                "calculation_rules": "年度计入工资总额按已确认的组合规则计算",
            },
        },
        uploads=[type("Upload", (), {
            "filename": source.name,
            "mime_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "content": source.read_bytes(),
        })()],
        document_font="宋体",
        document_font_size=14,
    )

    operations = result.artifacts[0].metadata["operation_results"]
    assert operations[0]["summary"].endswith("“年度计入工资总额”列填充 1 行公式")
    assert any(item["event"] == "spreadsheet.plan.repaired" for item in result.trace)
    assert result.interpretation.execution_mode.value == "execute"
    assert result.interpretation.reason_code == "explicit_new_task_brief"
    assert len(gateway.calls) == 3


def test_policy_followup_reuses_existing_answer_without_starting_new_task(tmp_path):
    gateway = RecordingGateway([
        route_json("policy_question"), "现有政策证据表明该事项需要人工审批。",
        route_json("policy_question", turn_action="explain_existing"),
        "这里的“人工审批”是指材料提交后由有权限的人员复核确认。",
    ])
    settings = Settings(data_dir=tmp_path, database_path=tmp_path / "unihr.db")
    repository = ConversationRepository(settings.database_path)
    service = WorkbenchService(
        settings, repository, model_gateway=gateway, workflow_checkpointer=InMemorySaver(),
    )

    first = service.handle_turn(
        message="这个事项按政策是否需要审批？",
        conversation_id=None,
        task_mode="政策问答",
        uploads=[],
        document_font="宋体",
        document_font_size=14,
    )
    followup = service.handle_turn(
        message="你说的人工审批具体指什么？",
        conversation_id=first.conversation.id,
        task_mode="政策问答",
        uploads=[],
        document_font="宋体",
        document_font_size=14,
    )

    assert followup.artifacts == []
    assert "人工审批" in followup.assistant_message.content
    assert any(item["event"] == "workflow.result_reused" for item in followup.trace)


def test_policy_retrieval_uses_contextual_rewrite_and_query_variants(tmp_path):
    router_result = json.dumps({
        "primary_intent": "policy_question",
        "secondary_intents": [],
        "context_mode": "uses_conversation",
        "missing_inputs": [],
        "risk_level": "medium",
        "data_class": "internal_hr",
        "required_capabilities": ["knowledge.retrieve", "evidence.answer"],
        "confidence": 0.96,
        "reason_summary": "补全了政策主题和地区",
        "turn_action": "start_new",
        "resolved_request": "浙江省博士后进站备案需要提交哪些材料",
        "retrieval_queries": ["浙江省 博士后 进站备案 材料", "博士后进站备案 管理办法"],
        "priority_constraints": ["适用地区：浙江省"],
    }, ensure_ascii=False)
    gateway = RecordingGateway([router_result, "根据当前索引，尚未找到可引用的有效依据。"])
    knowledge = RecordingKnowledgeIndex()
    settings = Settings(data_dir=tmp_path, database_path=tmp_path / "unihr.db")
    repository = ConversationRepository(settings.database_path)
    service = WorkbenchService(
        settings,
        repository,
        model_gateway=gateway,
        knowledge_index=knowledge,
        workflow_checkpointer=InMemorySaver(),
    )

    result = service.handle_turn(
        message="那浙江省需要哪些材料？",
        conversation_id=None,
        task_mode="政策问答",
        uploads=[],
        document_font="宋体",
        document_font_size=14,
    )

    assert "浙江省博士后进站备案需要提交哪些材料" in knowledge.queries
    assert "浙江省 博士后 进站备案 材料" in knowledge.queries
    assert any(item["event"] == "query.rewritten" for item in result.trace)


def test_explicit_slash_skill_is_injected_and_traced(tmp_path):
    gateway = RecordingGateway([route_json("recruitment_screening", missing_inputs=["objective"]), "已按显式 Skill 执行。"])
    settings = Settings(data_dir=tmp_path, database_path=tmp_path / "unihr.db")
    repository = ConversationRepository(settings.database_path)
    registry = SkillRegistry(settings.database_path)
    skill = registry.create(SkillDraft(
        name="verify-onboarding-files",
        display_name="入职材料核验",
        description="核验入职材料并输出缺失清单。",
        instructions="确认清单；逐项核验；输出缺失项。",
        trigger_phrases=["入职材料核验"],
        source_type="prompt",
        enabled=True,
        status="active",
        metadata={"resource_plan": [{
            "kind": "references",
            "name": "onboarding-checklist.md",
            "purpose": "入职材料标准清单",
            "required": False,
        }]},
    ))
    resources = SkillResourceStore(settings.database_path, tmp_path / "skill-resources")
    resources.add(
        skill, None,
        kind="references",
        filename="onboarding-checklist.md",
        content="身份证明、劳动合同和审批记录必须逐项核验。".encode("utf-8"),
        mime_type="text/markdown",
    )
    resources.add(
        skill, None,
        kind="scripts",
        filename="unsafe.py",
        content=b"raise RuntimeError('must never execute')\n",
        mime_type="text/x-python",
    )
    service = WorkbenchService(
        settings, repository, model_gateway=gateway, skill_registry=registry,
        workflow_checkpointer=InMemorySaver(),
    )

    result = service.handle_turn(
        message="请处理这件事",
        conversation_id=None,
        task_mode="自动选择",
        task_brief={
            "kind": "general",
            "objective": "处理入职材料",
            "selected_skill_ids": [skill.id],
        },
        uploads=[], document_font="宋体", document_font_size=14,
    )

    selected = next(item for item in result.trace if item["event"] == "skill.selected")
    assert selected["skill_id"] == skill.id
    assert selected["selection_mode"] == "explicit"
    answer_prompt = json.dumps(gateway.calls[-1][0], ensure_ascii=False)
    assert "逐项核验" in answer_prompt
    assert "身份证明、劳动合同和审批记录" in answer_prompt
    assert "must never execute" not in answer_prompt
    loaded = next(item for item in result.trace if item["event"] == "skill.resource.loaded")
    assert loaded["resource_name"] == "onboarding-checklist.md"
    quarantined = next(item for item in result.trace if item["event"] == "skill.script.quarantined")
    assert quarantined["scripts_executed"] is False
    assert any(item["kind"] == "skill_reference" for item in result.context_uses)
    assert "execute" in result.run.graph_path
    assert not any(item["event"] == "workflow.clarification_required" for item in result.trace)
    assert result.interpretation.business_intent.value == "conversation"
    assert result.interpretation.reason_code == "explicit_skill_domain_contract"
    assert len([item for item in result.trace if item["event"] == "skill.selected"]) == 1
    assert any(item["event"] == "skill.graph.started" for item in result.trace)
    assert any(item["event"] == "skill.graph.node.completed" for item in result.trace)
    assert any(item["event"] == "skill.graph.completed" for item in result.trace)


def test_explicit_data_skill_requires_real_business_input_instead_of_simulating(tmp_path):
    gateway = RecordingGateway([
        route_json("recruitment_screening"),
        route_json("conversation", turn_action="continue_with_tools"),
        "已根据真实合同数据完成检查。",
    ])
    settings = Settings(data_dir=tmp_path, database_path=tmp_path / "unihr.db")
    repository = ConversationRepository(settings.database_path)
    registry = SkillRegistry(settings.database_path)
    skill = registry.create(SkillDraft(
        name="check-contract-expiry",
        display_name="合同到期周检",
        description="读取员工合同数据并形成到期提醒。",
        instructions="读取真实数据；检查到期日；不得虚构人员；输出人工复核清单。",
        trigger_phrases=["合同到期周检"],
        capabilities=["file.extract", "document.compose", "knowledge.retrieve"],
        source_type="prompt",
        enabled=True,
        status="active",
    ))
    resources = SkillResourceStore(settings.database_path, tmp_path / "skill-resources")
    resources.add(
        skill, None,
        kind="references",
        filename="schema.md",
        content="字段：employee_id、employee_name、contract_end_date。".encode("utf-8"),
        mime_type="text/markdown",
    )
    service = WorkbenchService(
        settings, repository, model_gateway=gateway, skill_registry=registry,
        workflow_checkpointer=InMemorySaver(),
    )

    result = service.handle_turn(
        message="帮我看看这周谁的合同要到期",
        conversation_id=None,
        task_mode="自动选择",
        task_brief={
            "kind": "general",
            "objective": "执行合同到期周检",
            "selected_skill_ids": [skill.id],
        },
        uploads=[], document_font="宋体", document_font_size=14,
    )

    assert result.interpretation.business_intent.value == "conversation"
    assert "没有可处理的业务数据" in result.assistant_message.content
    assert "模拟数据" in result.assistant_message.content
    assert any(item["event"] == "skill.resource.loaded" for item in result.trace)
    assert any(item["event"] == "skill.input.required" for item in result.trace)
    assert len(gateway.calls) == 1

    continued = service.handle_turn(
        message="数据传好了，继续",
        conversation_id=result.conversation.id,
        task_mode="自动选择",
        task_brief={"kind": "general", "objective": "继续"},
        uploads=[UploadPayload(
            filename="contracts.csv",
            mime_type="text/csv",
            content="employee_id,employee_name,contract_end_date\nE001,测试员工,2026-08-12\n".encode("utf-8"),
        )],
        document_font="宋体", document_font_size=14,
    )

    assert "真实合同数据" in continued.assistant_message.content
    assert any(item["event"] == "skill.selection.resumed" for item in continued.trace)
    assert any(item["event"] == "skill.resource.loaded" for item in continued.trace)
    assert not any(item["event"] == "skill.input.required" for item in continued.trace)
