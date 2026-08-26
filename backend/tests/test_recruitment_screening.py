import json
import re
import sqlite3
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from app.domains.conversations.contracts import AttachmentRecord
from app.domains.recruitment import RecruitmentCapacityEvaluator, RecruitmentScreeningModule
from app.domains.recruitment.contracts import ResumeSourceProfile
from app.domains.recruitment.version_resolver import ResumeVersionResolver
from app.domains.runs import RunPayloadStore, RunPointer
from app.models.gateway import ModelReply
from app.observability import tool_progress_scope


class FakeGateway:
    def complete(self, messages, **kwargs):
        payload = {
            "job_title": "研发工程师",
            "criteria": [
                {"id": "python", "label": "Python经验", "weight": 60, "requirement": "3年以上"},
                {"id": "delivery", "label": "交付成果", "weight": 40, "requirement": "可验证项目"},
            ],
            "candidates": [
                {
                    "candidate_name": "候选人甲",
                    "source_name": "甲.pdf",
                    "criterion_scores": {"python": 90, "delivery": 80},
                    "score_evidence": {"python": "5年Python开发", "delivery": "负责两个已上线项目"},
                    "strengths": ["经历匹配"], "gaps": [], "verification_questions": ["核实项目角色"],
                },
                {
                    "candidate_name": "候选人乙",
                    "source_name": "乙.docx",
                    "criterion_scores": {"python": 60, "delivery": 50},
                    "score_evidence": {"python": "2年Python开发", "delivery": "未找到明确证据"},
                    "strengths": [], "gaps": ["年限不足"], "verification_questions": ["补充项目材料"],
                },
            ],
        }
        return ModelReply(json.dumps(payload, ensure_ascii=False), "fake-model", 12, {})


class BatchGateway:
    def __init__(self, *, omit_source: str = "", include_protected: bool = False):
        self.calls = []
        self.omit_source = omit_source
        self.include_protected = include_protected

    def complete(self, messages, **kwargs):
        self.calls.append((messages, kwargs))
        if kwargs.get("purpose") == "recruitment_rubric":
            payload = {
                "job_title": "批次测试岗位",
                "criteria": [
                    {"id": "required", "label": "必须条件", "weight": 40, "requirement": "满足岗位硬条件"},
                    {"id": "experience", "label": "相关经历", "weight": 30, "requirement": "岗位相关经验"},
                    {"id": "delivery", "label": "成果证据", "weight": 20, "requirement": "可验证成果"},
                    {"id": "evidence", "label": "证据完整性", "weight": 10, "requirement": "材料完整"},
                ],
            }
            return ModelReply(json.dumps(payload, ensure_ascii=False), "rubric-model", 5, {})
        content = messages[-1]["content"]
        sources = list(dict.fromkeys(re.findall(r"resume-\d+\.txt", content)))
        candidates = []
        for index, source in enumerate(sources):
            if source == self.omit_source:
                continue
            evidence = "女性，25岁，有3年相关经验" if self.include_protected and index == 0 else "有3年相关经验和可核验成果"
            candidates.append({
                "candidate_name": f"候选人-{Path(source).stem}",
                "source_name": source,
                "criterion_scores": {"required": 80, "experience": 80, "delivery": 70, "evidence": 80},
                "score_evidence": {
                    "required": evidence,
                    "experience": "有3年相关经验",
                    "delivery": "材料列出已交付项目",
                    "evidence": "材料包含项目和职责",
                },
                "strengths": [evidence],
                "gaps": [],
                "verification_questions": ["核实项目角色"],
            })
        candidates.append({
            "candidate_name": "虚构候选人", "source_name": "not-in-case.pdf",
            "criterion_scores": {}, "score_evidence": {}, "strengths": [], "gaps": [],
            "verification_questions": [],
        })
        return ModelReply(json.dumps({"candidates": candidates}, ensure_ascii=False), "screen-model", 7, {})


class LocationProxyGateway:
    def complete(self, messages, **kwargs):
        if kwargs.get("purpose") == "recruitment_rubric":
            return ModelReply(json.dumps({
                "job_title": "宁波实习生",
                "criteria": [
                    {"id": "location", "label": "宁波实习可行性", "weight": 50, "requirement": "可在宁波实习"},
                    {"id": "experience", "label": "相关经历", "weight": 30, "requirement": "岗位经历"},
                    {"id": "evidence", "label": "证据完整性", "weight": 20, "requirement": "材料完整"},
                ],
            }, ensure_ascii=False), "rubric-model", 5, {})
        return ModelReply(json.dumps({"candidates": [{
            "candidate_name": "候选人0", "source_name": "resume-0.txt",
            "criterion_scores": {"location": 100, "experience": 80, "evidence": 80},
            "score_evidence": {
                "location": "现居宁波，就读于宁波工程学院，具备地理便利性",
                "experience": "有相关实习经历", "evidence": "材料完整",
            },
            "strengths": [], "gaps": [], "verification_questions": [],
        }]}, ensure_ascii=False), "screen-model", 7, {})


class LocalizationGateway:
    def complete(self, messages, **kwargs):
        if kwargs.get("purpose") == "recruitment_rubric":
            return ModelReply(json.dumps({
                "job_title": "Partnership Intern",
                "criteria": [
                    {"id": "experience", "label": "Relevant experience", "weight": 50},
                    {"id": "delivery", "label": "Delivery results", "weight": 30},
                    {"id": "evidence", "label": "Evidence quality", "weight": 20},
                ],
            }), "rubric-model", 5, {})
        return ModelReply(json.dumps({"candidates": [{
            "candidate_name": "Candidate A",
            "source_name": "resume.pdf",
            "criterion_scores": {"experience": 90, "delivery": 85, "evidence": 80},
            "score_evidence": {
                "experience": "Led B2B partner development",
                "delivery": "closed three partnership projects",
                "evidence": "documented customer requirements",
            },
            "strengths": ["B2B partnership delivery"],
            "gaps": [],
            "verification_questions": ["Verify independent contribution"],
        }]}), "screen-model", 7, {})


class VersionGateway:
    def complete(self, messages, **kwargs):
        if kwargs.get("purpose") == "recruitment_rubric":
            return ModelReply(json.dumps({
                "job_title": "产业拓展实习生",
                "criteria": [
                    {"id": "experience", "label": "拓展经历", "weight": 50},
                    {"id": "delivery", "label": "结果证据", "weight": 30},
                    {"id": "evidence", "label": "材料完整", "weight": 20},
                ],
            }, ensure_ascii=False), "rubric-model", 5, {})
        content = messages[-1]["content"]
        candidates = []
        if "张三-v2.pdf" in content:
            candidates.append({
                "candidate_name": "张三", "source_name": "张三-v2.pdf",
                "criterion_scores": {"experience": 90, "delivery": 80, "evidence": 90},
                "score_evidence": {
                    "experience": "负责产业合作伙伴拓展",
                    "delivery": "完成三个合作项目",
                    "evidence": "材料包含项目职责和结果",
                },
                "strengths": ["有产业拓展经验"], "gaps": [],
                "verification_questions": ["核实独立贡献"],
            })
        if "李四.pdf" in content:
            candidates.append({
                "candidate_name": "李四", "source_name": "李四.pdf",
                "criterion_scores": {"experience": 70, "delivery": 60, "evidence": 80},
                "score_evidence": {
                    "experience": "参与客户沟通",
                    "delivery": "协助完成一个项目",
                    "evidence": "材料包含职责描述",
                },
                "strengths": ["有沟通经验"], "gaps": [],
                "verification_questions": ["核实拓展深度"],
            })
        return ModelReply(json.dumps({"candidates": candidates}, ensure_ascii=False), "screen-model", 7, {})


def _attachment(tmp_path, name, extension, text):
    path = tmp_path / name
    path.write_text(text, encoding="utf-8")
    return AttachmentRecord(
        id=f"att-{name}", conversation_id="conv", message_id="msg", display_name=name,
        mime_type="application/octet-stream", extension=extension, byte_size=len(text),
        saved_path=str(path), extracted_text=text, summary={},
        created_at=datetime.now(timezone.utc).isoformat(),
    )


def test_recruitment_screening_creates_evidence_workbook(tmp_path):
    service = RecruitmentScreeningModule(FakeGateway(), tmp_path, "fake-model")
    result = service.screen(
        "岗位：研发工程师；Python至少3年；有可验证交付成果",
        [
            _attachment(tmp_path, "甲.pdf", ".pdf", "候选人甲，5年Python开发，两个上线项目"),
            _attachment(tmp_path, "乙.docx", ".docx", "候选人乙，2年Python开发"),
        ],
    )

    assert [item.candidate_name for item in result.report.candidates] == ["候选人甲", "候选人乙"]
    assert result.report.candidates[0].total_score == 86.0
    assert result.report.candidates[1].total_score == 56.0
    assert result.artifact_path.exists()

    workbook = load_workbook(result.artifact_path, data_only=False)
    assert workbook.sheetnames == ["筛选说明", "候选人排名", "评分明细", "重复材料", "版本归并", "人工复核", "Harness"]
    assert "不构成录用" in workbook["筛选说明"]["B6"].value
    assert workbook["评分明细"]["E2"].value == "=C2*D2"
    assert workbook["人工复核"]["E2"].value == "待复核"
    workbook.close()


def test_recruitment_capacity_baseline_exercises_thirty_resume_graph(tmp_path):
    report = RecruitmentCapacityEvaluator(tmp_path / "capacity-eval").evaluate()

    assert report.status == "passed"
    assert report.source_count == 30
    assert report.unique_source_count == 30
    assert report.candidate_count == 30
    assert report.run_count == 2
    assert report.batch_count == 8
    assert report.model_call_count == 10
    assert report.source_coverage == 1.0
    assert report.ranking_consistent is True
    assert report.payload_offload_verified is True
    assert report.artifact_verified is True
    assert report.formula_error_count == 0
    assert all(report.checks.values())
    assert (tmp_path / "capacity-eval" / "latest.json").exists()


def test_recruitment_graph_offloads_and_hydrates_candidate_batches(tmp_path):
    database = tmp_path / "unihr.db"
    store = RunPayloadStore(database)
    pointer = RunPointer(run_id="run_recruitment", tenant_id="tenant_a", owner_user_id="user_a")
    service = RecruitmentScreeningModule(
        FakeGateway(), tmp_path, "fake-model", payload_store=store,
    )

    result = service.screen(
        "研发工程师；Python至少3年；有可验证交付成果",
        [
            _attachment(tmp_path, "甲.pdf", ".pdf", "候选人甲，5年Python开发，两个上线项目"),
            _attachment(tmp_path, "乙.docx", ".docx", "候选人乙，2年Python开发"),
        ],
        payload_scope=pointer,
    )

    events = [item["event"] for item in result.trace]
    assert "graph.payload.offloaded" in events
    assert "graph.payload.hydrated" in events
    stage_events = [
        item for item in result.trace
        if item["event"] == "recruitment.stage.completed"
    ]
    stage_ids = [item["stage_id"] for item in stage_events]
    assert stage_ids[:10] == [
        "prepare", "deduplicate", "version_resolve", "rubric_lock", "batch_map",
        "reduce", "fairness_guard", "evidence_localize", "harness", "quality_gate",
    ]
    if result.repair_count:
        assert "repair" in stage_ids
        repair_index = stage_ids.index("repair")
        assert stage_ids[repair_index + 1:repair_index + 5] == [
            "fairness_guard", "evidence_localize", "harness", "quality_gate",
        ]
    assert stage_ids[-2:] == ["artifact_create", "finalize"]
    assert all(item["output_ref"].get("payload_id") for item in stage_events)
    with sqlite3.connect(database) as db:
        rows = db.execute(
            "SELECT kind,byte_size FROM run_payloads WHERE run_id=?", (pointer.run_id,)
        ).fetchall()
    kinds = {row[0] for row in rows}
    assert "recruitment.candidate_batches" in kinds
    assert "professional.recruitment.stage.prepare" in kinds
    assert "professional.recruitment.stage.finalize" in kinds
    assert all(row[1] > 0 for row in rows)


def test_recruitment_deduplicates_identical_resume_content(tmp_path):
    service = RecruitmentScreeningModule(FakeGateway(), tmp_path, "fake-model")
    same_text = "候选人甲，5年Python开发，两个上线项目"
    result = service.screen(
        "岗位：研发工程师",
        [
            _attachment(tmp_path, "甲.pdf", ".pdf", same_text),
            _attachment(tmp_path, "甲-重复.pdf", ".pdf", same_text),
            _attachment(tmp_path, "乙.docx", ".docx", "候选人乙，2年Python开发"),
        ],
    )

    assert result.report.case is not None
    assert result.report.case.source_count == 3
    assert result.report.case.unique_source_count == 2
    assert result.report.case.duplicate_source_count == 1
    assert result.report.duplicate_groups[0].representative_source == "甲.pdf"
    assert result.report.duplicate_groups[0].duplicate_sources == ["甲-重复.pdf"]
    assert [item.source_name for item in result.report.candidates] == ["甲.pdf", "乙.docx"]


def test_recruitment_consolidates_same_candidate_resume_versions(tmp_path):
    service = RecruitmentScreeningModule(VersionGateway(), tmp_path, "fake-model")
    result = service.screen(
        "产业拓展实习生，建议两人进入人工复核",
        [
            _attachment(
                tmp_path, "张三-v1.pdf", ".pdf",
                "姓名：张三\n电话：13800138000\n负责产业合作伙伴拓展，完成两个合作项目。",
            ),
            _attachment(
                tmp_path, "张三-v2.pdf", ".pdf",
                "姓名：张三\n电话：13800138000\n负责产业合作伙伴拓展，完成三个合作项目，材料包含项目职责和结果。",
            ),
            _attachment(tmp_path, "李四.pdf", ".pdf", "姓名：李四\n参与客户沟通，协助完成一个项目。"),
        ],
        selection_count=2,
    )

    assert [item.source_name for item in result.report.candidates] == ["张三-v2.pdf", "李四.pdf"]
    assert result.report.case is not None
    assert result.report.case.source_count == 3
    assert result.report.case.unique_source_count == 2
    assert result.report.case.duplicate_source_count == 0
    assert len(result.report.version_groups) == 1
    group = result.report.version_groups[0]
    assert group.status == "merged"
    assert group.representative_source == "张三-v2.pdf"
    assert group.identity_basis == "shared_contact_anchor"
    assert result.report.harness is not None
    assert result.report.harness.metrics["semantic_merged_source_count"] == 1
    assert "同一候选人的 1 份旧版本已归并" in result.answer
    assert "version_resolve" in result.graph_path

    workbook = load_workbook(result.artifact_path, data_only=False)
    assert workbook["版本归并"]["B2"].value == "已自动归并"
    assert workbook["版本归并"]["C2"].value == "张三-v2.pdf"
    workbook.close()


def test_resume_version_resolver_keeps_ambiguous_same_name_for_review(tmp_path):
    resolver = ResumeVersionResolver()
    sources = [
        _attachment(
            tmp_path, "张三-申请版.pdf", ".pdf",
            "姓名：张三\nPython development systems design API testing delivery Zhejiang University.",
        ),
        _attachment(
            tmp_path, "张三-更新版.pdf", ".pdf",
            "姓名：张三\nPython development requirement analysis frontend customer communication Zhejiang University.",
        ),
    ]
    profiles = [ResumeSourceProfile(
        source_id=item.id,
        source_name=item.display_name,
        content_fingerprint=f"fingerprint-{index}",
        extracted_chars=len(item.extracted_text),
    ) for index, item in enumerate(sources)]

    resolution = resolver.resolve(sources, profiles)

    assert len(resolution.representatives) == 2
    assert len(resolution.groups) == 1
    assert resolution.groups[0].status == "review_required"
    assert all(item.version_status == "review_required" for item in resolution.source_profiles)


def test_resume_version_resolver_never_merges_different_names_from_template_similarity(tmp_path):
    resolver = ResumeVersionResolver()
    shared = "五年Python开发经验，负责系统设计、接口开发、自动化测试、项目交付。"
    sources = [
        _attachment(tmp_path, "张三.pdf", ".pdf", f"姓名：张三\n{shared}"),
        _attachment(tmp_path, "李四.pdf", ".pdf", f"姓名：李四\n{shared}"),
    ]
    profiles = [ResumeSourceProfile(
        source_id=item.id,
        source_name=item.display_name,
        content_fingerprint=f"fingerprint-{index}",
        extracted_chars=len(item.extracted_text),
    ) for index, item in enumerate(sources)]

    resolution = resolver.resolve(sources, profiles)

    assert len(resolution.representatives) == 2
    assert resolution.groups == []


def test_recruitment_locks_rubric_maps_batches_rejects_unknown_and_fills_missing(tmp_path):
    gateway = BatchGateway(omit_source="resume-8.txt")
    service = RecruitmentScreeningModule(gateway, tmp_path, "fake-model")
    attachments = [
        _attachment(tmp_path, f"resume-{index}.txt", ".txt", f"候选人{index}，3年Python经验，项目编号{index}")
        for index in range(9)
    ]
    result = service.screen("岗位：批次测试岗位", attachments)

    purposes = [call[1]["purpose"] for call in gateway.calls]
    assert purposes == ["recruitment_rubric", "recruitment_screening", "recruitment_screening"]
    assert len(result.report.candidates) == 9
    assert all(item.source_name != "not-in-case.pdf" for item in result.report.candidates)
    missing = next(item for item in result.report.candidates if item.source_name == "resume-8.txt")
    assert missing.total_score == 0
    assert missing.evidence_coverage == 0
    assert result.report.rubric is not None and result.report.rubric.locked
    assert "batch_map" in result.graph_path
    assert any(item["event"] == "recruitment.reduce" and item["rejected_unknown_sources"] == 2 for item in result.trace)
    screening_systems = [call[0][0]["content"] for call in gateway.calls if call[1]["purpose"] == "recruitment_screening"]
    assert len(set(re.findall(r"已锁定Rubric：(.+)", value)[0] for value in screening_systems)) == 1


def test_recruitment_streams_bounded_tool_progress_without_resume_identity(tmp_path):
    gateway = BatchGateway()
    service = RecruitmentScreeningModule(gateway, tmp_path, "fake-model")
    attachments = [
        _attachment(
            tmp_path,
            f"private-candidate-{index}.txt",
            ".txt",
            f"Confidential Candidate {index}; five years of relevant delivery experience",
        )
        for index in range(9)
    ]
    progress_events: list[tuple[str, str, dict]] = []

    with tool_progress_scope(
        lambda event, summary, details: progress_events.append((event, summary, details))
    ):
        result = service.screen("Batch screening role", attachments)

    progress_names = [event for event, _summary, _details in progress_events]
    expected = [
        "tool.recruitment.parse.started",
        "tool.recruitment.parse.completed",
        "tool.recruitment.deduplicate.started",
        "tool.recruitment.deduplicate.completed",
        "tool.recruitment.version.started",
        "tool.recruitment.version.completed",
        "tool.recruitment.rubric.started",
        "tool.recruitment.rubric.completed",
        "tool.recruitment.batch.started",
        "tool.recruitment.batch_item.started",
        "tool.recruitment.batch_item.completed",
        "tool.recruitment.batch_item.started",
        "tool.recruitment.batch_item.completed",
        "tool.recruitment.batch.completed",
        "tool.recruitment.reduce.started",
        "tool.recruitment.reduce.completed",
        "tool.recruitment.fairness.started",
        "tool.recruitment.fairness.completed",
        "tool.recruitment.evidence.started",
        "tool.recruitment.evidence.completed",
        "tool.recruitment.harness.started",
        "tool.recruitment.harness.completed",
        "tool.recruitment.quality.completed",
        "tool.recruitment.report.started",
        "tool.recruitment.report.completed",
        "tool.recruitment.finalize.completed",
    ]
    assert progress_names == expected
    assert result.report.case is not None
    assert result.report.case.source_count == 9
    assert next(
        details for event, _summary, details in progress_events
        if event == "tool.recruitment.batch.started"
    )["batch_count"] == 2
    serialized_progress = json.dumps(progress_events, ensure_ascii=False)
    assert "private-candidate" not in serialized_progress
    assert "Confidential Candidate" not in serialized_progress


def test_recruitment_uses_stable_tie_breaking_and_marks_near_scores_for_review(tmp_path):
    service = RecruitmentScreeningModule(BatchGateway(), tmp_path, "fake-model")
    result = service.screen(
        "岗位：批次测试岗位",
        [
            _attachment(tmp_path, "resume-2.txt", ".txt", "候选人2，3年相关经验和项目成果"),
            _attachment(tmp_path, "resume-1.txt", ".txt", "候选人1，3年相关经验和项目成果"),
        ],
    )

    # The model follows attachment order, while the final order is deterministic
    # after equal score/evidence values.  A tie remains a human-review signal.
    assert [item.source_name for item in result.report.candidates] == ["resume-1.txt", "resume-2.txt"]
    assert result.report.harness is not None
    metrics = result.report.harness.metrics
    assert metrics["ranking_stability"] == "review_required"
    assert metrics["near_score_groups"] == 1
    assert metrics["score_recalculation_mismatches"] == 0
    assert "排名稳定性：存在1 组近分" in result.answer


def test_recruitment_fairness_guard_removes_protected_attributes(tmp_path):
    gateway = BatchGateway(include_protected=True)
    service = RecruitmentScreeningModule(gateway, tmp_path, "fake-model")
    result = service.screen(
        "岗位：批次测试岗位",
        [_attachment(tmp_path, "resume-0.txt", ".txt", "候选人0，3年相关经验和项目成果")],
    )

    candidate_text = json.dumps(result.report.candidates[0].model_dump(mode="json"), ensure_ascii=False)
    assert "女性" not in candidate_text
    assert "25岁" not in candidate_text
    assert "[已排除受保护属性]" in candidate_text
    assert any(item["event"] == "fairness.guard" and item["removed_mentions"] >= 1 for item in result.trace)


def test_recruitment_does_not_infer_location_availability_from_residence_or_school(tmp_path):
    service = RecruitmentScreeningModule(LocationProxyGateway(), tmp_path, "fake-model")
    result = service.screen(
        "岗位：宁波实习生",
        [_attachment(tmp_path, "resume-0.txt", ".txt", "候选人0，有相关实习经历，就读于宁波高校")],
    )

    candidate = result.report.candidates[0]
    location = next(item for item in candidate.criterion_scores if item.criterion_id == "location")
    assert location.score == 0
    assert location.confidence == "low"
    assert "不作为评分依据" in location.evidence
    assert candidate.total_score == 40.0


def test_recruitment_localizes_native_pdf_evidence_to_original_page(tmp_path):
    service = RecruitmentScreeningModule(LocalizationGateway(), tmp_path, "fake-model")
    result = service.screen(
        "Select a partnership intern",
        [_attachment(
            tmp_path,
            "resume.pdf",
            ".pdf",
            "[第 1 页]\nCandidate A\n\n[第 2 页]\nLed B2B partner development; "
            "closed three partnership projects; documented customer requirements.",
        )],
    )

    locators = [row.locator for row in result.report.candidates[0].criterion_scores]
    assert all(locator is not None and locator.status == "located" for locator in locators)
    assert all(locator.page_number == 2 for locator in locators if locator is not None)
    assert "evidence_localize" in result.graph_path
    assert result.report.harness is not None
    assert result.report.harness.metrics["evidence_localization_rate"] == 1.0
    assert result.metadata["candidate_summary"][0]["evidence_locations"][0]["locator"]["page_number"] == 2
    assert "原文定位：3 条已定位" in result.answer

    workbook = load_workbook(result.artifact_path, data_only=False)
    headers = [cell.value for cell in workbook["评分明细"][1]]
    assert headers[-4:] == ["原文位置", "定位原文", "定位状态", "复核原因"]
    assert "第 2 页" in workbook["评分明细"]["H2"].value
    workbook.close()


def test_recruitment_marks_low_confidence_visual_bbox_for_human_review(tmp_path):
    service = RecruitmentScreeningModule(LocalizationGateway(), tmp_path, "fake-model")
    result = service.screen(
        "Select a partnership intern",
        [_attachment(tmp_path, "resume.pdf", ".pdf", "[第 1 页]\nCandidate A")],
        visual_evidence=[{
            "source_name": "resume.pdf",
            "page": 3,
            "text": "Led B2B partner development and closed three partnership projects; documented customer requirements.",
            "bbox": [120, 240, 880, 420],
            "model_confidence": 0.71,
            "needs_review": True,
            "review_reasons": ["low_model_confidence"],
        }],
    )

    locators = [row.locator for row in result.report.candidates[0].criterion_scores]
    assert all(locator is not None and locator.anchor_kind == "bbox" for locator in locators)
    assert all(locator.status == "review_required" for locator in locators if locator is not None)
    assert all(locator.page_number == 3 for locator in locators if locator is not None)
    assert all(locator.bbox == [120, 240, 880, 420] for locator in locators if locator is not None)
    assert result.report.harness is not None
    assert result.report.harness.metrics["review_evidence_rows"] == 3
