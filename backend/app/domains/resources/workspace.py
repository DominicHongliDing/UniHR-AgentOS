from __future__ import annotations

import csv
import hashlib
import io
import sqlite3
import re
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterator, Literal

from docx import Document
from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, range_boundaries
from pypdf import PdfReader
from pydantic import BaseModel, Field

from app.domains.conversations.repository import ConversationRepository
from app.domains.identity import Principal


class ResourceItem(BaseModel):
    id: str
    resource_type: Literal["source", "artifact"]
    message_id: str
    display_name: str
    mime_type: str
    byte_size: int
    created_at: str
    status: str
    download_url: str
    parent_resource_ids: list[str] = Field(default_factory=list)
    lineage_key: str = ""
    version_number: int = 1
    is_current: bool = True
    archived_at: str = ""
    preview_url: str = ""
    metadata: dict[str, Any] = Field(default_factory=dict)


class ResourceWorkspaceManifest(BaseModel):
    conversation_id: str
    source_count: int
    artifact_count: int
    total_bytes: int
    active_count: int = 0
    archived_count: int = 0
    items: list[ResourceItem]


class ResourceLocator(BaseModel):
    page_start: int | None = Field(default=None, ge=1)
    page_end: int | None = Field(default=None, ge=1)
    section: str = Field(default="", max_length=200)
    sheet: str = Field(default="", max_length=200)
    cell_range: str = Field(default="", max_length=100)
    columns: list[str] = Field(default_factory=list, max_length=50)
    row_start: int | None = Field(default=None, ge=1)
    row_end: int | None = Field(default=None, ge=1)
    line_start: int | None = Field(default=None, ge=1)
    line_end: int | None = Field(default=None, ge=1)
    query: str = Field(default="", max_length=500)
    max_chars: int = Field(default=12_000, ge=500, le=30_000)


class ResourceExcerpt(BaseModel):
    resource_id: str
    resource_type: Literal["source", "artifact"]
    display_name: str
    locator: ResourceLocator
    text: str
    characters: int
    truncated: bool
    content_hash: str
    metadata: dict[str, Any] = Field(default_factory=dict)


class ResourceWorkspace:
    """Expose conversation files and their lineage through one stable workspace interface."""

    def __init__(self, repository: ConversationRepository) -> None:
        self.repository = repository
        self._initialize()

    @contextmanager
    def _connect(self) -> Iterator[sqlite3.Connection]:
        """Own and close every SQLite handle; transaction context alone does not close it."""
        db = sqlite3.connect(self.repository.database_path)
        try:
            yield db
            db.commit()
        finally:
            db.close()

    def _initialize(self) -> None:
        with self._connect() as db:
            db.execute(
                """CREATE TABLE IF NOT EXISTS resource_workspace_states (
                    resource_id TEXT NOT NULL,
                    tenant_id TEXT NOT NULL,
                    owner_user_id TEXT NOT NULL,
                    archived_at TEXT NOT NULL DEFAULT '',
                    updated_at TEXT NOT NULL,
                    PRIMARY KEY(resource_id, tenant_id, owner_user_id)
                )"""
            )

    def build(
        self,
        conversation_id: str,
        principal: Principal,
        *,
        include_archived: bool = False,
    ) -> ResourceWorkspaceManifest:
        attachments = self.repository.list_attachments(conversation_id, principal)
        artifacts = self.repository.list_artifacts(conversation_id, principal)
        archived = self._archived_states(principal)
        source_ids_by_name: dict[str, list[str]] = {}
        items: list[ResourceItem] = []
        for attachment in attachments:
            source_ids_by_name.setdefault(attachment.display_name, []).append(attachment.id)
            items.append(ResourceItem(
                id=attachment.id,
                resource_type="source",
                message_id=attachment.message_id,
                display_name=attachment.display_name,
                mime_type=attachment.mime_type,
                byte_size=attachment.byte_size,
                created_at=attachment.created_at,
                status="ready",
                download_url=f"/api/v1/resources/attachments/{attachment.id}/download",
                archived_at=archived.get(attachment.id, ""),
                metadata={
                    "extension": attachment.extension,
                    "summary": attachment.summary,
                    "locators": _locator_hints(attachment.extension, attachment.summary),
                },
            ))
        for artifact in artifacts:
            source_names = list(artifact.metadata.get("source_files") or [])
            parents = [
                source_id
                for name in source_names
                for source_id in source_ids_by_name.get(str(name), [])
            ]
            harness = dict(artifact.metadata.get("artifact_harness") or {})
            details = dict(harness.get("details") or {})
            preview_files = list(details.get("preview_files") or artifact.metadata.get("preview_files") or [])
            items.append(ResourceItem(
                id=artifact.id,
                resource_type="artifact",
                message_id=artifact.message_id,
                display_name=artifact.display_name,
                mime_type=artifact.mime_type,
                byte_size=artifact.byte_size,
                created_at=artifact.created_at,
                status=str(harness.get("status") or "not_checked"),
                download_url=f"/api/v1/artifacts/{artifact.id}/download",
                parent_resource_ids=list(dict.fromkeys(parents)),
                archived_at=archived.get(artifact.id, ""),
                preview_url=(
                    f"/api/v1/artifacts/{artifact.id}/previews/0" if preview_files else ""
                ),
                metadata={
                    "kind": artifact.kind,
                    "harness": harness,
                    "validation": artifact.metadata.get("validation") or {},
                    "created_from_scratch": bool(artifact.metadata.get("created_from_scratch")),
                    "changed_cell_count": int(artifact.metadata.get("changed_cell_count") or 0),
                    "operation_count": len(artifact.metadata.get("operation_results") or []),
                },
            ))
        items = _assign_versions(items)
        archived_count = sum(bool(item.archived_at) for item in items)
        active_count = len(items) - archived_count
        if not include_archived:
            items = [item for item in items if not item.archived_at]
        items.sort(key=lambda item: item.created_at, reverse=True)
        return ResourceWorkspaceManifest(
            conversation_id=conversation_id,
            source_count=len(attachments),
            artifact_count=len(artifacts),
            total_bytes=sum(item.byte_size for item in items),
            active_count=active_count,
            archived_count=archived_count,
            items=items,
        )

    def set_archived(
        self,
        resource_id: str,
        archived: bool,
        principal: Principal,
    ) -> ResourceItem:
        try:
            source = self.repository.get_attachment(resource_id, principal)
            conversation_id = source.conversation_id
        except KeyError:
            artifact = self.repository.get_artifact(resource_id, principal)
            conversation_id = artifact.conversation_id
        archived_at = datetime.now(timezone.utc).isoformat() if archived else ""
        with self._connect() as db:
            db.execute(
                """INSERT INTO resource_workspace_states(
                    resource_id,tenant_id,owner_user_id,archived_at,updated_at
                ) VALUES(?,?,?,?,?)
                ON CONFLICT(resource_id,tenant_id,owner_user_id) DO UPDATE SET
                    archived_at=excluded.archived_at,updated_at=excluded.updated_at""",
                (resource_id, principal.tenant_id, principal.user_id, archived_at,
                 datetime.now(timezone.utc).isoformat()),
            )
        manifest = self.build(conversation_id, principal, include_archived=True)
        return next(item for item in manifest.items if item.id == resource_id)

    def _archived_states(self, principal: Principal) -> dict[str, str]:
        with self._connect() as db:
            rows = db.execute(
                """SELECT resource_id,archived_at FROM resource_workspace_states
                WHERE tenant_id=? AND owner_user_id=?""",
                (principal.tenant_id, principal.user_id),
            ).fetchall()
        return {str(row[0]): str(row[1]) for row in rows}

    def query(
        self,
        resource_id: str,
        locator: ResourceLocator,
        principal: Principal,
    ) -> ResourceExcerpt:
        """Read a bounded, addressable slice from one authorized workspace resource."""
        try:
            item = self.repository.get_attachment(resource_id, principal)
            resource_type: Literal["source", "artifact"] = "source"
            path = Path(item.saved_path)
            display_name = item.display_name
            extension = item.extension.casefold()
        except KeyError:
            artifact = self.repository.get_artifact(resource_id, principal)
            resource_type = "artifact"
            path = Path(artifact.path)
            display_name = artifact.display_name
            extension = path.suffix.casefold()
        if not path.is_file():
            raise KeyError(resource_id)

        text, metadata = _read_slice(path, extension, locator)
        full_length = len(text)
        bounded = text[:locator.max_chars]
        return ResourceExcerpt(
            resource_id=resource_id,
            resource_type=resource_type,
            display_name=display_name,
            locator=locator,
            text=bounded,
            characters=len(bounded),
            truncated=full_length > len(bounded),
            content_hash=hashlib.sha256(bounded.encode("utf-8")).hexdigest(),
            metadata=metadata,
        )


def _assign_versions(items: list[ResourceItem]) -> list[ResourceItem]:
    grouped: dict[str, list[ResourceItem]] = {}
    for item in items:
        key = _lineage_key(item)
        grouped.setdefault(key, []).append(item)
    result: list[ResourceItem] = []
    for key, family in grouped.items():
        ordered = sorted(family, key=lambda item: (item.created_at, item.id))
        current_id = ordered[-1].id
        for version, item in enumerate(ordered, start=1):
            result.append(item.model_copy(update={
                "lineage_key": key,
                "version_number": version,
                "is_current": item.id == current_id,
            }))
    return result


def _lineage_key(item: ResourceItem) -> str:
    stem = Path(item.display_name).stem.casefold()
    # Remove common copy/version suffixes while preserving the business title.
    stem = re.sub(r"(?:\s*[-_]?v\d+(?:\.\d+)*)$", "", stem, flags=re.IGNORECASE)
    stem = re.sub(r"(?:\s*\(\d+\)|\s*（\d+）)$", "", stem)
    normalized = re.sub(r"[^\w\u4e00-\u9fff]+", "", stem)
    return f"{item.resource_type}:{normalized or item.id}"


def _locator_hints(extension: str, summary: dict[str, Any]) -> dict[str, Any]:
    if extension == ".pdf":
        return {"type": "page", "pages": int(summary.get("pages") or 0)}
    if extension == ".docx":
        return {"type": "section", "paragraphs": int(summary.get("paragraphs") or 0)}
    if extension == ".xlsx":
        return {"type": "sheet_range", "sheets": summary.get("sheets") or []}
    return {"type": "line"}


def _read_slice(path: Path, extension: str, locator: ResourceLocator) -> tuple[str, dict[str, Any]]:
    if extension == ".pdf":
        return _read_pdf(path, locator)
    if extension == ".docx":
        return _read_docx(path, locator)
    if extension == ".xlsx":
        return _read_xlsx(path, locator)
    if extension in {".csv", ".tsv"}:
        return _read_delimited(path, locator, "\t" if extension == ".tsv" else ",")
    if extension in {".txt", ".md"}:
        return _read_text(path, locator)
    return "", {"unsupported_text_extraction": True, "extension": extension}


def _read_pdf(path: Path, locator: ResourceLocator) -> tuple[str, dict[str, Any]]:
    reader = PdfReader(path)
    total = len(reader.pages)
    start = min(locator.page_start or 1, max(total, 1))
    end = min(locator.page_end or start, total)
    if end < start:
        raise ValueError("page_end_before_page_start")
    blocks = []
    for number in range(start, end + 1):
        text = (reader.pages[number - 1].extract_text() or "").strip()
        blocks.append(f"[第 {number} 页]\n{text}")
    return "\n\n".join(blocks), {"page_start": start, "page_end": end, "total_pages": total}


def _read_docx(path: Path, locator: ResourceLocator) -> tuple[str, dict[str, Any]]:
    document = Document(path)
    paragraphs = [(index, paragraph) for index, paragraph in enumerate(document.paragraphs, start=1) if paragraph.text.strip()]
    selected = paragraphs
    matched_section = ""
    if locator.section.strip():
        needle = locator.section.strip().casefold()
        start_at = next((i for i, (_, p) in enumerate(paragraphs) if needle in p.text.casefold()), None)
        if start_at is not None:
            matched_section = paragraphs[start_at][1].text.strip()
            start_level = _heading_level(paragraphs[start_at][1].style.name)
            end_at = len(paragraphs)
            for offset in range(start_at + 1, len(paragraphs)):
                level = _heading_level(paragraphs[offset][1].style.name)
                if level and start_level and level <= start_level:
                    end_at = offset
                    break
            selected = paragraphs[start_at:end_at]
        else:
            selected = []
    elif locator.query.strip():
        terms = _terms(locator.query)
        hits = [i for i, (_, p) in enumerate(paragraphs) if any(term in p.text.casefold() for term in terms)]
        indexes = sorted({near for hit in hits for near in range(max(0, hit - 1), min(len(paragraphs), hit + 2))})
        selected = [paragraphs[i] for i in indexes]
    text = "\n".join(f"[段落 {index}] {paragraph.text.strip()}" for index, paragraph in selected)
    return text, {"matched_section": matched_section, "paragraph_count": len(selected)}


def _heading_level(style_name: str) -> int:
    match = re.search(r"(?:Heading|标题)\s*(\d+)", style_name or "", re.IGNORECASE)
    return int(match.group(1)) if match else 0


def _read_xlsx(path: Path, locator: ResourceLocator) -> tuple[str, dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        worksheet = workbook[locator.sheet] if locator.sheet else workbook.worksheets[0]
        if locator.cell_range:
            min_col, min_row, max_col, max_row = range_boundaries(locator.cell_range)
        else:
            min_row = locator.row_start or 1
            max_row = locator.row_end or min(worksheet.max_row or min_row, min_row + 99)
            if locator.columns:
                indexes = [_column_index(value) for value in locator.columns]
                min_col, max_col = min(indexes), max(indexes)
            else:
                min_col, max_col = 1, min(worksheet.max_column or 1, 40)
        rows = []
        for row in worksheet.iter_rows(
            min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col, values_only=True,
        ):
            rows.append(" | ".join("" if value is None else str(value) for value in row))
        return "\n".join(rows), {
            "sheet": worksheet.title,
            "range": f"{worksheet.cell(min_row, min_col).coordinate}:{worksheet.cell(max_row, max_col).coordinate}",
            "rows": max_row - min_row + 1,
            "columns": max_col - min_col + 1,
        }
    finally:
        workbook.close()


def _column_index(value: str) -> int:
    clean = str(value).strip()
    if clean.isdigit():
        return max(1, int(clean))
    return column_index_from_string(clean.upper())


def _read_delimited(path: Path, locator: ResourceLocator, delimiter: str) -> tuple[str, dict[str, Any]]:
    text = _decode_file(path)
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    start = locator.row_start or 1
    end = min(locator.row_end or start + 99, len(rows))
    selected = rows[start - 1:end]
    if locator.columns and rows:
        header = rows[0]
        indexes = [header.index(name) for name in locator.columns if name in header]
        if indexes:
            selected = [[row[index] if index < len(row) else "" for index in indexes] for row in selected]
    return "\n".join(" | ".join(row) for row in selected), {"row_start": start, "row_end": end}


def _read_text(path: Path, locator: ResourceLocator) -> tuple[str, dict[str, Any]]:
    lines = _decode_file(path).splitlines()
    if locator.query.strip():
        terms = _terms(locator.query)
        hits = [i for i, line in enumerate(lines) if any(term in line.casefold() for term in terms)]
        indexes = sorted({near for hit in hits for near in range(max(0, hit - 2), min(len(lines), hit + 3))})
        selected = [(index + 1, lines[index]) for index in indexes]
    else:
        start = locator.line_start or 1
        end = min(locator.line_end or start + 199, len(lines))
        selected = list(enumerate(lines[start - 1:end], start=start))
    return "\n".join(f"[{number}] {line}" for number, line in selected), {"line_count": len(selected)}


def _terms(query: str) -> list[str]:
    return [term.casefold() for term in re.findall(r"[\w\u4e00-\u9fff]{2,}", query)][:20]


def _decode_file(path: Path) -> str:
    content = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")
