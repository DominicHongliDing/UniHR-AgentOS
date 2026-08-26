from __future__ import annotations

import hashlib
import json
import math
import re
import sqlite3
from difflib import SequenceMatcher
from collections import Counter
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Callable, Sequence

from docx import Document
from pypdf import PdfReader

from app.models.embeddings import EmbeddingGateway

from .catalog import SourceCatalog
from .ingestion import IngestionQualityReport, KnowledgeDocumentExtractor
from .contracts import (
    IngestionDisposition,
    KnowledgeAccessContext,
    PublicationStatus,
    SourceVersionRecord,
)


POLICY_METADATA_VERSION = "policy-metadata-v3"


@dataclass(frozen=True)
class ChunkDraft:
    text: str
    page_number: int | None = None
    section: str = ""
    anchor_kind: str = "section"
    anchor_start: int | None = None
    anchor_end: int | None = None
    bbox: tuple[int, int, int, int] | None = None
    model_confidence: float | None = None
    needs_review: bool = False
    evidence_kind: str = "text"


@dataclass(frozen=True)
class SearchHit:
    chunk_id: str
    source_id: str
    title: str
    collection: str
    text: str
    page_number: int | None
    section: str
    score: float
    vector_score: float = 0.0
    lexical_score: float = 0.0
    retrieval_method: str = "vector"
    source_version: str | None = None
    publication_date: str | None = None
    effective_date: str | None = None
    expiry_date: str | None = None
    policy_status: str = "unknown"
    authority_level: str = "unknown"
    version_id: str | None = None
    publication_id: str | None = None
    anchor_kind: str = "section"
    anchor_start: int | None = None
    anchor_end: int | None = None
    rerank_score: float = 0.0
    rerank_model: str = ""

    @property
    def citation(self) -> str:
        if self.page_number:
            location = f"第 {self.page_number} 页"
        elif self.anchor_kind == "paragraph" and self.anchor_start:
            end = self.anchor_end or self.anchor_start
            location = f"第 {self.anchor_start}–{end} 段"
        elif self.anchor_kind == "line" and self.anchor_start:
            end = self.anchor_end or self.anchor_start
            location = f"第 {self.anchor_start}–{end} 行"
        else:
            location = self.section or "文档正文"
        if self.section and self.section not in location:
            location = f"{location} · {self.section}"
        return f"《{self.title}》{location}"


class KnowledgeVectorIndex:
    """Small local vector store for approved policy candidates.

    Only records pre-classified as ``vector`` are read and embedded. Candidate,
    compensation, and other case-only material never crosses this boundary.
    """

    def __init__(
        self,
        database_path: str | Path,
        embedder: EmbeddingGateway,
        default_tenant_id: str = "*",
        document_extractor: KnowledgeDocumentExtractor | None = None,
    ) -> None:
        self.database_path = Path(database_path)
        self.embedder = embedder
        self.default_tenant_id = default_tenant_id
        self.document_extractor = document_extractor
        self._ensure_schema()

    def index(
        self,
        root: str | Path,
        *,
        tenant_id: str | None = None,
        actor_id: str = "system",
        auto_publish: bool = True,
        progress_callback: Callable[[dict[str, object]], None] | None = None,
        only_relative_paths: Sequence[str] = (),
    ) -> dict[str, object]:
        root_path = Path(root)
        catalog_candidates = [
            (path, record)
            for path, record in SourceCatalog().records_with_paths(root_path)
            if record.disposition == IngestionDisposition.VECTOR
        ]
        if only_relative_paths:
            allowed_paths = {value.replace("\\", "/") for value in only_relative_paths}
            catalog_candidates = [
                (path, record) for path, record in catalog_candidates
                if path.relative_to(root_path).as_posix() in allowed_paths
            ]
        supported_extensions = {".pdf", ".docx", ".md", ".txt"}
        if self.document_extractor is not None:
            supported_extensions.update({".doc", ".png", ".jpg", ".jpeg", ".webp"})
        candidates = [
            (path, record) for path, record in catalog_candidates
            if path.suffix.casefold() in supported_extensions
        ]
        indexed_sources = 0
        indexed_chunks = 0
        skipped_sources = 0
        failed_sources = len(catalog_candidates) - len(candidates)
        failed_by_reason: dict[str, int] = {}
        processed_files = 0
        total_files = len(catalog_candidates)
        if progress_callback:
            progress_callback({"stage": "catalogued", "total": total_files, "processed": 0, "current_file": ""})
        for path, record in catalog_candidates:
            if path.suffix.casefold() in supported_extensions:
                continue
            failed_by_reason["unsupported_legacy_format"] = (
                failed_by_reason.get("unsupported_legacy_format", 0) + 1
            )
            self._record_failure(
                relative_path=path.relative_to(root_path).as_posix(),
                source_id=record.source_id,
                collection=record.collection,
                reason="unsupported_legacy_format",
            )
            processed_files += 1
            if progress_callback:
                progress_callback({"stage": "unsupported", "total": total_files,
                                   "processed": processed_files,
                                   "current_file": path.relative_to(root_path).as_posix()})
        for path, record in candidates:
            relative_path = path.relative_to(root_path).as_posix()
            if progress_callback:
                progress_callback({"stage": "extracting", "total": total_files,
                                   "processed": processed_files, "current_file": relative_path})
            checksum = _sha256(path)
            if self._is_current(relative_path, checksum):
                skipped_sources += 1
                self._clear_failure(relative_path)
                processed_files += 1
                if progress_callback:
                    progress_callback({"stage": "skipped", "total": total_files,
                                       "processed": processed_files, "current_file": relative_path})
                continue
            # Failed sources are quarantined after the first attempt. They must not
            # make every normal reindex replay expensive OCR/conversion work.
            if self._has_failure(relative_path) and not only_relative_paths:
                skipped_sources += 1
                processed_files += 1
                if progress_callback:
                    progress_callback({"stage": "quarantined", "total": total_files,
                                       "processed": processed_files, "current_file": relative_path})
                continue
            try:
                quality = IngestionQualityReport(status="ready", score=100, extractor="native")
                if self.document_extractor is not None:
                    extraction = self.document_extractor.extract(path)
                    quality = extraction.quality
                    chunks = [ChunkDraft(
                        text=item.text, page_number=item.page_number, section=item.section,
                        anchor_kind=item.anchor_kind, anchor_start=item.anchor_start,
                        anchor_end=item.anchor_end, bbox=item.bbox,
                        model_confidence=item.model_confidence, needs_review=item.needs_review,
                        evidence_kind=item.evidence_kind,
                    ) for item in extraction.chunks]
                else:
                    chunks = _extract_chunks(path)
                if not chunks:
                    failed_sources += 1
                    failed_by_reason["empty_or_scanned"] = failed_by_reason.get("empty_or_scanned", 0) + 1
                    self._record_failure(
                        relative_path=relative_path,
                        source_id=record.source_id,
                        collection=record.collection,
                        reason="empty_or_scanned",
                    )
                    processed_files += 1
                    if progress_callback:
                        progress_callback({"stage": "failed", "total": total_files,
                                           "processed": processed_files, "current_file": relative_path})
                    continue
                vectors = self.embedder.embed([chunk.text for chunk in chunks])
                policy_metadata = _infer_policy_metadata(path, relative_path, chunks)
                stable_source_id = "src_" + hashlib.sha256(relative_path.encode("utf-8")).hexdigest()[:16]
                version_id = self._stage_source(
                    source_id=stable_source_id,
                    relative_path=relative_path,
                    title=path.name,
                    collection=record.collection,
                    checksum=checksum,
                    chunks=chunks,
                    vectors=vectors,
                    policy_metadata=policy_metadata,
                    quality=quality,
                )
                if auto_publish:
                    self.publish_version(
                        version_id,
                        tenant_id=tenant_id or self.default_tenant_id,
                        actor_id=actor_id,
                        allowed_roles=("admin", "user"),
                    )
                indexed_sources += 1
                indexed_chunks += len(chunks)
                self._clear_failure(relative_path)
                processed_files += 1
                if progress_callback:
                    progress_callback({"stage": "staged", "total": total_files,
                                       "processed": processed_files, "current_file": relative_path})
            except Exception as exc:
                failed_sources += 1
                reason = type(exc).__name__
                failed_by_reason[reason] = failed_by_reason.get(reason, 0) + 1
                self._record_failure(
                    relative_path=relative_path,
                    source_id=record.source_id,
                    collection=record.collection,
                    reason=reason,
                )
                processed_files += 1
                if progress_callback:
                    progress_callback({"stage": "failed", "total": total_files,
                                       "processed": processed_files, "current_file": relative_path})
        stats = self.stats()
        return {
            **stats,
            "indexed_now": indexed_sources,
            "chunks_now": indexed_chunks,
            "skipped_now": skipped_sources,
            "failed_now": failed_sources,
            "failed_by_reason": failed_by_reason,
            "candidate_sources": len(catalog_candidates),
            "embedding_model": self.embedder.model_name,
        }

    def stats(self) -> dict[str, object]:
        with self._connect() as connection:
            source_count = connection.execute("SELECT COUNT(*) FROM knowledge_sources").fetchone()[0]
            chunk_count = connection.execute("SELECT COUNT(*) FROM knowledge_chunks").fetchone()[0]
            failure_rows = connection.execute(
                "SELECT reason, COUNT(*) FROM knowledge_ingestion_failures GROUP BY reason"
            ).fetchall()
            version_count = connection.execute(
                "SELECT COUNT(*) FROM knowledge_source_versions"
            ).fetchone()[0]
            pending_count = connection.execute(
                """SELECT COUNT(*) FROM knowledge_source_versions v
                WHERE NOT EXISTS(
                    SELECT 1 FROM knowledge_publications p
                    WHERE p.version_id=v.version_id AND p.status='published'
                )"""
            ).fetchone()[0]
        return {
            "indexed_sources": int(source_count),
            "indexed_chunks": int(chunk_count),
            "failed_sources": sum(int(row[1]) for row in failure_rows),
            "failure_reasons": {str(row[0]): int(row[1]) for row in failure_rows},
            "source_versions": int(version_count),
            "pending_versions": int(pending_count),
            "embedding_model": self.embedder.model_name,
        }

    def search(
        self,
        query: str,
        limit: int = 6,
        *,
        access: KnowledgeAccessContext | None = None,
        include_withdrawn: bool = False,
    ) -> list[SearchHit]:
        if not query.strip():
            return []
        query_vectors = self.embedder.embed([query])
        if not query_vectors:
            return []
        query_vector = _normalize(query_vectors[0])
        resolved_access = access or KnowledgeAccessContext(
            tenant_id=self.default_tenant_id, user_id="system", role="admin",
        )
        allowed_statuses = ("published", "withdrawn") if include_withdrawn else ("published",)
        status_placeholders = ",".join("?" for _ in allowed_statuses)
        with self._connect() as connection:
            rows = connection.execute(
                f"""
                SELECT DISTINCT c.chunk_id, c.source_id, s.title, s.collection, c.content, c.page_number,
                       c.section, c.embedding, s.source_version, s.publication_date,
                       s.effective_date, s.expiry_date, s.policy_status, s.authority_level,
                       s.version_id, p.publication_id, c.anchor_kind, c.anchor_start, c.anchor_end
                FROM knowledge_version_chunks c
                JOIN knowledge_source_versions s ON s.version_id = c.version_id
                JOIN knowledge_publications p ON p.version_id=s.version_id
                WHERE c.embedding_model = ? AND p.status IN ({status_placeholders})
                  AND p.tenant_id IN (?, '*')
                  AND EXISTS(
                    SELECT 1 FROM knowledge_publication_acl a
                    WHERE a.publication_id=p.publication_id AND a.permission='retrieve'
                      AND (
                        (a.subject_type='organization' AND a.subject_id IN (?, '*'))
                        OR (a.subject_type='role' AND a.subject_id=?)
                        OR (a.subject_type='user' AND a.subject_id=?)
                      )
                  )
                """,
                (
                    self.embedder.model_name,
                    *allowed_statuses,
                    resolved_access.tenant_id,
                    resolved_access.tenant_id,
                    resolved_access.role,
                    resolved_access.user_id,
                ),
            ).fetchall()
        query_terms = _lexical_terms(query)
        document_terms = [
            Counter(_lexical_terms(f"{row[2]} {row[6] or ''} {row[4]}"))
            for row in rows
        ]
        document_frequency = Counter(term for terms in document_terms for term in terms)
        average_length = (
            sum(sum(terms.values()) for terms in document_terms) / len(document_terms)
            if document_terms else 1.0
        )
        raw: list[tuple[sqlite3.Row, float, float]] = []
        for row, terms in zip(rows, document_terms, strict=True):
            vector = json.loads(row[7])
            vector_score = sum(
                left * right for left, right in zip(query_vector, vector, strict=False)
            )
            lexical_score = _bm25_score(
                query_terms,
                terms,
                document_frequency,
                document_count=len(rows),
                average_length=average_length,
            )
            raw.append((row, float(vector_score), lexical_score))

        vector_order = sorted(
            range(len(raw)), key=lambda index: raw[index][1], reverse=True,
        )
        lexical_order = sorted(
            (index for index in range(len(raw)) if raw[index][2] > 0),
            key=lambda index: raw[index][2], reverse=True,
        )
        fused: dict[int, float] = {}
        for rank, index in enumerate(vector_order, start=1):
            fused[index] = fused.get(index, 0.0) + 1.0 / (60 + rank)
        for rank, index in enumerate(lexical_order, start=1):
            fused[index] = fused.get(index, 0.0) + 1.25 / (60 + rank)
        max_rrf = 2.25 / 61
        scored: list[SearchHit] = []
        for index, fused_score in fused.items():
            row, vector_score, lexical_score = raw[index]
            scored.append(SearchHit(
                chunk_id=row[0], source_id=row[1], title=row[2], collection=row[3],
                text=row[4], page_number=row[5], section=row[6] or "",
                score=min(1.0, fused_score / max_rrf),
                vector_score=vector_score,
                lexical_score=lexical_score,
                retrieval_method="hybrid" if lexical_score > 0 else "vector",
                source_version=row[8], publication_date=row[9], effective_date=row[10],
                expiry_date=row[11], policy_status=row[12] or "unknown",
                authority_level=row[13] or "unknown",
                version_id=row[14], publication_id=row[15], anchor_kind=row[16] or "section",
                anchor_start=row[17], anchor_end=row[18],
            ))
        return sorted(scored, key=lambda candidate: candidate.score, reverse=True)[
            :max(1, min(limit, 100))
        ]

    def _is_current(self, relative_path: str, checksum: str) -> bool:
        with self._connect() as connection:
            row = connection.execute(
                """SELECT checksum, embedding_model, metadata_version
                FROM knowledge_source_versions
                WHERE relative_path=? ORDER BY created_at DESC LIMIT 1""",
                (relative_path,),
            ).fetchone()
        return bool(
            row
            and row[0] == checksum
            and row[1] == self.embedder.model_name
            and row[2] == POLICY_METADATA_VERSION
        )

    def _has_failure(self, relative_path: str) -> bool:
        with self._connect() as connection:
            row = connection.execute(
                "SELECT 1 FROM knowledge_ingestion_failures WHERE relative_path=?", (relative_path,),
            ).fetchone()
        return row is not None

    def list_failures(self) -> list[dict[str, object]]:
        with self._connect() as connection:
            rows = connection.execute(
                """SELECT relative_path,source_id,collection,reason,last_attempted_at,
                retry_count,last_error FROM knowledge_ingestion_failures
                ORDER BY last_attempted_at DESC"""
            ).fetchall()
        return [{
            "relative_path": row[0], "source_id": row[1], "collection": row[2],
            "reason": row[3], "last_attempted_at": row[4],
            "retry_count": int(row[5] or 0), "last_error": row[6] or "",
        } for row in rows]

    def prepare_failure_retry(self, source_id: str) -> str:
        with self._connect() as connection:
            row = connection.execute(
                "SELECT relative_path FROM knowledge_ingestion_failures WHERE source_id=?", (source_id,),
            ).fetchone()
            if row is None:
                raise KeyError("knowledge_ingestion_failure_not_found")
            relative_path = str(row[0])
            connection.execute(
                "UPDATE knowledge_ingestion_failures SET retry_count=retry_count+1,last_attempted_at=CURRENT_TIMESTAMP WHERE source_id=?",
                (source_id,),
            )
        return relative_path

    def _stage_source(
        self,
        *,
        source_id: str,
        relative_path: str,
        title: str,
        collection: str,
        checksum: str,
        chunks: Sequence[ChunkDraft],
        vectors: Sequence[Sequence[float]],
        policy_metadata: dict[str, str | None],
        quality: IngestionQualityReport,
    ) -> str:
        version_digest = hashlib.sha256(
            f"{source_id}:{checksum}:{self.embedder.model_name}:{POLICY_METADATA_VERSION}".encode("utf-8")
        ).hexdigest()[:24]
        version_id = f"ver_{version_digest}"
        with self._connect() as connection:
            connection.execute(
                """
                INSERT OR IGNORE INTO knowledge_source_versions
                    (version_id, source_id, relative_path, title, collection, checksum, embedding_model,
                     source_version, publication_date, effective_date, expiry_date,
                     policy_status, authority_level, metadata_version, ingestion_status,
                     quality_score, quality_review_status, quality_report_json)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    version_id, source_id, relative_path, title, collection, checksum, self.embedder.model_name,
                    policy_metadata["source_version"], policy_metadata["publication_date"],
                    policy_metadata["effective_date"], policy_metadata["expiry_date"],
                    policy_metadata["policy_status"], policy_metadata["authority_level"],
                    POLICY_METADATA_VERSION,
                    quality.status, quality.score,
                    "pending" if quality.status == "review_required" else "not_required",
                    json.dumps(quality.as_dict(), ensure_ascii=False),
                ),
            )
            connection.execute(
                "DELETE FROM knowledge_version_chunks WHERE version_id=?", (version_id,),
            )
            for index, (chunk, vector) in enumerate(zip(chunks, vectors, strict=True)):
                chunk_id = hashlib.sha256(f"{version_id}:{index}:{chunk.text}".encode("utf-8")).hexdigest()[:24]
                connection.execute(
                    """
                    INSERT INTO knowledge_version_chunks
                        (chunk_id, version_id, source_id, chunk_index, page_number, section,
                         anchor_kind, anchor_start, anchor_end, content, embedding, embedding_model,
                         bbox_json, model_confidence, needs_review, evidence_kind)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        f"chk_{chunk_id}", version_id, source_id, index, chunk.page_number, chunk.section,
                        chunk.anchor_kind, chunk.anchor_start, chunk.anchor_end,
                        chunk.text, json.dumps(_normalize(vector), separators=(",", ":")),
                        self.embedder.model_name,
                        json.dumps(chunk.bbox) if chunk.bbox else None,
                        chunk.model_confidence, int(chunk.needs_review), chunk.evidence_kind,
                    ),
                )
        return version_id

    def publish_version(
        self,
        version_id: str,
        *,
        tenant_id: str,
        actor_id: str,
        allowed_roles: Sequence[str] = ("admin", "user"),
        allowed_user_ids: Sequence[str] = (),
    ) -> SourceVersionRecord:
        with self._connect() as connection:
            version = connection.execute(
                "SELECT * FROM knowledge_source_versions WHERE version_id=?", (version_id,),
            ).fetchone()
            if version is None:
                raise KeyError("source_version_not_found")
            columns = [row[1] for row in connection.execute(
                "PRAGMA table_info(knowledge_source_versions)"
            ).fetchall()]
            data = dict(zip(columns, version, strict=True))
            if data.get("ingestion_status") == "failed":
                raise ValueError("knowledge_ingestion_failed")
            if data.get("ingestion_status") == "review_required" and data.get("quality_review_status") != "approved":
                raise ValueError("knowledge_quality_review_required")
            publication_id = "pub_" + hashlib.sha256(
                f"{version_id}:{tenant_id}".encode("utf-8")
            ).hexdigest()[:24]
            connection.execute(
                """UPDATE knowledge_publications SET status='withdrawn', withdrawn_at=CURRENT_TIMESTAMP
                WHERE tenant_id=? AND status='published' AND version_id IN (
                    SELECT version_id FROM knowledge_source_versions WHERE source_id=?
                )""",
                (tenant_id, data["source_id"]),
            )
            connection.execute(
                """INSERT INTO knowledge_publications(
                    publication_id,version_id,tenant_id,status,published_by,published_at
                ) VALUES(?,?,?,'published',?,CURRENT_TIMESTAMP)
                ON CONFLICT(publication_id) DO UPDATE SET
                    status='published',published_by=excluded.published_by,
                    published_at=CURRENT_TIMESTAMP,withdrawn_at=NULL""",
                (publication_id, version_id, tenant_id, actor_id),
            )
            connection.execute("DELETE FROM knowledge_publication_acl WHERE publication_id=?", (publication_id,))
            for role in dict.fromkeys(role for role in allowed_roles if role):
                connection.execute(
                    "INSERT OR IGNORE INTO knowledge_publication_acl(publication_id,subject_type,subject_id,permission) VALUES(?,?,?,'retrieve')",
                    (publication_id, "role", role),
                )
            for user_id in dict.fromkeys(user_id for user_id in allowed_user_ids if user_id):
                connection.execute(
                    "INSERT OR IGNORE INTO knowledge_publication_acl(publication_id,subject_type,subject_id,permission) VALUES(?,?,?,'retrieve')",
                    (publication_id, "user", user_id),
                )
            connection.execute("DELETE FROM knowledge_chunks WHERE source_id=?", (data["source_id"],))
            connection.execute("DELETE FROM knowledge_sources WHERE source_id=?", (data["source_id"],))
            connection.execute(
                """INSERT INTO knowledge_sources(
                    source_id,relative_path,title,collection,checksum,embedding_model,
                    source_version,publication_date,effective_date,expiry_date,policy_status,
                    authority_level,metadata_version,version_id,publication_id,tenant_id
                ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    data["source_id"], data["relative_path"], data["title"], data["collection"],
                    data["checksum"], data["embedding_model"], data["source_version"],
                    data["publication_date"], data["effective_date"], data["expiry_date"],
                    data["policy_status"], data["authority_level"], data["metadata_version"],
                    version_id, publication_id, tenant_id,
                ),
            )
            rows = connection.execute(
                """SELECT chunk_id,source_id,chunk_index,page_number,section,anchor_kind,
                anchor_start,anchor_end,content,embedding,embedding_model
                FROM knowledge_version_chunks WHERE version_id=? ORDER BY chunk_index""",
                (version_id,),
            ).fetchall()
            connection.executemany(
                """INSERT INTO knowledge_chunks(
                    chunk_id,source_id,chunk_index,page_number,section,anchor_kind,
                    anchor_start,anchor_end,content,embedding,embedding_model
                ) VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
                rows,
            )
        return self.get_version(version_id, tenant_id=tenant_id)

    def withdraw_publication(self, publication_id: str, *, tenant_id: str) -> None:
        with self._connect() as connection:
            row = connection.execute(
                "SELECT version_id FROM knowledge_publications WHERE publication_id=? AND tenant_id=?",
                (publication_id, tenant_id),
            ).fetchone()
            if row is None:
                raise KeyError("publication_not_found")
            connection.execute(
                "UPDATE knowledge_publications SET status='withdrawn',withdrawn_at=CURRENT_TIMESTAMP WHERE publication_id=?",
                (publication_id,),
            )
            connection.execute("DELETE FROM knowledge_sources WHERE publication_id=?", (publication_id,))

    def get_version(self, version_id: str, *, tenant_id: str) -> SourceVersionRecord:
        records = self.list_versions(tenant_id=tenant_id, include_unpublished=True, version_id=version_id)
        if not records:
            raise KeyError("source_version_not_found")
        return records[0]

    def list_versions(
        self,
        *,
        tenant_id: str,
        include_unpublished: bool,
        version_id: str | None = None,
    ) -> list[SourceVersionRecord]:
        clauses = ["(p.tenant_id=? OR p.tenant_id IS NULL)"] if include_unpublished else ["p.tenant_id=? AND p.status='published'"]
        params: list[object] = [tenant_id]
        if version_id:
            clauses.append("v.version_id=?")
            params.append(version_id)
        with self._connect() as connection:
            rows = connection.execute(
                f"""SELECT v.version_id,v.source_id,v.title,v.collection,v.checksum,
                v.source_version,v.publication_date,v.effective_date,v.expiry_date,
                v.policy_status,v.authority_level,v.created_at,p.status,p.publication_id,p.published_at,
                v.ingestion_status,v.quality_score,v.quality_review_status,v.quality_report_json,
                (SELECT COUNT(*) FROM knowledge_version_chunks c WHERE c.version_id=v.version_id)
                FROM knowledge_source_versions v
                LEFT JOIN knowledge_publications p ON p.version_id=v.version_id AND p.tenant_id=?
                WHERE {' AND '.join(clauses)} ORDER BY v.created_at DESC""",
                [tenant_id, *params],
            ).fetchall()
            result: list[SourceVersionRecord] = []
            for row in rows:
                roles = []
                if row[13]:
                    roles = [item[0] for item in connection.execute(
                        "SELECT subject_id FROM knowledge_publication_acl WHERE publication_id=? AND subject_type='role' ORDER BY subject_id",
                        (row[13],),
                    ).fetchall()]
                result.append(SourceVersionRecord(
                    version_id=row[0], source_id=row[1], title=row[2], collection=row[3], checksum=row[4],
                    source_version=row[5], publication_date=row[6], effective_date=row[7], expiry_date=row[8],
                    policy_status=row[9], authority_level=row[10], created_at=row[11],
                    lifecycle_status=PublicationStatus(row[12] or "draft"), publication_id=row[13],
                    published_at=row[14], ingestion_status=row[15] or "ready",
                    quality_score=int(row[16] or 0), quality_review_status=row[17] or "not_required",
                    quality_report=json.loads(row[18] or "{}"), chunk_count=int(row[19]), allowed_roles=roles,
                ))
        return result

    def review_quality(self, version_id: str, *, actor_id: str, approved: bool, comment: str = "") -> SourceVersionRecord:
        with self._connect() as connection:
            row = connection.execute(
                "SELECT ingestion_status FROM knowledge_source_versions WHERE version_id=?", (version_id,),
            ).fetchone()
            if row is None:
                raise KeyError("source_version_not_found")
            if row[0] != "review_required":
                raise ValueError("knowledge_quality_review_not_required")
            connection.execute(
                """UPDATE knowledge_source_versions SET quality_review_status=?, quality_reviewed_by=?,
                quality_reviewed_at=CURRENT_TIMESTAMP, quality_review_comment=? WHERE version_id=?""",
                ("approved" if approved else "rejected", actor_id, comment[:1000], version_id),
            )
        return self.get_version(version_id, tenant_id=self.default_tenant_id)

    def evidence(self, version_id: str) -> dict[str, object]:
        with self._connect() as connection:
            version = connection.execute(
                "SELECT title,relative_path,quality_report_json FROM knowledge_source_versions WHERE version_id=?",
                (version_id,),
            ).fetchone()
            if version is None:
                raise KeyError("source_version_not_found")
            rows = connection.execute(
                """SELECT page_number,content,bbox_json,model_confidence,needs_review,evidence_kind
                FROM knowledge_version_chunks WHERE version_id=? AND bbox_json IS NOT NULL ORDER BY page_number,chunk_index""",
                (version_id,),
            ).fetchall()
        return {"version_id": version_id, "title": version[0], "relative_path": version[1],
                "quality_report": json.loads(version[2] or "{}"), "blocks": [
                    {"page": row[0], "text": row[1], "bbox": json.loads(row[2]),
                     "model_confidence": row[3], "needs_review": bool(row[4]), "kind": row[5]}
                    for row in rows]}

    def visual_evidence_by_title(self, titles: Sequence[str]) -> dict[str, list[dict[str, object]]]:
        result: dict[str, list[dict[str, object]]] = {}
        with self._connect() as connection:
            for title in dict.fromkeys(titles):
                version = connection.execute(
                    """SELECT version_id FROM knowledge_source_versions WHERE title=?
                    ORDER BY created_at DESC LIMIT 1""", (title,),
                ).fetchone()
                if version is None:
                    result[title] = []
                    continue
                rows = connection.execute(
                    """SELECT page_number,content,bbox_json,model_confidence,needs_review,evidence_kind
                    FROM knowledge_version_chunks WHERE version_id=? AND bbox_json IS NOT NULL
                    ORDER BY page_number,chunk_index""", (version[0],),
                ).fetchall()
                result[title] = [{
                    "page": row[0], "text": row[1], "bbox": json.loads(row[2]),
                    "model_confidence": row[3], "needs_review": bool(row[4]), "kind": row[5],
                } for row in rows]
        return result

    def save_visual_evaluation(self, report: dict[str, object], *, actor_id: str, tenant_id: str) -> str:
        evaluation_id = "veval_" + hashlib.sha256(
            f"{tenant_id}:{actor_id}:{datetime.now(timezone.utc).isoformat()}".encode("utf-8")
        ).hexdigest()[:24]
        with self._connect() as connection:
            connection.execute(
                "INSERT INTO knowledge_visual_evaluations(evaluation_id,tenant_id,actor_id,report_json) VALUES(?,?,?,?)",
                (evaluation_id, tenant_id, actor_id, json.dumps(report, ensure_ascii=False)),
            )
        return evaluation_id

    def latest_visual_evaluation(self, *, tenant_id: str) -> dict[str, object] | None:
        with self._connect() as connection:
            row = connection.execute(
                """SELECT evaluation_id,report_json,created_at FROM knowledge_visual_evaluations
                WHERE tenant_id=? ORDER BY created_at DESC LIMIT 1""", (tenant_id,),
            ).fetchone()
        if row is None:
            return None
        return {"evaluation_id": row[0], "created_at": row[2], **json.loads(row[1])}

    def compare_versions(
        self,
        left_version_id: str,
        right_version_id: str,
        *,
        tenant_id: str,
    ) -> dict[str, object]:
        left = self.get_version(left_version_id, tenant_id=tenant_id)
        right = self.get_version(right_version_id, tenant_id=tenant_id)
        same_source = left.source_id == right.source_id
        same_policy_family = _normalized_policy_title(left.title) == _normalized_policy_title(right.title)
        if not same_source and not same_policy_family:
            raise ValueError("knowledge_versions_not_same_source")
        with self._connect() as connection:
            left_chunks = connection.execute(
                "SELECT content FROM knowledge_version_chunks WHERE version_id=? ORDER BY chunk_index",
                (left_version_id,),
            ).fetchall()
            right_chunks = connection.execute(
                "SELECT content FROM knowledge_version_chunks WHERE version_id=? ORDER BY chunk_index",
                (right_version_id,),
            ).fetchall()
        left_texts = [str(row[0]) for row in left_chunks]
        right_texts = [str(row[0]) for row in right_chunks]
        left_fingerprints = {_normalized_fingerprint(text): text for text in left_texts}
        right_fingerprints = {_normalized_fingerprint(text): text for text in right_texts}
        added = [right_fingerprints[key] for key in right_fingerprints.keys() - left_fingerprints.keys()]
        removed = [left_fingerprints[key] for key in left_fingerprints.keys() - right_fingerprints.keys()]
        left_full = "\n".join(left_texts)
        right_full = "\n".join(right_texts)
        metadata_fields = (
            "source_version", "publication_date", "effective_date", "expiry_date",
            "policy_status", "authority_level",
        )
        metadata_changes = {
            field: {"before": getattr(left, field), "after": getattr(right, field)}
            for field in metadata_fields
            if getattr(left, field) != getattr(right, field)
        }
        return {
            "source_id": left.source_id if same_source else f"{left.source_id}|{right.source_id}",
            "title": right.title,
            "left_version_id": left_version_id,
            "right_version_id": right_version_id,
            "similarity": round(SequenceMatcher(None, left_full, right_full).ratio(), 4),
            "left_chunk_count": len(left_texts),
            "right_chunk_count": len(right_texts),
            "added_chunk_count": len(added),
            "removed_chunk_count": len(removed),
            "added_samples": [text[:800] for text in added[:5]],
            "removed_samples": [text[:800] for text in removed[:5]],
            "metadata_changes": metadata_changes,
        }

    def _record_failure(
        self,
        *,
        relative_path: str,
        source_id: str,
        collection: str,
        reason: str,
    ) -> None:
        with self._connect() as connection:
            connection.execute(
                """INSERT INTO knowledge_ingestion_failures(
                    relative_path,source_id,collection,reason,last_attempted_at
                ) VALUES(?,?,?,?,CURRENT_TIMESTAMP)
                ON CONFLICT(relative_path) DO UPDATE SET
                    source_id=excluded.source_id,
                    collection=excluded.collection,
                    reason=excluded.reason,
                    last_attempted_at=CURRENT_TIMESTAMP""",
                (relative_path, source_id, collection, reason),
            )

    def _clear_failure(self, relative_path: str) -> None:
        with self._connect() as connection:
            connection.execute(
                "DELETE FROM knowledge_ingestion_failures WHERE relative_path=?",
                (relative_path,),
            )

    def _ensure_schema(self) -> None:
        self.database_path.parent.mkdir(parents=True, exist_ok=True)
        with self._connect() as connection:
            connection.executescript(
                """
                CREATE TABLE IF NOT EXISTS knowledge_sources (
                    source_id TEXT PRIMARY KEY,
                    relative_path TEXT NOT NULL UNIQUE,
                    title TEXT NOT NULL,
                    collection TEXT NOT NULL,
                    checksum TEXT NOT NULL,
                    embedding_model TEXT NOT NULL,
                    source_version TEXT,
                    publication_date TEXT,
                    effective_date TEXT,
                    expiry_date TEXT,
                    policy_status TEXT NOT NULL DEFAULT 'unknown',
                    authority_level TEXT NOT NULL DEFAULT 'unknown',
                    metadata_version TEXT NOT NULL DEFAULT '',
                    indexed_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
                );
                CREATE TABLE IF NOT EXISTS knowledge_chunks (
                    chunk_id TEXT PRIMARY KEY,
                    source_id TEXT NOT NULL REFERENCES knowledge_sources(source_id) ON DELETE CASCADE,
                    chunk_index INTEGER NOT NULL,
                    page_number INTEGER,
                    section TEXT NOT NULL DEFAULT '',
                    anchor_kind TEXT NOT NULL DEFAULT 'section',
                    anchor_start INTEGER,
                    anchor_end INTEGER,
                    content TEXT NOT NULL,
                    embedding TEXT NOT NULL,
                    embedding_model TEXT NOT NULL
                );
                CREATE INDEX IF NOT EXISTS idx_knowledge_chunks_source ON knowledge_chunks(source_id);
                CREATE INDEX IF NOT EXISTS idx_knowledge_chunks_model ON knowledge_chunks(embedding_model);
                CREATE TABLE IF NOT EXISTS knowledge_source_versions (
                    version_id TEXT PRIMARY KEY,
                    source_id TEXT NOT NULL,
                    relative_path TEXT NOT NULL,
                    title TEXT NOT NULL,
                    collection TEXT NOT NULL,
                    checksum TEXT NOT NULL,
                    embedding_model TEXT NOT NULL,
                    source_version TEXT,
                    publication_date TEXT,
                    effective_date TEXT,
                    expiry_date TEXT,
                    policy_status TEXT NOT NULL DEFAULT 'unknown',
                    authority_level TEXT NOT NULL DEFAULT 'unknown',
                    metadata_version TEXT NOT NULL DEFAULT '',
                    ingestion_status TEXT NOT NULL DEFAULT 'ready',
                    quality_score INTEGER NOT NULL DEFAULT 100,
                    quality_review_status TEXT NOT NULL DEFAULT 'not_required',
                    quality_report_json TEXT NOT NULL DEFAULT '{}',
                    quality_reviewed_by TEXT,
                    quality_reviewed_at TEXT,
                    quality_review_comment TEXT NOT NULL DEFAULT '',
                    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    UNIQUE(source_id, checksum, embedding_model, metadata_version)
                );
                CREATE INDEX IF NOT EXISTS idx_knowledge_versions_source
                    ON knowledge_source_versions(source_id, created_at);
                CREATE TABLE IF NOT EXISTS knowledge_version_chunks (
                    chunk_id TEXT PRIMARY KEY,
                    version_id TEXT NOT NULL REFERENCES knowledge_source_versions(version_id) ON DELETE CASCADE,
                    source_id TEXT NOT NULL,
                    chunk_index INTEGER NOT NULL,
                    page_number INTEGER,
                    section TEXT NOT NULL DEFAULT '',
                    anchor_kind TEXT NOT NULL DEFAULT 'section',
                    anchor_start INTEGER,
                    anchor_end INTEGER,
                    content TEXT NOT NULL,
                    embedding TEXT NOT NULL,
                    embedding_model TEXT NOT NULL
                    ,bbox_json TEXT
                    ,model_confidence REAL
                    ,needs_review INTEGER NOT NULL DEFAULT 0
                    ,evidence_kind TEXT NOT NULL DEFAULT 'text'
                );
                CREATE INDEX IF NOT EXISTS idx_knowledge_version_chunks_version
                    ON knowledge_version_chunks(version_id, chunk_index);
                CREATE TABLE IF NOT EXISTS knowledge_publications (
                    publication_id TEXT PRIMARY KEY,
                    version_id TEXT NOT NULL REFERENCES knowledge_source_versions(version_id),
                    tenant_id TEXT NOT NULL,
                    status TEXT NOT NULL CHECK(status IN ('draft','published','withdrawn')),
                    published_by TEXT NOT NULL,
                    published_at TEXT,
                    withdrawn_at TEXT,
                    UNIQUE(version_id, tenant_id)
                );
                CREATE INDEX IF NOT EXISTS idx_knowledge_publications_access
                    ON knowledge_publications(tenant_id, status);
                CREATE TABLE IF NOT EXISTS knowledge_publication_acl (
                    publication_id TEXT NOT NULL REFERENCES knowledge_publications(publication_id) ON DELETE CASCADE,
                    subject_type TEXT NOT NULL CHECK(subject_type IN ('organization','role','user')),
                    subject_id TEXT NOT NULL,
                    permission TEXT NOT NULL CHECK(permission IN ('retrieve')),
                    PRIMARY KEY(publication_id, subject_type, subject_id, permission)
                );
                CREATE TABLE IF NOT EXISTS knowledge_ingestion_failures (
                    relative_path TEXT PRIMARY KEY,
                    source_id TEXT NOT NULL,
                    collection TEXT NOT NULL,
                    reason TEXT NOT NULL,
                    last_attempted_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    retry_count INTEGER NOT NULL DEFAULT 0,
                    last_error TEXT NOT NULL DEFAULT ''
                );
                CREATE TABLE IF NOT EXISTS knowledge_visual_evaluations (
                    evaluation_id TEXT PRIMARY KEY,
                    tenant_id TEXT NOT NULL,
                    actor_id TEXT NOT NULL,
                    report_json TEXT NOT NULL,
                    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
                );
                """
            )
            _ensure_source_metadata_columns(connection)
            _ensure_chunk_anchor_columns(connection)
            _ensure_knowledge_quality_columns(connection)
            _ensure_ingestion_failure_columns(connection)
            _bootstrap_legacy_sources(connection, self.default_tenant_id)

    def _connect(self) -> sqlite3.Connection:
        connection = sqlite3.connect(self.database_path)
        connection.execute("PRAGMA foreign_keys = ON")
        return connection


def _ensure_source_metadata_columns(connection: sqlite3.Connection) -> None:
    existing = {
        str(row[1])
        for row in connection.execute("PRAGMA table_info(knowledge_sources)").fetchall()
    }
    definitions = {
        "source_version": "TEXT",
        "publication_date": "TEXT",
        "effective_date": "TEXT",
        "expiry_date": "TEXT",
        "policy_status": "TEXT NOT NULL DEFAULT 'unknown'",
        "authority_level": "TEXT NOT NULL DEFAULT 'unknown'",
        "metadata_version": "TEXT NOT NULL DEFAULT ''",
        "version_id": "TEXT",
        "publication_id": "TEXT",
        "tenant_id": "TEXT NOT NULL DEFAULT '*'",
    }
    for name, definition in definitions.items():
        if name not in existing:
            connection.execute(
                f"ALTER TABLE knowledge_sources ADD COLUMN {name} {definition}"
            )


def _ensure_chunk_anchor_columns(connection: sqlite3.Connection) -> None:
    existing = {
        str(row[1]) for row in connection.execute("PRAGMA table_info(knowledge_chunks)").fetchall()
    }
    definitions = {
        "anchor_kind": "TEXT NOT NULL DEFAULT 'section'",
        "anchor_start": "INTEGER",
        "anchor_end": "INTEGER",
    }
    for name, definition in definitions.items():
        if name not in existing:
            connection.execute(f"ALTER TABLE knowledge_chunks ADD COLUMN {name} {definition}")


def _ensure_knowledge_quality_columns(connection: sqlite3.Connection) -> None:
    versions = {str(row[1]) for row in connection.execute(
        "PRAGMA table_info(knowledge_source_versions)"
    ).fetchall()}
    for name, definition in {
        "ingestion_status": "TEXT NOT NULL DEFAULT 'ready'",
        "quality_score": "INTEGER NOT NULL DEFAULT 100",
        "quality_review_status": "TEXT NOT NULL DEFAULT 'not_required'",
        "quality_report_json": "TEXT NOT NULL DEFAULT '{}'",
        "quality_reviewed_by": "TEXT",
        "quality_reviewed_at": "TEXT",
        "quality_review_comment": "TEXT NOT NULL DEFAULT ''",
    }.items():
        if name not in versions:
            connection.execute(f"ALTER TABLE knowledge_source_versions ADD COLUMN {name} {definition}")
    chunks = {str(row[1]) for row in connection.execute(
        "PRAGMA table_info(knowledge_version_chunks)"
    ).fetchall()}
    for name, definition in {
        "bbox_json": "TEXT", "model_confidence": "REAL",
        "needs_review": "INTEGER NOT NULL DEFAULT 0",
        "evidence_kind": "TEXT NOT NULL DEFAULT 'text'",
    }.items():
        if name not in chunks:
            connection.execute(f"ALTER TABLE knowledge_version_chunks ADD COLUMN {name} {definition}")


def _ensure_ingestion_failure_columns(connection: sqlite3.Connection) -> None:
    existing = {str(row[1]) for row in connection.execute(
        "PRAGMA table_info(knowledge_ingestion_failures)"
    ).fetchall()}
    for name, definition in {
        "retry_count": "INTEGER NOT NULL DEFAULT 0",
        "last_error": "TEXT NOT NULL DEFAULT ''",
    }.items():
        if name not in existing:
            connection.execute(f"ALTER TABLE knowledge_ingestion_failures ADD COLUMN {name} {definition}")


def _bootstrap_legacy_sources(connection: sqlite3.Connection, tenant_id: str) -> None:
    rows = connection.execute(
        """SELECT source_id,relative_path,title,collection,checksum,embedding_model,
        source_version,publication_date,effective_date,expiry_date,policy_status,
        authority_level,metadata_version FROM knowledge_sources WHERE version_id IS NULL"""
    ).fetchall()
    for row in rows:
        version_id = "ver_" + hashlib.sha256(
            f"legacy:{row[0]}:{row[4]}:{row[5]}".encode("utf-8")
        ).hexdigest()[:24]
        publication_id = "pub_" + hashlib.sha256(
            f"{version_id}:{tenant_id}".encode("utf-8")
        ).hexdigest()[:24]
        connection.execute(
            """INSERT OR IGNORE INTO knowledge_source_versions(
                version_id,source_id,relative_path,title,collection,checksum,embedding_model,
                source_version,publication_date,effective_date,expiry_date,policy_status,
                authority_level,metadata_version
            ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (version_id, *row),
        )
        legacy_chunks = connection.execute(
            """SELECT chunk_id,source_id,chunk_index,page_number,section,anchor_kind,
            anchor_start,anchor_end,content,embedding,embedding_model
            FROM knowledge_chunks WHERE source_id=?""",
            (row[0],),
        ).fetchall()
        connection.executemany(
            """INSERT OR IGNORE INTO knowledge_version_chunks(
                chunk_id,version_id,source_id,chunk_index,page_number,section,anchor_kind,
                anchor_start,anchor_end,content,embedding,embedding_model
            ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)""",
            [(item[0], version_id, *item[1:]) for item in legacy_chunks],
        )
        connection.execute(
            """INSERT OR IGNORE INTO knowledge_publications(
                publication_id,version_id,tenant_id,status,published_by,published_at
            ) VALUES(?,?,?,'published','legacy-migration',CURRENT_TIMESTAMP)""",
            (publication_id, version_id, tenant_id),
        )
        for subject_type, subject_id in (("role", "admin"), ("role", "user")):
            connection.execute(
                """INSERT OR IGNORE INTO knowledge_publication_acl(
                publication_id,subject_type,subject_id,permission
                ) VALUES(?,?,?,'retrieve')""",
                (publication_id, subject_type, subject_id),
            )
        connection.execute(
            "UPDATE knowledge_sources SET version_id=?,publication_id=?,tenant_id=? WHERE source_id=?",
            (version_id, publication_id, tenant_id, row[0]),
        )


def _infer_policy_metadata(
    path: Path,
    relative_path: str,
    chunks: Sequence[ChunkDraft],
) -> dict[str, str | None]:
    title = path.stem
    leading_text = "\n".join(chunk.text for chunk in chunks[:4])[:8_000]
    searchable = f"{title}\n{leading_text}"
    source_version = _first_match(
        title,
        (
            r"((?:[\u3400-\u9fff]{2,20})?发[〔【\[](?:19|20)\d{2}[〕】\]]\s*\d+\s*号)",
            r"(?:v|V|版本)\s*(\d+(?:\.\d+)+)",
            r"((?:19|20)\d{2}年(?:修订|版|试行))",
            r"((?:修订|修正)版)",
        ),
    )
    publication_date = _date_after_markers(
        searchable,
        ("发布日期", "印发日期", "公布日期", "发布于", "印发"),
    ) or _first_policy_date(searchable)
    effective_date = _date_after_markers(
        searchable,
        ("施行", "执行", "生效", "实施"),
        allow_marker_after_date=True,
    )
    expiry_date = _date_after_markers(
        searchable,
        ("有效期至", "失效日期", "废止日期", "截止"),
    )
    title_status = title.casefold()
    if any(marker in title_status for marker in ("废止", "失效", "作废", "已停止执行")):
        policy_status = "repealed"
    elif "试行" in title_status:
        policy_status = "trial"
    elif re.search(r"本(?:办法|规定|通知|文件).{0,12}(?:废止|失效|停止执行)", leading_text):
        policy_status = "repealed"
    elif any(marker in searchable for marker in ("自本通知印发之日起施行", "自发布之日起施行", "现予印发")):
        policy_status = "active"
    else:
        policy_status = "unknown"
    authority_level = _authority_level(f"{relative_path}\n{searchable[:2_000]}")
    return {
        "source_version": source_version,
        "publication_date": publication_date,
        "effective_date": effective_date,
        "expiry_date": expiry_date,
        "policy_status": policy_status,
        "authority_level": authority_level,
    }


def _first_match(value: str, patterns: Sequence[str]) -> str | None:
    for pattern in patterns:
        match = re.search(pattern, value)
        if match:
            return match.group(1)
    return None


def _date_after_markers(
    value: str,
    markers: Sequence[str],
    *,
    allow_marker_after_date: bool = False,
) -> str | None:
    date_pattern = r"((?:19|20)\d{2})[年./-](\d{1,2})[月./-](\d{1,2})日?"
    for marker in markers:
        after = re.search(rf"{re.escape(marker)}.{{0,24}}?{date_pattern}", value)
        if after:
            return _normalize_date(after.group(1), after.group(2), after.group(3))
        if allow_marker_after_date:
            before = re.search(rf"{date_pattern}.{{0,12}}?{re.escape(marker)}", value)
            if before:
                return _normalize_date(before.group(1), before.group(2), before.group(3))
    return None


def _first_policy_date(value: str) -> str | None:
    match = re.search(
        r"((?:19|20)\d{2})[年./-](\d{1,2})[月./-](\d{1,2})日?",
        value,
    )
    return _normalize_date(match.group(1), match.group(2), match.group(3)) if match else None


def _normalize_date(year: str, month: str, day: str) -> str:
    return f"{int(year):04d}-{int(month):02d}-{int(day):02d}"


def _authority_level(value: str) -> str:
    if any(marker in value for marker in ("国务院", "全国人民代表大会", "人力资源社会保障部", "财政部", "教育部")):
        return "national"
    if any(marker in value for marker in ("浙江省", "省人民政府", "省人力资源", "省财政")):
        return "provincial"
    if any(marker in value for marker in ("宁波市", "市人民政府", "市人力资源", "市财政")):
        return "municipal"
    if any(marker in value for marker in ("大学", "学院", "学校", "实验室", "研究院")):
        return "institution"
    if any(marker in value for marker in ("公司", "集团", "学院", "学校", "医院")):
        return "organization"
    return "unknown"


def _extract_chunks(path: Path) -> list[ChunkDraft]:
    extension = path.suffix.casefold()
    if extension == ".pdf":
        reader = PdfReader(path)
        if reader.is_encrypted:
            return []
        result: list[ChunkDraft] = []
        for page_number, page in enumerate(reader.pages, start=1):
            result.extend(_split_text(
                page.extract_text() or "",
                page_number=page_number,
                anchor_kind="page",
                fixed_anchor_start=page_number,
                fixed_anchor_end=page_number,
            ))
        return result
    if extension == ".docx":
        document = Document(path)
        parts: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                parts.append(text)
        return _split_text("\n".join(parts), anchor_kind="paragraph")
    if extension in {".md", ".txt"}:
        raw = path.read_bytes()
        for encoding in ("utf-8-sig", "utf-8", "gb18030"):
            try:
                return _split_text(raw.decode(encoding), anchor_kind="line")
            except UnicodeDecodeError:
                continue
    return []


def _split_text(
    text: str,
    *,
    page_number: int | None = None,
    anchor_kind: str = "section",
    fixed_anchor_start: int | None = None,
    fixed_anchor_end: int | None = None,
    max_chars: int = 900,
    overlap: int = 120,
) -> list[ChunkDraft]:
    clean = re.sub(r"[ \t]+", " ", text.replace("\r", "\n"))
    paragraphs = [item.strip() for item in re.split(r"\n+", clean) if item.strip()]
    if not paragraphs:
        return []
    chunks: list[ChunkDraft] = []
    current: list[str] = []
    current_length = 0
    current_start = 1
    section = ""
    for paragraph_index, paragraph in enumerate(paragraphs, start=1):
        if _looks_like_heading(paragraph):
            section = paragraph[:120]
        if current and current_length + len(paragraph) + 1 > max_chars:
            joined = "\n".join(current)
            chunks.append(ChunkDraft(
                joined, page_number, section, anchor_kind,
                fixed_anchor_start or current_start,
                fixed_anchor_end or max(current_start, paragraph_index - 1),
            ))
            tail = joined[-overlap:] if overlap else ""
            current = [tail, paragraph] if tail else [paragraph]
            current_length = sum(len(item) for item in current) + len(current) - 1
            current_start = paragraph_index
        else:
            if not current:
                current_start = paragraph_index
            current.append(paragraph)
            current_length += len(paragraph) + 1
    if current:
        chunks.append(ChunkDraft(
            "\n".join(current), page_number, section, anchor_kind,
            fixed_anchor_start or current_start,
            fixed_anchor_end or len(paragraphs),
        ))
    return [chunk for chunk in chunks if len(chunk.text.strip()) >= 40]


def _looks_like_heading(value: str) -> bool:
    return len(value) <= 120 and bool(re.match(
        r"^(第[一二三四五六七八九十百0-9]+[编章节条]|[一二三四五六七八九十]+、|\d+[.、]|#+\s*)",
        value,
    ))


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _normalized_fingerprint(value: str) -> str:
    normalized = re.sub(r"\s+", "", value.casefold())
    return hashlib.sha256(normalized.encode("utf-8")).hexdigest()


def _normalized_policy_title(value: str) -> str:
    """Identify archived snapshots that represent the same named policy family."""

    normalized = Path(value).stem.casefold()
    normalized = re.sub(r"(?:19|20)\d{2}[年./_-]?\d{0,2}[月./_-]?\d{0,2}日?", "", normalized)
    normalized = re.sub(r"[（(](?:修订|修正版|草案|征求意见稿|第\d+版)[）)]", "", normalized)
    return re.sub(r"[\s._-]+", "", normalized)


def _normalize(values: Sequence[float]) -> list[float]:
    norm = math.sqrt(sum(float(value) ** 2 for value in values))
    if not norm:
        return [float(value) for value in values]
    return [float(value) / norm for value in values]


def _lexical_terms(value: str) -> list[str]:
    """Tokenize Chinese policy language and exact identifiers without external services."""

    normalized = re.sub(r"\s+", "", value.casefold())
    terms = re.findall(r"[a-z0-9][a-z0-9._/-]*", normalized)
    for sequence in re.findall(r"[\u3400-\u9fff]+", normalized):
        if len(sequence) <= 8:
            terms.append(sequence)
        terms.extend(sequence[index:index + 2] for index in range(max(0, len(sequence) - 1)))
        terms.extend(sequence[index:index + 3] for index in range(max(0, len(sequence) - 2)))
    return terms


def _bm25_score(
    query_terms: list[str],
    document_terms: Counter[str],
    document_frequency: Counter[str],
    *,
    document_count: int,
    average_length: float,
) -> float:
    if not query_terms or not document_terms or document_count <= 0:
        return 0.0
    length = sum(document_terms.values())
    score = 0.0
    k1 = 1.2
    b = 0.75
    for term in set(query_terms):
        frequency = document_terms.get(term, 0)
        if not frequency:
            continue
        frequency_docs = document_frequency.get(term, 0)
        inverse_frequency = math.log(
            1 + (document_count - frequency_docs + 0.5) / (frequency_docs + 0.5)
        )
        denominator = frequency + k1 * (1 - b + b * length / max(average_length, 1.0))
        score += inverse_frequency * frequency * (k1 + 1) / denominator
    return score
