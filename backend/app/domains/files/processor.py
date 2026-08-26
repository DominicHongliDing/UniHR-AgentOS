from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from pathlib import Path
from uuid import uuid4

from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader

from .legacy_office import LegacyOfficeConverter, LegacyOfficeConversionError


class UnsupportedFileError(ValueError):
    pass


@dataclass(frozen=True)
class ProcessedUpload:
    display_name: str
    mime_type: str
    extension: str
    byte_size: int
    saved_path: Path
    extracted_text: str
    summary: dict[str, object]


class FileProcessor:
    allowed_extensions = {
        ".txt", ".md", ".csv", ".tsv", ".pdf", ".doc", ".docx", ".xls", ".xlsx",
        ".png", ".jpg", ".jpeg", ".webp",
    }

    def __init__(
        self,
        upload_dir: str | Path,
        max_bytes: int,
        legacy_converter: LegacyOfficeConverter | None = None,
    ) -> None:
        self.upload_dir = Path(upload_dir)
        self.upload_dir.mkdir(parents=True, exist_ok=True)
        self.max_bytes = max_bytes
        self.legacy_converter = legacy_converter or LegacyOfficeConverter(
            self.upload_dir / "converted",
        )

    def process(self, filename: str, mime_type: str, content: bytes) -> ProcessedUpload:
        display_name = Path(filename or "upload").name
        extension = Path(display_name).suffix.casefold()
        if extension not in self.allowed_extensions:
            raise UnsupportedFileError(f"暂不支持 {extension or '无扩展名'} 文件")
        if not content:
            raise UnsupportedFileError("文件内容为空")
        if len(content) > self.max_bytes:
            raise UnsupportedFileError(f"单个文件不能超过 {self.max_bytes // 1024 // 1024} MB")
        saved_path = self.upload_dir / f"upl_{uuid4().hex}{extension}"
        saved_path.write_bytes(content)
        effective_extension = extension
        effective_mime_type = mime_type or "application/octet-stream"
        try:
            if extension in {".doc", ".xls"}:
                converted = self.legacy_converter.convert(saved_path)
                converted_content = converted.path.read_bytes()
                extracted, summary = self._extract(converted.extension, converted_content)
                summary = {
                    **summary,
                    "original_extension": extension,
                    "converted_format": converted.extension,
                    "conversion_engine": converted.engine,
                    "application_version": converted.application_version,
                    "macros_removed": True,
                }
                saved_path.unlink(missing_ok=True)
                saved_path = converted.path
                effective_extension = converted.extension
                effective_mime_type = converted.mime_type
            else:
                extracted, summary = self._extract(extension, content)
        except LegacyOfficeConversionError as exc:
            saved_path.unlink(missing_ok=True)
            raise UnsupportedFileError(
                "旧版 Word/Excel 转换失败，请先在 Office 中另存为 .docx 或 .xlsx 后重新上传"
            ) from exc
        except Exception:
            saved_path.unlink(missing_ok=True)
            raise
        return ProcessedUpload(
            display_name=display_name,
            mime_type=effective_mime_type,
            extension=effective_extension,
            byte_size=len(content),
            saved_path=saved_path,
            extracted_text=extracted[:80_000],
            summary=summary,
        )

    def extract_content(self, filename: str, content: bytes) -> tuple[str, str, dict[str, object]]:
        """Extract an existing governed file without copying it into the upload workspace."""
        extension = Path(filename).suffix.casefold()
        if extension not in self.allowed_extensions or extension in {".doc", ".xls"}:
            raise UnsupportedFileError(f"暂不支持读取 {extension or '无扩展名'} Skill 资源")
        if not content:
            raise UnsupportedFileError("Skill 资源内容为空")
        if len(content) > self.max_bytes:
            raise UnsupportedFileError(f"Skill 资源不能超过 {self.max_bytes // 1024 // 1024} MB")
        extracted, summary = self._extract(extension, content)
        return extension, extracted[:80_000], summary

    def _extract(self, extension: str, content: bytes) -> tuple[str, dict[str, object]]:
        if extension in {".txt", ".md"}:
            text = _decode_text(content)
            return text, {"kind": "text", "characters": len(text)}
        if extension in {".csv", ".tsv"}:
            return _extract_delimited(content, "\t" if extension == ".tsv" else ",")
        if extension == ".pdf":
            return _extract_pdf(content)
        if extension == ".docx":
            return _extract_docx(content)
        if extension == ".xlsx":
            return _extract_xlsx(content)
        if extension in {".png", ".jpg", ".jpeg", ".webp"}:
            return _inspect_image(content)
        raise UnsupportedFileError(f"暂不支持 {extension} 文件")


def image_data_url(path: str | Path, mime_type: str) -> str:
    import base64

    data = Path(path).read_bytes()
    return f"data:{mime_type};base64,{base64.b64encode(data).decode('ascii')}"


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _extract_delimited(content: bytes, delimiter: str) -> tuple[str, dict[str, object]]:
    text = _decode_text(content)
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    width = max((len(row) for row in rows), default=0)
    preview = rows[:80]
    rendered = "\n".join(" | ".join(cell[:300] for cell in row) for row in preview)
    return rendered, {
        "kind": "table",
        "rows": len(rows),
        "columns": width,
        "headers": rows[0][:50] if rows else [],
        "preview_truncated": len(rows) > len(preview),
    }


def _extract_pdf(content: bytes) -> tuple[str, dict[str, object]]:
    try:
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            raise UnsupportedFileError("PDF 已加密，需要先解密后上传")
        pages: list[str] = []
        low_text_pages: list[int] = []
        for index, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if len(text) < 30:
                low_text_pages.append(index)
            pages.append(f"[第 {index} 页]\n{text}")
        return "\n\n".join(pages), {
            "kind": "pdf",
            "pages": len(reader.pages),
            "low_text_pages": low_text_pages,
            "requires_vision": bool(low_text_pages),
        }
    except UnsupportedFileError:
        raise
    except Exception as exc:
        raise UnsupportedFileError("PDF 无法解析，请检查文件是否损坏") from exc


def _extract_docx(content: bytes) -> tuple[str, dict[str, object]]:
    try:
        document = Document(io.BytesIO(content))
        parts: list[str] = []
        for index, paragraph in enumerate(document.paragraphs, start=1):
            text = " ".join(paragraph.text.split())
            if text:
                parts.append(f"[段落 {index}] {text}")
        for table_index, table in enumerate(document.tables, start=1):
            parts.append(f"[表格 {table_index}]")
            for row in table.rows[:100]:
                parts.append(" | ".join(" ".join(cell.text.split()) for cell in row.cells))
        return "\n".join(parts), {
            "kind": "document",
            "paragraphs": len(document.paragraphs),
            "tables": len(document.tables),
        }
    except Exception as exc:
        raise UnsupportedFileError("Word 文档无法解析") from exc


def _extract_xlsx(content: bytes) -> tuple[str, dict[str, object]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        parts: list[str] = []
        sheets: list[dict[str, object]] = []
        formula_count = 0
        for worksheet in workbook.worksheets:
            max_rows = min(worksheet.max_row or 0, 100)
            max_cols = min(worksheet.max_column or 0, 40)
            sheet_formulas = 0
            parts.append(f"[Sheet: {worksheet.title}]")
            for row in worksheet.iter_rows(min_row=1, max_row=max_rows, max_col=max_cols, values_only=False):
                values: list[str] = []
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        sheet_formulas += 1
                    values.append("" if cell.value is None else str(cell.value)[:300])
                if any(values):
                    parts.append(" | ".join(values))
            formula_count += sheet_formulas
            sheets.append({
                "name": worksheet.title,
                "rows": worksheet.max_row or 0,
                "columns": worksheet.max_column or 0,
                "preview_rows": max_rows,
                "formulas_in_preview": sheet_formulas,
            })
        workbook.close()
        return "\n".join(parts), {
            "kind": "workbook",
            "sheets": sheets,
            "formula_count_in_preview": formula_count,
            "values_truncated": any(int(item["rows"]) > int(item["preview_rows"]) for item in sheets),
        }
    except Exception as exc:
        raise UnsupportedFileError("Excel 工作簿无法解析") from exc


def _inspect_image(content: bytes) -> tuple[str, dict[str, object]]:
    try:
        with Image.open(io.BytesIO(content)) as image:
            width, height = image.size
            mode = image.mode
            image_format = image.format or "unknown"
        return "", {
            "kind": "image",
            "width": width,
            "height": height,
            "mode": mode,
            "format": image_format,
            "requires_vision": True,
        }
    except Exception as exc:
        raise UnsupportedFileError("图片无法解析") from exc
