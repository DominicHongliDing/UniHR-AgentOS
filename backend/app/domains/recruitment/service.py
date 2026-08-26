from __future__ import annotations

import hashlib
import json
import operator
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Annotated, Any, Literal, TypedDict
from uuid import uuid4

from langgraph.graph import END, START, StateGraph
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

from app.domains.conversations.contracts import AttachmentRecord
from app.models.gateway import ModelGateway, parse_json_object
from app.domains.runs import RunPointer, RunPayloadStore
from app.observability import emit_tool_progress

from .contracts import (
    CandidateEvaluation,
    CriterionEvidence,
    DuplicateSourceGroup,
    EvidenceLocator,
    RecruitmentCaseSnapshot,
    RecruitmentHarnessReport,
    RecruitmentReport,
    RecruitmentRequest,
    ResumeSourceProfile,
    ResumeVersionGroup,
    RubricSnapshot,
    ScreeningCriterion,
)
from .harness import RecruitmentHarness, remove_protected_attributes
from .evidence_locator import RecruitmentEvidenceLocator
from .version_resolver import ResumeVersionResolver


_RESUME_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".xlsx", ".csv", ".tsv"}
_FAIRNESS_NOTE = (
    "本报告仅用于辅助人工复核，不构成录用、淘汰或自动化就业决定。"
    "评分不得使用性别、年龄、民族、婚育、照片、籍贯、住址等与岗位胜任力无关的属性。"
)
_LOCATION_CRITERION_PATTERN = re.compile(r"(?:宁波|工作地点|实习地点|到岗|地点可行性)")
_EXPLICIT_LOCATION_INTENT_PATTERN = re.compile(r"(?:期望城市|意向城市|可在.{0,8}(?:实习|工作)|愿意.{0,8}(?:实习|工作)|能够.{0,8}(?:实习|工作)|可到岗|接受.{0,8}地点)")


class _RecruitmentState(TypedDict, total=False):
    request: dict[str, Any]
    sources: list[dict[str, Any]]
    source_profiles: list[dict[str, Any]]
    duplicate_groups: list[dict[str, Any]]
    version_groups: list[dict[str, Any]]
    representative_sources: list[dict[str, Any]]
    job_title: str
    criteria: list[dict[str, Any]]
    rubric_hash: str
    rubric_candidate_payload: dict[str, Any]
    candidate_payloads: list[dict[str, Any]]
    candidate_payloads_ref: dict[str, Any]
    visual_evidence: list[dict[str, Any]]
    visual_evidence_ref: dict[str, Any]
    payload_scope: dict[str, Any]
    report: dict[str, Any]
    harness: dict[str, Any]
    artifact_path: str
    repair_count: int
    model_names: list[str]
    latency_ms: int
    graph_path: Annotated[list[str], operator.add]
    trace: Annotated[list[dict[str, Any]], operator.add]


@dataclass(frozen=True)
class RecruitmentScreeningResult:
    answer: str
    report: RecruitmentReport
    artifact_path: Path
    display_name: str
    model: str
    latency_ms: int
    metadata: dict[str, Any]
    trace: list[dict[str, Any]]
    graph_path: list[str]
    repair_count: int = 0


class RecruitmentScreeningModule:
    """Deep module from one locked recruitment request to a reviewable report."""

    def __init__(
        self,
        gateway: ModelGateway,
        output_dir: str | Path,
        model: str,
        payload_store: RunPayloadStore | None = None,
    ) -> None:
        self.gateway = gateway
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.model = model
        self.payload_store = payload_store
        self.harness = RecruitmentHarness()
        self.evidence_locator = RecruitmentEvidenceLocator()
        self.version_resolver = ResumeVersionResolver()
        self.graph = self._build_graph()

    def screen(
        self,
        job_requirements: str,
        attachments: list[AttachmentRecord],
        *,
        skill_instructions: str = "",
        job_title: str = "",
        must_haves: str = "",
        nice_to_haves: str = "",
        selection_count: int = 3,
        payload_scope: RunPointer | None = None,
        visual_evidence: list[dict[str, Any]] | None = None,
    ) -> RecruitmentScreeningResult:
        eligible = [item for item in attachments if item.extension in _RESUME_EXTENSIONS]
        if len(eligible) > 30:
            raise ValueError("resume_batch_limit_30")
        return self.run(RecruitmentRequest(
            job_requirements=job_requirements,
            attachments=eligible,
            job_title=job_title,
            must_haves=must_haves,
            nice_to_haves=nice_to_haves,
            selection_count=selection_count,
            skill_instructions=skill_instructions,
        ), payload_scope=payload_scope, visual_evidence=visual_evidence)

    def run(
        self,
        request: RecruitmentRequest,
        *,
        payload_scope: RunPointer | None = None,
        visual_evidence: list[dict[str, Any]] | None = None,
    ) -> RecruitmentScreeningResult:
        visual_state, visual_trace = self._prepare_visual_evidence(
            visual_evidence or [], payload_scope,
        )
        state = self.graph.invoke({
            "request": request.model_dump(mode="json"),
            "repair_count": 0,
            "model_names": [],
            "latency_ms": 0,
            "graph_path": [],
            "trace": visual_trace,
            "payload_scope": payload_scope.model_dump(mode="json") if payload_scope else {},
            **visual_state,
        })
        report = RecruitmentReport.model_validate(state["report"])
        artifact_path = Path(state["artifact_path"])
        graph_path = list(state.get("graph_path") or [])
        repair_count = int(state.get("repair_count") or 0)
        model_names = list(dict.fromkeys(state.get("model_names") or []))
        metadata = {
            "job_title": report.job_title,
            "candidate_count": len(report.candidates),
            "source_count": len(report.source_profiles),
            "unique_source_count": report.case.unique_source_count if report.case else len(report.source_profiles),
            "duplicate_source_count": report.case.duplicate_source_count if report.case else 0,
            "selection_count": report.case.selection_count if report.case else request.selection_count,
            "rubric": report.rubric.model_dump(mode="json") if report.rubric else {},
            "candidate_summary": [
                {
                    "candidate_id": _candidate_id(
                        report.case.case_id if report.case else "recruitment_case",
                        candidate.source_name,
                        candidate.candidate_name,
                    ),
                    "rank": index,
                    "candidate_name": candidate.candidate_name,
                    "source_name": candidate.source_name,
                    "total_score": candidate.total_score,
                    "recommendation": candidate.recommendation,
                    "evidence_coverage": candidate.evidence_coverage,
                    "strengths": candidate.strengths[:5],
                    "gaps": candidate.gaps[:3],
                    "verification_questions": candidate.verification_questions[:5],
                    "evidence_locations": [
                        {
                            "criterion_label": row.criterion_label,
                            "evidence": row.evidence,
                            "locator": row.locator.model_dump(mode="json") if row.locator else {},
                        }
                        for row in candidate.criterion_scores
                    ],
                }
                for index, candidate in enumerate(report.candidates, start=1)
            ],
            "duplicate_groups": [item.model_dump(mode="json") for item in report.duplicate_groups],
            "version_groups": [item.model_dump(mode="json") for item in report.version_groups],
            "semantic_merged_source_count": sum(
                max(0, len(item.version_sources) - 1)
                for item in report.version_groups if item.status == "merged"
            ),
            "version_review_group_count": sum(
                item.status == "review_required" for item in report.version_groups
            ),
            "recruitment_harness": report.harness.model_dump(mode="json") if report.harness else {},
            "recruitment_case": report.case.model_dump(mode="json") if report.case else {},
            "recruitment_graph_path": graph_path,
            "repair_count": repair_count,
            "human_review_required": True,
            "automated_employment_decision": False,
            "fairness_note": report.fairness_note,
        }
        return RecruitmentScreeningResult(
            answer=_render_answer(report),
            report=report,
            artifact_path=artifact_path,
            display_name=f"候选人筛选报告-{_safe_name(report.job_title)}.xlsx",
            model=" + ".join(model_names) or self.model,
            latency_ms=int(state.get("latency_ms") or 0),
            metadata=metadata,
            trace=list(state.get("trace") or []),
            graph_path=graph_path,
            repair_count=repair_count,
        )

    def _build_graph(self):
        builder = StateGraph(_RecruitmentState)
        builder.add_node("prepare", self._prepare)
        builder.add_node("deduplicate", self._deduplicate)
        builder.add_node("version_resolve", self._version_resolve)
        builder.add_node("rubric_lock", self._rubric_lock)
        builder.add_node("batch_map", self._batch_map)
        builder.add_node("reduce", self._reduce)
        builder.add_node("evidence_localize", self._evidence_localize)
        builder.add_node("fairness_guard", self._fairness_guard)
        builder.add_node("recruitment_harness", self._recruitment_harness)
        builder.add_node("quality_decision", self._quality_decision)
        builder.add_node("repair_once", self._repair_once)
        builder.add_node("artifact_create", self._artifact_create)
        builder.add_node("finalize", self._finalize)
        builder.add_edge(START, "prepare")
        builder.add_edge("prepare", "deduplicate")
        builder.add_edge("deduplicate", "version_resolve")
        builder.add_edge("version_resolve", "rubric_lock")
        builder.add_edge("rubric_lock", "batch_map")
        builder.add_edge("batch_map", "reduce")
        builder.add_edge("reduce", "fairness_guard")
        builder.add_edge("fairness_guard", "evidence_localize")
        builder.add_edge("evidence_localize", "recruitment_harness")
        builder.add_edge("recruitment_harness", "quality_decision")
        builder.add_conditional_edges(
            "quality_decision",
            self._after_quality,
            {"repair_once": "repair_once", "artifact_create": "artifact_create"},
        )
        builder.add_edge("repair_once", "fairness_guard")
        builder.add_edge("artifact_create", "finalize")
        builder.add_edge("finalize", END)
        return builder.compile(name="unihr-recruitment-agent")

    def _stage_trace(
        self,
        state: _RecruitmentState,
        stage_id: str,
        payload: dict[str, Any],
        *,
        status: str = "completed",
        summary: str,
        output_ref: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        """Persist one real recruitment phase behind a compact stage reference."""
        resolved_ref = dict(output_ref or {})
        scope = dict(state.get("payload_scope") or {})
        if not resolved_ref and self.payload_store is not None and scope:
            reference = self.payload_store.put_json(
                RunPointer.model_validate(scope), payload,
                kind=f"professional.recruitment.stage.{stage_id}",
            )
            resolved_ref = reference.model_dump(mode="json")
        return _trace(
            "recruitment.stage.completed", status, summary,
            stage_id=stage_id, output_ref=resolved_ref,
        )

    def _prepare(self, state: _RecruitmentState) -> dict[str, Any]:
        request = RecruitmentRequest.model_validate(state["request"])
        emit_tool_progress(
            "tool.recruitment.parse.started",
            "正在读取并解析候选人材料",
            attachment_count=len(request.attachments),
        )
        sources = [
            item for item in request.attachments
            if item.extension in _RESUME_EXTENSIONS and item.extracted_text.strip()
        ]
        if not sources:
            emit_tool_progress(
                "tool.recruitment.parse.failed",
                "没有发现可解析的候选人材料",
                status="failed",
                error_code="resume_files_required",
            )
            raise ValueError("resume_files_required")
        limited_count = sum(len(item.extracted_text.strip()) < 200 for item in sources)
        emit_tool_progress(
            "tool.recruitment.parse.completed",
            f"已读取 {len(sources)} 份候选人材料",
            status="completed",
            source_count=len(sources),
            limited_source_count=limited_count,
        )
        source_payload = [item.model_dump(mode="json") for item in sources]
        return {
            "sources": source_payload,
            "graph_path": ["prepare"],
            "trace": [
                _trace("resume.parse", "completed", f"读取 {len(sources)} 份可解析候选人材料"),
                self._stage_trace(
                    state, "prepare", {"sources": source_payload},
                    summary=f"已读取 {len(sources)} 份可解析候选人材料",
                ),
            ],
        }

    def _deduplicate(self, state: _RecruitmentState) -> dict[str, Any]:
        sources = [AttachmentRecord.model_validate(item) for item in state["sources"]]
        emit_tool_progress(
            "tool.recruitment.deduplicate.started",
            "正在识别完全重复的候选人材料",
            source_count=len(sources),
        )
        first_by_fingerprint: dict[str, AttachmentRecord] = {}
        duplicate_map: dict[str, list[str]] = {}
        profiles: list[ResumeSourceProfile] = []
        representatives: list[AttachmentRecord] = []
        for source in sources:
            fingerprint = _content_fingerprint(source.extracted_text)
            representative = first_by_fingerprint.get(fingerprint)
            duplicate_of = representative.display_name if representative else ""
            profiles.append(ResumeSourceProfile(
                source_id=source.id,
                source_name=source.display_name,
                content_fingerprint=fingerprint,
                extracted_chars=len(source.extracted_text),
                parse_status="parsed" if len(source.extracted_text.strip()) >= 200 else "limited",
                duplicate_of=duplicate_of,
            ))
            if representative is None:
                first_by_fingerprint[fingerprint] = source
                representatives.append(source)
            else:
                duplicate_map.setdefault(representative.display_name, []).append(source.display_name)
        groups = [
            DuplicateSourceGroup(representative_source=name, duplicate_sources=duplicates)
            for name, duplicates in duplicate_map.items()
        ]
        duplicate_count = len(sources) - len(representatives)
        emit_tool_progress(
            "tool.recruitment.deduplicate.completed",
            f"去重完成，保留 {len(representatives)} 份唯一材料，识别 {duplicate_count} 份重复材料",
            status="completed",
            source_count=len(sources),
            unique_source_count=len(representatives),
            duplicate_source_count=duplicate_count,
        )
        profiles_payload = [item.model_dump(mode="json") for item in profiles]
        groups_payload = [item.model_dump(mode="json") for item in groups]
        representatives_payload = [item.model_dump(mode="json") for item in representatives]
        return {
            "source_profiles": profiles_payload,
            "duplicate_groups": groups_payload,
            "representative_sources": representatives_payload,
            "graph_path": ["deduplicate"],
            "trace": [
                _trace(
                    "resume.deduplicate", "completed",
                    f"保留 {len(representatives)} 份唯一材料，识别 {len(sources) - len(representatives)} 份完全重复材料",
                ),
                self._stage_trace(
                    state, "deduplicate",
                    {
                        "source_profiles": profiles_payload,
                        "duplicate_groups": groups_payload,
                        "representative_sources": representatives_payload,
                    },
                    summary=f"识别 {duplicate_count} 份完全重复材料，保留 {len(representatives)} 份唯一材料",
                ),
            ],
        }

    def _version_resolve(self, state: _RecruitmentState) -> dict[str, Any]:
        sources = [
            AttachmentRecord.model_validate(item)
            for item in state["representative_sources"]
        ]
        profiles = [
            ResumeSourceProfile.model_validate(item)
            for item in state["source_profiles"]
        ]
        emit_tool_progress(
            "tool.recruitment.version.started",
            "正在识别同一候选人的不同简历版本",
            unique_source_count=len(sources),
        )
        resolution = self.version_resolver.resolve(sources, profiles)
        merged = sum(
            max(0, len(group.version_sources) - 1)
            for group in resolution.groups if group.status == "merged"
        )
        review_groups = sum(group.status == "review_required" for group in resolution.groups)
        summary = (
            f"归并 {merged} 份同一候选人的旧版本，保留 {len(resolution.representatives)} 份代表材料"
        )
        if review_groups:
            summary += f"；{review_groups} 组疑似版本关系待人工确认"
        emit_tool_progress(
            "tool.recruitment.version.completed",
            summary,
            status="waiting" if review_groups else "completed",
            representative_count=len(resolution.representatives),
            merged_source_count=merged,
            review_group_count=review_groups,
        )
        representatives_payload = [
            item.model_dump(mode="json") for item in resolution.representatives
        ]
        profiles_payload = [
            item.model_dump(mode="json") for item in resolution.source_profiles
        ]
        groups_payload = [item.model_dump(mode="json") for item in resolution.groups]
        return {
            "representative_sources": representatives_payload,
            "source_profiles": profiles_payload,
            "version_groups": groups_payload,
            "graph_path": ["version_resolve"],
            "trace": [
                _trace(
                    "resume.version.resolve",
                    "warning" if review_groups else "completed",
                    summary,
                    semantic_merged_sources=merged,
                    version_groups=len(resolution.groups),
                    review_groups=review_groups,
                ),
                self._stage_trace(
                    state, "version_resolve",
                    {
                        "representative_sources": representatives_payload,
                        "source_profiles": profiles_payload,
                        "version_groups": groups_payload,
                    },
                    status="completed",
                    summary=summary,
                ),
            ],
        }

    def _rubric_lock(self, state: _RecruitmentState) -> dict[str, Any]:
        request = RecruitmentRequest.model_validate(state["request"])
        emit_tool_progress(
            "tool.recruitment.rubric.started",
            "正在根据岗位要求设计并锁定评分 Rubric",
        )
        system = (
            "你是高校人事招聘的岗位评分标准设计助手。只输出严格 JSON 对象，字段为 job_title 和 criteria。"
            "criteria 必须是3到8项，每项包含 id,label,weight,requirement，权重合计100。"
            "必须条件应获得最高权重；不得把性别、年龄、民族、婚育、照片、颜值、籍贯或住址作为标准。"
            "本步骤只设计Rubric，不评价候选人。"
        )
        reply = self.gateway.complete(
            [{"role": "system", "content": system}, {"role": "user", "content": _rubric_prompt(request)}],
            purpose="recruitment_rubric",
            model=self.model,
            structured=True,
            max_tokens=2_000,
        )
        payload = parse_json_object(reply.content)
        criteria = _normalize_criteria(payload.get("criteria"))
        rubric_hash = _rubric_hash(criteria)
        job_title = _clean_text(request.job_title or str(payload.get("job_title") or "未命名岗位"), 100)
        emit_tool_progress(
            "tool.recruitment.rubric.completed",
            f"评分 Rubric 已锁定，共 {len(criteria)} 个岗位相关维度",
            status="completed",
            criterion_count=len(criteria),
            rubric_hash=rubric_hash[:12],
            latency_ms=reply.latency_ms,
        )
        criteria_payload = [item.model_dump(mode="json") for item in criteria]
        return {
            "job_title": job_title,
            "criteria": criteria_payload,
            "rubric_hash": rubric_hash,
            "rubric_candidate_payload": payload if isinstance(payload.get("candidates"), list) else {},
            "model_names": [reply.model],
            "latency_ms": int(state.get("latency_ms") or 0) + reply.latency_ms,
            "graph_path": ["rubric_lock"],
            "trace": [
                _trace(
                    "recruitment.rubric.locked", "completed",
                    f"锁定 {len(criteria)} 项评分维度，Rubric {rubric_hash[:10]}",
                    model=reply.model, latency_ms=reply.latency_ms, rubric_hash=rubric_hash,
                ),
                self._stage_trace(
                    state, "rubric_lock",
                    {"job_title": job_title, "criteria": criteria_payload, "rubric_hash": rubric_hash},
                    summary=f"已锁定 {len(criteria)} 项岗位相关评分维度",
                ),
            ],
        }

    def _batch_map(self, state: _RecruitmentState) -> dict[str, Any]:
        request = RecruitmentRequest.model_validate(state["request"])
        criteria = [ScreeningCriterion.model_validate(item) for item in state["criteria"]]
        sources = [AttachmentRecord.model_validate(item) for item in state["representative_sources"]]
        visual_evidence = self._hydrate_visual_evidence(state)
        payloads: list[dict[str, Any]] = []
        model_names: list[str] = []
        latency_ms = int(state.get("latency_ms") or 0)
        batch_size = request.map_batch_size
        batch_count = (len(sources) + batch_size - 1) // batch_size
        emit_tool_progress(
            "tool.recruitment.batch.started",
            f"正在把 {len(sources)} 份唯一材料分为 {batch_count} 个批次评分",
            candidate_count=len(sources),
            batch_count=batch_count,
            batch_size=batch_size,
        )
        compatibility_payload = dict(state.get("rubric_candidate_payload") or {})
        if compatibility_payload:
            payloads = [{
                    "allowed_sources": [item.display_name for item in sources],
                    "payload": compatibility_payload,
                }]
            offloaded = self._offload_candidate_payloads(state, payloads)
            emit_tool_progress(
                "tool.recruitment.batch.completed",
                f"已完成 {len(sources)} 份唯一材料的结构化评分",
                status="completed",
                candidate_count=len(sources),
                batch_count=1,
                compatibility_mode=True,
            )
            stage_ref = dict(offloaded.get("candidate_payloads_ref") or {})
            return {
                **offloaded,
                "latency_ms": latency_ms,
                "graph_path": ["batch_map"],
                "trace": [
                    *(offloaded.get("trace") or []),
                    _trace(
                        "recruitment.evaluate", "completed",
                        f"复用结构化响应评估 {len(sources)} 份唯一材料",
                        batch_size=len(sources), compatibility_mode=True,
                    ),
                    self._stage_trace(
                        state, "batch_map", {"candidate_payloads": payloads},
                        output_ref=stage_ref,
                        summary=f"已形成 {len(sources)} 份材料的结构化评分批次",
                    ),
                ],
            }
        for offset in range(0, len(sources), batch_size):
            batch = sources[offset:offset + batch_size]
            batch_index = offset // batch_size + 1
            emit_tool_progress(
                "tool.recruitment.batch_item.started",
                f"正在评分第 {batch_index}/{batch_count} 批候选人材料",
                batch_index=batch_index,
                batch_count=batch_count,
                candidate_count=len(batch),
            )
            reply = self.gateway.complete(
                [
                    {"role": "system", "content": _evaluation_system(criteria, request.skill_instructions)},
                    {"role": "user", "content": _batch_prompt(
                        request, batch, criteria, visual_evidence=visual_evidence,
                    )},
                ],
                purpose="recruitment_screening",
                model=self.model,
                structured=True,
                max_tokens=6_000,
            )
            payloads.append({
                "allowed_sources": [item.display_name for item in batch],
                "payload": parse_json_object(reply.content),
            })
            model_names.append(reply.model)
            latency_ms += reply.latency_ms
            emit_tool_progress(
                "tool.recruitment.batch_item.completed",
                f"第 {batch_index}/{batch_count} 批评分完成",
                status="completed",
                batch_index=batch_index,
                batch_count=batch_count,
                candidate_count=len(batch),
                latency_ms=reply.latency_ms,
            )
        offloaded = self._offload_candidate_payloads(state, payloads)
        emit_tool_progress(
            "tool.recruitment.batch.completed",
            f"全部 {batch_count} 个评分批次已完成",
            status="completed",
            candidate_count=len(sources),
            batch_count=batch_count,
            total_latency_ms=latency_ms,
            payload_offloaded=bool(offloaded.get("candidate_payloads_ref")),
        )
        stage_ref = dict(offloaded.get("candidate_payloads_ref") or {})
        return {
            **offloaded,
            "model_names": model_names,
            "latency_ms": latency_ms,
            "graph_path": ["batch_map"],
            "trace": [
                *(offloaded.get("trace") or []),
                _trace(
                    "recruitment.evaluate", "completed",
                    f"按 {len(payloads)} 个批次评估 {len(sources)} 份唯一材料",
                    batch_size=batch_size,
                ),
                self._stage_trace(
                    state, "batch_map", {"candidate_payloads": payloads},
                    output_ref=stage_ref,
                    summary=f"已按 {len(payloads)} 个批次完成 {len(sources)} 份材料评分",
                ),
            ],
        }

    def _offload_candidate_payloads(
        self, state: _RecruitmentState, payloads: list[dict[str, Any]]
    ) -> dict[str, Any]:
        scope_data = dict(state.get("payload_scope") or {})
        if self.payload_store is None or not scope_data:
            return {"candidate_payloads": payloads}
        reference = self.payload_store.put_json(
            RunPointer.model_validate(scope_data),
            payloads,
            kind="recruitment.candidate_batches",
        )
        return {
            "candidate_payloads": [],
            "candidate_payloads_ref": reference.model_dump(mode="json"),
            "trace": [_trace(
                "graph.payload.offloaded",
                "completed",
                "批量候选人评分原文已外置，Graph 仅保留引用",
                kind=reference.kind,
                byte_size=reference.byte_size,
                payload_version=reference.version,
            )],
        }

    def _prepare_visual_evidence(
        self,
        visual_evidence: list[dict[str, Any]],
        payload_scope: RunPointer | None,
    ) -> tuple[dict[str, Any], list[dict[str, Any]]]:
        if not visual_evidence:
            return {"visual_evidence": [], "visual_evidence_ref": {}}, []
        if self.payload_store is None or payload_scope is None:
            return {"visual_evidence": visual_evidence, "visual_evidence_ref": {}}, []
        reference = self.payload_store.put_json(
            payload_scope,
            visual_evidence,
            kind="recruitment.visual_evidence",
        )
        return {
            "visual_evidence": [],
            "visual_evidence_ref": reference.model_dump(mode="json"),
        }, [_trace(
            "graph.payload.offloaded",
            "completed",
            "招聘视觉证据已外置，Graph 仅保留受控引用",
            kind=reference.kind,
            byte_size=reference.byte_size,
            payload_version=reference.version,
        )]

    def _hydrate_visual_evidence(self, state: _RecruitmentState) -> list[dict[str, Any]]:
        payload = list(state.get("visual_evidence") or [])
        reference = dict(state.get("visual_evidence_ref") or {})
        scope = dict(state.get("payload_scope") or {})
        if payload or not reference or not scope or self.payload_store is None:
            return payload
        hydrated = self.payload_store.get_json(
            RunPointer.model_validate(scope), reference,
        )
        return [dict(item) for item in hydrated if isinstance(item, dict)]

    def _reduce(self, state: _RecruitmentState) -> dict[str, Any]:
        request = RecruitmentRequest.model_validate(state["request"])
        criteria = [ScreeningCriterion.model_validate(item) for item in state["criteria"]]
        sources = [AttachmentRecord.model_validate(item) for item in state["representative_sources"]]
        emit_tool_progress(
            "tool.recruitment.reduce.started",
            "正在合并各批评分并建立稳定排名输入",
            representative_count=len(sources),
        )
        raw_candidates: list[dict[str, Any]] = []
        rejected_sources = 0
        payloads = list(state.get("candidate_payloads") or [])
        hydrated_trace: list[dict[str, Any]] = []
        reference = dict(state.get("candidate_payloads_ref") or {})
        scope_data = dict(state.get("payload_scope") or {})
        if not payloads and reference and scope_data and self.payload_store is not None:
            payloads = list(self.payload_store.get_json(
                RunPointer.model_validate(scope_data), reference,
            ))
            hydrated_trace.append(_trace(
                "graph.payload.hydrated",
                "completed",
                "按引用恢复批量候选人评分原文",
                kind=reference.get("kind", ""),
                byte_size=int(reference.get("byte_size") or 0),
            ))
        for batch in payloads:
            allowed = list(batch.get("allowed_sources") or [])
            payload = dict(batch.get("payload") or {})
            items = payload.get("candidates") if isinstance(payload.get("candidates"), list) else []
            for item in items:
                if not isinstance(item, dict):
                    continue
                resolved = _resolve_source_name(str(item.get("source_name") or ""), allowed)
                if not resolved:
                    rejected_sources += 1
                    continue
                raw_candidates.append({**item, "source_name": resolved})
        report = _normalize_report(
            {"job_title": state.get("job_title"), "candidates": raw_candidates},
            request.job_requirements,
            sources,
            criteria=criteria,
            source_profiles=[ResumeSourceProfile.model_validate(item) for item in state["source_profiles"]],
            duplicate_groups=[DuplicateSourceGroup.model_validate(item) for item in state["duplicate_groups"]],
            version_groups=[ResumeVersionGroup.model_validate(item) for item in state.get("version_groups", [])],
        )
        emit_tool_progress(
            "tool.recruitment.reduce.completed",
            f"批次结果已合并，共形成 {len(report.candidates)} 位候选人的评分记录",
            status="completed",
            candidate_count=len(report.candidates),
            rejected_unknown_source_count=rejected_sources,
        )
        report_payload = report.model_dump(mode="json")
        return {
            "report": report_payload,
            "graph_path": ["reduce"],
            "trace": [
                *hydrated_trace,
                _trace(
                    "recruitment.reduce", "completed",
                    f"合并 {len(report.candidates)} 位候选人并拒绝 {rejected_sources} 个未知来源结果",
                    rejected_unknown_sources=rejected_sources,
                ),
                self._stage_trace(
                    state, "reduce", {"report": report_payload},
                    summary=f"已合并 {len(report.candidates)} 位候选人的稳定排名输入",
                ),
            ],
        }

    def _evidence_localize(self, state: _RecruitmentState) -> dict[str, Any]:
        report = RecruitmentReport.model_validate(state["report"])
        sources = [
            AttachmentRecord.model_validate(item)
            for item in state["representative_sources"]
        ]
        evidence_row_count = sum(len(item.criterion_scores) for item in report.candidates)
        emit_tool_progress(
            "tool.recruitment.evidence.started",
            "正在把评分依据定位回简历原文、页码或段落",
            evidence_row_count=evidence_row_count,
        )
        result = self.evidence_locator.localize(
            report,
            sources,
            self._hydrate_visual_evidence(state),
        )
        emit_tool_progress(
            "tool.recruitment.evidence.completed",
            f"证据定位完成：{result.located} 条已定位，{result.review_required} 条需核对",
            status="waiting" if result.review_required or result.unlocated else "completed",
            located_count=result.located,
            review_required_count=result.review_required,
            unlocated_count=result.unlocated,
        )
        report_payload = result.report.model_dump(mode="json")
        return {
            "report": report_payload,
            "graph_path": ["evidence_localize"],
            "trace": [
                _trace(
                    "recruitment.evidence.localized",
                    "completed" if not result.review_required else "warning",
                    f"定位 {result.located} 条原文证据，{result.review_required} 条需核对，{result.unlocated} 条未定位",
                    located=result.located,
                    review_required=result.review_required,
                    unlocated=result.unlocated,
                ),
                self._stage_trace(
                    state, "evidence_localize", {"report": report_payload},
                    status="completed",
                    summary=f"已定位 {result.located} 条原文证据，{result.review_required + result.unlocated} 条待核对",
                ),
            ],
        }

    def _fairness_guard(self, state: _RecruitmentState) -> dict[str, Any]:
        report = RecruitmentReport.model_validate(state["report"])
        emit_tool_progress(
            "tool.recruitment.fairness.started",
            "正在排除与岗位胜任力无关的受保护属性",
            candidate_count=len(report.candidates),
        )
        removed = 0
        cleaned: list[CandidateEvaluation] = []
        for candidate in report.candidates:
            rows: list[CriterionEvidence] = []
            for row in candidate.criterion_scores:
                row_update: dict[str, Any] = {}
                if (
                    _LOCATION_CRITERION_PATTERN.search(row.criterion_label)
                    and not _EXPLICIT_LOCATION_INTENT_PATTERN.search(row.evidence)
                ):
                    row_update = {
                        "score": 0,
                        "evidence": "未找到明确的工作地点意愿或可到岗承诺；学校所在地、现居地、籍贯和住址不作为评分依据",
                        "confidence": "low",
                    }
                    removed += 1
                evidence, changed = remove_protected_attributes(str(row_update.get("evidence") or row.evidence))
                removed += int(changed)
                rows.append(row.model_copy(update={**row_update, "evidence": evidence}))
            strengths, count = _sanitize_list(candidate.strengths)
            removed += count
            gaps, count = _sanitize_list(candidate.gaps)
            removed += count
            questions, count = _sanitize_list(candidate.verification_questions)
            removed += count
            total = round(sum(row.score * row.weight / 100 for row in rows), 1)
            recommendation = (
                "建议优先人工复核" if total >= 80
                else "建议重点核实后复核" if total >= 60
                else "证据不足，建议补充材料后复核"
            )
            evidence_coverage = round(sum(
                row.confidence != "low" and "未找到明确" not in row.evidence for row in rows
            ) / max(1, len(rows)), 4)
            cleaned.append(candidate.model_copy(update={
                "criterion_scores": rows,
                "strengths": strengths,
                "gaps": gaps,
                "verification_questions": questions,
                "total_score": total,
                "recommendation": recommendation,
                "evidence_coverage": evidence_coverage,
            }))
        cleaned = _stable_candidate_sort(cleaned)
        emit_tool_progress(
            "tool.recruitment.fairness.completed",
            f"公平性检查完成，清理 {removed} 处不应参与评分的表述",
            status="completed",
            candidate_count=len(cleaned),
            removed_mention_count=removed,
        )
        report_payload = report.model_copy(update={"candidates": cleaned}).model_dump(mode="json")
        return {
            "report": report_payload,
            "graph_path": ["fairness_guard"],
            "trace": [
                _trace(
                    "fairness.guard", "completed",
                    f"排除受保护属性，清理 {removed} 处不应参与评分的表述",
                    removed_mentions=removed,
                ),
                self._stage_trace(
                    state, "fairness_guard", {"report": report_payload, "removed_mentions": removed},
                    summary=f"已完成公平性检查并清理 {removed} 处不相关表述",
                ),
            ],
        }

    def _recruitment_harness(self, state: _RecruitmentState) -> dict[str, Any]:
        report = RecruitmentReport.model_validate(state["report"])
        emit_tool_progress(
            "tool.recruitment.harness.started",
            "正在检查材料覆盖、证据质量、公平性和排名稳定性",
            candidate_count=len(report.candidates),
        )
        harness = self.harness.inspect(report)
        report = report.model_copy(update={"harness": harness})
        emit_tool_progress(
            "tool.recruitment.harness.completed",
            f"招聘 Harness {harness.status}，质量得分 {harness.score:.1f}",
            status="completed" if harness.status != "failed" else "failed",
            harness_status=harness.status,
            harness_score=harness.score,
            issue_count=len(harness.issues),
            warning_count=len(harness.warnings),
            evidence_coverage=harness.metrics.get("evidence_coverage"),
            ranking_stability=harness.metrics.get("ranking_stability"),
            near_score_group_count=harness.metrics.get("near_score_groups"),
        )
        return {
            "report": report.model_dump(mode="json"),
            "harness": harness.model_dump(mode="json"),
            "graph_path": ["recruitment_harness"],
            "trace": [
                _trace(
                    "recruitment.harness", harness.status,
                    f"Recruitment Harness {harness.status}，得分 {harness.score:.1f}",
                    issues=harness.issues, warnings=harness.warnings, metrics=harness.metrics,
                ),
                self._stage_trace(
                    state, "harness", {"harness": harness.model_dump(mode="json")},
                    status=harness.status,
                    summary=f"招聘质量检查 {harness.status}，得分 {harness.score:.1f}",
                ),
            ],
        }

    def _quality_decision(self, state: _RecruitmentState) -> dict[str, Any]:
        harness = RecruitmentHarnessReport.model_validate(state["harness"])
        emit_tool_progress(
            "tool.recruitment.quality.completed",
            "招聘质量门通过" if harness.status != "failed" else "质量门要求一次确定性修复",
            status="completed" if harness.status != "failed" else "waiting",
            harness_status=harness.status,
            issue_count=len(harness.issues),
        )
        return {
            "graph_path": ["quality_decision"],
            "trace": [
                _trace(
                    "recruitment.quality_gate", "completed" if harness.status != "failed" else "warning",
                    "质量门通过" if harness.status != "failed" else "质量门要求一次确定性修复",
                ),
                self._stage_trace(
                    state, "quality_gate", {"harness": harness.model_dump(mode="json")},
                    status="completed" if harness.status != "failed" else "waiting",
                    summary="质量门通过" if harness.status != "failed" else "质量门要求一次确定性修复",
                ),
            ],
        }

    @staticmethod
    def _after_quality(state: _RecruitmentState) -> Literal["repair_once", "artifact_create"]:
        harness = RecruitmentHarnessReport.model_validate(state["harness"])
        if harness.status == "failed" and int(state.get("repair_count") or 0) < 1:
            return "repair_once"
        return "artifact_create"

    def _repair_once(self, state: _RecruitmentState) -> dict[str, Any]:
        report = RecruitmentReport.model_validate(state["report"])
        emit_tool_progress(
            "tool.recruitment.repair.started",
            "正在补齐漏评材料并移除未知来源结果",
            candidate_count=len(report.candidates),
        )
        allowed = {
            item.source_name for item in report.source_profiles
            if not item.duplicate_of and item.version_status != "superseded"
        }
        candidates = [item for item in report.candidates if item.source_name in allowed]
        covered = {item.source_name for item in candidates}
        for source_name in sorted(allowed - covered):
            candidates.append(_missing_candidate(source_name, report.criteria))
        candidates = _stable_candidate_sort(candidates)
        repair_count = int(state.get("repair_count") or 0) + 1
        emit_tool_progress(
            "tool.recruitment.repair.completed",
            "确定性修复完成，正在重新运行公平性与质量检查",
            status="completed",
            repair_count=repair_count,
            candidate_count=len(candidates),
        )
        return {
            "report": report.model_copy(update={"candidates": candidates}).model_dump(mode="json"),
            "repair_count": repair_count,
            "graph_path": ["repair_once"],
            "trace": [
                _trace("recruitment.repair", "completed", "补齐漏评材料并移除未知来源，重新执行质量门"),
                self._stage_trace(
                    state, "repair", {"report": report.model_copy(update={"candidates": candidates}).model_dump(mode="json")},
                    summary="已完成一次确定性修复并准备重新质量检查",
                ),
            ],
        }

    def _artifact_create(self, state: _RecruitmentState) -> dict[str, Any]:
        request = RecruitmentRequest.model_validate(state["request"])
        report = RecruitmentReport.model_validate(state["report"])
        emit_tool_progress(
            "tool.recruitment.report.started",
            "正在生成候选人排名、评分明细、证据索引和人工复核 Excel",
            candidate_count=len(report.candidates),
        )
        case_id = f"recruitment_case_{uuid4().hex}"
        case = RecruitmentCaseSnapshot(
            case_id=case_id,
            rubric_hash=str(state["rubric_hash"]),
            source_count=len(report.source_profiles),
            unique_source_count=sum(
                not item.duplicate_of and item.version_status != "superseded"
                for item in report.source_profiles
            ),
            duplicate_source_count=sum(
                bool(item.duplicate_of)
                for item in report.source_profiles
            ),
            selection_count=min(request.selection_count, max(1, len(report.candidates))),
            graph_path=[*state.get("graph_path", []), "artifact_create", "finalize"],
        )
        rubric = RubricSnapshot(
            rubric_hash=str(state["rubric_hash"]),
            criteria=report.criteria,
        )
        report = report.model_copy(update={"rubric": rubric, "case": case})
        path = self._build_workbook(report)
        emit_tool_progress(
            "tool.recruitment.report.completed",
            "七表招聘复核 Excel 已生成",
            status="completed",
            candidate_count=len(report.candidates),
            worksheet_count=7,
            selection_count=case.selection_count,
        )
        return {
            "report": report.model_dump(mode="json"),
            "artifact_path": str(path),
            "graph_path": ["artifact_create"],
            "trace": [
                _trace("artifact.create", "completed", f"生成七表 Excel：候选人筛选报告-{_safe_name(report.job_title)}.xlsx"),
                self._stage_trace(
                    state, "artifact_create",
                    {"artifact_path": str(path), "case": case.model_dump(mode="json")},
                    summary="已生成七表招聘复核 Excel 和招聘 Case 快照",
                ),
            ],
        }

    def _finalize(self, state: _RecruitmentState) -> dict[str, Any]:
        emit_tool_progress(
            "tool.recruitment.finalize.completed",
            "筛选结果已进入人工复核，不触发自动录用或淘汰",
            status="completed",
            repair_count=int(state.get("repair_count") or 0),
        )
        return {
            "graph_path": ["finalize"],
            "trace": [
                _trace("recruitment.finalize", "completed", "结果进入人工复核，不触发自动录用或淘汰"),
                self._stage_trace(
                    state, "finalize",
                    {
                        "artifact_path": str(state.get("artifact_path") or ""),
                        "case": dict(RecruitmentReport.model_validate(state["report"]).case.model_dump(mode="json"))
                        if RecruitmentReport.model_validate(state["report"]).case else {},
                    },
                    summary="已固定招聘报告、Case 快照和人工复核边界",
                ),
            ],
        }

    def _build_workbook(self, report: RecruitmentReport) -> Path:
        output = self.output_dir / f"recruitment_{uuid4().hex}.xlsx"
        workbook = Workbook()
        overview = workbook.active
        overview.title = "筛选说明"
        ranking = workbook.create_sheet("候选人排名")
        details = workbook.create_sheet("评分明细")
        duplicates = workbook.create_sheet("重复材料")
        versions = workbook.create_sheet("版本归并")
        review = workbook.create_sheet("人工复核")
        harness_sheet = workbook.create_sheet("Harness")

        overview.sheet_view.showGridLines = False
        overview.column_dimensions["A"].width = 24
        overview.column_dimensions["B"].width = 86
        overview.append(["UniHR 批量简历筛选报告", ""])
        overview.merge_cells("A1:B1")
        overview.append(["岗位", report.job_title])
        overview.append(["岗位要求", report.job_requirements[:4_000]])
        overview.append(["候选人数", len(report.candidates)])
        overview.append(["评分标准", "；".join(f"{item.label} {item.weight:.1f}%" for item in report.criteria)])
        overview.append(["公平性与责任", report.fairness_note])
        overview.append(["生成时间", report.generated_at])
        overview.append(["使用说明", "请逐条核对评分证据；排名相近或证据不足的候选人必须进一步人工复核。"])
        overview.append(["材料统计", f"共 {len(report.source_profiles)} 份；唯一材料 {report.case.unique_source_count if report.case else len(report.source_profiles)} 份；重复 {report.case.duplicate_source_count if report.case else 0} 份"])
        overview.append(["Rubric 版本", f"{report.rubric.version} · {report.rubric.rubric_hash}" if report.rubric else "未记录"])
        _style_overview(overview)

        ranking.append(["人工复核排名", "候选人", "来源文件", "综合分", "复核建议", "主要优势", "证据缺口", "待核实问题"])
        for index, candidate in enumerate(report.candidates, start=1):
            ranking.append([
                index,
                candidate.candidate_name,
                candidate.source_name,
                candidate.total_score,
                candidate.recommendation,
                "；".join(candidate.strengths),
                "；".join(candidate.gaps),
                "；".join(candidate.verification_questions),
            ])
        _style_table_sheet(ranking, [13, 18, 30, 12, 22, 42, 42, 46], "CandidateRanking")
        for row in ranking.iter_rows(min_row=2, min_col=4, max_col=4):
            row[0].number_format = "0.0"

        details.append([
            "候选人", "评分维度", "权重", "维度得分", "加权得分", "简历证据", "证据置信度",
            "原文位置", "定位原文", "定位状态", "复核原因",
        ])
        row_index = 2
        for candidate in report.candidates:
            for score in candidate.criterion_scores:
                details.append([
                    candidate.candidate_name,
                    score.criterion_label,
                    score.weight / 100,
                    score.score,
                    f"=C{row_index}*D{row_index}",
                    score.evidence,
                    score.confidence,
                    _locator_label(score.locator),
                    score.locator.quote if score.locator else "",
                    score.locator.status if score.locator else "unlocated",
                    "；".join(score.locator.review_reasons) if score.locator else "evidence_not_localized",
                ])
                row_index += 1
        _style_table_sheet(details, [18, 24, 12, 14, 14, 58, 16, 34, 58, 18, 36], "ScoreDetails")
        for row in details.iter_rows(min_row=2, min_col=3, max_col=5):
            row[0].number_format = "0.0%"
            row[1].number_format = "0.0"
            row[2].number_format = "0.0"

        duplicates.append(["代表材料", "重复材料", "处理方式"])
        for group in report.duplicate_groups:
            for duplicate in group.duplicate_sources:
                duplicates.append([group.representative_source, duplicate, "仅保留代表材料参与评分，重复文件不重复排名"])
        if not report.duplicate_groups:
            duplicates.append(["—", "—", "未发现内容完全重复的材料"])
        _style_table_sheet(duplicates, [38, 38, 56], "DuplicateSources")

        versions.append(["版本组", "处理状态", "代表材料", "组内材料", "判定依据", "相似度", "处理说明"])
        for group in report.version_groups:
            versions.append([
                group.group_id,
                "已自动归并" if group.status == "merged" else "待人工确认",
                group.representative_source or "—",
                "；".join(group.version_sources),
                group.identity_basis,
                "；".join(f"{key}: {value:.1%}" for key, value in group.similarities.items()),
                group.selection_reason,
            ])
        if not report.version_groups:
            versions.append(["—", "未发现", "—", "—", "—", "—", "未发现同一候选人的多版本关系"])
        _style_table_sheet(versions, [26, 16, 38, 62, 30, 58, 62], "ResumeVersions")

        review.append(["人工复核排名", "候选人", "来源文件", "AI综合分", "复核结论", "复核意见", "复核人", "复核时间"])
        for index, candidate in enumerate(report.candidates, start=1):
            review.append([index, candidate.candidate_name, candidate.source_name, candidate.total_score, "待复核", "", "", ""])
        _style_table_sheet(review, [13, 18, 34, 12, 18, 48, 18, 22], "HumanReview")
        for row in review.iter_rows(min_row=2, min_col=4, max_col=4):
            row[0].number_format = "0.0"

        harness_sheet.append(["Recruitment Harness", "结果"])
        if report.harness:
            harness_sheet.append(["状态", report.harness.status])
            harness_sheet.append(["得分", report.harness.score])
            harness_sheet.append(["问题", "；".join(report.harness.issues) or "无"])
            harness_sheet.append(["警告", "；".join(report.harness.warnings) or "无"])
            for key, value in report.harness.metrics.items():
                harness_sheet.append([key, value])
        else:
            harness_sheet.append(["状态", "未运行"])
        _style_overview(harness_sheet)

        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.save(output)
        workbook.close()
        return output


def _stable_candidate_sort(candidates: list[CandidateEvaluation]) -> list[CandidateEvaluation]:
    """Produce a reproducible order without treating a model's return order as a tie-breaker."""

    return sorted(
        candidates,
        key=lambda item: (
            -item.total_score,
            -item.evidence_coverage,
            item.candidate_name.casefold(),
            item.source_name.casefold(),
        ),
    )


def _locator_label(locator: EvidenceLocator | None) -> str:
    if locator is None or locator.status == "unlocated":
        return "未定位"
    parts = [locator.source_name]
    if locator.page_number:
        parts.append(f"第 {locator.page_number} 页")
    if locator.paragraph_number:
        parts.append(f"段落 {locator.paragraph_number}")
    if locator.table_number:
        parts.append(f"表格 {locator.table_number}")
    if locator.sheet_name:
        parts.append(f"Sheet {locator.sheet_name}")
    if locator.bbox:
        parts.append("bbox=" + ",".join(str(value) for value in locator.bbox))
    return " · ".join(part for part in parts if part)


def _normalize_report(
    payload: dict[str, Any],
    job_requirements: str,
    sources: list[AttachmentRecord],
    *,
    criteria: list[ScreeningCriterion] | None = None,
    source_profiles: list[ResumeSourceProfile] | None = None,
    duplicate_groups: list[DuplicateSourceGroup] | None = None,
    version_groups: list[ResumeVersionGroup] | None = None,
) -> RecruitmentReport:
    criteria = criteria or _normalize_criteria(payload.get("criteria"))
    raw_candidates = payload.get("candidates") if isinstance(payload.get("candidates"), list) else []
    candidates: list[CandidateEvaluation] = []
    allowed_sources = {item.display_name for item in sources}
    seen_candidates: set[tuple[str, str]] = set()
    for index, item in enumerate(raw_candidates):
        if not isinstance(item, dict):
            continue
        source_name = str(item.get("source_name") or "").strip()
        if not source_name:
            source_name = sources[min(index, len(sources) - 1)].display_name
        if source_name not in allowed_sources:
            continue
        candidate_name = str(item.get("candidate_name") or Path(source_name).stem or f"候选人{index + 1}").strip()
        candidate_key = (source_name.casefold(), candidate_name.casefold())
        if candidate_key in seen_candidates:
            continue
        seen_candidates.add(candidate_key)
        raw_scores = item.get("criterion_scores") if isinstance(item.get("criterion_scores"), dict) else {}
        raw_evidence = item.get("score_evidence") if isinstance(item.get("score_evidence"), dict) else {}
        evidence_rows: list[CriterionEvidence] = []
        for criterion in criteria:
            score_value = raw_scores.get(criterion.id, raw_scores.get(criterion.label, 0))
            if isinstance(score_value, dict):
                score = _number(score_value.get("score"), 0)
                evidence = str(score_value.get("evidence") or raw_evidence.get(criterion.id) or "未找到明确证据")
                confidence = str(score_value.get("confidence") or "medium")
            else:
                score = _number(score_value, 0)
                evidence = str(raw_evidence.get(criterion.id) or raw_evidence.get(criterion.label) or "未找到明确证据")
                confidence = "low" if "未找到" in evidence else "medium"
            evidence_rows.append(CriterionEvidence(
                criterion_id=criterion.id,
                criterion_label=criterion.label,
                score=max(0, min(100, score)),
                weight=criterion.weight,
                evidence=_clean_text(evidence, 240),
                confidence=confidence if confidence in {"low", "medium", "high"} else "medium",
            ))
        total = round(sum(row.score * row.weight / 100 for row in evidence_rows), 1)
        recommendation = (
            "建议优先人工复核" if total >= 80
            else "建议重点核实后复核" if total >= 60
            else "证据不足，建议补充材料后复核"
        )
        evidence_coverage = round(sum(row.confidence != "low" and "未找到明确证据" not in row.evidence for row in evidence_rows) / max(1, len(evidence_rows)), 4)
        candidates.append(CandidateEvaluation(
            candidate_name=_clean_text(candidate_name, 80),
            source_name=_clean_text(source_name, 160),
            total_score=total,
            criterion_scores=evidence_rows,
            strengths=_string_list(item.get("strengths"), 5),
            gaps=_string_list(item.get("gaps"), 5),
            verification_questions=_string_list(item.get("verification_questions"), 5),
            recommendation=recommendation,
            evidence_coverage=evidence_coverage,
        ))
    covered_sources = {item.source_name for item in candidates}
    for source in sources:
        if source.display_name not in covered_sources:
            candidates.append(_missing_candidate(source.display_name, criteria))
    candidates.sort(key=lambda item: (-item.total_score, item.candidate_name.casefold()))
    return RecruitmentReport(
        job_title=_clean_text(str(payload.get("job_title") or "未命名岗位"), 100),
        job_requirements=job_requirements.strip()[:8_000],
        criteria=criteria,
        candidates=candidates,
        fairness_note=_FAIRNESS_NOTE,
        generated_at=datetime.now(timezone.utc).isoformat(),
        source_profiles=source_profiles or [
            ResumeSourceProfile(
                source_id=item.id,
                source_name=item.display_name,
                content_fingerprint=_content_fingerprint(item.extracted_text),
                extracted_chars=len(item.extracted_text),
                parse_status="parsed" if len(item.extracted_text.strip()) >= 200 else "limited",
            ) for item in sources
        ],
        duplicate_groups=duplicate_groups or [],
        version_groups=version_groups or [],
    )


def _normalize_criteria(raw: Any) -> list[ScreeningCriterion]:
    items = raw if isinstance(raw, list) else []
    prepared: list[tuple[str, str, float, str]] = []
    for index, item in enumerate(items[:8], start=1):
        if not isinstance(item, dict):
            continue
        label = _clean_text(str(item.get("label") or ""), 60)
        if not label:
            continue
        prepared.append((
            _clean_text(str(item.get("id") or f"c{index}"), 32),
            label,
            max(0, _number(item.get("weight"), 0)),
            _clean_text(str(item.get("requirement") or ""), 240),
        ))
    if not prepared:
        prepared = [
            ("essential", "必备条件匹配", 40, "用户明确列出的必要条件"),
            ("experience", "相关经历", 30, "与岗位职责直接相关的项目或工作经历"),
            ("skills", "专业能力", 20, "岗位所需技能及可验证成果"),
            ("evidence", "证据完整性", 10, "材料是否包含可核验的事实依据"),
        ]
    total = sum(item[2] for item in prepared) or len(prepared)
    return [ScreeningCriterion(
        id=item[0] or f"c{index + 1}",
        label=item[1],
        weight=round((item[2] or 1) * 100 / total, 2),
        requirement=item[3],
    ) for index, item in enumerate(prepared)]


def _render_answer(report: RecruitmentReport) -> str:
    duplicate_count = sum(len(item.duplicate_sources) for item in report.duplicate_groups)
    semantic_merged_count = sum(
        max(0, len(item.version_sources) - 1)
        for item in report.version_groups if item.status == "merged"
    )
    version_review_count = sum(item.status == "review_required" for item in report.version_groups)
    harness_text = (
        f"Recruitment Harness：{report.harness.status}（{report.harness.score:.1f}/100）"
        if report.harness else "Recruitment Harness：未运行"
    )
    ranking_stability_text = _ranking_stability_message(report)
    evidence_localization_text = _evidence_localization_message(report)
    lines = [
        f"已完成 **{report.job_title}** 的批量简历辅助筛选，共分析 {len(report.candidates)} 位候选人。"
        + (f"另识别 {duplicate_count} 份完全重复材料，未重复计入排名。" if duplicate_count else "")
        + (f"同一候选人的 {semantic_merged_count} 份旧版本已归并，不重复占用排名名额。" if semantic_merged_count else "")
        + (f"另有 {version_review_count} 组疑似版本关系证据不足，已保留并标记人工确认。" if version_review_count else ""),
        "",
        "| 人工复核排名 | 候选人 | 综合分 | 建议 | 主要证据 |",
        "|---:|---|---:|---|---|",
    ]
    for index, candidate in enumerate(report.candidates[:15], start=1):
        evidence = next((item.evidence for item in candidate.criterion_scores if "未找到" not in item.evidence), "证据待补充")
        lines.append(
            f"| {index} | {_escape_markdown(candidate.candidate_name)} | {candidate.total_score:.1f} | "
            f"{_escape_markdown(candidate.recommendation)} | {_escape_markdown(evidence[:90])} |"
        )
    lines.extend([
        "",
        "评分维度：" + "、".join(f"{item.label}（{item.weight:.1f}%）" for item in report.criteria),
        harness_text,
        ranking_stability_text,
        evidence_localization_text,
        "",
        f"> {report.fairness_note}",
        "",
        "完整的逐项证据、重复材料、版本归并、人工复核表和 Harness 结果已写入 Excel 报告。",
    ])
    return "\n".join(lines)


def _ranking_stability_message(report: RecruitmentReport) -> str:
    """Expose deterministic ranking checks as a decision aid, never an auto-decision."""
    metrics = report.harness.metrics if report.harness else {}
    near_score_groups = int(metrics.get("near_score_groups", 0) or 0)
    tied_candidates = int(metrics.get("tied_candidates", 0) or 0)
    if near_score_groups or tied_candidates:
        details = []
        if near_score_groups:
            details.append(f"{near_score_groups} 组近分")
        if tied_candidates:
            details.append(f"{tied_candidates} 位并列候选人")
        return "排名稳定性：存在" + "、".join(details) + "，建议逐项人工复核后再确认最终人选。"
    return "排名稳定性：结构稳定；仍请结合原始证据与业务判断完成人工确认。"


def _evidence_localization_message(report: RecruitmentReport) -> str:
    metrics = report.harness.metrics if report.harness else {}
    located = int(metrics.get("located_evidence_rows", 0) or 0)
    review = int(metrics.get("review_evidence_rows", 0) or 0)
    unlocated = int(metrics.get("unlocated_evidence_rows", 0) or 0)
    return (
        f"原文定位：{located} 条已定位，{review} 条需核对，{unlocated} 条未定位；"
        "只有可回到原始材料的证据才可作为人工复核依据。"
    )


def _rubric_prompt(request: RecruitmentRequest) -> str:
    return "\n".join([
        f"岗位名称：{request.job_title or '请根据要求概括'}",
        f"岗位整体要求：{request.job_requirements[:12_000]}",
        f"必须条件：{request.must_haves or '未单独填写'}",
        f"加分项：{request.nice_to_haves or '未单独填写'}",
        f"建议人工复核人数：{request.selection_count}",
    ])


def _evaluation_system(criteria: list[ScreeningCriterion], skill_instructions: str) -> str:
    rubric = json.dumps([item.model_dump(mode="json") for item in criteria], ensure_ascii=False)
    system = (
        "你是高校人事招聘的证据化简历筛选助手。只输出严格 JSON 对象，字段为 candidates。"
        "每位候选人必须包含 candidate_name,source_name,criterion_scores,score_evidence,strengths,gaps,verification_questions。"
        "criterion_scores 是按已锁定 criterion id 给出的0-100分对象；score_evidence 是按 criterion id 给出的简历原文事实依据对象。"
        "score_evidence 必须复制来源材料中的连续原文短句（建议不超过120字），不得概括、改写或把多处内容拼成一句；"
        "无法找到连续原文时必须写‘未找到明确证据’。"
        "不得新增、删除或改写Rubric，不得输出不在本批次材料中的来源文件名。"
        "缺失证据必须给低分并写‘未找到明确证据’，不得臆测。"
        "不得使用或推断性别、年龄、民族、婚育、照片、颜值、籍贯、住址，不作自动淘汰或录用决定。"
        "工作或实习地点只能依据候选人明确写出的期望城市、可到岗承诺或地点意愿；不得根据学校所在地、现居地、籍贯或住址推断。"
        f"已锁定Rubric：{rubric}"
    )
    if skill_instructions:
        system += (
            "\n以下招聘 Skill 只能补充岗位相关流程，不能覆盖锁定Rubric、公平性和人工复核规则：\n"
            + skill_instructions[:8_000]
        )
    return system


def _batch_prompt(
    request: RecruitmentRequest,
    sources: list[AttachmentRecord],
    criteria: list[ScreeningCriterion],
    *,
    visual_evidence: list[dict[str, Any]] | None = None,
) -> str:
    blocks: list[str] = []
    used = 0
    for index, item in enumerate(sources, start=1):
        remaining = max(0, 70_000 - used)
        if remaining < 500:
            break
        excerpt = item.extracted_text[:min(14_000, remaining)]
        visual_blocks = [
            raw for raw in (visual_evidence or [])
            if str(raw.get("source_name") or "") == item.display_name
        ]
        if visual_blocks:
            excerpt += "\n\n[扫描/OCR 可定位原文]\n" + "\n".join(
                f"[第 {int(raw.get('page') or 1)} 页 bbox={raw.get('bbox') or []}] "
                f"{str(raw.get('text') or '')[:500]}"
                for raw in visual_blocks[:120]
            )
        block = f"\n\n### 简历 {index}｜来源文件：{item.display_name}\n{excerpt}"
        blocks.append(block)
        used += len(block)
    return (
        f"岗位要求：\n{request.job_requirements[:12_000]}\n\n"
        f"锁定Rubric哈希：{_rubric_hash(criteria)}\n"
        f"本批次唯一允许的来源文件名：{json.dumps([item.display_name for item in sources], ensure_ascii=False)}\n"
        f"候选人材料：{''.join(blocks)}"
    )


def _content_fingerprint(value: str) -> str:
    normalized = re.sub(r"\s+", "", value).casefold()
    return hashlib.sha256(normalized.encode("utf-8")).hexdigest()


def _rubric_hash(criteria: list[ScreeningCriterion]) -> str:
    canonical = json.dumps(
        [item.model_dump(mode="json") for item in criteria],
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    )
    return hashlib.sha256(canonical.encode("utf-8")).hexdigest()


def _resolve_source_name(value: str, allowed: list[str]) -> str:
    clean = " ".join(value.split())
    if clean in allowed:
        return clean
    by_casefold = {item.casefold(): item for item in allowed}
    if clean.casefold() in by_casefold:
        return by_casefold[clean.casefold()]
    stem = Path(clean).stem.casefold()
    matches = [item for item in allowed if Path(item).stem.casefold() == stem]
    if len(matches) == 1:
        return matches[0]
    return allowed[0] if not clean and len(allowed) == 1 else ""


def _missing_candidate(source_name: str, criteria: list[ScreeningCriterion]) -> CandidateEvaluation:
    return CandidateEvaluation(
        candidate_name=Path(source_name).stem,
        source_name=source_name,
        total_score=0,
        criterion_scores=[CriterionEvidence(
            criterion_id=criterion.id,
            criterion_label=criterion.label,
            score=0,
            weight=criterion.weight,
            evidence="模型未返回可核验评分，请人工复核原始材料",
            confidence="low",
        ) for criterion in criteria],
        gaps=["缺少可解析的结构化评分结果"],
        verification_questions=["请人工核对原始简历并补充岗位证据"],
        recommendation="证据不足，建议补充材料后复核",
        evidence_coverage=0,
    )


def _sanitize_list(values: list[str]) -> tuple[list[str], int]:
    cleaned: list[str] = []
    removed = 0
    for value in values:
        safe, changed = remove_protected_attributes(value)
        cleaned.append(safe)
        removed += int(changed)
    return cleaned, removed


def _trace(event: str, status: str, summary: str, **details: Any) -> dict[str, Any]:
    return {"event": event, "status": status, "summary": summary, **details}


def _style_overview(sheet: Any) -> None:
    sheet["A1"].font = Font(name="微软雅黑", size=18, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="2563EB")
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 34
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, min_col=1, max_col=2):
        row[0].font = Font(name="微软雅黑", size=11, bold=True, color="1F2937")
        row[0].fill = PatternFill("solid", fgColor="EFF6FF")
        for cell in row:
            cell.font = Font(name="微软雅黑", size=11, bold=cell.column == 1)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=Side(style="thin", color="DDE3EA"))
        sheet.row_dimensions[row[0].row].height = 30
    sheet.freeze_panes = "A2"


def _style_table_sheet(sheet: Any, widths: list[int], table_name: str) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    for cell in sheet[1]:
        cell.font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 30
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
        for cell in row:
            cell.font = Font(name="微软雅黑", size=10, color="111827")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.row_dimensions[row[0].row].height = 36
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    if sheet.max_row > 1:
        table = Table(displayName=table_name, ref=f"A1:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False)
        sheet.add_table(table)


def _number(value: Any, default: float) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _string_list(value: Any, limit: int) -> list[str]:
    if not isinstance(value, list):
        return []
    return [_clean_text(str(item), 180) for item in value[:limit] if str(item).strip()]


def _clean_text(value: str, limit: int) -> str:
    return " ".join(value.split())[:limit]


def _safe_name(value: str) -> str:
    clean = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", value).strip(" .")
    return (clean or "未命名岗位")[:60]


def _candidate_id(case_id: str, source_name: str, candidate_name: str) -> str:
    digest = hashlib.sha256(f"{case_id}|{source_name}|{candidate_name}".encode("utf-8")).hexdigest()
    return f"candidate_{digest[:20]}"


def _escape_markdown(value: str) -> str:
    return value.replace("|", "\\|").replace("\n", " ")
