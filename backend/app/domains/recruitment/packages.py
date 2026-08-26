from __future__ import annotations

import json
import re
import sqlite3
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterator
from uuid import uuid4
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from app.domains.artifacts.harness import ArtifactHarness, attach_harness_report
from app.domains.conversations.contracts import ArtifactRecord
from app.domains.conversations.repository import ConversationRepository
from app.domains.documents.word_verification import WordAutomationVerifier
from app.domains.identity import Principal, RUN_APPROVE

from .case_management import RecruitmentCaseModule
from .contracts import (
    CandidateReviewRecord,
    RecruitmentCaseRecord,
    RecruitmentPackageFile,
    RecruitmentPackagePreview,
    RecruitmentPackageRecord,
    RecruitmentPackageType,
)


DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PACKAGE_TITLES = {
    "interview_evaluation": "结构化面试评价包",
    "final_decision": "最终人工决策证据包",
}
COMPLIANCE_NOTE = (
    "本材料仅用于辅助人工复核，不构成录用、淘汰或自动化就业决定。"
    "不得使用性别、年龄、民族、婚育、照片、籍贯、住址等与岗位胜任力无关的属性。"
)


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _safe_name(value: str) -> str:
    clean = re.sub(r'[\\/:*?"<>|]+', "-", value).strip().strip(".")
    return clean[:80] or "招聘Case"


class RecruitmentEvidencePackageModule:
    """Build, verify and persist governed recruitment evidence packages."""

    def __init__(
        self,
        database_path: str | Path,
        output_dir: str | Path,
        cases: RecruitmentCaseModule,
        repository: ConversationRepository,
        *,
        artifact_harness: ArtifactHarness | None = None,
        word_verifier: WordAutomationVerifier | None = None,
    ) -> None:
        self.database_path = Path(database_path)
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.cases = cases
        self.repository = repository
        self.artifact_harness = artifact_harness or ArtifactHarness()
        self.word_verifier = word_verifier
        self._initialize()

    @contextmanager
    def _connect(self) -> Iterator[sqlite3.Connection]:
        connection = sqlite3.connect(self.database_path)
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA foreign_keys = ON")
        try:
            yield connection
            connection.commit()
        finally:
            connection.close()

    def _initialize(self) -> None:
        with self._connect() as db:
            db.executescript(
                """
                CREATE TABLE IF NOT EXISTS recruitment_case_packages (
                    id TEXT PRIMARY KEY,
                    case_id TEXT NOT NULL REFERENCES recruitment_cases(id) ON DELETE CASCADE,
                    tenant_id TEXT NOT NULL,
                    owner_user_id TEXT NOT NULL,
                    package_type TEXT NOT NULL,
                    title TEXT NOT NULL,
                    files_json TEXT NOT NULL DEFAULT '[]',
                    harness_json TEXT NOT NULL DEFAULT '{}',
                    paths_json TEXT NOT NULL DEFAULT '{}',
                    created_by TEXT NOT NULL,
                    created_at TEXT NOT NULL
                );
                CREATE INDEX IF NOT EXISTS idx_recruitment_packages_case_created
                    ON recruitment_case_packages(case_id, created_at DESC);
                """
            )

    def preview(
        self, case_id: str, package_type: RecruitmentPackageType, principal: Principal,
    ) -> RecruitmentPackagePreview:
        case = self.cases.get(case_id, principal)
        blockers: list[str] = []
        can_manage = case.owner_user_id == principal.user_id or principal.can(RUN_APPROVE)
        if not can_manage:
            blockers.append("recruitment_case_owner_required")
        if package_type == "final_decision" and case.status != "finalized":
            blockers.append("recruitment_case_not_finalized")
        if package_type == "interview_evaluation" and not case.candidates:
            blockers.append("recruitment_candidates_required")
        feedback_count = sum(len(session.feedback) for session in case.interview_sessions)
        sections = ["Case 与岗位摘要", "候选人证据总览", "结构化核实问题", "面试评价表", "合规说明"]
        if package_type == "final_decision":
            sections = ["最终人工建议", "候选人证据总览", "多人面试反馈", "风险与待核实项", "审计信息", "合规说明"]
        return RecruitmentPackagePreview(
            case_id=case.id,
            package_type=package_type,
            title=PACKAGE_TITLES[package_type],
            can_generate=not blockers,
            blockers=blockers,
            sections=sections,
            candidate_count=len(case.candidates),
            interview_session_count=len(case.interview_sessions),
            feedback_count=feedback_count,
            final_candidate_count=len(case.final_candidate_ids),
        )

    def list(self, case_id: str, principal: Principal) -> list[RecruitmentPackageRecord]:
        self.cases.get(case_id, principal)
        with self._connect() as db:
            rows = db.execute(
                "SELECT * FROM recruitment_case_packages WHERE case_id=? ORDER BY created_at DESC",
                (case_id,),
            ).fetchall()
        return [self._record(row) for row in rows]

    def generate(
        self, case_id: str, package_type: RecruitmentPackageType, principal: Principal,
    ) -> RecruitmentPackageRecord:
        readiness = self.preview(case_id, package_type, principal)
        if readiness.blockers:
            raise ValueError(readiness.blockers[0])
        case = self.cases.get(case_id, principal)
        package_id = f"recruitment_package_{uuid4().hex}"
        created_at = _now()
        stem = f"{_safe_name(case.case_name)}-{PACKAGE_TITLES[package_type]}-{created_at[:10]}"
        package_dir = self.output_dir / case.id / package_id
        package_dir.mkdir(parents=True, exist_ok=True)
        docx_path = package_dir / f"{stem}.docx"
        xlsx_path = package_dir / f"{stem}.xlsx"

        self._build_docx(case, package_type, principal, created_at, docx_path)
        self._build_xlsx(case, package_type, principal, created_at, xlsx_path)

        docx_metadata: dict[str, object] = {
            "package_id": package_id,
            "recruitment_case_id": case.id,
            "package_type": package_type,
            "preset": "standard_business_brief",
            "named_override": "recruitment_evidence_a4_zh_cn",
            "requires_human_review": True,
        }
        if self.word_verifier is not None:
            word_report = self.word_verifier.verify(docx_path)
            docx_metadata["word_automation"] = word_report.model_dump(mode="json")
            if word_report.preview_files:
                docx_metadata["preview_files"] = word_report.preview_files
        docx_report = self.artifact_harness.inspect(
            docx_path, kind="recruitment_evidence_package", mime_type=DOCX_MIME,
            metadata=docx_metadata,
        )
        xlsx_report = self.artifact_harness.inspect(
            xlsx_path, kind="recruitment_evidence_package", mime_type=XLSX_MIME,
            metadata={"package_id": package_id, "recruitment_case_id": case.id},
        )
        docx_metadata = attach_harness_report(docx_metadata, docx_report)
        xlsx_metadata = attach_harness_report({
            "package_id": package_id,
            "recruitment_case_id": case.id,
            "package_type": package_type,
            "requires_human_review": True,
        }, xlsx_report)

        docx_artifact_id = self._register_artifact(
            case, principal, docx_path, DOCX_MIME, "recruitment_evidence_docx", docx_metadata,
        )
        xlsx_artifact_id = self._register_artifact(
            case, principal, xlsx_path, XLSX_MIME, "recruitment_evidence_xlsx", xlsx_metadata,
        )
        files = [
            RecruitmentPackageFile(
                format="docx", display_name=docx_path.name, byte_size=docx_path.stat().st_size,
                artifact_id=docx_artifact_id,
            ),
            RecruitmentPackageFile(
                format="xlsx", display_name=xlsx_path.name, byte_size=xlsx_path.stat().st_size,
                artifact_id=xlsx_artifact_id,
            ),
        ]
        harness = {
            "passed": docx_report.passed and xlsx_report.passed,
            "docx": docx_report.model_dump(mode="json"),
            "xlsx": xlsx_report.model_dump(mode="json"),
        }
        with self._connect() as db:
            db.execute(
                """INSERT INTO recruitment_case_packages(
                id,case_id,tenant_id,owner_user_id,package_type,title,files_json,harness_json,
                paths_json,created_by,created_at
                ) VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    package_id, case.id, case.tenant_id, case.owner_user_id, package_type,
                    PACKAGE_TITLES[package_type],
                    json.dumps([item.model_dump(mode="json") for item in files], ensure_ascii=False),
                    json.dumps(harness, ensure_ascii=False),
                    json.dumps({"docx": str(docx_path), "xlsx": str(xlsx_path)}, ensure_ascii=False),
                    principal.display_name, created_at,
                ),
            )
            row = db.execute(
                "SELECT * FROM recruitment_case_packages WHERE id=?", (package_id,),
            ).fetchone()
        if row is None:  # pragma: no cover
            raise RuntimeError("recruitment_package_persist_failed")
        return self._record(row)

    def download(
        self, case_id: str, package_id: str, file_format: str, principal: Principal,
    ) -> tuple[Path, str, str]:
        self.cases.get(case_id, principal)
        if file_format not in {"docx", "xlsx"}:
            raise ValueError("recruitment_package_format_invalid")
        with self._connect() as db:
            row = db.execute(
                "SELECT * FROM recruitment_case_packages WHERE id=? AND case_id=?",
                (package_id, case_id),
            ).fetchone()
        if row is None:
            raise KeyError("recruitment_package_not_found")
        paths = json.loads(row["paths_json"] or "{}")
        path = Path(str(paths.get(file_format) or "")).resolve()
        root = self.output_dir.resolve()
        if not path.is_file() or root not in path.parents:
            raise KeyError("recruitment_package_file_not_found")
        mime = DOCX_MIME if file_format == "docx" else XLSX_MIME
        return path, mime, path.name

    def _register_artifact(
        self,
        case: RecruitmentCaseRecord,
        principal: Principal,
        path: Path,
        mime_type: str,
        kind: str,
        metadata: dict[str, object],
    ) -> str:
        # Conversation artifacts remain strictly owner scoped. Tenant admins can still
        # generate and download through the governed Case package endpoint.
        if case.owner_user_id != principal.user_id:
            return ""
        source = self.repository.get_artifact(case.artifact_id, principal)
        artifact_id = f"artifact_{uuid4().hex}"
        self.repository.add_artifact(ArtifactRecord(
            id=artifact_id,
            conversation_id=case.conversation_id,
            message_id=source.message_id,
            kind=kind,
            display_name=path.name,
            path=str(path),
            mime_type=mime_type,
            byte_size=path.stat().st_size,
            metadata=metadata,
            created_at=_now(),
        ), principal)
        return artifact_id

    @staticmethod
    def _record(row: sqlite3.Row) -> RecruitmentPackageRecord:
        return RecruitmentPackageRecord(
            id=row["id"], case_id=row["case_id"], tenant_id=row["tenant_id"],
            owner_user_id=row["owner_user_id"], package_type=row["package_type"],
            title=row["title"],
            files=[RecruitmentPackageFile.model_validate(item) for item in json.loads(row["files_json"])],
            harness=json.loads(row["harness_json"] or "{}"),
            created_by=row["created_by"], created_at=row["created_at"],
        )

    @classmethod
    def _build_docx(
        cls,
        case: RecruitmentCaseRecord,
        package_type: RecruitmentPackageType,
        principal: Principal,
        created_at: str,
        output: Path,
    ) -> None:
        document = Document()
        section = document.sections[0]
        section.start_type = WD_SECTION_START.NEW_PAGE
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.1)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)
        section.header_distance = Cm(1.1)
        section.footer_distance = Cm(1.1)

        styles = document.styles
        cls._set_style(styles["Normal"], "Microsoft YaHei", 10.5, after=6, line=1.1)
        cls._set_style(styles["Heading 1"], "Microsoft YaHei", 16, bold=True, color="2E74B5", before=16, after=8)
        cls._set_style(styles["Heading 2"], "Microsoft YaHei", 13, bold=True, color="2E74B5", before=12, after=6)
        cls._set_style(styles["Heading 3"], "Microsoft YaHei", 12, bold=True, color="1F4D78", before=8, after=4)

        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        cls._run(header, "UniHR AgentOS · 招聘证据包", 8.5, color="6B7280")
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cls._run(footer, "仅供授权招聘人员内部复核 · 需人工确认", 8.5, color="6B7280")

        kicker = document.add_paragraph()
        kicker.paragraph_format.space_after = Pt(4)
        cls._run(kicker, "RECRUITMENT EVIDENCE PACKAGE", 9, bold=True, color="2E74B5")
        title = document.add_paragraph()
        title.paragraph_format.space_after = Pt(5)
        cls._run(title, PACKAGE_TITLES[package_type], 23, bold=True, color="0B2545")
        subtitle = document.add_paragraph()
        subtitle.paragraph_format.space_after = Pt(14)
        cls._run(subtitle, f"{case.case_name} · {case.job_title}", 12.5, color="4B5563")

        cls._label_table(document, [
            ("Case ID", case.id), ("状态", cls._case_status(case.status)),
            ("Rubric 快照", case.rubric_hash[:20]), ("生成时间", cls._display_datetime(created_at)),
            ("生成操作人", principal.display_name), ("版本边界", "生成时 Case 快照；后续修改需重新生成"),
        ])
        cls._callout(document, COMPLIANCE_NOTE, "FFF8E8", "7A5A00")

        if package_type == "final_decision":
            document.add_heading("1. 最终人工建议", level=1)
            selected = [item for item in case.candidates if item.candidate_id in case.final_candidate_ids]
            cls._label_table(document, [
                ("建议人选", "、".join(item.candidate_name for item in selected) or "未填写"),
                ("人工理由", case.final_rationale),
                ("确认人", case.finalized_by),
                ("确认时间", cls._display_datetime(case.finalized_at)),
            ])
            first_section = 2
        else:
            document.add_heading("1. 面试使用说明", level=1)
            document.add_paragraph("面试官应记录候选人的原始回答、本人贡献和可核实结果，再给出 1-5 分维度评价；不要用笼统印象替代证据。")
            first_section = 2

        document.add_heading(f"{first_section}. 候选人证据总览", level=1)
        overview = document.add_table(rows=1, cols=6)
        overview.alignment = WD_TABLE_ALIGNMENT.CENTER
        overview.autofit = False
        headings = ["排名", "候选人", "AI 辅助分", "证据覆盖", "人工复核", "面试结论"]
        for index, value in enumerate(headings):
            cls._cell(overview.rows[0].cells[index], value, bold=True, fill="E8EEF5", align="center")
        widths = [1.0, 2.8, 2.0, 2.0, 3.0, 3.2]
        for candidate in case.candidates:
            row = overview.add_row().cells
            values = [
                str(candidate.rank), candidate.candidate_name, f"{candidate.ai_score:.1f}",
                f"{candidate.evidence_coverage:.0%}", cls._review_status(candidate.review_status),
                cls._interview_outcome(candidate.interview_outcome),
            ]
            for index, value in enumerate(values):
                cls._cell(row[index], value, align="center" if index != 1 else "left")
        cls._set_table_widths(overview, widths)

        document.add_heading(f"{first_section + 1}. 候选人逐项复核", level=1)
        for candidate in case.candidates:
            document.add_heading(f"{candidate.rank}. {candidate.candidate_name}", level=2)
            cls._label_table(document, [
                ("来源文件", candidate.source_name),
                ("AI 建议", candidate.ai_recommendation),
                ("人工记录", candidate.reviewer_note or "尚未填写"),
                ("优势", "；".join(candidate.strengths) or "未提取"),
                ("缺口/风险", "；".join(candidate.gaps) or "未提取"),
            ])
            if candidate.evidence_locations:
                evidence_table = document.add_table(rows=1, cols=3)
                evidence_table.autofit = False
                for index, value in enumerate(["评价维度", "简历事实证据", "原文定位"]):
                    cls._cell(evidence_table.rows[0].cells[index], value, bold=True, fill="F2F4F7")
                for item in candidate.evidence_locations:
                    row = evidence_table.add_row().cells
                    locator = dict(item.get("locator") or {})
                    location = cls._locator_text(locator)
                    cls._cell(row[0], str(item.get("criterion_label") or "综合"))
                    cls._cell(row[1], str(item.get("evidence") or "证据待补充"))
                    cls._cell(row[2], location)
                cls._set_table_widths(evidence_table, [3.0, 8.2, 3.8])
            if candidate.verification_questions:
                document.add_heading("建议核实问题", level=3)
                for question in candidate.verification_questions:
                    paragraph = document.add_paragraph(style="List Bullet")
                    cls._run(paragraph, question, 10.5)

        document.add_heading(f"{first_section + 2}. 面试场次与多人反馈", level=1)
        if not case.interview_sessions:
            document.add_paragraph("尚未安排正式面试场次。")
        for session in case.interview_sessions:
            candidate = cls._candidate(case, session.candidate_id)
            document.add_heading(f"{candidate.candidate_name} · {session.title}", level=2)
            cls._label_table(document, [
                ("时间", f"{cls._display_datetime(session.start_at)} - {cls._display_datetime(session.end_at, time_only=True)}"),
                ("面试官", "、".join(item.display_name for item in session.interviewers)),
                ("评价维度", "、".join(session.focus_dimensions)),
                ("场次状态", session.status),
            ])
            feedback_by_user = {item.interviewer_user_id: item for item in session.feedback}
            for interviewer in session.interviewers:
                feedback = feedback_by_user.get(interviewer.user_id)
                if feedback:
                    score_text = "；".join(f"{label} {score}/5" for label, score in feedback.dimension_scores.items())
                    cls._label_table(document, [
                        ("面试官", feedback.interviewer_name),
                        ("总体建议", cls._recommendation(feedback.recommendation)),
                        ("维度评分", score_text),
                        ("事实依据", feedback.evidence_note),
                        ("风险/待核实", feedback.risk_note or "无补充"),
                    ])
                else:
                    cls._label_table(document, [
                        ("面试官", interviewer.display_name), ("总体建议", "□ 强烈建议  □ 建议  □ 保留  □ 不建议  □ 明确不建议"),
                        ("事实依据", "\n\n"), ("风险/待核实", "\n"),
                    ])

        document.add_heading(f"{first_section + 3}. 审计与签字", level=1)
        cls._label_table(document, [
            ("AI 与人工边界", "AI 分数仅作为复核线索；候选人状态、面试反馈和最终建议均按操作人独立留痕。"),
            ("Case 负责人签字", "________________________"),
            ("复核日期", "________年____月____日"),
        ])
        document.core_properties.title = f"{case.case_name}-{PACKAGE_TITLES[package_type]}"
        document.core_properties.subject = "UniHR AgentOS 招聘人工复核证据"
        document.core_properties.author = "UniHR AgentOS"
        document.save(output)

    @classmethod
    def _build_xlsx(
        cls,
        case: RecruitmentCaseRecord,
        package_type: RecruitmentPackageType,
        principal: Principal,
        created_at: str,
        output: Path,
    ) -> None:
        workbook = Workbook()
        workbook.remove(workbook.active)
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = "auto"
        sheets = {
            name: workbook.create_sheet(name)
            for name in ["使用说明", "候选人总览", "面试评价模板", "多人反馈", "最终人工建议", "证据索引", "审计信息"]
        }
        package_data = workbook.create_sheet("PackageData")
        package_data.sheet_state = "hidden"
        package_data.cell(1, 1, "final_candidate_id")
        for row_index, candidate_id in enumerate(case.final_candidate_ids, start=2):
            package_data.cell(row_index, 1, candidate_id)
        for sheet in sheets.values():
            sheet.sheet_view.showGridLines = False
            sheet.freeze_panes = "A4"

        instruction = sheets["使用说明"]
        cls._sheet_title(instruction, PACKAGE_TITLES[package_type], f"{case.case_name} · {created_at[:10]}", 8)
        rows = [
            ("使用边界", COMPLIANCE_NOTE),
            ("蓝色单元格", "来自 Case 的系统快照，不建议直接修改。"),
            ("浅黄色单元格", "供面试官补充评价；评分范围 1-5。"),
            ("版本规则", "如 Case 中候选人、反馈或最终建议变化，请重新生成新版本。"),
            ("Package 类型", PACKAGE_TITLES[package_type]),
        ]
        for row_index, (label, value) in enumerate(rows, start=4):
            instruction.cell(row_index, 1, label)
            instruction.cell(row_index, 2, value)
            instruction.merge_cells(start_row=row_index, start_column=2, end_row=row_index, end_column=8)
            cls._style_label(instruction.cell(row_index, 1))
            cls._style_value(instruction.cell(row_index, 2))
        instruction.column_dimensions["A"].width = 18
        for letter in "BCDEFGH":
            instruction.column_dimensions[letter].width = 15

        overview = sheets["候选人总览"]
        cls._sheet_title(overview, "候选人总览", case.case_name, 10)
        overview_headers = ["候选人ID", "排名", "候选人", "来源文件", "AI辅助分", "证据覆盖率", "人工复核", "面试结论", "是否最终建议", "人工复核备注"]
        cls._write_headers(overview, 4, overview_headers)
        for row_index, candidate in enumerate(case.candidates, start=5):
            values = [candidate.candidate_id, candidate.rank, candidate.candidate_name, candidate.source_name, candidate.ai_score, candidate.evidence_coverage, cls._review_status(candidate.review_status), cls._interview_outcome(candidate.interview_outcome), None, candidate.reviewer_note]
            for column, value in enumerate(values, start=1):
                overview.cell(row_index, column, value)
            overview.cell(row_index, 9, f'=IF(COUNTIF(\'PackageData\'!$A$2:$A$31,A{row_index})>0,"是","否")')
            overview.cell(row_index, 5).number_format = "0.0"
            overview.cell(row_index, 6).number_format = "0%"
        cls._format_data_sheet(overview, 10, max(5, 4 + len(case.candidates)))
        for row_index in range(5, max(5, 4 + len(case.candidates)) + 1):
            overview.row_dimensions[row_index].height = 50
        overview.auto_filter.ref = f"A4:J{max(5, 4 + len(case.candidates))}"

        evaluation = sheets["面试评价模板"]
        cls._sheet_title(evaluation, "结构化面试评价模板", "黄色列由面试官填写；事实依据为必填", 12)
        eval_headers = ["场次ID", "候选人", "面试官", "评价维度", "评分(1-5)", "事实依据", "风险/待核实", "总体建议", "是否已提交", "面试时间", "方式", "地点/链接"]
        cls._write_headers(evaluation, 4, eval_headers)
        eval_rows: list[list[object]] = []
        for session in case.interview_sessions:
            candidate = cls._candidate(case, session.candidate_id)
            feedback_map = {item.interviewer_user_id: item for item in session.feedback}
            for interviewer in session.interviewers:
                feedback = feedback_map.get(interviewer.user_id)
                for dimension in session.focus_dimensions:
                    eval_rows.append([
                        session.id, candidate.candidate_name, interviewer.display_name, dimension,
                        feedback.dimension_scores.get(dimension) if feedback else None,
                        feedback.evidence_note if feedback else "",
                        feedback.risk_note if feedback else "",
                        cls._recommendation(feedback.recommendation) if feedback else "",
                        "是" if feedback else "否", cls._display_datetime(session.start_at),
                        session.mode, session.meeting_url or session.location,
                    ])
        if not eval_rows:
            for candidate in case.candidates:
                for dimension in cls._dimensions(case):
                    eval_rows.append(["", candidate.candidate_name, "", dimension, None, "", "", "", "否", "", "", ""])
        for row_index, row in enumerate(eval_rows, start=5):
            for column, value in enumerate(row, start=1):
                evaluation.cell(row_index, column, value)
        eval_end = max(5, 4 + len(eval_rows))
        score_validation = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
        recommendation_validation = DataValidation(type="list", formula1='"强烈建议,建议进入后续,保留观察,不建议,明确不建议"', allow_blank=True)
        evaluation.add_data_validation(score_validation)
        evaluation.add_data_validation(recommendation_validation)
        score_validation.add(f"E5:E{eval_end}")
        recommendation_validation.add(f"H5:H{eval_end}")
        for row in range(5, eval_end + 1):
            for column in range(5, 9):
                evaluation.cell(row, column).fill = PatternFill("solid", fgColor="FFF8E8")
        cls._format_data_sheet(evaluation, 12, eval_end)
        for row_index in range(5, eval_end + 1):
            evaluation.row_dimensions[row_index].height = 92
        evaluation.auto_filter.ref = f"A4:L{eval_end}"

        feedback_sheet = sheets["多人反馈"]
        cls._sheet_title(feedback_sheet, "多人面试反馈", "每位面试官独立记录，不覆盖他人结论", 9)
        cls._write_headers(feedback_sheet, 4, ["候选人", "场次", "面试官", "总体建议", "维度评分", "事实依据", "风险/待核实", "提交时间", "状态"])
        feedback_rows = []
        for session in case.interview_sessions:
            candidate = cls._candidate(case, session.candidate_id)
            for feedback in session.feedback:
                feedback_rows.append([
                    candidate.candidate_name, session.title, feedback.interviewer_name,
                    cls._recommendation(feedback.recommendation),
                    "；".join(f"{key}:{value}/5" for key, value in feedback.dimension_scores.items()),
                    feedback.evidence_note, feedback.risk_note, cls._display_datetime(feedback.updated_at),
                    "已提交",
                ])
        for row_index, row in enumerate(feedback_rows or [["", "", "", "", "", "尚无已提交反馈", "", "", "待提交"]], start=5):
            for column, value in enumerate(row, start=1):
                feedback_sheet.cell(row_index, column, value)
        cls._format_data_sheet(feedback_sheet, 9, max(5, 4 + len(feedback_rows)))
        for row_index in range(5, max(5, 4 + len(feedback_rows)) + 1):
            feedback_sheet.row_dimensions[row_index].height = 76

        final_sheet = sheets["最终人工建议"]
        cls._sheet_title(final_sheet, "最终人工建议", "该页只呈现 Case 负责人/管理员的人工结论", 7)
        cls._write_headers(final_sheet, 4, ["候选人ID", "候选人", "人工结论", "人工理由", "确认人", "确认时间", "AI辅助分"])
        selected = [item for item in case.candidates if item.candidate_id in case.final_candidate_ids]
        rows_to_write = selected if package_type == "final_decision" else []
        for row_index, candidate in enumerate(rows_to_write, start=5):
            row = [candidate.candidate_id, candidate.candidate_name, "最终建议人选", case.final_rationale, case.finalized_by, cls._display_datetime(case.finalized_at), candidate.ai_score]
            for column, value in enumerate(row, start=1):
                final_sheet.cell(row_index, column, value)
        if not rows_to_write:
            final_sheet.cell(5, 3, "尚未形成最终人工建议" if case.status != "finalized" else "本包为面试评价版本")
        cls._format_data_sheet(final_sheet, 7, max(5, 4 + len(rows_to_write)))
        for row_index in range(5, max(5, 4 + len(rows_to_write)) + 1):
            rationale = str(final_sheet.cell(row_index, 4).value or "")
            final_sheet.row_dimensions[row_index].height = cls._wrapped_row_height(
                rationale,
                visual_chars_per_line=24,
                minimum=50,
            )

        evidence = sheets["证据索引"]
        cls._sheet_title(evidence, "简历证据索引", "从评分项返回原始简历的可复核定位", 9)
        cls._write_headers(evidence, 4, ["候选人", "评价维度", "事实证据", "来源文件", "定位类型", "页/段/表/Sheet", "原文摘录", "定位状态", "匹配度"])
        evidence_rows = []
        for candidate in case.candidates:
            for item in candidate.evidence_locations:
                locator = dict(item.get("locator") or {})
                evidence_rows.append([
                    candidate.candidate_name, item.get("criterion_label") or "综合", item.get("evidence") or "",
                    locator.get("source_name") or candidate.source_name, locator.get("anchor_kind") or "unlocated",
                    cls._locator_text(locator), locator.get("quote") or "", locator.get("status") or "unlocated",
                    locator.get("match_score") or 0,
                ])
        for row_index, row in enumerate(evidence_rows or [["", "", "无结构化证据定位", "", "", "", "", "unlocated", 0]], start=5):
            for column, value in enumerate(row, start=1):
                evidence.cell(row_index, column, value)
            evidence.cell(row_index, 9).number_format = "0%"
        cls._format_data_sheet(evidence, 9, max(5, 4 + len(evidence_rows)))
        for row_index in range(5, max(5, 4 + len(evidence_rows)) + 1):
            evidence.row_dimensions[row_index].height = 60
        evidence.auto_filter.ref = f"A4:I{max(5, 4 + len(evidence_rows))}"

        audit = sheets["审计信息"]
        cls._sheet_title(audit, "审计信息", "Package 与 Case 的快照关系", 6)
        audit_rows = [
            ("Case ID", case.id), ("Conversation ID", case.conversation_id), ("Run ID", case.run_id),
            ("原筛选产物 ID", case.artifact_id), ("Rubric Hash", case.rubric_hash),
            ("Case 状态", cls._case_status(case.status)), ("生成类型", PACKAGE_TITLES[package_type]),
            ("生成操作人", principal.display_name), ("生成时间 UTC", created_at),
            ("人工最终确认人", case.finalized_by or "未形成"), ("人工最终确认时间", case.finalized_at or "未形成"),
            ("合规声明", COMPLIANCE_NOTE),
        ]
        for row_index, (label, value) in enumerate(audit_rows, start=4):
            audit.cell(row_index, 1, label)
            audit.cell(row_index, 2, value)
            audit.merge_cells(start_row=row_index, start_column=2, end_row=row_index, end_column=6)
            cls._style_label(audit.cell(row_index, 1))
            cls._style_value(audit.cell(row_index, 2))
        audit.column_dimensions["A"].width = 22
        for letter in "BCDEF":
            audit.column_dimensions[letter].width = 18

        for sheet in sheets.values():
            sheet.page_setup.orientation = "landscape" if sheet.max_column >= 9 else "portrait"
            sheet.page_setup.fitToWidth = 1
            sheet.page_setup.fitToHeight = 0
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet.print_title_rows = "1:4"
        workbook.save(output)

    @staticmethod
    def _set_style(style, font: str, size: float, *, bold: bool = False, color: str = "111827", before: float = 0, after: float = 0, line: float = 1.0) -> None:
        style.font.name = font
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = line

    @staticmethod
    def _run(paragraph, text: str, size: float, *, bold: bool = False, color: str = "111827"):
        run = paragraph.add_run(text)
        run.font.name = "Microsoft YaHei"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        return run

    @classmethod
    def _cell(cls, cell, text: str, *, bold: bool = False, fill: str = "", align: str = "left") -> None:
        cell.text = ""
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if align == "center" else WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(0)
        cls._run(paragraph, text, 9, bold=bold, color="0B2545" if bold else "111827")
        margins = cell._tc.get_or_add_tcPr()
        tc_mar = margins.first_child_found_in("w:tcMar")
        if tc_mar is None:
            tc_mar = OxmlElement("w:tcMar")
            margins.append(tc_mar)
        for edge in ("top", "bottom", "start", "end"):
            node = OxmlElement(f"w:{edge}")
            node.set(qn("w:w"), "100" if edge in {"top", "bottom"} else "120")
            node.set(qn("w:type"), "dxa")
            tc_mar.append(node)
        if fill:
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), fill)
            margins.append(shading)

    @classmethod
    def _label_table(cls, document: Document, rows: list[tuple[str, str]]) -> None:
        table = document.add_table(rows=0, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        for label, value in rows:
            cells = table.add_row().cells
            cls._cell(cells[0], label, bold=True, fill="F2F4F7")
            cls._cell(cells[1], value)
        cls._set_table_widths(table, [3.2, 11.8])
        document.add_paragraph().paragraph_format.space_after = Pt(2)

    @classmethod
    def _callout(cls, document: Document, text: str, fill: str, color: str) -> None:
        table = document.add_table(rows=1, cols=1)
        table.autofit = False
        cls._cell(table.cell(0, 0), text, fill=fill)
        for run in table.cell(0, 0).paragraphs[0].runs:
            run.font.color.rgb = RGBColor.from_string(color)
        cls._set_table_widths(table, [15.0])
        document.add_paragraph().paragraph_format.space_after = Pt(2)

    @staticmethod
    def _set_table_widths(table, widths_cm: list[float]) -> None:
        for row in table.rows:
            for index, width in enumerate(widths_cm):
                row.cells[index].width = Cm(width)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

    @staticmethod
    def _sheet_title(sheet, title: str, subtitle: str, columns: int) -> None:
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
        sheet.cell(1, 1, title)
        sheet.cell(1, 1).font = Font(name="Microsoft YaHei", size=18, bold=True, color="0B2545")
        sheet.cell(1, 1).alignment = Alignment(vertical="center")
        sheet.row_dimensions[1].height = 34
        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=columns)
        sheet.cell(2, 1, subtitle)
        sheet.cell(2, 1).font = Font(name="Microsoft YaHei", size=10, color="6B7280")
        sheet.row_dimensions[2].height = 24

    @staticmethod
    def _write_headers(sheet, row: int, headers: list[str]) -> None:
        for column, value in enumerate(headers, start=1):
            cell = sheet.cell(row, column, value)
            cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2E74B5")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.row_dimensions[row].height = 30

    @staticmethod
    def _format_data_sheet(sheet, columns: int, end_row: int) -> None:
        thin = Side(style="thin", color="DCE3EA")
        widths = [22, 10, 16, 28, 14, 18, 20, 20, 18, 34, 16, 30]
        for column in range(1, columns + 1):
            sheet.column_dimensions[chr(64 + column)].width = widths[column - 1]
        for row in sheet.iter_rows(min_row=5, max_row=end_row, min_col=1, max_col=columns):
            for cell in row:
                cell.font = Font(name="Microsoft YaHei", size=9, color="1F2937")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = Border(bottom=thin)
            sheet.row_dimensions[cell.row].height = 42
        if end_row >= 5:
            sheet.conditional_formatting.add(
                f"E5:E{end_row}",
                CellIsRule(operator="greaterThanOrEqual", formula=["80"], fill=PatternFill("solid", fgColor="E6F4EA")),
            )

    @staticmethod
    def _wrapped_row_height(
        value: str,
        *,
        visual_chars_per_line: int,
        minimum: float,
        line_height: float = 16,
    ) -> float:
        """Reserve enough row height for wrapped Chinese/Latin text in exported workbooks."""
        visual_width = sum(2 if ord(character) > 127 else 1 for character in value)
        explicit_lines = max(1, value.count("\n") + 1)
        wrapped_lines = max(1, (visual_width + visual_chars_per_line - 1) // visual_chars_per_line)
        return max(minimum, line_height * max(explicit_lines, wrapped_lines))

    @staticmethod
    def _style_label(cell) -> None:
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="0B2545")
        cell.fill = PatternFill("solid", fgColor="E8EEF5")
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    @staticmethod
    def _style_value(cell) -> None:
        cell.font = Font(name="Microsoft YaHei", size=10, color="1F2937")
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    @staticmethod
    def _candidate(case: RecruitmentCaseRecord, candidate_id: str) -> CandidateReviewRecord:
        return next(item for item in case.candidates if item.candidate_id == candidate_id)

    @staticmethod
    def _dimensions(case: RecruitmentCaseRecord) -> list[str]:
        values = [dimension for session in case.interview_sessions for dimension in session.focus_dimensions]
        values.extend(
            str(item.get("criterion_label") or "").strip()
            for candidate in case.candidates for item in candidate.evidence_locations
        )
        result = list(dict.fromkeys(value for value in values if value))
        return result[:8] or ["岗位动机", "相关经验", "问题解决", "沟通协作"]

    @staticmethod
    def _locator_text(locator: dict[str, object]) -> str:
        parts = []
        if locator.get("page_number"):
            parts.append(f"第{locator['page_number']}页")
        if locator.get("paragraph_number"):
            parts.append(f"第{locator['paragraph_number']}段")
        if locator.get("table_number"):
            parts.append(f"表{locator['table_number']}")
        if locator.get("sheet_name"):
            parts.append(f"Sheet:{locator['sheet_name']}")
        return " / ".join(parts) or str(locator.get("status") or "未定位")

    @staticmethod
    def _case_status(value: str) -> str:
        return {"reviewing": "逐人复核中", "interviewing": "面试复核中", "finalized": "已形成最终人工建议"}.get(value, value)

    @staticmethod
    def _review_status(value: str) -> str:
        return {"pending": "待复核", "shortlisted": "进入面试", "hold": "待核实", "not_advanced": "本轮不进入"}.get(value, value)

    @staticmethod
    def _interview_outcome(value: str) -> str:
        return {"not_scheduled": "未安排", "scheduled": "已安排", "passed": "建议进入后续", "hold": "继续观察", "not_advanced": "不进入后续"}.get(value, value)

    @staticmethod
    def _recommendation(value: str) -> str:
        return {"strong_yes": "强烈建议", "yes": "建议进入后续", "hold": "保留观察", "no": "不建议", "strong_no": "明确不建议"}.get(value, value)

    @staticmethod
    def _display_datetime(value: str, *, time_only: bool = False) -> str:
        if not value:
            return ""
        try:
            parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
            if parsed.tzinfo is None:
                parsed = parsed.replace(tzinfo=ZoneInfo("Asia/Shanghai"))
            local = parsed.astimezone(ZoneInfo("Asia/Shanghai"))
            return local.strftime("%H:%M") if time_only else local.strftime("%Y-%m-%d %H:%M")
        except (TypeError, ValueError):
            return value
