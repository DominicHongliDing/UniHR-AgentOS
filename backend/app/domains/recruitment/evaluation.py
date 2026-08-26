from __future__ import annotations

import json
import re
import sqlite3
from datetime import datetime, timezone
from pathlib import Path
from time import perf_counter

from openpyxl import load_workbook

from app.domains.conversations.contracts import AttachmentRecord
from app.domains.runs import RunPayloadStore, RunPointer
from app.models.gateway import ModelReply

from .contracts import RecruitmentCapacityReport, RecruitmentRequest
from .service import RecruitmentScreeningModule


_DATASET_VERSION = "recruitment-capacity-v1"
_SOURCE_PATTERN = re.compile(r"capacity-resume-(\d{2})\.txt")


class RecruitmentCapacityEvaluator:
    """Exercise the real recruitment Graph through one small evaluation interface.

    The internal deterministic model adapter isolates Graph capacity and stability
    from provider latency and model drift. This report therefore does not claim
    semantic screening accuracy.
    """

    def __init__(self, output_dir: str | Path) -> None:
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def evaluate(self, *, source_count: int = 30, map_batch_size: int = 8) -> RecruitmentCapacityReport:
        if source_count != 30:
            raise ValueError("recruitment_capacity_fixture_requires_30_sources")
        started = perf_counter()
        run_root = self.output_dir / datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
        fixture_dir = run_root / "fixtures"
        fixture_dir.mkdir(parents=True, exist_ok=True)
        database_path = run_root / "capacity.sqlite3"
        payload_store = RunPayloadStore(database_path)
        attachments = self._fixtures(fixture_dir, source_count)

        results = []
        gateways = []
        for index, ordered in enumerate((attachments, list(reversed(attachments))), start=1):
            gateway = _DeterministicCapacityGateway()
            gateways.append(gateway)
            module = RecruitmentScreeningModule(
                gateway,
                run_root / f"run-{index}",
                "deterministic-capacity-adapter",
                payload_store=payload_store,
            )
            request = RecruitmentRequest(
                job_requirements="筛选产业拓展实习生，必须具备沟通协作能力，优先可核验项目成果。",
                job_title="产业拓展实习生容量基线",
                must_haves="沟通协作；能够说明本人贡献",
                nice_to_haves="B2B项目；量化成果；材料证据完整",
                selection_count=3,
                map_batch_size=map_batch_size,
                attachments=ordered,
            )
            results.append(module.run(
                request,
                payload_scope=RunPointer(
                    tenant_id="tenant-capacity-eval",
                    owner_user_id="capacity-evaluator",
                    run_id=f"capacity-run-{index}",
                ),
            ))

        first, second = results
        first_signature = [item.source_name for item in first.report.candidates]
        second_signature = [item.source_name for item in second.report.candidates]
        ranking_consistent = first_signature == second_signature
        first_metrics = dict(first.report.harness.metrics if first.report.harness else {})
        batch_count = sum(item.screening_calls for item in gateways)
        model_call_count = sum(item.calls for item in gateways)
        offload_verified = all(
            any(row.get("event") == "graph.payload.offloaded" for row in result.trace)
            for result in results
        ) and self._payload_count(database_path) >= len(results)
        artifact_paths = [str(item.artifact_path) for item in results]
        formula_error_count, artifact_verified = self._verify_workbooks(
            [item.artifact_path for item in results]
        )
        graph_path = list(first.graph_path)
        required_nodes = {
            "prepare", "deduplicate", "version_resolve", "rubric_lock", "batch_map",
            "reduce", "fairness_guard", "evidence_localize", "recruitment_harness",
            "quality_decision", "artifact_create", "finalize",
        }
        checks = {
            "thirty_sources_accepted": len(first.report.source_profiles) == source_count,
            "all_sources_unique": int(first_metrics.get("unique_source_count", 0)) == source_count,
            "all_candidates_returned": len(first.report.candidates) == source_count,
            "source_coverage_complete": float(first_metrics.get("source_coverage", 0)) == 1.0,
            "ranking_consistent_when_input_reversed": ranking_consistent,
            "expected_batch_split": batch_count == 8,
            "payload_offload_verified": offload_verified,
            "artifact_verified": artifact_verified,
            "formula_errors_absent": formula_error_count == 0,
            "graph_path_complete": required_nodes.issubset(set(graph_path)),
            "harness_has_no_failures": all(
                result.report.harness is not None and result.report.harness.status != "failed"
                for result in results
            ),
        }
        issues = [label for label, passed in checks.items() if not passed]
        duration_ms = max(0, int((perf_counter() - started) * 1000))
        warnings = []
        if duration_ms > 10_000:
            warnings.append("固定容量评测超过10秒，请检查文件生成或 Graph 编排性能")
        status = "failed" if issues else "warning" if warnings else "passed"
        report = RecruitmentCapacityReport(
            dataset_version=_DATASET_VERSION,
            execution_mode="deterministic_graph_capacity",
            status=status,
            source_count=source_count,
            unique_source_count=int(first_metrics.get("unique_source_count", 0)),
            candidate_count=len(first.report.candidates),
            run_count=2,
            batch_count=batch_count,
            model_call_count=model_call_count,
            duration_ms=duration_ms,
            source_coverage=float(first_metrics.get("source_coverage", 0)),
            evidence_coverage=float(first_metrics.get("evidence_coverage", 0)),
            ranking_consistent=ranking_consistent,
            payload_offload_verified=offload_verified,
            artifact_verified=artifact_verified,
            formula_error_count=formula_error_count,
            graph_path=graph_path,
            checks=checks,
            issues=issues,
            warnings=warnings,
            artifact_paths=artifact_paths,
            generated_at=datetime.now(timezone.utc).isoformat(),
        )
        (self.output_dir / "latest.json").write_text(
            report.model_dump_json(indent=2), encoding="utf-8"
        )
        return report

    @staticmethod
    def _fixtures(directory: Path, count: int) -> list[AttachmentRecord]:
        now = datetime.now(timezone.utc).isoformat()
        attachments = []
        for index in range(1, count + 1):
            name = f"capacity-resume-{index:02d}.txt"
            text = (
                f"候选人{index:02d}\n"
                f"岗位相关经验：参与产业合作与客户沟通项目 {index}。\n"
                f"必须条件：能够说明本人贡献并与团队协作 {index}。\n"
                f"项目成果：完成 {index % 5 + 1} 项可核验交付。\n"
                f"材料证据：包含职责、行动与结果记录 {index}。\n"
                "联系方式仅用于本次容量测试，不包含真实个人信息。"
            )
            path = directory / name
            path.write_text(text, encoding="utf-8")
            attachments.append(AttachmentRecord(
                id=f"capacity-att-{index:02d}",
                conversation_id="capacity-evaluation",
                message_id="capacity-message",
                display_name=name,
                mime_type="text/plain",
                extension=".txt",
                byte_size=len(text.encode("utf-8")),
                saved_path=str(path),
                extracted_text=text,
                summary={},
                created_at=now,
            ))
        return attachments

    @staticmethod
    def _payload_count(database_path: Path) -> int:
        with sqlite3.connect(database_path) as db:
            return int(db.execute(
                "SELECT COUNT(*) FROM run_payloads WHERE kind='recruitment.candidate_batches'"
            ).fetchone()[0])

    @staticmethod
    def _verify_workbooks(paths: list[Path]) -> tuple[int, bool]:
        errors = 0
        verified = True
        expected = {"筛选说明", "候选人排名", "评分明细", "重复材料", "版本归并", "人工复核", "Harness"}
        for path in paths:
            if not path.exists() or path.stat().st_size <= 0:
                verified = False
                continue
            workbook = load_workbook(path, data_only=False, read_only=False)
            verified = verified and expected.issubset(set(workbook.sheetnames))
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows():
                    for cell in row:
                        if isinstance(cell.value, str) and cell.value.startswith("#"):
                            errors += 1
            workbook.close()
        return errors, verified


class _DeterministicCapacityGateway:
    """Internal adapter used only to isolate Graph capacity from provider drift."""

    def __init__(self) -> None:
        self.calls = 0
        self.screening_calls = 0

    def complete(self, messages, **kwargs) -> ModelReply:
        self.calls += 1
        if kwargs.get("purpose") == "recruitment_rubric":
            payload = {
                "job_title": "产业拓展实习生容量基线",
                "criteria": [
                    {"id": "required", "label": "必须条件", "weight": 40, "requirement": "沟通协作与本人贡献"},
                    {"id": "experience", "label": "相关经历", "weight": 30, "requirement": "产业合作或客户沟通"},
                    {"id": "delivery", "label": "成果证据", "weight": 20, "requirement": "可核验交付"},
                    {"id": "evidence", "label": "材料完整", "weight": 10, "requirement": "职责行动结果完整"},
                ],
            }
            return ModelReply(json.dumps(payload, ensure_ascii=False), "capacity-rubric", 1, {})
        self.screening_calls += 1
        prompt = str(messages[-1]["content"])
        candidates = []
        for token in dict.fromkeys(_SOURCE_PATTERN.findall(prompt)):
            index = int(token)
            source = f"capacity-resume-{index:02d}.txt"
            candidates.append({
                "candidate_name": f"候选人{index:02d}",
                "source_name": source,
                "criterion_scores": {
                    "required": 70 + index % 21,
                    "experience": 60 + index * 3 % 31,
                    "delivery": 65 + index * 5 % 26,
                    "evidence": 75 + index * 7 % 16,
                },
                "score_evidence": {
                    "required": f"必须条件：能够说明本人贡献并与团队协作 {index}。",
                    "experience": f"岗位相关经验：参与产业合作与客户沟通项目 {index}。",
                    "delivery": f"项目成果：完成 {index % 5 + 1} 项可核验交付。",
                    "evidence": f"材料证据：包含职责、行动与结果记录 {index}。",
                },
                "strengths": ["具有岗位相关经历和可核验结果"],
                "gaps": ["需在面试中核实独立贡献"],
                "verification_questions": ["请说明一个项目中由本人独立完成的部分"],
            })
        return ModelReply(
            json.dumps({"candidates": candidates}, ensure_ascii=False),
            "capacity-screening", 1, {},
        )
