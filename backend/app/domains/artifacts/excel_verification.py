from __future__ import annotations

import json
import subprocess
from pathlib import Path
from shutil import copy2
from typing import Any
from uuid import uuid4

from openpyxl import load_workbook
from pydantic import BaseModel, Field
from pypdf import PdfReader

from app.observability import emit_tool_progress


class ExcelVerificationReport(BaseModel):
    passed: bool
    checks: dict[str, bool] = Field(default_factory=dict)
    metrics: dict[str, int | float | str | bool] = Field(default_factory=dict)
    issues: list[str] = Field(default_factory=list)
    preview_files: list[str] = Field(default_factory=list)


class ExcelAutomationVerifier:
    """Recalculate with desktop Excel and prove that the workbook can render to PDF."""

    def __init__(
        self,
        preview_dir: str | Path,
        *,
        excel_executable: str | Path = r"C:\Program Files\Microsoft Office\root\Office16\EXCEL.EXE",
        timeout_seconds: int = 90,
    ) -> None:
        self.preview_dir = Path(preview_dir)
        self.preview_dir.mkdir(parents=True, exist_ok=True)
        self.excel_executable = Path(excel_executable)
        self.timeout_seconds = timeout_seconds

    @property
    def available(self) -> bool:
        return self.excel_executable.is_file()

    def verify(self, workbook_path: str | Path) -> ExcelVerificationReport:
        workbook = Path(workbook_path).resolve()
        checks: dict[str, bool] = {"excel_automation_available": self.available}
        metrics: dict[str, int | float | str | bool] = {"calculation_engine": "Microsoft Excel COM"}
        issues: list[str] = []
        if not self.available:
            emit_tool_progress(
                "tool.spreadsheet.recalculate.failed",
                "当前环境无法调用桌面 Excel 进行公式重算",
                status="failed",
                error_code="excel_automation_unavailable",
            )
            return ExcelVerificationReport(
                passed=False, checks=checks, metrics=metrics, issues=["excel_automation_unavailable"],
            )
        validation_id = uuid4().hex
        preview = self.preview_dir / f"excel-preview-{validation_id}.pdf"
        validation_copy = self.preview_dir / f"excel-validation-{validation_id}.xlsx"
        copy2(workbook, validation_copy)
        metrics["validation_workspace"] = "isolated_copy"
        script = self._powershell_script(validation_copy, preview)
        emit_tool_progress(
            "tool.spreadsheet.recalculate.started",
            "正在调用桌面 Excel 重算公式并刷新缓存",
        )
        try:
            completed = subprocess.run(
                [
                    "powershell.exe", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass",
                    "-Command", script,
                ],
                capture_output=True,
                text=True,
                timeout=self.timeout_seconds,
                check=False,
                encoding="utf-8",
                errors="replace",
            )
        except subprocess.TimeoutExpired:
            validation_copy.unlink(missing_ok=True)
            preview.unlink(missing_ok=True)
            emit_tool_progress(
                "tool.spreadsheet.recalculate.failed",
                "桌面 Excel 公式重算超时",
                status="failed",
                error_code="excel_automation_timeout",
            )
            return ExcelVerificationReport(
                passed=False, checks={**checks, "excel_recalculation_completed": False},
                metrics=metrics, issues=["excel_automation_timeout"],
            )
        checks["excel_recalculation_completed"] = completed.returncode == 0
        if completed.returncode != 0:
            emit_tool_progress(
                "tool.spreadsheet.recalculate.failed",
                "桌面 Excel 未能完成公式重算",
                status="failed",
                error_code="excel_automation_failed",
            )
            issues.append("excel_automation_failed")
            metrics["automation_error"] = (completed.stderr or completed.stdout)[-1_000:]
            validation_copy.unlink(missing_ok=True)
            preview.unlink(missing_ok=True)
            return ExcelVerificationReport(passed=False, checks=checks, metrics=metrics, issues=issues)

        try:
            automation = json.loads((completed.stdout.strip().splitlines() or ["{}"])[-1])
        except json.JSONDecodeError:
            automation = {}
            issues.append("excel_automation_output_invalid")
        metrics["excel_version"] = str(automation.get("excel_version") or "unknown")
        emit_tool_progress(
            "tool.spreadsheet.recalculate.completed",
            "桌面 Excel 已完成公式重算并保存结果",
            status="completed",
            engine="Microsoft Excel COM",
            excel_version=metrics["excel_version"],
        )
        checks["workbook_reopens_after_recalculation"] = self._inspect_formula_cache(
            validation_copy, checks, metrics, issues,
        )
        emit_tool_progress(
            "tool.spreadsheet.render.started",
            "正在检查公式缓存和 PDF 预览页",
        )
        checks["rendered_pdf_exists"] = preview.is_file() and preview.stat().st_size > 0
        if checks["rendered_pdf_exists"]:
            try:
                reader = PdfReader(preview)
                page_count = len(reader.pages)
                content_pages = sum(
                    1 for page in reader.pages
                    if (page.get_contents() is not None and len(page.get_contents().get_data()) > 0)
                )
                checks["rendered_pdf_has_pages"] = page_count > 0
                checks["rendered_pages_have_content"] = content_pages == page_count
                metrics["rendered_page_count"] = page_count
                metrics["rendered_content_page_count"] = content_pages
            except (OSError, ValueError, KeyError) as exc:
                checks["rendered_pdf_has_pages"] = False
                checks["rendered_pages_have_content"] = False
                issues.append(f"rendered_pdf_invalid:{type(exc).__name__}")
        else:
            checks["rendered_pdf_has_pages"] = False
            checks["rendered_pages_have_content"] = False
            issues.append("rendered_pdf_missing")
        passed = all(checks.values())
        if passed:
            try:
                validation_copy.replace(workbook)
                checks["validated_copy_committed"] = True
            except OSError as exc:
                checks["validated_copy_committed"] = False
                issues.append(f"validated_copy_commit_failed:{type(exc).__name__}")
                passed = False
        else:
            validation_copy.unlink(missing_ok=True)
        emit_tool_progress(
            "tool.spreadsheet.render.completed",
            (
                f"预览渲染完成，共 {int(metrics.get('rendered_page_count') or 0)} 页"
                if checks.get("rendered_pages_have_content")
                else "预览渲染检查发现问题"
            ),
            status="completed" if passed else "failed",
            formula_count=metrics.get("formula_count_after_recalculation"),
            formula_cache_coverage=metrics.get("formula_cache_coverage"),
            rendered_page_count=metrics.get("rendered_page_count"),
            rendered_content_page_count=metrics.get("rendered_content_page_count"),
        )
        return ExcelVerificationReport(
            passed=passed,
            checks=checks,
            metrics=metrics,
            issues=list(dict.fromkeys(issues)),
            preview_files=[str(preview)] if preview.is_file() else [],
        )

    @staticmethod
    def _inspect_formula_cache(
        path: Path,
        checks: dict[str, bool],
        metrics: dict[str, int | float | str | bool],
        issues: list[str],
    ) -> bool:
        formulas = load_workbook(path, data_only=False, keep_links=False, read_only=True)
        values = load_workbook(path, data_only=True, keep_links=False, read_only=True)
        formula_count = 0
        cached_count = 0
        error_count = 0
        errors = {"#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A"}
        try:
            for sheet_name in formulas.sheetnames:
                formula_sheet = formulas[sheet_name]
                value_sheet = values[sheet_name]
                for row_number, row in enumerate(formula_sheet.iter_rows(), start=1):
                    for column_number, cell in enumerate(row, start=1):
                        if not (isinstance(cell.value, str) and cell.value.startswith("=")):
                            continue
                        formula_count += 1
                        # Read-only worksheets can yield EmptyCell placeholders
                        # for ragged rows; those placeholders intentionally have
                        # no ``coordinate`` attribute.  Use the iteration position
                        # so validation remains robust for any valid workbook shape.
                        cached = value_sheet.cell(row=row_number, column=column_number).value
                        if cached is not None:
                            cached_count += 1
                        if isinstance(cached, str) and cached.upper() in errors:
                            error_count += 1
        finally:
            formulas.close()
            values.close()
        coverage = 1.0 if formula_count == 0 else cached_count / formula_count
        metrics["formula_count_after_recalculation"] = formula_count
        metrics["formula_cached_value_count"] = cached_count
        metrics["formula_cache_coverage"] = round(coverage, 4)
        metrics["formula_error_count_after_recalculation"] = error_count
        checks["formula_results_have_no_errors"] = error_count == 0
        checks["formula_cache_populated"] = formula_count == 0 or coverage >= 0.95
        if not checks["formula_cache_populated"]:
            issues.append("formula_cache_incomplete")
        if error_count:
            issues.append("formula_result_error")
        return True

    @staticmethod
    def _powershell_script(workbook: Path, preview: Path) -> str:
        workbook_literal = str(workbook).replace("'", "''")
        preview_literal = str(preview).replace("'", "''")
        return f"""
$ErrorActionPreference = 'Stop'
$excel = $null
$book = $null
try {{
  $excel = New-Object -ComObject Excel.Application
  $excel.Visible = $false
  $excel.DisplayAlerts = $false
  $excel.AskToUpdateLinks = $false
  $excel.AutomationSecurity = 3
  $book = $excel.Workbooks.Open('{workbook_literal}', 0, $false)
  $excel.Calculation = -4105
  $excel.CalculateFullRebuild()
  $book.Save()
  $book.ExportAsFixedFormat(0, '{preview_literal}', 0, $true, $false)
  [pscustomobject]@{{ excel_version = [string]$excel.Version; rendered = (Test-Path -LiteralPath '{preview_literal}') }} | ConvertTo-Json -Compress
}} finally {{
  if ($book -ne $null) {{ $book.Close($false); [void][Runtime.InteropServices.Marshal]::ReleaseComObject($book) }}
  if ($excel -ne $null) {{ $excel.Quit(); [void][Runtime.InteropServices.Marshal]::ReleaseComObject($excel) }}
  [GC]::Collect()
  [GC]::WaitForPendingFinalizers()
}}
"""
