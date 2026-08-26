from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from uuid import uuid4

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.domains.task_briefs import DocumentStyleSpec


@dataclass(frozen=True)
class DraftSection:
    heading: str
    paragraphs: list[str] = field(default_factory=list)


@dataclass(frozen=True)
class DraftTable:
    title: str
    headers: list[str]
    rows: list[list[str]] = field(default_factory=list)
    blank_rows: int = 0
    column_widths_cm: list[float] = field(default_factory=list)
    alignment: str = "center"
    style: str = "table_grid"


@dataclass(frozen=True)
class DocumentDraft:
    title: str
    document_type: str = "公文草稿"
    addressee: str = ""
    lead: str = ""
    sections: list[DraftSection] = field(default_factory=list)
    tables: list[DraftTable] = field(default_factory=list)
    closing: str = ""
    issuer: str = ""
    date_text: str = ""


class DocumentArtifactBuilder:
    """Create a reviewable Chinese business document behind one interface."""

    allowed_fonts = {
        "方正小标宋简体", "仿宋_GB2312", "仿宋", "宋体", "微软雅黑", "黑体", "楷体", "Arial", "Calibri"
    }

    def __init__(self, output_dir: str | Path) -> None:
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def build(
        self,
        draft: DocumentDraft,
        *,
        body_font: str = "仿宋_GB2312",
        body_size_pt: float = 16,
        style: DocumentStyleSpec | None = None,
    ) -> tuple[Path, dict[str, object]]:
        style = style or DocumentStyleSpec(body_font=body_font, body_size_pt=body_size_pt)
        font = style.body_font if style.body_font in self.allowed_fonts else "仿宋_GB2312"
        title_font = style.title_font if style.title_font in self.allowed_fonts else "方正小标宋简体"
        size = min(22.0, max(10.5, float(style.body_size_pt)))
        document = Document()
        section = document.sections[0]
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(style.top_margin_cm)
        section.bottom_margin = Cm(style.bottom_margin_cm)
        section.left_margin = Cm(style.left_margin_cm)
        section.right_margin = Cm(style.right_margin_cm)
        section.header_distance = Cm(1.5)
        section.footer_distance = Cm(1.5)

        styles = document.styles
        normal = styles["Normal"]
        _set_style_font(normal, font, size)
        normal.paragraph_format.first_line_indent = Cm(style.first_line_indent_cm)
        normal.paragraph_format.line_spacing = style.line_spacing
        normal.paragraph_format.space_before = Pt(0)
        normal.paragraph_format.space_after = Pt(0)
        normal.paragraph_format.alignment = _alignment(style.body_alignment)

        for style_name, style_font, style_size in (
            ("Heading 1", "黑体", min(18, size)),
            ("Heading 2", "楷体", min(17, size)),
        ):
            heading_style = styles[style_name]
            _set_style_font(heading_style, style_font, style_size, bold=True)
            heading_style.font.color.rgb = RGBColor(0, 0, 0)
            heading_style.paragraph_format.space_before = Pt(8)
            heading_style.paragraph_format.space_after = Pt(4)
            heading_style.paragraph_format.keep_with_next = True

        if "Document Meta" not in [style.name for style in styles]:
            meta_style = styles.add_style("Document Meta", WD_STYLE_TYPE.PARAGRAPH)
        else:
            meta_style = styles["Document Meta"]
        _set_style_font(meta_style, "楷体", 11)
        meta_style.font.color.rgb = RGBColor(89, 96, 105)
        meta_style.paragraph_format.space_after = Pt(10)

        title = document.add_paragraph()
        title.alignment = _alignment(style.title_alignment)
        title.paragraph_format.space_before = Pt(0)
        title.paragraph_format.space_after = Pt(18)
        title.paragraph_format.keep_with_next = True
        _add_run(title, draft.title.strip() or "未命名公文", title_font, style.title_size_pt, bold=False)

        if draft.addressee.strip():
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.first_line_indent = Cm(0)
            _add_run(paragraph, f"{draft.addressee.strip()}：", font, size)
        if draft.lead.strip():
            _body_paragraph(document, draft.lead, font, size, style)
        for item in draft.sections:
            if item.heading.strip():
                document.add_paragraph(item.heading.strip(), style="Heading 1")
            for text in item.paragraphs:
                if text.strip():
                    _body_paragraph(document, text, font, size, style)
        for item in draft.tables:
            _add_table(document, item, font, size)
        if draft.closing.strip():
            _body_paragraph(document, draft.closing, font, size, style)
        if draft.issuer.strip() or draft.date_text.strip():
            signature = document.add_paragraph()
            signature.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            signature.paragraph_format.space_before = Pt(18)
            signature.paragraph_format.first_line_indent = Cm(0)
            if draft.issuer.strip():
                _add_run(signature, draft.issuer.strip(), font, size)
            if draft.date_text.strip():
                signature.add_run("\n")
                _add_run(signature, draft.date_text.strip(), font, size)

        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_run(footer, style.footer_text, "微软雅黑", 9)
        footer.runs[0].font.color.rgb = RGBColor(128, 128, 128)

        output = self.output_dir / f"doc_{uuid4().hex}.docx"
        document.save(output)
        validation = _validate_document(output, style, body_font=font, title_font=title_font)
        return output, {
            "preset": style.preset,
            "named_override": "chinese_official_a4",
            "body_font": font,
            "body_size_pt": size,
            "title_font": title_font,
            "title_size_pt": style.title_size_pt,
            "body_alignment": style.body_alignment,
            "line_spacing": style.line_spacing,
            "page_size": "A4",
            "requires_human_review": False,
            "review_recommended": True,
            "validation": validation,
            "table_count": len(document.tables),
        }


def _set_style_font(style: object, font: str, size: float, bold: bool = False) -> None:
    style.font.name = font  # type: ignore[attr-defined]
    style.font.size = Pt(size)  # type: ignore[attr-defined]
    style.font.bold = bold  # type: ignore[attr-defined]
    rpr = style.element.get_or_add_rPr()  # type: ignore[attr-defined]
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(key), font)


def _add_run(paragraph: object, text: str, font: str, size: float, bold: bool = False):
    run = paragraph.add_run(text)  # type: ignore[attr-defined]
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
    return run


def _body_paragraph(
    document: Document,
    text: str,
    font: str,
    size: float,
    style: DocumentStyleSpec,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(style.first_line_indent_cm)
    paragraph.paragraph_format.line_spacing = style.line_spacing
    paragraph.alignment = _alignment(style.body_alignment)
    _add_run(paragraph, " ".join(text.split()), font, size)


def _add_table(
    document: Document,
    draft: DraftTable,
    font: str,
    size: float,
) -> None:
    headers = [value.strip() for value in draft.headers if value.strip()]
    if not headers:
        return
    if draft.title.strip():
        title = document.add_paragraph()
        title.paragraph_format.first_line_indent = Cm(0)
        title.paragraph_format.space_before = Pt(8)
        title.paragraph_format.space_after = Pt(4)
        _add_run(title, draft.title.strip(), "黑体", min(16, size), bold=True)

    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid" if draft.style == "table_grid" else "Light Grid Accent 1"
    table.alignment = {
        "left": WD_TABLE_ALIGNMENT.LEFT,
        "right": WD_TABLE_ALIGNMENT.RIGHT,
    }.get(draft.alignment, WD_TABLE_ALIGNMENT.CENTER)
    usable_width_cm = section_usable_width_cm(document.sections[-1])
    fitted_widths = _fit_column_widths(
        draft.column_widths_cm,
        column_count=len(headers),
        usable_width_cm=usable_width_cm,
    )
    table.autofit = not bool(fitted_widths)
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "D9EAF7")
        cell._tc.get_or_add_tcPr().append(shading)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Cm(0)
        _add_run(paragraph, header, "黑体", min(size, 12), bold=True)

    normalized_rows = [
        (list(row[:len(headers)]) + [""] * len(headers))[:len(headers)]
        for row in draft.rows
    ]
    normalized_rows.extend([[""] * len(headers) for _ in range(draft.blank_rows)])
    for values in normalized_rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            _add_run(paragraph, str(value), font, min(size, 12))

    if fitted_widths:
        _apply_fixed_table_widths(table, fitted_widths)
    document.add_paragraph()


def section_usable_width_cm(section: object) -> float:
    return max(
        0.0,
        section.page_width.cm - section.left_margin.cm - section.right_margin.cm,
    )


def table_width_cm(table: object) -> float:
    if not table.rows:
        return 0.0
    widths = [cell.width.cm for cell in table.rows[0].cells if cell.width is not None]
    return sum(widths) if len(widths) == len(table.rows[0].cells) else 0.0


def _fit_column_widths(
    requested: list[float],
    *,
    column_count: int,
    usable_width_cm: float,
) -> list[float]:
    if column_count <= 0 or len(requested) != column_count or usable_width_cm <= 0:
        return []
    widths = [max(0.5, min(8.0, float(width))) for width in requested]
    target_width = max(0.5, usable_width_cm - 0.05)
    total = sum(widths)
    if total <= target_width:
        return widths
    scale = target_width / total
    return [width * scale for width in widths]


def _apply_fixed_table_widths(table: object, widths_cm: list[float]) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    total_twips = sum(Cm(width).twips for width in widths_cm)
    tbl_width = tbl_pr.find(qn("w:tblW"))
    if tbl_width is None:
        tbl_width = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_width)
    tbl_width.set(qn("w:type"), "dxa")
    tbl_width.set(qn("w:w"), str(total_twips))

    grid_columns = list(table._tbl.tblGrid.gridCol_lst)
    for index, width in enumerate(widths_cm):
        width_value = Cm(width)
        table.columns[index].width = width_value
        if index < len(grid_columns):
            grid_columns[index].set(qn("w:w"), str(width_value.twips))
        for row in table.rows:
            row.cells[index].width = width_value


def _alignment(value: str):
    return {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }.get(value, WD_ALIGN_PARAGRAPH.JUSTIFY)


def _validate_document(
    path: Path,
    style: DocumentStyleSpec,
    *,
    body_font: str,
    title_font: str,
) -> dict[str, object]:
    document = Document(path)
    section = document.sections[0]
    normal = document.styles["Normal"]
    title = document.paragraphs[0] if document.paragraphs else None
    title_run = title.runs[0] if title and title.runs else None
    usable_width_cm = section_usable_width_cm(section)
    table_widths_cm = [table_width_cm(table) for table in document.tables]
    checks = {
        "file_reopens": True,
        "page_is_a4": round(section.page_width.cm, 1) == 21.0 and round(section.page_height.cm, 1) == 29.7,
        "margins_match": all(abs(actual - expected) <= 0.11 for actual, expected in (
            (section.top_margin.cm, style.top_margin_cm),
            (section.bottom_margin.cm, style.bottom_margin_cm),
            (section.left_margin.cm, style.left_margin_cm),
            (section.right_margin.cm, style.right_margin_cm),
        )),
        "title_present": bool(title and title.text.strip()),
        "title_alignment_matches": bool(
            title and title.alignment == _alignment(style.title_alignment)
        ),
        "title_font_matches": bool(title_run and title_run.font.name == title_font),
        "title_size_matches": bool(
            title_run and title_run.font.size
            and abs(title_run.font.size.pt - style.title_size_pt) <= 0.1
        ),
        "body_font_matches": normal.font.name == body_font,
        "body_size_matches": bool(
            normal.font.size and abs(normal.font.size.pt - style.body_size_pt) <= 0.1
        ),
        "body_alignment_matches": normal.paragraph_format.alignment == _alignment(style.body_alignment),
        "line_spacing_matches": bool(
            isinstance(normal.paragraph_format.line_spacing, (int, float))
            and abs(float(normal.paragraph_format.line_spacing) - style.line_spacing) <= 0.01
        ),
        "tables_within_page_width": all(
            not width or width <= usable_width_cm + 0.05
            for width in table_widths_cm
        ),
    }
    return {
        "passed": all(checks.values()),
        "checks": checks,
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "usable_width_cm": round(usable_width_cm, 3),
        "max_table_width_cm": round(max(table_widths_cm, default=0.0), 3),
    }
