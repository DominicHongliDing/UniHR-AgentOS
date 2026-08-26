from __future__ import annotations

import re
import unicodedata
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formula.translate import Translator, TranslatorError
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.utils.datetime import from_excel

from .spreadsheet_plans import (
    AddChartOperation,
    AddDataValidationOperation,
    AddFormulaColumnOperation,
    CopySheetOperation,
    CreateSheetOperation,
    CreateSummaryOperation,
    FilterDateRangeOperation,
    FilterRowsOperation,
    FormatRangeOperation,
    MapColumnOperation,
    RenameSheetOperation,
    ReplaceYearTextOperation,
    SetCellsOperation,
    SetAutoFilterOperation,
    SortRowsOperation,
    SpreadsheetPlan,
)


_FORMULA_FUNCTION = re.compile(r"(?<![A-Z0-9_.])([A-Z][A-Z0-9_.]*)\s*\(", re.IGNORECASE)
_FORMULA_SHEET_REFERENCE = re.compile(r"(?:'([^']+)'|([A-Za-z0-9_\u3400-\u9fff]+))!")
_FORBIDDEN_FORMULA_MARKERS = (
    "[",
    "]",
    "HTTP:",
    "HTTPS:",
    "WEBSERVICE(",
    "HYPERLINK(",
    "RTD(",
    "DDE",
    "CALL(",
    "REGISTER.ID(",
    "EXEC(",
    "SHELL(",
    "INDIRECT(",
    "OFFSET(",
    "#REF!",
)
_ALLOWED_FORMULA_FUNCTIONS = {
    "ABS",
    "AND",
    "AVERAGE",
    "AVERAGEIF",
    "COUNT",
    "COUNTA",
    "COUNTIF",
    "COUNTIFS",
    "DATE",
    "DAY",
    "IF",
    "IFERROR",
    "IFNA",
    "INDEX",
    "ISBLANK",
    "ISERROR",
    "LEFT",
    "LEN",
    "LOWER",
    "MATCH",
    "MAX",
    "MAXIFS",
    "MID",
    "MIN",
    "MINIFS",
    "MONTH",
    "NA",
    "NOT",
    "OR",
    "RIGHT",
    "ROUND",
    "ROUNDDOWN",
    "ROUNDUP",
    "SUM",
    "SUMIF",
    "SUMIFS",
    "TEXT",
    "TRIM",
    "UPPER",
    "VALUE",
    "VLOOKUP",
    "XLOOKUP",
    "YEAR",
}
_FORMULA_ERROR_VALUES = {"#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A"}


class SpreadsheetPlanError(ValueError):
    pass


@dataclass(frozen=True)
class SpreadsheetExecutionResult:
    operation_results: list[dict[str, Any]]
    exceptions: list[dict[str, Any]]
    validation: dict[str, Any]
    changed_cell_count: int


class SpreadsheetPlanExecutor:
    """Execute a small allowlisted workbook plan and expose a reviewable audit trail."""

    def execute(self, workbook: Workbook, plan: SpreadsheetPlan) -> SpreadsheetExecutionResult:
        self._validate_plan_against_workbook(workbook, plan)
        results: list[dict[str, Any]] = []
        exceptions: list[dict[str, Any]] = []
        changed_cells = 0

        for index, operation in enumerate(plan.operations, start=1):
            if isinstance(operation, CreateSheetOperation):
                affected, summary = self._create_sheet(workbook, operation)
            elif isinstance(operation, CopySheetOperation):
                affected, summary = self._copy_sheet(workbook, operation)
            elif isinstance(operation, RenameSheetOperation):
                affected, summary = self._rename_sheet(workbook, operation)
            elif isinstance(operation, ReplaceYearTextOperation):
                affected, summary = self._replace_year_text(workbook, operation)
            elif isinstance(operation, SetCellsOperation):
                affected, summary = self._set_cells(workbook, operation)
            elif isinstance(operation, MapColumnOperation):
                affected, summary, mapping_exceptions = self._map_column(workbook, operation, index)
                exceptions.extend(mapping_exceptions)
            elif isinstance(operation, AddFormulaColumnOperation):
                affected, summary = self._add_formula_column(workbook, operation)
            elif isinstance(operation, FormatRangeOperation):
                affected, summary = self._format_range(workbook, operation)
            elif isinstance(operation, SetAutoFilterOperation):
                affected, summary = self._set_auto_filter(workbook, operation)
            elif isinstance(operation, FilterRowsOperation):
                affected, summary = self._filter_rows(workbook, operation)
            elif isinstance(operation, FilterDateRangeOperation):
                affected, summary, date_exceptions = self._filter_date_range(
                    workbook, operation, index,
                )
                exceptions.extend(date_exceptions)
            elif isinstance(operation, SortRowsOperation):
                affected, summary = self._sort_rows(workbook, operation)
            elif isinstance(operation, AddDataValidationOperation):
                affected, summary = self._add_data_validation(workbook, operation)
            elif isinstance(operation, CreateSummaryOperation):
                affected, summary = self._create_summary(workbook, operation)
            elif isinstance(operation, AddChartOperation):
                affected, summary = self._add_chart(workbook, operation)
            else:  # pragma: no cover - Pydantic's discriminated union prevents this.
                raise SpreadsheetPlanError("spreadsheet_operation_not_supported")
            changed_cells += affected
            results.append({
                "index": index,
                "type": operation.type,
                "status": "completed",
                "affected_cells": affected,
                "summary": summary,
            })

        exceptions.extend(self._collect_data_quality_exceptions(workbook, plan))
        self._write_exception_sheet(workbook, exceptions)
        self._write_operation_audit_sheet(workbook, results)
        validation = self._validate_output(workbook, plan, exceptions)
        return SpreadsheetExecutionResult(results, exceptions, validation, changed_cells)

    def _validate_plan_against_workbook(self, workbook: Workbook, plan: SpreadsheetPlan) -> None:
        known_sheets = set(workbook.sheetnames)
        simulated_sheets = set(known_sheets)
        summary_targets: dict[str, tuple[str, str, int]] = {}
        for operation in plan.operations:
            if isinstance(operation, CreateSheetOperation):
                if operation.sheet in simulated_sheets:
                    raise SpreadsheetPlanError("spreadsheet_target_sheet_exists")
                simulated_sheets.add(operation.sheet)
            elif isinstance(operation, CopySheetOperation):
                if operation.source_sheet not in simulated_sheets:
                    raise SpreadsheetPlanError("spreadsheet_source_sheet_not_found")
                if operation.target_sheet in simulated_sheets:
                    raise SpreadsheetPlanError("spreadsheet_target_sheet_exists")
                simulated_sheets.add(operation.target_sheet)
            elif isinstance(operation, RenameSheetOperation):
                if operation.source_sheet not in simulated_sheets:
                    raise SpreadsheetPlanError("spreadsheet_source_sheet_not_found")
                if operation.target_sheet in simulated_sheets:
                    raise SpreadsheetPlanError("spreadsheet_target_sheet_exists")
                simulated_sheets.remove(operation.source_sheet)
                simulated_sheets.add(operation.target_sheet)
            elif isinstance(operation, MapColumnOperation):
                if operation.source_sheet not in simulated_sheets or operation.target_sheet not in simulated_sheets:
                    raise SpreadsheetPlanError("spreadsheet_mapping_sheet_not_found")
            elif isinstance(operation, CreateSummaryOperation):
                if operation.source_sheet not in simulated_sheets:
                    raise SpreadsheetPlanError("spreadsheet_source_sheet_not_found")
                if operation.target_sheet in simulated_sheets:
                    signature = summary_targets.get(operation.target_sheet)
                    expected = (operation.source_sheet, operation.group_header, operation.source_header_row)
                    if signature != expected:
                        raise SpreadsheetPlanError("spreadsheet_target_sheet_exists")
                else:
                    simulated_sheets.add(operation.target_sheet)
                    summary_targets[operation.target_sheet] = (
                        operation.source_sheet, operation.group_header, operation.source_header_row,
                    )
            else:
                sheet_name = getattr(operation, "sheet", "")
                if sheet_name not in simulated_sheets:
                    raise SpreadsheetPlanError("spreadsheet_operation_sheet_not_found")

            if isinstance(operation, AddFormulaColumnOperation):
                self._validate_formula(operation.formula_template, simulated_sheets)

    @staticmethod
    def _create_sheet(workbook: Workbook, operation: CreateSheetOperation) -> tuple[int, str]:
        sheet = workbook.create_sheet(operation.sheet)
        sheet.append(operation.headers)
        for row in operation.rows:
            sheet.append(row)
        affected = len(operation.headers) + sum(
            1 for row in operation.rows for value in row if value is not None
        )
        return affected, f"创建工作表 {operation.sheet}，包含 {len(operation.headers)} 列、{len(operation.rows)} 行数据"

    @staticmethod
    def _copy_sheet(workbook: Workbook, operation: CopySheetOperation) -> tuple[int, str]:
        copied = workbook.copy_worksheet(workbook[operation.source_sheet])
        copied.title = operation.target_sheet
        return 0, f"复制工作表 {operation.source_sheet} 为 {operation.target_sheet}"

    @staticmethod
    def _rename_sheet(workbook: Workbook, operation: RenameSheetOperation) -> tuple[int, str]:
        workbook[operation.source_sheet].title = operation.target_sheet
        return 0, f"重命名工作表 {operation.source_sheet} 为 {operation.target_sheet}"

    @staticmethod
    def _replace_year_text(workbook: Workbook, operation: ReplaceYearTextOperation) -> tuple[int, str]:
        sheet = workbook[operation.sheet]
        source_year = str(operation.from_year)
        target_year = str(operation.to_year)
        changed = 0
        for row in sheet.iter_rows():
            for cell in row:
                if not isinstance(cell.value, str):
                    continue
                if cell.value.startswith("=") and not operation.include_formulas:
                    continue
                if source_year in cell.value:
                    cell.value = cell.value.replace(source_year, target_year)
                    changed += 1
        return changed, f"将 {operation.sheet} 中的 {source_year} 年文本更新为 {target_year} 年"

    @staticmethod
    def _set_cells(workbook: Workbook, operation: SetCellsOperation) -> tuple[int, str]:
        sheet = workbook[operation.sheet]
        for item in operation.cells:
            cell = sheet[item.cell]
            cell.value = item.value
            if item.number_format:
                cell.number_format = item.number_format
        return len(operation.cells), f"在 {operation.sheet} 写入 {len(operation.cells)} 个已确认参数单元格"

    def _map_column(
        self,
        workbook: Workbook,
        operation: MapColumnOperation,
        operation_index: int,
    ) -> tuple[int, str, list[dict[str, Any]]]:
        source = workbook[operation.source_sheet]
        target = workbook[operation.target_sheet]
        source_headers = self._header_columns(source, operation.source_header_row)
        target_headers = self._header_columns(target, operation.target_header_row)
        source_key_column = self._require_header(source_headers, operation.source_key_header)
        source_value_column = self._require_header(source_headers, operation.source_value_header)
        target_key_column = self._require_header(target_headers, operation.target_key_header)
        target_value_column = target_headers.get(operation.target_header)
        if target_value_column is None:
            target_value_column = target.max_column + 1
            target.cell(operation.target_header_row, target_value_column, operation.target_header)

        values_by_key: dict[str, Any] = {}
        duplicate_keys: set[str] = set()
        for row in range(operation.source_header_row + 1, source.max_row + 1):
            raw_key = source.cell(row, source_key_column).value
            key = self._normalize_key(raw_key, operation.normalize_key)
            if not key:
                continue
            if key in values_by_key:
                duplicate_keys.add(key)
                continue
            values_by_key[key] = source.cell(row, source_value_column).value

        exceptions: list[dict[str, Any]] = []
        for key in sorted(duplicate_keys):
            exceptions.append(self._exception(
                operation_index,
                "duplicate_source_key",
                operation.source_sheet,
                key,
                f"来源键在字段“{operation.source_key_header}”中重复，未自动映射",
            ))

        changed = 0
        for row in range(operation.target_header_row + 1, target.max_row + 1):
            raw_key = target.cell(row, target_key_column).value
            key = self._normalize_key(raw_key, operation.normalize_key)
            if not key:
                continue
            if key in duplicate_keys:
                continue
            if key not in values_by_key:
                exceptions.append(self._exception(
                    operation_index,
                    "unmatched_target_key",
                    operation.target_sheet,
                    key,
                    f"目标键未在 {operation.source_sheet} 中找到",
                ))
                if operation.missing_behavior == "blank_and_exception":
                    target.cell(row, target_value_column).value = None
                continue
            target.cell(row, target_value_column).value = values_by_key[key]
            changed += 1
        return changed, f"按 {operation.target_key_header} 映射 {operation.target_header}，成功 {changed} 行", exceptions

    def _add_formula_column(
        self,
        workbook: Workbook,
        operation: AddFormulaColumnOperation,
    ) -> tuple[int, str]:
        sheet = workbook[operation.sheet]
        headers = self._header_columns(sheet, operation.header_row)
        column = headers.get(operation.header)
        if column is None:
            column = sheet.max_column + 1
            sheet.cell(operation.header_row, column, operation.header)
        start_row = operation.start_row or operation.header_row + 1
        end_row = operation.end_row or sheet.max_row
        if end_row < start_row:
            return 0, f"{operation.sheet} 没有可填充的数据行"
        if end_row - start_row + 1 > 100_000:
            raise SpreadsheetPlanError("spreadsheet_formula_fill_too_large")
        if not operation.overwrite_existing:
            for row in range(start_row, end_row + 1):
                if sheet.cell(row, column).value not in (None, ""):
                    raise SpreadsheetPlanError("spreadsheet_formula_target_column_has_data")
        changed = 0
        for row in range(start_row, end_row + 1):
            formula = operation.formula_template.replace("{row}", str(row))
            self._validate_formula(formula, set(workbook.sheetnames))
            cell = sheet.cell(row, column)
            cell.value = formula
            if operation.number_format:
                cell.number_format = operation.number_format
            changed += 1
        return changed, f"在 {operation.sheet} 的“{operation.header}”列填充 {changed} 行公式"

    @staticmethod
    def _format_range(workbook: Workbook, operation: FormatRangeOperation) -> tuple[int, str]:
        sheet = workbook[operation.sheet]
        cells = sheet[operation.range]
        matrix = cells if isinstance(cells, tuple) else ((cells,),)
        flat = [cell for row in matrix for cell in (row if isinstance(row, tuple) else (row,))]
        if len(flat) > 100_000:
            raise SpreadsheetPlanError("spreadsheet_format_range_too_large")
        for cell in flat:
            font = copy(cell.font)
            if operation.font_name is not None:
                font.name = operation.font_name
            if operation.font_size is not None:
                font.sz = operation.font_size
            if operation.bold is not None:
                font.b = operation.bold
            if operation.italic is not None:
                font.i = operation.italic
            if operation.font_color is not None:
                font.color = operation.font_color
            cell.font = font
            if operation.fill_color is not None:
                cell.fill = PatternFill("solid", fgColor=operation.fill_color)
            if operation.number_format is not None:
                cell.number_format = operation.number_format
            alignment = copy(cell.alignment)
            if operation.horizontal_alignment is not None:
                alignment.horizontal = operation.horizontal_alignment
            if operation.vertical_alignment is not None:
                alignment.vertical = operation.vertical_alignment
            if operation.wrap_text is not None:
                alignment.wrap_text = operation.wrap_text
            cell.alignment = alignment
        return len(flat), f"设置 {operation.sheet}!{operation.range} 的版式与数字格式"

    @staticmethod
    def _set_auto_filter(workbook: Workbook, operation: SetAutoFilterOperation) -> tuple[int, str]:
        sheet = workbook[operation.sheet]
        sheet.auto_filter.ref = operation.range
        return 0, f"为 {operation.sheet}!{operation.range} 启用筛选"

    def _filter_rows(self, workbook: Workbook, operation: FilterRowsOperation) -> tuple[int, str]:
        sheet = workbook[operation.sheet]
        headers = self._header_columns(sheet, operation.header_row)
        column = self._require_header(headers, operation.header)
        allowed = {
            self._filter_key(value, operation.case_sensitive)
            for value in operation.allowed_values
        }
        changed = 0
        for row in range(operation.header_row + 1, sheet.max_row + 1):
            should_hide = self._filter_key(sheet.cell(row, column).value, operation.case_sensitive) not in allowed
            if bool(sheet.row_dimensions[row].hidden) != should_hide:
                sheet.row_dimensions[row].hidden = should_hide
                changed += 1
        sheet.auto_filter.ref = f"A{operation.header_row}:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}"
        return changed, f"按“{operation.header}”筛选 {operation.sheet}，显示指定值对应的数据行"

    def _filter_date_range(
        self,
        workbook: Workbook,
        operation: FilterDateRangeOperation,
        operation_index: int,
    ) -> tuple[int, str, list[dict[str, Any]]]:
        sheet = workbook[operation.sheet]
        headers = self._header_columns(sheet, operation.header_row)
        column = self._require_header(headers, operation.header)
        changed = 0
        matched = 0
        invalid = 0
        exceptions: list[dict[str, Any]] = []
        for row in range(operation.header_row + 1, sheet.max_row + 1):
            raw_value = sheet.cell(row, column).value
            parsed = self._coerce_date(raw_value, workbook.epoch)
            if parsed is None:
                invalid += 1
                should_hide = operation.invalid_date_policy == "exclude_and_report"
                exceptions.append(self._exception(
                    operation_index,
                    "date_filter_invalid_date",
                    operation.sheet,
                    f"第{row}行",
                    f"“{operation.header}”的值“{str(raw_value)[:120]}”无法识别为日期，已"
                    + ("排除" if should_hide else "保留")
                    + "并等待人工复核",
                ))
            else:
                in_range = operation.start_date <= parsed <= operation.end_date
                matched += int(in_range)
                should_hide = not in_range
            if bool(sheet.row_dimensions[row].hidden) != should_hide:
                sheet.row_dimensions[row].hidden = should_hide
                changed += sheet.max_column
        sheet.auto_filter.ref = (
            f"A{operation.header_row}:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}"
        )
        summary = (
            f"按“{operation.header}”筛选 {operation.sheet}："
            f"{operation.start_date.isoformat()} 至 {operation.end_date.isoformat()}（含首尾），"
            f"匹配 {matched} 行，日期异常 {invalid} 行，其余行已隐藏"
        )
        return changed, summary, exceptions

    @staticmethod
    def _coerce_date(value: Any, epoch: datetime) -> date | None:
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            try:
                converted = from_excel(value, epoch)
                return converted.date() if isinstance(converted, datetime) else converted
            except (TypeError, ValueError, OverflowError):
                return None
        text = unicodedata.normalize("NFKC", str(value if value is not None else "")).strip()
        if not text:
            return None
        text = re.sub(r"[年月/]", "-", text).replace("日", "")
        text = text.split("T", 1)[0].split(" ", 1)[0].strip()
        for pattern in ("%Y-%m-%d", "%Y-%m", "%Y%m%d"):
            try:
                return datetime.strptime(text, pattern).date()
            except ValueError:
                continue
        return None

    @staticmethod
    def _filter_key(value: Any, case_sensitive: bool) -> str:
        text = unicodedata.normalize("NFKC", str(value if value is not None else "")).strip()
        return text if case_sensitive else text.casefold()

    def _sort_rows(self, workbook: Workbook, operation: SortRowsOperation) -> tuple[int, str]:
        sheet = workbook[operation.sheet]
        headers = self._header_columns(sheet, operation.header_row)
        column = self._require_header(headers, operation.header)
        start_row = operation.start_row or operation.header_row + 1
        end_row = operation.end_row or sheet.max_row
        if end_row < start_row:
            return 0, f"{operation.sheet} 没有可排序的数据行"
        if end_row - start_row + 1 > 100_000:
            raise SpreadsheetPlanError("spreadsheet_sort_range_too_large")
        if any(
            merged.min_row <= end_row and merged.max_row >= start_row
            for merged in sheet.merged_cells.ranges
        ):
            raise SpreadsheetPlanError("spreadsheet_sort_merged_cells_not_supported")
        rows: list[dict[str, Any]] = []
        for row_number in range(start_row, end_row + 1):
            cells = [sheet.cell(row_number, col) for col in range(1, sheet.max_column + 1)]
            sort_value = sheet.cell(row_number, column).value
            if isinstance(sort_value, str) and sort_value.startswith("="):
                sort_value = self._sortable_formula_value(sheet, row_number, sort_value)
            rows.append({
                "key": self._sort_key(sort_value),
                "cells": [self._snapshot_cell(cell) for cell in cells],
                "height": sheet.row_dimensions[row_number].height,
                "hidden": sheet.row_dimensions[row_number].hidden,
            })
        populated = [item for item in rows if item["key"][0] != 2]
        blanks = [item for item in rows if item["key"][0] == 2]
        populated.sort(key=lambda item: item["key"], reverse=operation.direction == "descending")
        rows = [*populated, *blanks]
        for row_number, snapshot in zip(range(start_row, end_row + 1), rows, strict=True):
            for column_number, cell_snapshot in enumerate(snapshot["cells"], start=1):
                self._restore_cell(sheet.cell(row_number, column_number), cell_snapshot)
            sheet.row_dimensions[row_number].height = snapshot["height"]
            sheet.row_dimensions[row_number].hidden = snapshot["hidden"]
        changed = len(rows) * sheet.max_column
        return changed, f"按“{operation.header}”{operation.direction}排序 {len(rows)} 行"

    @staticmethod
    def _sort_key(value: Any) -> tuple[int, Any]:
        if value in (None, ""):
            return (2, "")
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            return (0, float(value))
        return (1, unicodedata.normalize("NFKC", str(value)).casefold())

    @staticmethod
    def _sortable_formula_value(sheet: Worksheet, row_number: int, formula: str) -> float:
        match = re.fullmatch(
            r"=\s*\$?([A-Z]{1,3})\$?(\d+)\s*([+\-*/])\s*\$?([A-Z]{1,3})\$?(\d+)\s*",
            formula.upper(),
        )
        if match is None or int(match.group(2)) != row_number or int(match.group(5)) != row_number:
            raise SpreadsheetPlanError("spreadsheet_sort_formula_rows_not_supported")
        left = sheet[f"{match.group(1)}{row_number}"].value
        right = sheet[f"{match.group(4)}{row_number}"].value
        if left in (None, ""):
            left = 0
        if right in (None, ""):
            right = 0
        if not isinstance(left, (int, float)) or isinstance(left, bool):
            raise SpreadsheetPlanError("spreadsheet_sort_formula_key_not_numeric")
        if not isinstance(right, (int, float)) or isinstance(right, bool):
            raise SpreadsheetPlanError("spreadsheet_sort_formula_key_not_numeric")
        operator = match.group(3)
        if operator == "+":
            return float(left + right)
        if operator == "-":
            return float(left - right)
        if operator == "*":
            return float(left * right)
        if right == 0:
            raise SpreadsheetPlanError("spreadsheet_sort_formula_division_by_zero")
        return float(left / right)

    @staticmethod
    def _snapshot_cell(cell: Any) -> dict[str, Any]:
        return {
            "coordinate": cell.coordinate,
            "value": cell.value,
            "style": copy(cell._style),
            "number_format": cell.number_format,
            "hyperlink": copy(cell.hyperlink),
            "comment": copy(cell.comment),
        }

    @staticmethod
    def _restore_cell(cell: Any, snapshot: dict[str, Any]) -> None:
        value = snapshot["value"]
        if isinstance(value, str) and value.startswith("="):
            try:
                value = Translator(value, origin=snapshot["coordinate"]).translate_formula(cell.coordinate)
            except TranslatorError as exc:
                raise SpreadsheetPlanError("spreadsheet_sort_formula_translation_failed") from exc
        cell.value = value
        cell._style = copy(snapshot["style"])
        cell.number_format = snapshot["number_format"]
        cell._hyperlink = copy(snapshot["hyperlink"])
        cell.comment = copy(snapshot["comment"])

    def _add_data_validation(
        self, workbook: Workbook, operation: AddDataValidationOperation,
    ) -> tuple[int, str]:
        sheet = workbook[operation.sheet]
        headers = self._header_columns(sheet, operation.header_row)
        column = self._require_header(headers, operation.header)
        start_row = operation.start_row or operation.header_row + 1
        end_row = operation.end_row or max(sheet.max_row, start_row + 99)
        if end_row - start_row + 1 > 100_000:
            raise SpreadsheetPlanError("spreadsheet_validation_range_too_large")
        if operation.validation_type == "list":
            escaped = [str(item).replace('"', '""') for item in operation.values]
            formula = f'"{",".join(escaped)}"'
            if len(formula) > 255:
                raise SpreadsheetPlanError("spreadsheet_validation_list_too_long")
            validation = DataValidation(type="list", formula1=formula, allow_blank=operation.allow_blank)
        else:
            if operation.minimum is not None and operation.maximum is not None:
                operator, formula1, formula2 = "between", str(operation.minimum), str(operation.maximum)
            elif operation.minimum is not None:
                operator, formula1, formula2 = "greaterThanOrEqual", str(operation.minimum), None
            else:
                operator, formula1, formula2 = "lessThanOrEqual", str(operation.maximum), None
            validation = DataValidation(
                type=operation.validation_type,
                operator=operator,
                formula1=formula1,
                formula2=formula2,
                allow_blank=operation.allow_blank,
            )
        validation.error = "输入值不符合本列规则"
        validation.errorTitle = "UniHR 数据校验"
        validation.prompt = operation.prompt or "请按本列要求填写"
        validation.promptTitle = operation.header
        validation.showErrorMessage = True
        validation.showInputMessage = True
        sheet.add_data_validation(validation)
        from openpyxl.utils import get_column_letter
        target = f"{get_column_letter(column)}{start_row}:{get_column_letter(column)}{end_row}"
        validation.add(target)
        return end_row - start_row + 1, f"为 {operation.sheet} 的“{operation.header}”添加数据校验"

    def _create_summary(self, workbook: Workbook, operation: CreateSummaryOperation) -> tuple[int, str]:
        source = workbook[operation.source_sheet]
        headers = self._header_columns(source, operation.source_header_row)
        group_column = self._require_header(headers, operation.group_header)
        value_column = (
            self._require_header(headers, operation.value_header)
            if operation.aggregation != "count"
            else None
        )
        categories: list[Any] = []
        seen: set[str] = set()
        for row in range(operation.source_header_row + 1, source.max_row + 1):
            value = source.cell(row, group_column).value
            key = self._normalize_key(value, True)
            if key and key not in seen:
                seen.add(key)
                categories.append(value)
        if len(categories) > 5_000:
            raise SpreadsheetPlanError("spreadsheet_summary_too_many_groups")
        if operation.target_sheet in workbook.sheetnames:
            target = workbook[operation.target_sheet]
            target_headers = self._header_columns(target, 1)
            if operation.group_header not in target_headers:
                raise SpreadsheetPlanError("spreadsheet_summary_group_header_mismatch")
            if operation.result_header in target_headers:
                raise SpreadsheetPlanError("spreadsheet_summary_result_header_exists")
            existing_categories = [
                self._normalize_key(target.cell(row, target_headers[operation.group_header]).value, True)
                for row in range(2, target.max_row + 1)
            ]
            if existing_categories != [self._normalize_key(value, True) for value in categories]:
                raise SpreadsheetPlanError("spreadsheet_summary_categories_mismatch")
            result_column = target.max_column + 1
            target.cell(1, result_column, operation.result_header)
        else:
            target = workbook.create_sheet(operation.target_sheet)
            target.sheet_view.showGridLines = False
            target.append([operation.group_header, operation.result_header])
            result_column = 2
        from openpyxl.utils import get_column_letter
        group_letter = get_column_letter(group_column)
        group_range = f"'{operation.source_sheet}'!${group_letter}${operation.source_header_row + 1}:${group_letter}${source.max_row}"
        value_range = ""
        if value_column is not None:
            value_letter = get_column_letter(value_column)
            value_range = f"'{operation.source_sheet}'!${value_letter}${operation.source_header_row + 1}:${value_letter}${source.max_row}"
        result_number_format = operation.number_format
        if value_column is not None and result_number_format == "0.00":
            inherited_format = source.cell(operation.source_header_row + 1, value_column).number_format
            if inherited_format and inherited_format != "General":
                result_number_format = inherited_format
        function = {
            "count": "COUNTIF",
            "sum": "SUMIF",
            "average": "AVERAGEIF",
            "min": "MINIFS",
            "max": "MAXIFS",
        }[operation.aggregation]
        for row_number, category in enumerate(categories, start=2):
            if target.cell(row_number, 1).value in (None, ""):
                target.cell(row_number, 1, category)
            if operation.aggregation == "count":
                formula = f"={function}({group_range},A{row_number})"
            elif operation.aggregation in {"min", "max"}:
                formula = f"={function}({value_range},{group_range},A{row_number})"
            else:
                formula = f"={function}({group_range},A{row_number},{value_range})"
            self._validate_formula(formula, set(workbook.sheetnames))
            target.cell(row_number, result_column, formula)
            target.cell(row_number, result_column).number_format = result_number_format
        for cell in target[1]:
            cell.font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2563EB")
        target.freeze_panes = "A2"
        from openpyxl.utils import get_column_letter
        target.auto_filter.ref = f"A1:{get_column_letter(target.max_column)}{max(2, target.max_row)}"
        target.column_dimensions["A"].width = 24
        target.column_dimensions[get_column_letter(result_column)].width = 18
        return len(categories) + 1, f"按“{operation.group_header}”生成 {operation.target_sheet} 的“{operation.result_header}”指标"

    def _add_chart(self, workbook: Workbook, operation: AddChartOperation) -> tuple[int, str]:
        sheet = workbook[operation.sheet]
        headers = self._header_columns(sheet, operation.header_row)
        category_column = self._require_header(headers, operation.category_header)
        value_columns = [self._require_header(headers, header) for header in operation.value_headers]
        if operation.chart_type == "pie" and len(value_columns) != 1:
            raise SpreadsheetPlanError("spreadsheet_pie_chart_requires_one_value_column")
        chart: Any
        if operation.chart_type == "line":
            chart = LineChart()
        elif operation.chart_type == "pie":
            chart = PieChart()
        else:
            chart = BarChart()
            chart.type = "bar" if operation.chart_type == "bar" else "col"
        chart.title = operation.title
        chart.style = 10
        chart.height = 8
        chart.width = 14
        categories = Reference(
            sheet, min_col=category_column, min_row=operation.header_row + 1, max_row=sheet.max_row,
        )
        for column in value_columns:
            data = Reference(
                sheet, min_col=column, min_row=operation.header_row,
                max_row=sheet.max_row,
            )
            chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        sheet.add_chart(chart, operation.position)
        return len(value_columns), f"在 {operation.sheet} 添加“{operation.title}”图表"

    @staticmethod
    def _header_columns(sheet: Worksheet, header_row: int) -> dict[str, int]:
        headers: dict[str, int] = {}
        duplicates: set[str] = set()
        for column in range(1, sheet.max_column + 1):
            value = sheet.cell(header_row, column).value
            if value in (None, ""):
                continue
            header = " ".join(str(value).split())
            if header in headers:
                duplicates.add(header)
            else:
                headers[header] = column
        if duplicates:
            raise SpreadsheetPlanError("spreadsheet_duplicate_headers")
        return headers

    @staticmethod
    def _require_header(headers: dict[str, int], header: str) -> int:
        if header not in headers:
            raise SpreadsheetPlanError(f"spreadsheet_header_not_found:{header}")
        return headers[header]

    @staticmethod
    def _normalize_key(value: Any, enabled: bool) -> str:
        if value in (None, ""):
            return ""
        if isinstance(value, float) and value.is_integer():
            text = str(int(value))
        else:
            text = str(value)
        if not enabled:
            return text
        return unicodedata.normalize("NFKC", text).strip()

    @staticmethod
    def _exception(index: int, code: str, sheet: str, key: str, details: str) -> dict[str, Any]:
        return {
            "operation_index": index,
            "code": code,
            "sheet": sheet,
            "key": key,
            "details": details,
        }

    def _collect_data_quality_exceptions(
        self, workbook: Workbook, plan: SpreadsheetPlan,
    ) -> list[dict[str, Any]]:
        issues: list[dict[str, Any]] = []
        for requirement in plan.validations.unique_keys:
            if requirement.sheet not in workbook.sheetnames:
                continue
            sheet = workbook[requirement.sheet]
            headers = self._header_columns(sheet, requirement.header_row)
            column = headers.get(requirement.header)
            if column is None:
                continue
            rows_by_key: dict[str, list[int]] = {}
            for row in range(requirement.header_row + 1, sheet.max_row + 1):
                key = self._normalize_key(sheet.cell(row, column).value, True)
                if not key:
                    if not requirement.allow_blank:
                        issues.append(self._exception(
                            0, "unique_key_blank", requirement.sheet, f"第{row}行",
                            f"“{requirement.header}”不能为空，请人工补充后复核",
                        ))
                    continue
                rows_by_key.setdefault(key, []).append(row)
            for key, row_numbers in rows_by_key.items():
                if len(row_numbers) > 1:
                    issues.append(self._exception(
                        0, "unique_key_duplicate", requirement.sheet, key,
                        f"“{requirement.header}”在第 {', '.join(map(str, row_numbers))} 行重复",
                    ))

        for index, operation in enumerate(plan.operations, start=1):
            if not isinstance(operation, AddDataValidationOperation) or operation.validation_type != "list":
                continue
            if operation.sheet not in workbook.sheetnames:
                continue
            sheet = workbook[operation.sheet]
            headers = self._header_columns(sheet, operation.header_row)
            column = headers.get(operation.header)
            if column is None:
                continue
            allowed = {self._normalize_key(value, True) for value in operation.values}
            start_row = operation.start_row or operation.header_row + 1
            end_row = operation.end_row or sheet.max_row
            for row in range(start_row, min(end_row, sheet.max_row) + 1):
                value = self._normalize_key(sheet.cell(row, column).value, True)
                if value and value not in allowed:
                    issues.append(self._exception(
                        index, "data_validation_existing_value_invalid", operation.sheet, f"第{row}行",
                        f"“{operation.header}”当前值“{value}”不在允许列表中",
                    ))
        return issues

    @staticmethod
    def _write_exception_sheet(workbook: Workbook, exceptions: list[dict[str, Any]]) -> None:
        name = "UniHR异常清单"
        if name in workbook.sheetnames:
            del workbook[name]
        sheet = workbook.create_sheet(name)
        sheet.append(["操作序号", "异常代码", "工作表", "匹配键", "说明"])
        for item in exceptions:
            sheet.append([
                item["operation_index"], item["code"], item["sheet"], item["key"], item["details"],
            ])
        SpreadsheetPlanExecutor._style_audit_sheet(sheet)
        if not exceptions:
            sheet.append(["", "无", "", "", "本次执行未发现字段映射或数据质量异常"])

    @staticmethod
    def _write_operation_audit_sheet(workbook: Workbook, results: list[dict[str, Any]]) -> None:
        name = "UniHR执行记录"
        if name in workbook.sheetnames:
            del workbook[name]
        sheet = workbook.create_sheet(name)
        sheet.append(["操作序号", "操作类型", "状态", "影响单元格", "说明"])
        for item in results:
            sheet.append([
                item["index"], item["type"], item["status"], item["affected_cells"], item["summary"],
            ])
        SpreadsheetPlanExecutor._style_audit_sheet(sheet)
        if not results:
            sheet.append(["", "inspect_only", "completed", 0, "未执行数据修改，仅保留源文件并生成审计记录"])

    @staticmethod
    def _style_audit_sheet(sheet: Worksheet) -> None:
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:E{max(sheet.max_row, 2)}"
        for cell in sheet[1]:
            cell.font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0F766E")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for column, width in {"A": 12, "B": 42, "C": 22, "D": 18, "E": 72}.items():
            sheet.column_dimensions[column].width = width

    def _validate_output(
        self,
        workbook: Workbook,
        plan: SpreadsheetPlan,
        exceptions: list[dict[str, Any]],
    ) -> dict[str, Any]:
        checks: dict[str, bool] = {
            "no_mapping_exceptions": not any(
                item["code"] in {"duplicate_source_key", "unmatched_target_key"} for item in exceptions
            ),
            "no_data_quality_exceptions": not any(
                item["code"].startswith(("unique_key_", "data_validation_", "date_filter_"))
                for item in exceptions
            ),
            "execution_audit_present": "UniHR执行记录" in workbook.sheetnames,
            "exception_sheet_present": "UniHR异常清单" in workbook.sheetnames,
        }
        for sheet_name in plan.validations.required_sheets:
            checks[f"required_sheet:{sheet_name}"] = sheet_name in workbook.sheetnames
        for requirement in plan.validations.required_headers:
            if requirement.sheet not in workbook.sheetnames:
                checks[f"required_headers:{requirement.sheet}"] = False
                continue
            headers = self._header_columns(workbook[requirement.sheet], requirement.header_row)
            checks[f"required_headers:{requirement.sheet}"] = all(
                header in headers for header in requirement.headers
            )
        for requirement in plan.validations.unique_keys:
            key = f"unique_key:{requirement.sheet}.{requirement.header}"
            if requirement.sheet not in workbook.sheetnames:
                checks[key] = False
                continue
            sheet = workbook[requirement.sheet]
            headers = self._header_columns(sheet, requirement.header_row)
            column = headers.get(requirement.header)
            if column is None:
                checks[key] = False
                continue
            values: list[str] = []
            blank_seen = False
            for row in range(requirement.header_row + 1, sheet.max_row + 1):
                value = self._normalize_key(sheet.cell(row, column).value, True)
                if not value:
                    blank_seen = True
                else:
                    values.append(value)
            checks[key] = len(values) == len(set(values)) and (requirement.allow_blank or not blank_seen)
        if plan.validations.reject_formula_errors:
            checks["no_visible_formula_errors"] = not any(
                isinstance(cell.value, str) and cell.value.upper() in _FORMULA_ERROR_VALUES
                for sheet in workbook.worksheets
                for row in sheet.iter_rows()
                for cell in row
            )
        return {"passed": all(checks.values()), "checks": checks}

    @staticmethod
    def _validate_formula(formula: str, workbook_sheets: set[str]) -> None:
        upper = formula.upper()
        if len(formula) > 2_000 or not formula.startswith("="):
            raise SpreadsheetPlanError("spreadsheet_formula_invalid")
        if any(marker in upper for marker in _FORBIDDEN_FORMULA_MARKERS):
            raise SpreadsheetPlanError("spreadsheet_formula_unsafe")
        functions = {match.group(1).upper() for match in _FORMULA_FUNCTION.finditer(formula)}
        unsupported = functions - _ALLOWED_FORMULA_FUNCTIONS
        if unsupported:
            raise SpreadsheetPlanError(f"spreadsheet_formula_function_not_allowed:{sorted(unsupported)[0]}")
        for match in _FORMULA_SHEET_REFERENCE.finditer(formula):
            sheet_name = match.group(1) or match.group(2)
            if sheet_name not in workbook_sheets:
                raise SpreadsheetPlanError(f"spreadsheet_formula_sheet_not_found:{sheet_name}")
