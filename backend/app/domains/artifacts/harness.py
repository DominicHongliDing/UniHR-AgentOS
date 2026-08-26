from __future__ import annotations

from pathlib import Path
from typing import Any, Literal
from zipfile import BadZipFile, ZipFile

from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pydantic import BaseModel, Field
from pypdf import PdfReader

from .excel_verification import ExcelAutomationVerifier


_FORMULA_ERRORS = {"#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A"}


class ArtifactHarnessReport(BaseModel):
    """File-backed artifact verification result used by the runtime quality gate."""

    passed: bool
    status: Literal["passed", "warning", "failed"]
    checks: dict[str, bool] = Field(default_factory=dict)
    metrics: dict[str, int | float | str | bool] = Field(default_factory=dict)
    issues: list[str] = Field(default_factory=list)
    details: dict[str, Any] = Field(default_factory=dict)


class ArtifactHarness:
    """Open the actual generated file and verify its package and domain structure."""

    def __init__(self, excel_verifier: ExcelAutomationVerifier | None = None) -> None:
        self.excel_verifier = excel_verifier

    def inspect(
        self,
        path: str | Path,
        *,
        kind: str,
        mime_type: str,
        metadata: dict[str, object] | None = None,
    ) -> ArtifactHarnessReport:
        artifact = Path(path)
        checks = {
            "file_exists": artifact.is_file(),
            "file_nonempty": artifact.is_file() and artifact.stat().st_size > 0,
        }
        metrics: dict[str, int | float | str | bool] = {
            "kind": kind,
            "mime_type": mime_type,
            "byte_size": artifact.stat().st_size if artifact.is_file() else 0,
        }
        issues: list[str] = []
        details: dict[str, Any] = {}
        if not all(checks.values()):
            return ArtifactHarnessReport(
                passed=False, status="failed", checks=checks, metrics=metrics,
                issues=[name for name, passed in checks.items() if not passed],
            )
        try:
            suffix = artifact.suffix.casefold()
            if suffix == ".xlsx":
                self._inspect_xlsx(artifact, checks, metrics, issues)
                if self.excel_verifier is not None:
                    excel_report = self.excel_verifier.verify(artifact)
                    checks.update(excel_report.checks)
                    metrics.update(excel_report.metrics)
                    issues.extend(excel_report.issues)
                    details["preview_files"] = excel_report.preview_files
            elif suffix == ".docx":
                self._inspect_docx(artifact, checks, metrics)
                preview_files = [
                    str(item) for item in (metadata or {}).get("preview_files", [])
                    if str(item).strip()
                ]
                if preview_files:
                    preview_path = Path(preview_files[-1])
                    checks["document_preview_exists"] = (
                        preview_path.is_file() and preview_path.stat().st_size > 0
                    )
                    if checks["document_preview_exists"]:
                        reader = PdfReader(preview_path)
                        checks["document_preview_has_pages"] = len(reader.pages) > 0
                        metrics["document_preview_page_count"] = len(reader.pages)
                        details["preview_files"] = preview_files
                    else:
                        issues.append("document_preview_missing")
            elif suffix == ".pdf":
                self._inspect_pdf(artifact, checks, metrics)
            elif suffix in {".png", ".jpg", ".jpeg", ".webp"}:
                self._inspect_image(artifact, checks, metrics)
            else:
                checks["package_opens"] = True
        except (BadZipFile, OSError, ValueError, KeyError) as exc:
            checks["package_opens"] = False
            issues.append(f"artifact_open_failed:{type(exc).__name__}")

        declared_validation = dict((metadata or {}).get("validation") or {})
        if declared_validation:
            checks["generation_validation_passed"] = bool(declared_validation.get("passed"))
            if not checks["generation_validation_passed"]:
                issues.append("generation_validation_failed")
        passed = all(checks.values())
        return ArtifactHarnessReport(
            passed=passed,
            status="passed" if passed and not issues else "warning" if passed else "failed",
            checks=checks,
            metrics=metrics,
            issues=list(dict.fromkeys(issues)),
            details=details,
        )

    @staticmethod
    def _inspect_xlsx(
        path: Path,
        checks: dict[str, bool],
        metrics: dict[str, int | float | str | bool],
        issues: list[str],
    ) -> None:
        with ZipFile(path, "r") as package:
            bad_member = package.testzip()
            checks["package_opens"] = bad_member is None
            if bad_member:
                issues.append(f"xlsx_corrupt_member:{bad_member}")
        workbook = load_workbook(path, data_only=False, keep_links=False, read_only=True)
        sheet_count = len(workbook.sheetnames)
        formula_count = 0
        formula_errors = 0
        styled_cells = 0
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if isinstance(value, str) and value.startswith("="):
                        formula_count += 1
                    if isinstance(value, str) and value.upper() in _FORMULA_ERRORS:
                        formula_errors += 1
                    if getattr(cell, "style_id", 0) or getattr(cell, "has_style", False):
                        styled_cells += 1
        calc = workbook.calculation
        checks["workbook_has_sheets"] = sheet_count > 0
        checks["formula_cells_have_no_literal_errors"] = formula_errors == 0
        metrics.update({
            "sheet_count": sheet_count,
            "formula_count": formula_count,
            "literal_formula_error_count": formula_errors,
            "styled_cell_count": styled_cells,
            "full_calc_on_load": bool(calc and calc.fullCalcOnLoad),
        })
        workbook.close()

    @staticmethod
    def _inspect_docx(path: Path, checks: dict[str, bool], metrics: dict[str, int | float | str | bool]) -> None:
        document = Document(path)
        checks["package_opens"] = True
        checks["document_has_content"] = any(paragraph.text.strip() for paragraph in document.paragraphs)
        metrics["paragraph_count"] = len(document.paragraphs)

    @staticmethod
    def _inspect_pdf(path: Path, checks: dict[str, bool], metrics: dict[str, int | float | str | bool]) -> None:
        reader = PdfReader(path)
        checks["package_opens"] = True
        checks["pdf_has_pages"] = len(reader.pages) > 0
        metrics["page_count"] = len(reader.pages)

    @staticmethod
    def _inspect_image(path: Path, checks: dict[str, bool], metrics: dict[str, int | float | str | bool]) -> None:
        with Image.open(path) as image:
            image.verify()
        with Image.open(path) as image:
            checks["package_opens"] = True
            checks["image_has_dimensions"] = image.width > 0 and image.height > 0
            metrics["width"] = image.width
            metrics["height"] = image.height


def attach_harness_report(
    metadata: dict[str, object], report: ArtifactHarnessReport,
) -> dict[str, object]:
    return {**metadata, "artifact_harness": report.model_dump(mode="json")}
