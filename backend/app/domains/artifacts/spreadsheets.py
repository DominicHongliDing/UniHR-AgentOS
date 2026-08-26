from __future__ import annotations

import csv
import io
import re
import unicodedata
from copy import copy
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from shutil import copy2
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.table import Table, TableStyleInfo

from app.domains.task_briefs import SpreadsheetStyleSpec

from .spreadsheet_engine import SpreadsheetExecutionResult, SpreadsheetPlanExecutor
from .spreadsheet_diff import WorkbookDiffEngine, WorkbookSnapshot
from .spreadsheet_plans import SpreadsheetPlan


@dataclass(frozen=True)
class SpreadsheetBuildResult:
    path: Path
    display_name: str
    metadata: dict[str, object]


class SpreadsheetArtifactBuilder:
    """Create a new reviewable workbook without overwriting the source."""

    def __init__(self, output_dir: str | Path) -> None:
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.executor = SpreadsheetPlanExecutor()
        self.diff_engine = WorkbookDiffEngine()

    def inspect_source(self, source_path: str | Path, *, header_row: int = 1) -> dict[str, object]:
        """Return planning-safe workbook structure without exposing row-level HR data."""
        return self.inspect_sources([(source_path, Path(source_path).name)], header_row=header_row)

    def inspect_sources(
        self,
        sources: list[tuple[str | Path, str]],
        *,
        header_row: int = 1,
    ) -> dict[str, object]:
        """Return the projected merged-workbook schema used by the planner."""
        if not sources:
            raise ValueError("spreadsheet_source_required")
        used_names: set[str] = set()
        inspected_sheets: list[dict[str, object]] = []
        source_files: list[str] = []
        for source_index, (source_path, display_name) in enumerate(sources):
            source = Path(source_path)
            source_files.append(display_name)
            extension = source.suffix.casefold()
            if extension == ".xlsx":
                workbook = load_workbook(source, data_only=False, read_only=True, keep_links=False)
                for sheet in workbook.worksheets:
                    projected_name = _project_sheet_name(
                        display_name, sheet.title, used_names, primary=source_index == 0,
                    )
                    inspected_sheets.append({
                        **_inspect_sheet(sheet),
                        "name": projected_name,
                        "source_sheet": sheet.title,
                        "source_file": display_name,
                        "header_row": header_row,
                        "headers": [
                            " ".join(str(cell.value).split())
                            for cell in sheet[header_row]
                            if cell.value not in (None, "")
                        ] if sheet.max_row >= header_row else [],
                        "header_columns": {
                            " ".join(str(cell.value).split()): cell.column_letter
                            for cell in sheet[header_row]
                            if cell.value not in (None, "")
                        } if sheet.max_row >= header_row else {},
                    })
                workbook.close()
                continue
            if extension in {".csv", ".tsv"}:
                rows = csv.reader(
                    io.StringIO(_decode(source.read_bytes())),
                    delimiter="\t" if extension == ".tsv" else ",",
                )
                header: list[str] = []
                header_columns: dict[str, str] = {}
                row_count = 0
                max_columns = 0
                for row_count, row in enumerate(rows, start=1):
                    max_columns = max(max_columns, len(row))
                    if row_count == header_row:
                        header = [" ".join(str(value).split()) for value in row if value not in (None, "")]
                        header_columns = {
                            " ".join(str(value).split()): get_column_letter(index)
                            for index, value in enumerate(row, start=1)
                            if value not in (None, "")
                        }
                projected_name = _project_sheet_name(
                    display_name, "输入数据", used_names, primary=source_index == 0,
                )
                inspected_sheets.append({
                    "name": projected_name, "source_sheet": "输入数据", "source_file": display_name,
                    "rows": row_count, "columns": max_columns, "formula_count": 0,
                    "merged_range_count": 0, "hidden": False,
                    "header_row": header_row, "headers": header,
                    "header_columns": header_columns,
                })
                continue
            raise ValueError("spreadsheet_source_required")
        return {
            "format": "xlsx_bundle" if len(sources) > 1 else Path(sources[0][0]).suffix.removeprefix("."),
            "source_files": source_files,
            "sheets": inspected_sheets,
        }

    def process(
        self,
        source_path: str | Path,
        display_name: str,
        request: str,
        *,
        style: SpreadsheetStyleSpec | None = None,
        plan: SpreadsheetPlan | None = None,
        additional_sources: list[tuple[str | Path, str]] | None = None,
        validate_output: bool = True,
    ) -> SpreadsheetBuildResult:
        source = Path(source_path)
        requested_style = style
        style = style or SpreadsheetStyleSpec()
        extension = source.suffix.casefold()
        output = self.output_dir / f"sheet_{uuid4().hex}.xlsx"
        if extension == ".xlsx":
            copy2(source, output)
            workbook = load_workbook(output, data_only=False, keep_links=False)
            source_sheets = list(workbook.sheetnames)
        elif extension in {".csv", ".tsv"}:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "输入数据"
            text = _decode(source.read_bytes())
            rows = list(csv.reader(io.StringIO(text), delimiter="\t" if extension == ".tsv" else ","))
            for row in rows:
                sheet.append(row)
            _style_data_sheet(sheet, style)
            source_sheets = [sheet.title]
        else:
            raise ValueError("spreadsheet_source_required")

        source_sheet_map: list[dict[str, str]] = [
            {"source_file": display_name, "source_sheet": name, "workbook_sheet": name}
            for name in source_sheets
        ]
        if additional_sources:
            imported_names, imported_map = _import_sources(workbook, additional_sources)
            source_sheets.extend(imported_names)
            source_sheet_map.extend(imported_map)

        source_inspection = [_inspect_sheet(workbook[name]) for name in source_sheets]
        source_snapshot = self.diff_engine.snapshot(workbook)
        execution: SpreadsheetExecutionResult | None = None
        if plan is not None:
            execution = self.executor.execute(workbook, plan)
        return self._finalize_workbook(
            workbook,
            output,
            request=request,
            output_display_name=f"UniHR处理结果-{Path(display_name).stem}.xlsx",
            source_sheets=source_sheets,
            source_files=[display_name] + [name for _, name in (additional_sources or [])],
            source_sheet_map=source_sheet_map,
            source_inspection=source_inspection,
            source_snapshot=source_snapshot,
            style=style,
            apply_style=requested_style is not None,
            plan=plan,
            execution=execution,
            created_from_scratch=False,
            validate_output=validate_output,
        )

    def create(
        self,
        request: str,
        *,
        display_name: str,
        plan: SpreadsheetPlan,
        style: SpreadsheetStyleSpec | None = None,
        validate_output: bool = True,
    ) -> SpreadsheetBuildResult:
        """Create a reviewable workbook from a validated plan without a source attachment."""
        if not any(operation.type == "create_sheet" for operation in plan.operations):
            raise SpreadsheetPlanError("spreadsheet_create_sheet_required")
        workbook = Workbook()
        workbook.remove(workbook.active)
        output = self.output_dir / f"sheet_{uuid4().hex}.xlsx"
        execution = self.executor.execute(workbook, plan)
        applied_style = style or SpreadsheetStyleSpec(preserve_existing_style=False)
        safe_name = Path(display_name).name
        if not safe_name.casefold().endswith(".xlsx"):
            safe_name += ".xlsx"
        return self._finalize_workbook(
            workbook,
            output,
            request=request,
            output_display_name=safe_name,
            source_sheets=[],
            source_files=[],
            source_sheet_map=[],
            source_inspection=[],
            source_snapshot=WorkbookSnapshot(sheets={}, scanned_cells=0, truncated=False),
            style=applied_style,
            apply_style=True,
            plan=plan,
            execution=execution,
            created_from_scratch=True,
            validate_output=validate_output,
        )

    def validate(self, result: SpreadsheetBuildResult) -> SpreadsheetBuildResult:
        """Reopen and validate a built workbook as an independent stage."""
        metadata = result.metadata
        style = SpreadsheetStyleSpec.model_validate(metadata.get("style_spec") or {})
        target_sheets = [str(item) for item in metadata.get("style_applied_to") or []]
        validation = _validate_workbook(
            result.path,
            target_sheets,
            style,
            validate_style=bool(target_sheets),
        )
        execution_validation = dict(metadata.get("execution_validation") or {})
        if execution_validation:
            validation = {
                "passed": bool(validation["passed"] and execution_validation.get("passed")),
                "checks": {
                    **dict(validation["checks"]),
                    **dict(execution_validation.get("checks") or {}),
                },
            }
        metadata["validation"] = validation
        metadata["validation_status"] = "completed"
        return result

    def _finalize_workbook(
        self,
        workbook: Workbook,
        output: Path,
        *,
        request: str,
        output_display_name: str,
        source_sheets: list[str],
        source_files: list[str],
        source_sheet_map: list[dict[str, str]],
        source_inspection: list[dict[str, object]],
        source_snapshot: WorkbookSnapshot,
        style: SpreadsheetStyleSpec,
        apply_style: bool,
        plan: SpreadsheetPlan | None,
        execution: SpreadsheetExecutionResult | None,
        created_from_scratch: bool,
        validate_output: bool = True,
    ) -> SpreadsheetBuildResult:
        """Finalize styling, audit metadata, calculation flags, and validation in one seam."""
        data_sheets = [name for name in workbook.sheetnames if not name.startswith("UniHR")]
        target_sheets = _resolve_target_sheets(workbook, data_sheets, style.target_sheet)
        if apply_style:
            for target in target_sheets:
                _apply_style_spec(target, style)
        output_inspection = [_inspect_sheet(workbook[name]) for name in data_sheets]
        workbook_diff = self.diff_engine.compare(source_snapshot, self.diff_engine.snapshot(workbook))

        if "UniHR处理说明" in workbook.sheetnames:
            del workbook["UniHR处理说明"]
        note = workbook.create_sheet("UniHR处理说明", 0)
        note.sheet_view.showGridLines = False
        note.column_dimensions["A"].width = 22
        note.column_dimensions["B"].width = 72
        operation_summary = (
            f"已执行 {len(execution.operation_results)} 个白名单操作，改动 {execution.changed_cell_count} 个单元格"
            if execution is not None
            else "仅生成准备稿，未执行数据变换"
        )
        exception_summary = (
            f"发现 {len(execution.exceptions)} 项，详见“UniHR异常清单”"
            if execution is not None and execution.exceptions
            else "未发现字段映射或数据质量异常"
        )
        rows = [
            ("处理状态", "已从纯文字生成可复核工作簿" if created_from_scratch else "已生成可复核草稿，未覆盖原文件"),
            ("用户要求", request.strip()[:2_000]),
            ("来源文件", "无（纯文字创建）" if created_from_scratch else "、".join(source_files)),
            ("来源工作表", "无" if created_from_scratch else "、".join(source_sheets)),
            ("执行计划", operation_summary),
            ("异常结果", exception_summary),
            ("样式执行", "、".join(sheet.title for sheet in target_sheets) if apply_style else "沿用来源样式"),
            ("表头设置", f"第 {style.header_row} 行；冻结={'是' if style.freeze_header else '否'}；筛选={'是' if style.enable_filter else '否'}"),
            ("生成时间", datetime.now(timezone.utc).isoformat()),
            ("安全说明", "模型只生成受限计划；工作表结构、字段、公式和参数均由代码校验并留下审计记录。"),
        ]
        note.append(["UniHR 表格处理记录", ""])
        note.merge_cells("A1:B1")
        note["A1"].font = Font(name="微软雅黑", size=18, bold=True, color="FFFFFF")
        note["A1"].fill = PatternFill("solid", fgColor="2563EB")
        note["A1"].alignment = Alignment(vertical="center")
        note.row_dimensions[1].height = 32
        for label, value in rows:
            note.append([label, value])
        for row in note.iter_rows(min_row=2, max_row=note.max_row, min_col=1, max_col=2):
            row[0].font = Font(name="微软雅黑", size=11, bold=True, color="1F2937")
            row[0].fill = PatternFill("solid", fgColor="EFF6FF")
            row[1].font = Font(name="微软雅黑", size=11, color="111827")
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = Border(bottom=Side(style="thin", color="DDE3EA"))
            note.row_dimensions[row[0].row].height = max(24, 18 * ((_visual_width(str(row[1].value or "")) // 70) + 1))
        note.freeze_panes = "A2"
        for sheet in workbook.worksheets:
            _configure_print_layout(sheet)
        if workbook.calculation is None:
            workbook.calculation = CalcProperties()
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.save(output)
        workbook.close()
        execution_validation = execution.validation if execution is not None else {}
        validation: dict[str, object] = {
            "passed": False,
            "checks": {},
            "status": "pending",
        }
        result = SpreadsheetBuildResult(
            path=output,
            display_name=output_display_name,
            metadata={
                "source_sheets": source_sheets,
                "source_files": source_files,
                "source_sheet_map": source_sheet_map,
                "source_preserved": True,
                "created_from_scratch": created_from_scratch,
                "requires_human_review": True,
                "calculation_on_open": True,
                "inspection": source_inspection,
                "output_inspection": output_inspection,
                "workbook_diff": workbook_diff,
                "style_applied_to": [sheet.title for sheet in target_sheets] if apply_style else [],
                "style_spec": style.model_dump(mode="json"),
                "operation_plan": plan.model_dump(mode="json") if plan is not None else None,
                "operation_results": execution.operation_results if execution is not None else [],
                "exceptions": execution.exceptions if execution is not None else [],
                "changed_cell_count": execution.changed_cell_count if execution is not None else 0,
                "execution_validation": execution_validation,
                "validation": validation,
                "validation_status": "pending",
            },
        )
        return self.validate(result) if validate_output else result


def _configure_print_layout(sheet) -> None:
    """Keep generated previews readable without changing business cell values."""
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(
        fitToPage=True,
        autoPageBreaks=False,
    )
    sheet.page_setup.orientation = "landscape" if sheet.max_column >= 8 else "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_options.horizontalCentered = True
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.45
    sheet.page_margins.bottom = 0.45
    sheet.print_area = f"A1:{get_column_letter(max(1, sheet.max_column))}{max(1, sheet.max_row)}"


def _import_sources(
    workbook: Workbook,
    sources: list[tuple[str | Path, str]],
) -> tuple[list[str], list[dict[str, str]]]:
    imported_names: list[str] = []
    imported_map: list[dict[str, str]] = []
    used_names = set(workbook.sheetnames)
    for source_path, display_name in sources:
        source = Path(source_path)
        extension = source.suffix.casefold()
        if extension == ".xlsx":
            incoming = load_workbook(source, data_only=False, keep_links=False)
            name_map = {
                sheet.title: _project_sheet_name(display_name, sheet.title, used_names, primary=False)
                for sheet in incoming.worksheets
            }
            for sheet in incoming.worksheets:
                target_name = name_map[sheet.title]
                target = workbook.create_sheet(target_name)
                _copy_sheet_contents(sheet, target, name_map)
                imported_names.append(target_name)
                imported_map.append({
                    "source_file": display_name,
                    "source_sheet": sheet.title,
                    "workbook_sheet": target_name,
                })
            incoming.close()
            continue
        if extension in {".csv", ".tsv"}:
            target_name = _project_sheet_name(display_name, "输入数据", used_names, primary=False)
            target = workbook.create_sheet(target_name)
            rows = csv.reader(
                io.StringIO(_decode(source.read_bytes())),
                delimiter="\t" if extension == ".tsv" else ",",
            )
            for row in rows:
                target.append(row)
            imported_names.append(target_name)
            imported_map.append({
                "source_file": display_name,
                "source_sheet": "输入数据",
                "workbook_sheet": target_name,
            })
            continue
        raise ValueError("spreadsheet_source_required")
    return imported_names, imported_map


def _copy_sheet_contents(
    source: object,
    target: object,
    sheet_name_map: dict[str, str],
) -> None:
    for row in source.iter_rows():  # type: ignore[attr-defined]
        for source_cell in row:
            target_cell = target.cell(source_cell.row, source_cell.column)  # type: ignore[attr-defined]
            value = source_cell.value
            if isinstance(value, str) and value.startswith("="):
                value = _rewrite_internal_sheet_references(value, sheet_name_map)
            target_cell.value = value
            if source_cell.has_style:
                target_cell._style = copy(source_cell._style)
            if source_cell.number_format:
                target_cell.number_format = source_cell.number_format
            if source_cell.hyperlink:
                target_cell._hyperlink = copy(source_cell.hyperlink)
            if source_cell.comment:
                target_cell.comment = copy(source_cell.comment)
    for key, dimension in source.column_dimensions.items():  # type: ignore[attr-defined]
        target.column_dimensions[key] = copy(dimension)  # type: ignore[attr-defined]
    for key, dimension in source.row_dimensions.items():  # type: ignore[attr-defined]
        target.row_dimensions[key] = copy(dimension)  # type: ignore[attr-defined]
    for merged_range in source.merged_cells.ranges:  # type: ignore[attr-defined]
        target.merge_cells(str(merged_range))  # type: ignore[attr-defined]
    target.freeze_panes = source.freeze_panes  # type: ignore[attr-defined]
    target.sheet_view.showGridLines = source.sheet_view.showGridLines  # type: ignore[attr-defined]
    target.sheet_state = source.sheet_state  # type: ignore[attr-defined]
    target.auto_filter.ref = source.auto_filter.ref  # type: ignore[attr-defined]
    target.sheet_format = copy(source.sheet_format)  # type: ignore[attr-defined]
    target.sheet_properties = copy(source.sheet_properties)  # type: ignore[attr-defined]
    target.page_margins = copy(source.page_margins)  # type: ignore[attr-defined]
    target.page_setup = copy(source.page_setup)  # type: ignore[attr-defined]
    target.print_options = copy(source.print_options)  # type: ignore[attr-defined]


def _rewrite_internal_sheet_references(formula: str, sheet_name_map: dict[str, str]) -> str:
    rewritten = formula
    for source_name, target_name in sorted(sheet_name_map.items(), key=lambda item: len(item[0]), reverse=True):
        quoted_source = "'" + source_name.replace("'", "''") + "'!"
        quoted_target = "'" + target_name.replace("'", "''") + "'!"
        rewritten = rewritten.replace(quoted_source, quoted_target)
        if re.fullmatch(r"[A-Za-z0-9_.\u3400-\u9fff]+", source_name):
            rewritten = re.sub(
                rf"(?<!['A-Za-z0-9_\u3400-\u9fff]){re.escape(source_name)}!",
                quoted_target,
                rewritten,
            )
    return rewritten


def _project_sheet_name(
    display_name: str,
    source_sheet: str,
    used_names: set[str],
    *,
    primary: bool,
) -> str:
    if primary:
        base = source_sheet
    else:
        prefix = re.sub(r"[\\/*?:\[\]]", "", Path(display_name).stem).strip()[:10] or "附件"
        clean_sheet = re.sub(r"[\\/*?:\[\]]", "", source_sheet).strip() or "工作表"
        base = f"{prefix}-{clean_sheet}"
    base = base[:31] or "工作表"
    candidate = base
    suffix = 2
    while candidate in used_names:
        marker = f"-{suffix}"
        candidate = base[: 31 - len(marker)] + marker
        suffix += 1
    used_names.add(candidate)
    return candidate


def _style_data_sheet(sheet: object, style: SpreadsheetStyleSpec | None = None) -> None:
    style = style or SpreadsheetStyleSpec()
    if sheet.max_row < 1 or sheet.max_column < 1:  # type: ignore[attr-defined]
        return
    sheet.freeze_panes = f"A{style.header_row + 1}" if style.freeze_header else None  # type: ignore[attr-defined]
    sheet.sheet_view.showGridLines = False  # type: ignore[attr-defined]
    header = sheet[style.header_row]  # type: ignore[index]
    for cell in header:
        cell.font = Font(name=style.header_font, size=11, bold=True, color=style.header_text_color)
        cell.fill = PatternFill("solid", fgColor=style.header_fill)
        cell.alignment = Alignment(horizontal=style.horizontal_alignment, vertical=style.vertical_alignment)
    for column in sheet.columns:  # type: ignore[attr-defined]
        values = [str(cell.value or "") for cell in column[:80]]
        width = min(44, max(12, max((_visual_width(value) for value in values), default=8) + 4))
        sheet.column_dimensions[column[0].column_letter].width = width  # type: ignore[attr-defined]
    if sheet.max_row > style.header_row:  # type: ignore[attr-defined]
        ref = f"A{style.header_row}:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}"  # type: ignore[attr-defined]
        table = Table(displayName="InputDataTable", ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False,
        )
        sheet.add_table(table)  # type: ignore[attr-defined]


def _resolve_target_sheets(workbook: object, source_sheets: list[str], target_sheet: str) -> list[object]:
    if target_sheet.strip():
        if target_sheet not in workbook.sheetnames:  # type: ignore[attr-defined]
            raise ValueError("spreadsheet_target_sheet_not_found")
        return [workbook[target_sheet]]  # type: ignore[index]
    return [workbook[name] for name in source_sheets]  # type: ignore[index]


def _apply_style_spec(sheet: object, style: SpreadsheetStyleSpec) -> None:
    if sheet.max_row < style.header_row or sheet.max_column < 1:  # type: ignore[attr-defined]
        return
    sheet.sheet_view.showGridLines = False  # type: ignore[attr-defined]
    sheet.freeze_panes = f"A{style.header_row + 1}" if style.freeze_header else None  # type: ignore[attr-defined]
    if style.enable_filter and not sheet.tables:  # type: ignore[attr-defined]
        # An Excel Table already owns its AutoFilter.  Adding a second worksheet
        # AutoFilter over the same range creates a package that openpyxl can
        # reopen but desktop Excel rejects (Workbooks.Open 0x800A03EC).
        sheet.auto_filter.ref = f"A{style.header_row}:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}"  # type: ignore[attr-defined]
    elif style.enable_filter:
        sheet.auto_filter.ref = None  # type: ignore[attr-defined]
    else:
        sheet.auto_filter.ref = None  # type: ignore[attr-defined]
    thin = Side(style="thin", color="DDE3EA")
    for cell in sheet[style.header_row]:  # type: ignore[index]
        cell.font = Font(name=style.header_font, size=11, bold=True, color=style.header_text_color)
        cell.fill = PatternFill("solid", fgColor=style.header_fill)
        cell.alignment = Alignment(
            horizontal=style.horizontal_alignment,
            vertical=style.vertical_alignment,
            wrap_text=True,
        )
        cell.border = Border(bottom=thin)
    for row in sheet.iter_rows(min_row=style.header_row + 1):  # type: ignore[attr-defined]
        for cell in row:
            cell.font = Font(name=style.body_font, size=11, color="111827")
            cell.alignment = Alignment(
                horizontal=style.horizontal_alignment,
                vertical=style.vertical_alignment,
                wrap_text=True,
            )
            if not style.preserve_existing_style:
                cell.border = Border(bottom=Side(style="hair", color="E5E7EB"))
    if style.auto_width:
        for column in sheet.columns:  # type: ignore[attr-defined]
            values = [str(cell.value or "") for cell in column[:100]]
            width = min(48, max(10, max((_visual_width(value) for value in values), default=8) + 4))
            sheet.column_dimensions[column[0].column_letter].width = width  # type: ignore[attr-defined]


def _inspect_sheet(sheet: object) -> dict[str, object]:
    max_row = getattr(sheet, "max_row", None)
    max_column = getattr(sheet, "max_column", None)
    if max_row is None or max_column is None:
        calculate_dimension = getattr(sheet, "calculate_dimension", None)
        if callable(calculate_dimension):
            try:
                calculate_dimension(force=True)
            except TypeError:
                calculate_dimension()
        max_row = getattr(sheet, "max_row", None)
        max_column = getattr(sheet, "max_column", None)
    formula_count = sum(
        1
        for row in sheet.iter_rows()  # type: ignore[attr-defined]
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    )
    merged_ranges = getattr(getattr(sheet, "merged_cells", None), "ranges", ())
    return {
        "name": sheet.title,  # type: ignore[attr-defined]
        "rows": int(max_row or 0),
        "columns": int(max_column or 0),
        "formula_count": formula_count,
        "merged_range_count": len(merged_ranges),
        "hidden": sheet.sheet_state != "visible",  # type: ignore[attr-defined]
    }


def _validate_workbook(
    path: Path,
    target_sheets: list[str],
    style: SpreadsheetStyleSpec,
    *,
    validate_style: bool = True,
) -> dict[str, object]:
    workbook = load_workbook(path, data_only=False, keep_links=False)
    checks: dict[str, bool] = {
        "file_reopens": True,
        "audit_sheet_present": "UniHR处理说明" in workbook.sheetnames,
        "target_sheets_present": all(name in workbook.sheetnames for name in target_sheets),
    }
    for name in target_sheets:
        sheet = workbook[name]
        if sheet.max_row >= style.header_row:
            checks[f"{name}:header_exists"] = any(
                cell.value not in (None, "") for cell in sheet[style.header_row]
            )
            if validate_style:
                checks[f"{name}:freeze_matches"] = (
                    str(sheet.freeze_panes or "") == (f"A{style.header_row + 1}" if style.freeze_header else "")
                )
    workbook.close()
    return {"passed": all(checks.values()), "checks": checks}


def _decode(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _visual_width(value: str) -> int:
    return sum(2 if unicodedata.east_asian_width(character) in {"W", "F", "A"} else 1 for character in value)
