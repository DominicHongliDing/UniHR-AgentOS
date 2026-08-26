from __future__ import annotations

import re
from difflib import SequenceMatcher
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.domains.artifacts.documents import section_usable_width_cm, table_width_cm
from app.domains.task_briefs import DocumentStyleSpec

from .contracts import (
    DocumentBrief,
    DocumentFact,
    DocumentPlan,
    DocumentValidationReport,
    EvidenceSource,
    FactStatus,
)
from .word_verification import WordAutomationVerifier


_CRITICAL_TOKEN = re.compile(
    r"(?<!\d)(?:20\d{2}[年./-](?:0?[1-9]|1[0-2])(?:[月./-](?:0?[1-9]|[12]\d|3[01])日?)?|"
    r"\d+(?:\.\d+)?%|\d+(?:\.\d+)?(?:万元|元|人|项|份|天|小时|分钟|周|月|年|次|场|页|个|倍|时))(?!\d)"
)


class DocumentContentHarness:
    """Validate source attribution and brief compliance without judging writing style."""

    def validate(
        self,
        plan: DocumentPlan,
        brief: DocumentBrief,
        evidence_sources: Iterable[EvidenceSource],
    ) -> tuple[DocumentPlan, DocumentValidationReport]:
        source_map = {source.name: source.text for source in evidence_sources}
        verified_facts = [self._verify_fact(fact, source_map) for fact in plan.facts]
        plan = plan.model_copy(update={"facts": verified_facts})
        text = plan.full_text()
        table_headers = {
            _normalize(header)
            for table in plan.tables
            for header in table.headers
            if header.strip()
        }
        checks = {
            "title_present": bool(plan.title.strip()),
            "body_present": bool(plan.lead.strip() or any(section.paragraphs for section in plan.sections)),
            "required_points_covered": all(_covered(point, text) for point in brief.required_points),
            "prohibited_content_absent": all(item not in text for item in brief.prohibited_content),
            "fact_evidence_verified": all(
                fact.status in {FactStatus.SUPPORTED, FactStatus.MISSING}
                for fact in verified_facts
            ),
            "required_tables_present": not brief.requires_table or bool(plan.tables),
            "required_table_headers_present": not brief.required_table_headers or all(
                _normalize(header) in table_headers for header in brief.required_table_headers
            ),
            "required_blank_rows_present": not brief.required_blank_rows or any(
                table.blank_rows >= brief.required_blank_rows for table in plan.tables
            ),
        }
        missing = [
            fact.label or fact.key or fact.value
            for fact in verified_facts
            if fact.required and fact.status != FactStatus.SUPPORTED
        ]
        unsupported_tokens = self._unsupported_critical_tokens(text, source_map, verified_facts)
        substantive_facts = [
            fact for fact in verified_facts
            if fact.key not in {
                "document_purpose", "table_headers", "required_points", "tone", "table_data",
            }
        ]
        # Older model fixtures and imported legacy drafts may not have a fact ledger yet.
        # Brief-contract facts (purpose, tone, table schema) make requirements visible
        # but must not by themselves switch on the strict value-attribution gate.
        # Only substantive business facts do that.
        checks["critical_values_attributed"] = not unsupported_tokens or not substantive_facts
        issues: list[str] = []
        if not checks["title_present"]:
            issues.append("document_title_missing")
        if not checks["body_present"]:
            issues.append("document_body_missing")
        if not checks["required_points_covered"]:
            issues.append("required_points_not_covered")
        if not checks["prohibited_content_absent"]:
            issues.append("prohibited_content_present")
        if not checks["fact_evidence_verified"]:
            issues.append("fact_evidence_unverified")
        if not checks["required_tables_present"]:
            issues.append("required_table_missing")
        if not checks["required_table_headers_present"]:
            issues.append("required_table_headers_missing")
        if not checks["required_blank_rows_present"]:
            issues.append("required_blank_rows_missing")
        if unsupported_tokens:
            issues.append("critical_values_without_source")
        placeholders = text.count("[待补充：") + text.count("【待补充：")
        warnings = [f"missing_required_fact:{item}" for item in missing]
        if unsupported_tokens and not substantive_facts:
            warnings.append("fact_ledger_missing_for_critical_values")
        if placeholders:
            warnings.append("document_contains_visible_placeholders")
        passed = all(checks.values())
        return plan, DocumentValidationReport(
            passed=passed,
            status="passed" if passed and not warnings else "warning" if passed else "failed",
            checks=checks,
            issues=issues,
            warnings=warnings,
            missing_facts=missing,
            metrics={
                "fact_count": len(verified_facts),
                "supported_fact_count": sum(fact.status == FactStatus.SUPPORTED for fact in verified_facts),
                "required_point_count": len(brief.required_points),
                "placeholder_count": placeholders,
                "unattributed_critical_value_count": len(unsupported_tokens),
                "table_count": len(plan.tables),
                "table_header_count": len(table_headers),
            },
            details={"unattributed_critical_values": unsupported_tokens},
        )

    @staticmethod
    def _verify_fact(fact: DocumentFact, source_map: dict[str, str]) -> DocumentFact:
        if fact.status == FactStatus.MISSING or not fact.value.strip():
            return fact.model_copy(update={"status": FactStatus.MISSING})
        source_name = fact.source_name
        source = source_map.get(source_name)
        resolved_alias = False
        if source is None:
            target = _normalize(fact.source_name)
            compatible = [
                (name, text) for name, text in source_map.items()
                if len(target) >= 8
                and (target in _normalize(name) or _normalize(name) in target)
            ]
            if compatible:
                source_name, source = compatible[0]
                resolved_alias = True
        if source is None:
            # Models occasionally name a handoff block (for example “用户补充”)
            # instead of the actual EvidenceSource (“用户需求卡”). Resolve that
            # alias from the quoted evidence/value itself. This keeps attribution
            # strict while avoiding futile model repair loops over a label only.
            excerpt = _normalize(fact.evidence_excerpt)
            value = _normalize(fact.value)
            evidence_matches = [
                (name, text) for name, text in source_map.items()
                if (excerpt and excerpt in _normalize(text))
                or (value and len(value) >= 2 and value in _normalize(text))
            ]
            if evidence_matches:
                source_name, source = evidence_matches[0]
                resolved_alias = True
        if source is None:
            return fact.model_copy(update={"status": FactStatus.UNSUPPORTED})
        excerpt = _normalize(fact.evidence_excerpt or fact.value)
        value = _normalize(fact.value)
        normalized_source = _normalize(source)
        supported = bool(
            (excerpt and excerpt in normalized_source)
            or (resolved_alias and value and value in normalized_source)
        )
        return fact.model_copy(update={
            "source_name": source_name,
            "status": FactStatus.SUPPORTED if supported else FactStatus.UNSUPPORTED,
        })

    @staticmethod
    def _unsupported_critical_tokens(
        text: str,
        source_map: dict[str, str],
        facts: list[DocumentFact],
    ) -> list[str]:
        allowed = _normalize("\n".join([*source_map.values(), *(fact.value for fact in facts)]))
        unsupported: list[str] = []
        for token in _CRITICAL_TOKEN.findall(text):
            if _normalize(token) not in allowed and token not in unsupported:
                unsupported.append(token)
        return unsupported


class DocumentLayoutHarness:
    """Prove DOCX structure and, when available, Word rendering."""

    def __init__(self, verifier: WordAutomationVerifier | None = None) -> None:
        self.verifier = verifier

    def validate(self, path: str | Path, style: DocumentStyleSpec) -> DocumentValidationReport:
        artifact = Path(path)
        issues: list[str] = []
        warnings: list[str] = []
        checks: dict[str, bool] = {
            "file_exists": artifact.is_file(),
            "file_nonempty": artifact.is_file() and artifact.stat().st_size > 0,
        }
        metrics: dict[str, int | float | str | bool] = {}
        details: dict[str, object] = {"preview_files": []}
        if not all(checks.values()):
            return DocumentValidationReport(
                passed=False, status="failed", checks=checks,
                issues=[key for key, value in checks.items() if not value],
            )
        try:
            document = Document(artifact)
            section = document.sections[0]
            title = document.paragraphs[0] if document.paragraphs else None
            normal = document.styles["Normal"]
            checks.update({
                "package_opens": True,
                "page_is_a4": round(section.page_width.cm, 1) == 21.0
                and round(section.page_height.cm, 1) == 29.7,
                "margins_match": all(abs(actual - expected) <= 0.11 for actual, expected in (
                    (section.top_margin.cm, style.top_margin_cm),
                    (section.bottom_margin.cm, style.bottom_margin_cm),
                    (section.left_margin.cm, style.left_margin_cm),
                    (section.right_margin.cm, style.right_margin_cm),
                )),
                "title_present": bool(title and title.text.strip()),
                "title_alignment_matches": bool(
                    title and title.alignment == _alignment(style.title_alignment)
                ),
                "body_size_matches": bool(
                    normal.font.size and abs(normal.font.size.pt - style.body_size_pt) <= 0.1
                ),
                "body_alignment_matches": normal.paragraph_format.alignment == _alignment(
                    style.body_alignment
                ),
            })
            metrics["paragraph_count"] = len(document.paragraphs)
            metrics["section_count"] = len(document.sections)
            metrics["table_count"] = len(document.tables)
            usable_width_cm = section_usable_width_cm(section)
            table_widths_cm = [table_width_cm(table) for table in document.tables]
            metrics["usable_width_cm"] = round(usable_width_cm, 3)
            metrics["max_table_width_cm"] = round(max(table_widths_cm, default=0.0), 3)
            table_structure_valid = all(
                len(table.rows) >= 1
                and len(table.columns) >= 1
                and all(len(row.cells) == len(table.columns) for row in table.rows)
                and all(cell.text.strip() for cell in table.rows[0].cells)
                for table in document.tables
            )
            checks["table_structure_valid"] = table_structure_valid
            checks["tables_within_page_width"] = all(
                not width or width <= usable_width_cm + 0.05
                for width in table_widths_cm
            )
        except (OSError, ValueError, KeyError) as exc:
            checks["package_opens"] = False
            issues.append(f"document_open_failed:{type(exc).__name__}")

        if self.verifier is None or not self.verifier.available:
            warnings.append("word_automation_unavailable")
            metrics["render_engine"] = "unavailable"
        else:
            rendered = self.verifier.verify(artifact)
            checks.update(rendered.checks)
            metrics.update(rendered.metrics)
            issues.extend(rendered.issues)
            details["preview_files"] = rendered.preview_files
        required_checks = {
            key: value for key, value in checks.items()
            if key != "word_automation_available"
        }
        passed = all(required_checks.values())
        return DocumentValidationReport(
            passed=passed,
            status="passed" if passed and not warnings else "warning" if passed else "failed",
            checks=checks,
            issues=list(dict.fromkeys(issues)),
            warnings=list(dict.fromkeys(warnings)),
            metrics=metrics,
            details=details,
        )


def merge_validation_reports(
    content: DocumentValidationReport,
    layout: DocumentValidationReport,
) -> DocumentValidationReport:
    passed = content.passed and layout.passed
    warnings = list(dict.fromkeys([*content.warnings, *layout.warnings]))
    return DocumentValidationReport(
        passed=passed,
        status="passed" if passed and not warnings else "warning" if passed else "failed",
        checks={
            **{f"content.{key}": value for key, value in content.checks.items()},
            **{f"layout.{key}": value for key, value in layout.checks.items()},
        },
        issues=list(dict.fromkeys([*content.issues, *layout.issues])),
        warnings=warnings,
        missing_facts=content.missing_facts,
        metrics={
            **{f"content.{key}": value for key, value in content.metrics.items()},
            **{f"layout.{key}": value for key, value in layout.metrics.items()},
        },
        details={"content": content.details, "layout": layout.details},
    )


def _normalize(value: str) -> str:
    return re.sub(r"[\s，。；：、,.!?！？:;（）()【】\[\]‘’“”\-_]", "", value).casefold()


def _covered(point: str, text: str) -> bool:
    expected = _normalize(point)
    actual = _normalize(text)
    if not expected:
        return True
    if expected in actual:
        return True
    # A formal draft will often rewrite “会议时间为……” as “会议定于……召开”.
    # Compare against individual sentences rather than the whole document, while
    # still requiring every date/amount/count anchor from the locked brief.
    critical_tokens = [_normalize(token) for token in _CRITICAL_TOKEN.findall(point)]
    if any(token and token not in actual for token in critical_tokens):
        return False
    segments = [
        _normalize(item)
        for item in re.split(r"[。！？!?；;\n\r]+", text)
        if _normalize(item)
    ]
    return any(
        SequenceMatcher(None, expected, segment).ratio() >= 0.56
        for segment in segments
    )


def _alignment(value: str):
    return {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }.get(value, WD_ALIGN_PARAGRAPH.JUSTIFY)
